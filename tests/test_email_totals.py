"""
Dollar totals in the proposal emails (the internal AE email, the client-facing email, the Step 08 revision, the Word
doc and the Gamma outline that carry them).

The figures are computed in ONE place — app/services/email_totals.py — from the CURRENT plan, equal the Excel export,
go into the prompts as the only amounts the model may state, and every dollar amount in the finished emails is checked
against them. What is pinned here:

  * the helper per billing period (week / month / quarter / full flight), several options, add-ons, customised per-period
    allocations (and merged periods), the agency fee, mixed period counts, added-value lines;
  * parity with the REAL workbook (generate_proposal -> openpyxl cells): the period multiplier, the ROUNDDOWNs, the add-on
    SUMIF, the gross mark-up and the per-period breakdown row all come out of the helper identically;
  * the prompts (first generation and the revision) carry those figures and the unit's own wording;
  * the checker: a wrong total is caught, correct text (per-period, per-option, add-on figures) is accepted, statistics /
    unit rates / odd formatting never raise a false alarm or a crash, the corrective pass repairs or warns and never
    blocks the email;
  * the endpoints and the SPA feed the CURRENT state (every option, add-ons, allocations, unit) into every generation /
    revision, and the SPA flags emails the plan has since moved away from.

No OpenAI call is made anywhere: the model is a stub. What a live model does with these instructions (does it obey the
authoritative block? does the corrective pass fix what it is shown?) can only be seen with a real key.
"""
from __future__ import annotations

import asyncio
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from types import SimpleNamespace

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import openpyxl
import pytest

from app import main
from app.services import ai_enricher
from app.services import email_totals as et
from app.services import monthly_allocation as mo
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import AddonItem, LineItem, generate_proposal

ROOT = Path(__file__).resolve().parent.parent
APP_JS = (ROOT / "app" / "static" / "app.js").read_text(encoding="utf-8")
INDEX_HTML = (ROOT / "app" / "templates" / "index.html").read_text(encoding="utf-8")

EMAIL = "Email Campaigns - Display Re-targeting"
IAB = "eDigital Network Display - Standard IAB"
GEOFENCE = "Display - Geo Fence"


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.catalog.resolve_product_alias", lambda name: None)
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


def _req(**kw) -> ProposalRequest:
    fields = dict(client_name="Acme", request_type="New Business", start_date="2026-10-01", end_date="2026-12-31",
                  total_months=3, agency_fee=None, requested_by="Pat Planner", salesperson_email="pat@example.com")
    fields.update(kw)
    return ProposalRequest(**fields)


def _li(name, budget, months=3, **kw) -> LineItem:
    return LineItem(product_name=name, monthly_budget=budget, months=months, id=kw.pop("id", name), **kw)


def _opt(request, lines, unit="month", addons=None, **tier):
    """The one option's figures for `lines`."""
    totals = et.compute_email_totals(request, [dict({"label": "A", "line_items": lines}, **tier)], addons, unit)
    return totals.options[0]


def _block(request, tiers, addons=None, unit="month"):
    return et.render_prompt_block(et.compute_email_totals(request, tiers, addons, unit))


# ===========================================================================================================
# The helper
# ===========================================================================================================

def test_month_totals_floor_the_per_period_total_then_multiply_like_the_workbook():
    # $2,000.40 + $1,500.80 = $3,501.20 a month: the sheet's TOTAL DIGITAL cell is ROUNDDOWN(.., 0) = 3,501, and the
    # campaign row multiplies THAT by the months cell: 10,503 (not 3 x 3,501.20 = 10,503.60, which prints as 10,504).
    opt = _opt(_req(), [_li("Display Alpha", 2000.40), _li("Search Beta", 1500.80)])
    assert (opt.unit, opt.n_periods) == ("month", 3)
    assert opt.per_period_net == 3501 and opt.media_net == 10503 and opt.total_net == 10503
    assert [round(l.total, 2) for l in opt.lines] == [6001.20, 4502.40]
    assert not opt.addons and opt.addons_net == 0 and opt.periods is None and not opt.custom_allocation


def test_week_totals_are_per_week_over_the_weeks_of_the_flight():
    weeks = len(mo.periods_between(mo.parse_flexible_date("2026-10-01"), mo.parse_flexible_date("2026-12-31"), "week"))
    assert weeks == 14
    lines = [_li("Display Alpha", 500.0, months=weeks), _li("Search Beta", 250.0, months=weeks)]
    opt = _opt(_req(total_months=3), lines, "week")
    assert (opt.unit, opt.n_periods) == ("week", 14)
    assert opt.per_period_net == 750 and opt.media_net == 10500 and opt.total_net == 10500
    block = _block(_req(total_months=3), [{"label": "A", "line_items": lines}], unit="week")
    assert "$500 per week × 14 weeks = $7,000" in block
    assert "All lines together, per week (net): $750" in block
    assert "Billing period: WEEKLY" in block and "ONE WEEK" in block
    assert "per month" not in block and "/month" not in block


def test_quarter_totals_are_per_quarter():
    lines = [_li("Display Alpha", 6000.0, months=2)]
    request = _req(start_date="2026-10-01", end_date="2027-03-31", total_months=6)
    opt = _opt(request, lines, "quarter")
    assert (opt.unit, opt.n_periods, opt.per_period_net, opt.media_net) == ("quarter", 2, 6000, 12000)
    block = _block(request, [{"label": "A", "line_items": lines}], unit="quarter")
    assert "$6,000 per quarter × 2 quarters = $12,000" in block and "Billing period: QUARTERLY" in block
    assert "per month" not in block


def test_full_flight_is_one_flat_flight_total_with_nothing_multiplied():
    stale = _li("Stale Gamma", 1000.0, months=3, monthly_allocations={"2026-10": 1.0, "2026-11": 1.0})   # a month-mode payload
    lines = [_li("Display Alpha", 5000.0, months=1), _li("Search Beta", 7750.0, months=1), stale]
    opt = _opt(_req(), lines, "full_flight")
    assert (opt.n_periods, opt.full_flight) == (1, True)
    assert stale.monthly_allocations is not None                              # the caller's objects are never mutated
    # the stale line is FOLDED into one period (1,000 x 3), exactly like proposal_generator.normalize_full_flight_lines
    assert [(l.per_period, l.periods, l.total) for l in opt.lines] == [(5000, 1, 5000), (7750, 1, 7750), (3000, 1, 3000)]
    assert opt.per_period_net == 15750 and opt.media_net == 15750 and opt.total_net == 15750
    assert opt.periods is None and not opt.custom_allocation                  # no per-period split exists
    block = et.render_prompt_block(et.compute_email_totals(_req(), [{"label": "A", "line_items": lines}], None, "full_flight"))
    assert "$5,000 flat for the whole flight" in block and "Billing period: FULL FLIGHT" in block
    assert "per month" not in block and "month-by-month" not in block.lower() and "×" not in block


def test_options_are_alternatives_each_with_its_own_totals():
    tiers = [
        {"label": "A", "name": "Independent", "line_items": [_li("Display Alpha", 2000.0)]},
        {"label": "B", "name": "", "line_items": [_li("Display Alpha", 4000.0), _li("Search Beta", 1000.0)]},
    ]
    totals = et.compute_email_totals(_req(), tiers, None, "month")
    a, b = totals.options
    assert totals.multi_option and (a.display_name, b.display_name) == ("Independent", "Option B")
    assert (a.total_net, b.total_net) == (6000, 15000)
    block = et.render_prompt_block(totals)
    assert "INDEPENDENT (option A)" in block and "\nOPTION B\n" in block
    assert "ALTERNATIVES" in block
    # the same product costs different amounts in each option — each option's lines say so
    assert "$2,000 per month × 3 months = $6,000" in block and "$4,000 per month × 3 months = $12,000" in block
    # nothing combines the options (6,000 + 15,000 = 21,000)
    assert 21000.0 not in et.allowed_amounts(totals)
    assert "$21,000" not in block


def test_add_ons_enter_the_campaign_total_but_free_ones_do_not():
    addons = [SimpleNamespace(product_name="Landing Page", amount=100.0), AddonItem("Call Tracking", 250.5),
              {"product_name": "Brand Lift", "amount": 0.0}]
    opt = _opt(_req(), [_li("Display Alpha", 2000.0)], addons=addons)
    assert [a.name for a in opt.addons] == ["Landing Page", "Call Tracking"] and opt.free_addons == ["Brand Lift"]
    assert opt.addons_net == pytest.approx(350.5)
    assert opt.media_net == 6000 and opt.total_net == 6350          # ROUNDDOWN(6,000 + 350.5)
    block = _block(_req(), [{"label": "A", "line_items": [_li("Display Alpha", 2000.0)]}], addons)
    assert "One-time add-ons (net, the same on every option): Landing Page $100; Call Tracking $250.50" in block
    assert "Add-ons at no charge: Brand Lift" in block
    assert "TOTAL INVESTMENT, net — the Excel's campaign grand-total row (media + one-time add-ons): $6,350" in block
    # every option carries the same add-ons
    two = et.compute_email_totals(_req(), [{"label": "A", "line_items": [_li("X", 1000.0)]},
                                           {"label": "B", "line_items": [_li("X", 2000.0)]}], addons, "month")
    assert [o.addons_net for o in two.options] == [pytest.approx(350.5)] * 2
    assert [o.total_net for o in two.options] == [3350, 6350]


def test_added_value_lines_cost_nothing_and_are_in_no_total():
    lines = [_li("Display Alpha", 2000.0), _li("Social Boost", 999.0, is_added_value=True)]
    opt = _opt(_req(), lines)
    assert opt.per_period_net == 2000 and opt.total_net == 6000
    assert [l.is_added_value for l in opt.lines] == [False, True]
    assert opt.lines[1].per_period == 0 and opt.lines[1].total == 0
    assert "Social Boost: added value — no cost to the client, and in no total" in et.render_prompt_block(
        et.compute_email_totals(_req(), [{"label": "A", "line_items": lines}], None, "month"))


def test_a_customised_breakdown_gives_every_period_its_own_amount():
    custom = _li("Display Alpha", 4000.0, monthly_allocations={"2026-10": 6000.0, "2026-11": 3000.0, "2026-12": 3000.0})
    plain = _li("Search Beta", 1500.0)                                        # never customised: its default (even) split
    opt = _opt(_req(), [custom, plain])
    assert opt.custom_allocation and opt.varies_by_period
    assert [(p.label, p.net) for p in opt.periods] == [("October 2026", 7500.0), ("November 2026", 4500.0), ("December 2026", 4500.0)]
    assert sum(p.net for p in opt.periods) == opt.media_net == 16500
    block = et.render_prompt_block(et.compute_email_totals(_req(), [{"label": "A", "line_items": [custom, plain]}], None, "month"))
    assert "Month-by-month total, all lines (net): October 2026 $7,500 · November 2026 $4,500 · December 2026 $4,500" in block
    assert "never present one as a flat monthly figure" in block
    assert "$12,000 for the flight over 3 months (the amount differs by month" in block      # not "$4,000 per month"
    assert "All lines together, per month" not in block
    assert "$4,000 per month" not in block


def test_the_default_split_honors_the_plans_distribution_mode_for_uncustomised_lines():
    request = _req(start_date="2026-10-15", end_date="2026-12-31")                    # a half-active first month
    custom = _li("Display Alpha", 1000.0, monthly_allocations={"2026-10": 1000.0, "2026-11": 1000.0, "2026-12": 1000.0})
    plain = _li("Search Beta", 3000.0)
    tiers = [{"label": "A", "line_items": [custom, plain]}]
    even = et.compute_email_totals(request, tiers, None, "month", "even").options[0]
    prorated = et.compute_email_totals(request, tiers, None, "month", "prorated").options[0]
    assert [p.net for p in even.periods] == [4000.0, 4000.0, 4000.0]
    assert prorated.periods[0].net < even.periods[0].net < prorated.periods[2].net
    assert sum(p.net for p in prorated.periods) == pytest.approx(12000.0)


def test_merged_periods_are_one_bucket():
    custom = _li("Display Alpha", 4000.0, monthly_allocations={"2026-10+2026-11": 9000.0, "2026-12": 3000.0})
    opt = _opt(_req(), [custom], period_merge_groups=[["2026-10", "2026-11"]])
    assert [(p.label, p.net) for p in opt.periods] == [("October–November 2026", 9000.0), ("December 2026", 3000.0)]


def test_stale_allocation_keys_are_never_quoted():
    stale = _li("Display Alpha", 4000.0, monthly_allocations={"W1-2026-10-01": 12000.0})      # week keys in a month plan
    opt = _opt(_req(), [stale])
    assert opt.periods is None and opt.custom_allocation and opt.notes
    block = et.render_prompt_block(et.compute_email_totals(_req(), [{"label": "A", "line_items": [stale]}], None, "month"))
    assert "can't be quoted; state only the totals below" in block
    assert "Month-by-month" not in block and "$12,000" in block                       # the line total is still stated
    # no readable dates: no per-period amounts either, and no crash
    undated = _opt(_req(start_date="", end_date=""), [_li("Display Alpha", 4000.0, monthly_allocations={"2026-10": 12000.0})])
    assert undated.periods is None and undated.total_net == 12000


def test_week_allocations_use_week_buckets():
    weeks = mo.periods_between(mo.parse_flexible_date("2026-10-01"), mo.parse_flexible_date("2026-10-21"), "week")
    assert len(weeks) == 3
    allocations = {weeks[0]["key"]: 100.0, weeks[1]["key"]: 200.0, weeks[2]["key"]: 300.0}
    line = _li("Display Alpha", 200.0, months=3, monthly_allocations=allocations)
    opt = _opt(_req(start_date="2026-10-01", end_date="2026-10-21", total_months=1), [line], "week")
    assert [(p.label, p.net) for p in opt.periods] == [("Week 1", 100.0), ("Week 2", 200.0), ("Week 3", 300.0)]
    block = et.render_prompt_block(et.compute_email_totals(
        _req(start_date="2026-10-01", end_date="2026-10-21", total_months=1), [{"label": "A", "line_items": [line]}], None, "week"))
    assert "Week-by-week total, all lines (net): Week 1 (Oct 1 – 7, 2026) $100 · Week 2 (Oct 8 – 14, 2026) $200" in block


def test_agency_fee_gross_figures_follow_the_gross_sheet_and_are_labelled_gross():
    addons = [SimpleNamespace(product_name="Landing Page", amount=100.0)]
    totals = et.compute_email_totals(_req(agency_fee=0.15), [{"label": "A", "line_items": [_li("Display Alpha", 2000.0)]}], addons, "month")
    o = totals.options[0]
    assert totals.agency_fee == 0.15
    assert o.per_period_gross == 2352 and o.media_gross == 7056                       # floor(2,000/.85) = 2,352 x 3
    assert o.addons_gross == pytest.approx(117.647, abs=1e-3)
    assert o.total_gross == 7173                                                      # floor(7,056 + 117.647)
    block = et.render_prompt_block(totals)
    assert "Gross, with the 15% agency fee (reference — quote only if the context calls for it, and always label it gross)" in block
    assert "TOTAL INVESTMENT, gross $7,173" in block
    for fee in (None, 0, 0.0, 1.0, 1.5, -0.1, "x"):
        # (a plain object: ProposalRequest itself already normalizes a typed 15 / 1.0 into a fraction)
        plain = SimpleNamespace(start_date="2026-10-01", end_date="2026-12-31", total_months=3, agency_fee=fee)
        t = et.compute_email_totals(plain, [{"label": "A", "line_items": [_li("Display Alpha", 2000.0)]}], None, "month")
        assert t.agency_fee is None and "agency fee" not in et.render_prompt_block(t)
        assert t.options[0].total_gross == t.options[0].total_net


def test_lines_with_different_period_counts_follow_the_calendar_span_like_the_workbook():
    opt = _opt(_req(), [_li("Display Alpha", 1000.0, months=2), _li("Search Beta", 1000.0, months=4)])
    assert opt.mixed_period_counts and opt.n_periods == 3 and opt.media_net == 6000     # the flight's 3 calendar months
    assert any("different period counts" in n for n in opt.notes)
    # no dates either: the request's own total_months, then 3
    assert _opt(_req(start_date="", end_date="", total_months=5), [_li("A", 100.0, months=2), _li("B", 100.0, months=4)]).n_periods == 5
    assert _opt(_req(start_date="", end_date="", total_months=None), [_li("A", 100.0, months=2), _li("B", 100.0, months=4)]).n_periods == 3


def test_products_the_workbook_would_skip_are_skipped_here_too():
    lines = [_li(EMAIL, 1000.0), _li("Not In The Catalog", 5000.0)]
    from app.catalog import by_name
    totals = et.compute_email_totals(_req(), [{"label": "A", "line_items": lines}], None, "month",
                                     known_product=lambda n: by_name(n) is not None)
    assert [l.product_name for l in totals.options[0].lines] == [EMAIL] and totals.options[0].total_net == 3000
    assert et.compute_email_totals(_req(), [{"label": "A", "line_items": lines}], None, "month").options[0].total_net == 18000


def test_odd_input_never_crashes():
    t = et.compute_email_totals(None, None, None, None)
    assert len(t.options) == 1 and t.options[0].total_net == 0 and t.unit == "month"
    odd = [SimpleNamespace(product_name=None, monthly_budget="abc", months=None, is_added_value=None, monthly_allocations={"x": "y"}),
           {"product_name": "Dict Line", "monthly_budget": 100, "months": "2"}]
    et.render_prompt_block(et.compute_email_totals(SimpleNamespace(), [{"label": "z", "line_items": odd}], [None, 3], "banana"))


def test_legacy_call_shape_gives_the_same_figures():
    lines = [_li("Display Alpha", 2000.0), _li("Search Beta", 1500.0)]
    legacy = et.totals_from_legacy(_req(), lines, None, "month")
    assert legacy.options[0].total_net == et.compute_email_totals(_req(), [{"label": "A", "line_items": lines}], None, "month").options[0].total_net == 10500
    multi = et.totals_from_legacy(_req(), lines, [{"label": "A", "name": "One", "line_items": lines[:1]},
                                                  {"label": "B", "name": "", "line_items": lines}], "month")
    assert [o.total_net for o in multi.options] == [6000, 10500]
    assert json.loads(json.dumps(multi.to_dict()))["options"][1]["total_net"] == 10500            # JSON-able for the API


# ===========================================================================================================
# Parity with the real workbook
# ===========================================================================================================

def _excel_figures(ws, *, gross: bool) -> dict:
    """What the sheet's own formulas come to, evaluated from the cells generate_proposal wrote (after checking each
    formula really has the shape this evaluation assumes)."""
    col = "N" if gross else "L"
    total_row = next(r for r in range(19, ws.max_row + 1)
                     if str(ws[f"C{r}"].value or "").startswith("TOTAL DIGITAL") and str(ws[f"L{r}"].value or "").startswith("=ROUNDDOWN(SUM(L19:L"))
    grand_row = next(r for r in range(total_row + 1, ws.max_row + 1) if str(ws[f"L{r}"].value or "").startswith(f"=ROUNDDOWN(L{total_row}"))
    fee = ws["I14"].value if gross else None
    factor = 1.0 / (1.0 - fee) if gross else 1.0
    lines = [float(ws[f"L{r}"].value) for r in range(19, total_row) if isinstance(ws[f"L{r}"].value, (int, float))]
    per_period = math.floor(math.fsum(v * factor for v in lines) + 1e-9)
    assert str(ws[f"{col}{total_row}"].value).startswith(f"=ROUNDDOWN(SUM({col}19:{col}")          # TOTAL DIGITAL = ROUNDDOWN(SUM(budgets))
    formula = ws[f"{col}{grand_row}"].value
    if "ISNUMBER(I10)" in formula:
        assert formula.startswith(f"=ROUNDDOWN({col}{total_row}*IF(ISNUMBER(I10),I10,3)")
        multiplier = ws["I10"].value
    else:
        assert formula.startswith(f"=ROUNDDOWN({col}{total_row} + SUMIF(")                     # Full Flight: nothing multiplied
        multiplier = 1
    addons = 0.0
    m = re.search(rf'SUMIF\({col}(\d+):{col}(\d+),">0"\)', formula)
    if m:
        for r in range(int(m.group(1)), int(m.group(2)) + 1):
            net = ws[f"L{r}"].value
            amount = net * factor if isinstance(net, (int, float)) else 0.0
            if amount > 0:
                addons += amount
    return {"per_period": per_period, "multiplier": multiplier, "grand": math.floor(per_period * multiplier + addons + 1e-9),
            "total_row": total_row}


def _workbook(tmp_path, request, tiers, addons, unit, mode="even"):
    out = tmp_path / "parity.xlsx"
    # addons=None would mean the generator's legacy hard-coded list; /api/generate always sends [] for "none picked"
    generate_proposal(request, tiers[0]["line_items"], out, tiers=tiers, addons=addons if addons is not None else [], time_unit=unit,
                      monthly_distribution_mode=mode, force_tabs={"net": True, "gross": True})
    return openpyxl.load_workbook(out)


def _compare(wb, totals: et.EmailTotals, label="A", *, gross=True):
    opt = next(o for o in totals.options if o.label == label)
    sheet = "Proposal A" if len(totals.options) == 1 else None
    ws = wb[sheet] if sheet else wb[f"{opt.display_name}"]
    net = _excel_figures(ws, gross=False)
    assert (opt.per_period_net, opt.n_periods, opt.total_net) == (net["per_period"], net["multiplier"], net["grand"])
    if gross and totals.agency_fee:
        ws_g = wb["Proposal A (Gross)"] if sheet else wb[f"{opt.display_name} (Gross)"]
        g = _excel_figures(ws_g, gross=True)
        assert (opt.per_period_gross, opt.total_gross) == (g["per_period"], g["grand"])
    return ws, net


def test_month_totals_equal_the_workbooks_including_add_ons_and_the_gross_sheet(tmp_path):
    request = _req(agency_fee=0.15, geo="Los Angeles")
    lines = [_li(EMAIL, 1666.67), _li(IAB, 2500.33), _li(GEOFENCE, 999.99)]
    addons = [AddonItem("Call Tracking", 100.0), AddonItem("Web Services - Landing Page", 250.5), AddonItem("Non Media Offering - Brand Lift", 0.0)]
    tiers = [{"label": "A", "line_items": lines, "avails_data": None}]
    totals = et.compute_email_totals(request, tiers, addons, "month")
    wb = _workbook(tmp_path, request, tiers, addons, "month")
    _compare(wb, totals)
    assert totals.options[0].n_periods == 3 and totals.options[0].addons_net == pytest.approx(350.5)


def test_full_flight_totals_equal_the_workbooks(tmp_path):
    request = _req(start_date="2026-11-24", end_date="2026-12-31", total_months=2, agency_fee=0.15)
    lines = [_li(EMAIL, 680.0, months=1), _li("Email Campaigns and/or Email Campaigns - Re-Drop", 3600.0, months=1),
             _li(IAB, 5470.0, months=1), _li(GEOFENCE, 3000.0, months=1)]
    addons = [AddonItem("Call Tracking", 100.0), AddonItem("Non Media Offering - Brand Lift", 0.0)]
    tiers = [{"label": "A", "line_items": lines, "avails_data": None}]
    totals = et.compute_email_totals(request, tiers, addons, "full_flight")
    wb = _workbook(tmp_path, request, tiers, addons, "full_flight")
    _compare(wb, totals)
    o = totals.options[0]
    assert (o.per_period_net, o.n_periods, o.media_net, o.total_net) == (12750, 1, 12750, 12850)
    assert (o.media_gross, o.total_gross) == (15000, 15117)                           # the reference export's $15,000 gross


def test_week_totals_equal_the_workbooks(tmp_path):
    request = _req(start_date="2026-10-01", end_date="2026-12-31", total_months=3)
    weeks = len(mo.periods_between(mo.parse_flexible_date("2026-10-01"), mo.parse_flexible_date("2026-12-31"), "week"))
    lines = [_li(EMAIL, 333.33, months=weeks), _li(IAB, 412.5, months=weeks)]
    tiers = [{"label": "A", "line_items": lines, "avails_data": None}]
    addons = [AddonItem("Call Tracking", 100.0)]
    totals = et.compute_email_totals(request, tiers, addons, "week")
    _, net = _compare(_workbook(tmp_path, request, tiers, addons, "week"), totals, gross=False)
    assert net["multiplier"] == weeks == totals.options[0].n_periods


def test_two_options_with_their_own_dates_equal_the_workbooks(tmp_path):
    request = _req(agency_fee=0.2)
    tiers = [
        {"label": "A", "name": "Independent", "line_items": [_li(EMAIL, 1000.0)], "avails_data": None},
        {"label": "B", "name": "Democrat", "start_date": "2026-10-01", "end_date": "2027-01-31",
         "line_items": [_li(EMAIL, 1200.0, months=4), _li(IAB, 800.0, months=4)], "avails_data": None},
    ]
    addons = [AddonItem("Call Tracking", 100.0)]
    totals = et.compute_email_totals(request, tiers, addons, "month")
    wb = _workbook(tmp_path, request, tiers, addons, "month")
    assert {"Independent", "Independent (Gross)", "Democrat", "Democrat (Gross)"} <= set(wb.sheetnames)
    for label in ("A", "B"):
        _compare(wb, totals, label)
    assert [o.n_periods for o in totals.options] == [3, 4]


def test_customised_per_period_amounts_equal_the_workbooks_breakdown_row(tmp_path):
    request = _req(agency_fee=0.15)
    custom = _li(EMAIL, 4000.0, monthly_allocations={"2026-10": 6000.0, "2026-11": 3000.0, "2026-12": 3000.0})
    plain = _li(IAB, 1500.0)
    tiers = [{"label": "A", "line_items": [custom, plain], "avails_data": None}]
    totals = et.compute_email_totals(request, tiers, None, "month")
    wb = _workbook(tmp_path, request, tiers, None, "month")
    ws, net = _compare(wb, totals)
    o = totals.options[0]
    net_row = [ws.cell(row=net["total_row"], column=19 + i).value for i in range(3)]          # S.. : the breakdown block
    assert net_row == pytest.approx([p.net for p in o.periods])
    assert [p.net for p in o.periods] == [7500.0, 4500.0, 4500.0]
    ws_g = wb["Proposal A (Gross)"]
    gross_row = [ws_g.cell(row=_excel_figures(ws_g, gross=True)["total_row"], column=21 + i).value for i in range(3)]   # U..
    assert gross_row == pytest.approx([p.gross for p in o.periods])


# ===========================================================================================================
# The prompts
# ===========================================================================================================

@pytest.fixture
def stub_catalog(monkeypatch):
    from app.catalog import Product  # noqa: F401  (only to prove the import works offline)
    catalog = {}
    monkeypatch.setattr(ai_enricher, "_catalog_by_name", lambda name: catalog.get(name))
    return catalog


def _prompt(request, lines, tiers=None, unit="month", **kw):
    return ai_enricher._build_prompt(request, lines, tiers=tiers, time_unit=unit, **kw)


def test_the_email_prompt_carries_the_authoritative_figures_and_the_grand_total(stub_catalog):
    lines = [_li("Display Alpha", 2000.0), _li("Search Beta", 1500.0)]
    addons = [SimpleNamespace(product_name="Landing Page", amount=100.0)]
    totals = et.compute_email_totals(_req(), [{"label": "A", "line_items": lines}], addons, "month")
    prompt = _prompt(_req(), lines, totals=totals)
    assert "## AUTHORITATIVE FIGURES — the ONLY dollar amounts you may state" in prompt
    assert "- Total Net Investment: $10,600 ($10,500 media + $100 one-time add-ons)" in prompt
    assert "TOTAL INVESTMENT, net — the Excel's campaign grand-total row (media + one-time add-ons): $10,600" in prompt
    assert prompt.index("## AUTHORITATIVE FIGURES") < prompt.index("## PRODUCTS IN THIS PROPOSAL")
    assert "- Dollar figures: state ONLY amounts listed in AUTHORITATIVE FIGURES" in prompt
    assert "copied exactly from the AUTHORITATIVE FIGURES block" in prompt
    # the JSON exemplar the model must return stays valid
    start = prompt.index('{\n  "campaign_name"')
    exemplar, _ = json.JSONDecoder().raw_decode(prompt[start:])
    assert {"internal_email_body", "client_email_body"} <= set(exemplar)


def test_a_second_line_for_the_same_product_is_in_the_totals_even_though_the_product_list_dedupes(stub_catalog):
    # main.py hands the enricher ONE line per product (for the blurbs); the figures come from the full option instead
    first, second = _li("Display Alpha", 2000.0, id="a1"), _li("Display Alpha", 1000.0, id="a2")
    totals = et.compute_email_totals(_req(), [{"label": "A", "line_items": [first, second]}], None, "month")
    prompt = _prompt(_req(), [first], totals=totals)
    assert "- Total Net Investment: $9,000" in prompt
    assert prompt.count("Display Alpha: $2,000 per month × 3 months = $6,000") == 1
    assert "Display Alpha: $1,000 per month × 3 months = $3,000" in prompt


def test_multi_option_prompt_has_no_combined_total_and_no_borrowed_product_amounts(stub_catalog):
    a = [_li("Display Alpha", 2000.0)]
    b = [_li("Display Alpha", 4000.0), _li("Search Beta", 1000.0)]
    tiers = [{"label": "A", "name": "Independent", "line_items": a}, {"label": "B", "name": "", "line_items": b}]
    union = [b[0], b[1]]
    totals = et.compute_email_totals(_req(), tiers, None, "month")
    prompt = _prompt(_req(), union, tiers=tiers, totals=totals)
    assert "- Total Net Investment: see each option below — they are ALTERNATIVES, never added together" in prompt
    assert "  - Independent: $2,000/month ($6,000 total flight) — Display Alpha" in prompt
    assert "  - Option B: $5,000/month ($15,000 total flight) — Display Alpha, Search Beta" in prompt
    # the union list can't show ONE option's amount for a product that costs different amounts per option
    assert "  - Display Alpha: (the amount differs by option — see AUTHORITATIVE FIGURES)" in prompt
    assert "$21,000" not in prompt


def test_week_and_quarter_prompts_use_their_own_units_everywhere(stub_catalog):
    for unit, noun in (("week", "week"), ("quarter", "quarter")):
        lines = [_li("Display Alpha", 500.0, months=4), _li("Search Beta", 250.0, months=4)]
        totals = et.compute_email_totals(_req(), [{"label": "A", "line_items": lines}], None, unit)
        prompt = _prompt(_req(), lines, unit=unit, totals=totals)
        assert f"$500/{noun} × 4 {noun}s = $2,000 total" in prompt
        assert f"All lines together, per {noun} (net): $750" in prompt
        assert f"- Billing period: this proposal is priced PER {noun.upper()}" in prompt
        assert "monthly" not in "\n".join(l for l in prompt.splitlines() if "impressions" in l or l.startswith("- Total Net Investment"))


def test_the_revision_prompt_uses_the_plan_as_it_is_now_not_the_stale_line_list(monkeypatch):
    captured = []
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    monkeypatch.setattr(ai_enricher, "_HAS_OPENAI", True)
    monkeypatch.setattr(ai_enricher, "_OpenAI", lambda api_key: object())
    monkeypatch.setattr(ai_enricher.llm_utils, "chat_text", lambda client, model, prompt, **k: captured.append(prompt) or json.dumps(
        {"internal_email_subject": "s", "internal_email_body": "b", "client_email_subject": "s", "client_email_body": "b"}))
    current = [{"label": "A", "name": "Independent", "line_items": [_li("Display Alpha", 5000.0)]},
               {"label": "B", "name": "Democrat", "line_items": [_li("Display Alpha", 8000.0)]}]
    totals = et.compute_email_totals(_req(), current, [SimpleNamespace(product_name="Landing Page", amount=100.0)], "month")
    stale_line_items = [_li("Display Alpha", 1000.0)]            # what the old call sent: the ACTIVE option, as it was
    result = ai_enricher.reprompt_emails(_req(), stale_line_items, "Holiday", "is", "ib", "cs", "cb", "shorter",
                                         scope="client", time_unit="month", totals=totals)
    assert result["error"] is None and result["figure_warnings"] == []
    prompt = captured[0]
    assert "$15,100" in prompt and "$24,100" in prompt           # both CURRENT options, add-ons included
    assert "Total Net Investment" not in prompt                  # two alternatives: no combined figure
    assert "$3,000" not in prompt                                # the stale 1,000 x 3 is nowhere
    assert "CRITICAL — DOLLAR FIGURES" in prompt and "replace any dollar amount that no longer matches" in prompt
    # without `totals` (old callers) the figures are derived from line_items alone, still unit-aware
    captured.clear()
    ai_enricher.reprompt_emails(_req(), stale_line_items, "Holiday", "is", "ib", "cs", "cb", "shorter")
    assert "TOTAL INVESTMENT, net" in captured[0] and "- Total Net Investment: $3,000" in captured[0]


# ===========================================================================================================
# The checker
# ===========================================================================================================

def _totals_30k():
    # 3 months x ($6,000 + $4,000) + a $100 add-on = $30,100; options: one
    return et.compute_email_totals(
        _req(), [{"label": "A", "line_items": [_li("Display Alpha", 6000.0), _li("Search Beta", 4000.0)]}],
        [SimpleNamespace(product_name="Landing Page", amount=100.0)], "month")


def _check(text, totals=None, **kw):
    totals = totals or _totals_30k()
    return et.check_texts({"client email": text}, totals, **kw)


def test_the_checker_accepts_every_legitimate_figure():
    text = ("Your total investment is $30,100: $30,000 across three months plus a one-time $100 landing page.\n"
            "Display Alpha runs $6,000 per month ($18,000 in total) and Search Beta $4,000 a month ($12,000 total).\n"
            "- Option budget: $10,000 per month.\n"
            "The all-in total comes to $30,100.00 net.")
    assert _check(text) == []


def test_the_checker_catches_a_wrong_total():
    issues = _check("Your total investment across the three months is $31,000.")
    assert [(i.amount, i.value) for i in issues] == [("$31,000", 31000.0)]
    assert "total investment" in issues[0].sentence
    # a monthly rate presented for a weekly plan, a multiplied figure, a per-period figure from the wrong unit
    assert _check("The plan is $10,000 per month for 12 months, $120,000 in all.")
    assert _check("Display Alpha is $3,400 per month.")
    assert _check("That is a $27,000 budget.")
    assert _check("Investment: $30K") == []            # shorthand within 1% of $30,100 is fine ...
    assert _check("Investment: $25K")                  # ... a rounded-away $5,100 is not


def test_the_checker_only_judges_amounts_that_are_about_the_plan():
    allowed_text = ("Households spend an average of $5,000 a year on streaming. The category is worth $1.2 billion and "
                    "$63 million in local ad spend. A $10 CPM keeps frequency affordable, at roughly $2.50 per click "
                    "or $15 per lead. Nielsen says listeners average $1,200 in annual podcast-driven purchases. "
                    "A $25 fee applies. Our minimum is just $20.")
    assert _check(allowed_text) == []
    # ... but the same amount in a sentence about THIS plan is judged
    assert _check("Your budget of $5,000 a month will do that.")
    # a product or option name makes a sentence about the plan, even without budget words
    assert _check("Search Beta gets $3,333 to start.")
    assert _check("Search Beta gets $4,000 to start.") == []


def test_the_checker_is_tolerant_of_formatting_and_never_raises():
    t = _totals_30k()
    for text in ("", None, "   ", "$", "$$", "$ ", "$abc", "$5,000,", "$ 30,100", "$30,100.", "(30,100)", "$1,2,3", "$-5",
                 " $30,100 ", "$" * 500, "\n".join(["Total $31,000"] * 200), "Total investment: $" + "9" * 400,
                 "💰 Total investment: $30,100 💰"):
        et.check_text(text, et.allowed_amounts(t))
    assert et.check_text("Total investment: $ 30,100", et.allowed_amounts(t)) == []
    assert et.check_text("Total investment: $30,100.00", et.allowed_amounts(t)) == []
    assert et.check_text(None, [1.0]) == []
    assert et.check_text("Total investment $31,000", None) == []         # a broken `allowed` is swallowed, not raised


def test_extra_allowed_amounts_cover_figures_the_context_itself_provided():
    text = "Your total investment is $5,000 per month at the minimum."
    assert _check(text)                                               # not a plan figure ...
    assert _check(text, extra_allowed=[5000.0]) == []                 # ... unless the prompt / planner supplied it
    assert et.amounts_in_text("a $1,500 minimum, $2.50 per click and $3 million") == [1500.0, 2.5, 3_000_000.0]


# ---- the corrective pass ----------------------------------------------------------------------------------

def _emails(**kw):
    base = {"internal_email_subject": "S", "internal_email_body": "Hi Pat,\n{{PROPOSAL_LINE}}\nTotal investment is $30,100.\n{{PLANNER_SIGNATURE}}",
            "client_email_subject": "CS", "client_email_body": "Your total investment is $30,100 net.\nDisplay Alpha runs $6,000 per month."}
    base.update(kw)
    return base


def _fake_client():
    return SimpleNamespace()


def _patch_repair(monkeypatch, reply):
    """Replace the corrective model call; `reply` is a dict (returned as JSON), an Exception (raised) or a callable."""
    calls = []

    def chat_text(client, model, prompt, **kw):
        calls.append(prompt)
        if isinstance(reply, Exception):
            raise reply
        return json.dumps(reply(prompt) if callable(reply) else reply)

    monkeypatch.setattr(ai_enricher.llm_utils, "chat_text", chat_text)
    return calls


def test_correct_emails_cost_no_second_call(monkeypatch):
    calls = _patch_repair(monkeypatch, {})
    emails, warnings = ai_enricher._check_and_repair_emails(_fake_client(), _totals_30k(), "", _emails(), "both")
    assert emails == _emails() and warnings == [] and calls == []


def test_a_wrong_total_gets_one_corrective_pass_that_quotes_the_right_figures(monkeypatch):
    wrong = _emails(client_email_body="Your total investment is $31,000 net.\nDisplay Alpha runs $6,000 per month.")
    fixed = _emails()
    calls = _patch_repair(monkeypatch, fixed)
    emails, warnings = ai_enricher._check_and_repair_emails(_fake_client(), _totals_30k(), "", wrong, "both")
    assert emails["client_email_body"] == fixed["client_email_body"] and warnings == []
    assert len(calls) == 1
    prompt = calls[0]
    assert "TOTAL INVESTMENT, net" in prompt and "$30,100" in prompt                # the right figures
    assert "$31,000 (client email)" in prompt                                      # the offending amount and where it is
    assert "Rewrite ONLY the sentences that contain a mismatching amount" in prompt
    assert "{{PLACEHOLDER}}" in prompt


def test_a_correction_that_still_fails_is_reported_not_shipped_silently(monkeypatch):
    wrong = _emails(client_email_body="Your total investment is $31,000 net.")
    _patch_repair(monkeypatch, _emails(client_email_body="Your total investment is $32,500 net."))
    emails, warnings = ai_enricher._check_and_repair_emails(_fake_client(), _totals_30k(), "", wrong, "both")
    assert len(warnings) == 1 and "$31,000" in warnings[0] and "(client email)" in warnings[0]
    # nothing is blanked: an email is always returned
    assert emails["client_email_body"] and emails["internal_email_body"] == wrong["internal_email_body"]


def test_a_failing_or_damaging_correction_never_blocks_the_email(monkeypatch):
    wrong = _emails(client_email_body="Your total investment is $31,000 net.")
    for reply in (RuntimeError("boom"), ai_enricher.llm_utils.ResponseTruncated("cut"), ValueError("no json")):
        _patch_repair(monkeypatch, reply)
        emails, warnings = ai_enricher._check_and_repair_emails(_fake_client(), _totals_30k(), "", wrong, "both")
        assert emails == wrong and len(warnings) == 1
    # a correction that loses a {{PLACEHOLDER}} (the internal email's link / signature slots) is discarded
    bad = _emails(internal_email_body="Hi Pat, total investment is $30,100.", client_email_body="Your total investment is $30,100 net.")
    _patch_repair(monkeypatch, bad)
    emails, warnings = ai_enricher._check_and_repair_emails(_fake_client(), _totals_30k(), "", wrong, "both")
    assert emails == wrong and warnings
    # no client at all (no key): warn, don't crash
    emails, warnings = ai_enricher._check_and_repair_emails(None, _totals_30k(), "", wrong, "both")
    assert emails == wrong and len(warnings) == 1


def test_scope_keeps_the_other_email_untouched_but_still_reports_its_stale_figures(monkeypatch):
    stale_client = _emails(client_email_body="Your total investment is $28,000 net.")
    # only the internal email may change; the corrective pass fixes the internal one and cannot touch the client one
    wrong = dict(stale_client, internal_email_body="Hi Pat,\n{{PROPOSAL_LINE}}\nTotal investment is $31,000.\n{{PLANNER_SIGNATURE}}")
    _patch_repair(monkeypatch, dict(_emails(), client_email_body="REWRITTEN BY THE MODEL anyway, total investment $30,100."))
    emails, warnings = ai_enricher._check_and_repair_emails(_fake_client(), _totals_30k(), "", wrong, "internal")
    assert emails["internal_email_body"] == _emails()["internal_email_body"]
    assert emails["client_email_body"] == stale_client["client_email_body"]                      # the guardrail held
    assert any("$28,000" in w and "client email" in w for w in warnings)                          # ... and the stale figure is surfaced


def test_a_products_rate_card_price_is_fair_to_quote(monkeypatch):
    emails = _emails(client_email_body="Search Beta has a $1,500 monthly minimum.")
    monkeypatch.setattr(ai_enricher, "_catalog_by_name",
                        lambda n: SimpleNamespace(minimum_spend=1500.0, base_rate=12.0) if n == "Search Beta" else None)
    assert ai_enricher._figure_issues(emails, _totals_30k(), "") == []
    monkeypatch.setattr(ai_enricher, "_catalog_by_name", lambda n: None)          # not on the rate card: judged like any figure
    assert [i.amount for i in ai_enricher._figure_issues(emails, _totals_30k(), "")] == ["$1,500"]
    monkeypatch.setattr(ai_enricher, "_catalog_by_name", lambda n: (_ for _ in ()).throw(RuntimeError("catalog down")))
    assert ai_enricher._figure_issues(emails, _totals_30k(), "")                   # a catalog failure never breaks the check


def test_warnings_are_capped_and_the_planners_own_figures_are_respected():
    many = _emails(client_email_body="\n".join(f"Total investment option {i}: ${31000 + i}." for i in range(12)))
    _, warnings = ai_enricher._check_and_repair_emails(None, _totals_30k(), "", many, "both")
    assert len(warnings) == ai_enricher._MAX_FIGURE_WARNINGS + 1 and warnings[-1].startswith("…and ")
    typed = _emails(client_email_body="As you asked, the total investment is $27,500 net.")
    _, warnings = ai_enricher._check_and_repair_emails(None, _totals_30k(), "make the total $27,500", typed, "both")
    assert warnings == []


# ---- enrich_proposal end to end (stub model) ------------------------------------------------------------------

class _FakeResponses:
    def __init__(self, text):
        self._text, self.calls = text, []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(output_text=self._text, status="completed", incomplete_details=None, usage=None)


def _enrich(monkeypatch, first_reply, repair_reply, **kw):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    monkeypatch.setattr(ai_enricher, "_HAS_OPENAI", True)
    fake = SimpleNamespace(responses=_FakeResponses(json.dumps(first_reply)))
    monkeypatch.setattr(ai_enricher, "_OpenAI", lambda api_key: fake)
    monkeypatch.setattr(ai_enricher, "_catalog_by_name", lambda name: None)
    repair_calls = _patch_repair(monkeypatch, repair_reply)
    lines = [_li("Display Alpha", 6000.0), _li("Search Beta", 4000.0)]
    result = ai_enricher.enrich_proposal(_req(), lines, "EVC-1", totals=_totals_30k(), **kw)
    return result, fake, repair_calls


def _model_reply(client_body):
    return {"campaign_name": "Holiday Push", "product_blurbs": [], "internal_email_subject": "s",
            "internal_email_body": "Hi Pat,\n{{PROPOSAL_LINE}}\nThe plan totals $30,100.\n{{PLANNER_SIGNATURE}}",
            "client_email_subject": "cs", "client_email_body": client_body}


def test_enrich_proposal_sends_the_figures_and_leaves_a_correct_answer_alone(monkeypatch):
    result, fake, repair_calls = _enrich(monkeypatch, _model_reply("Your total investment is $30,100 net."), {})
    assert result.error is None and result.figure_warnings == [] and repair_calls == []
    assert "TOTAL INVESTMENT, net — the Excel's campaign grand-total row (media + one-time add-ons): $30,100" in fake.responses.calls[0]["input"]


def test_enrich_proposal_repairs_a_model_that_did_its_own_arithmetic(monkeypatch):
    result, _, repair_calls = _enrich(monkeypatch, _model_reply("Your total investment is $36,000 net (12 months)."),
                                      _model_reply("Your total investment is $30,100 net."))
    assert result.client_email_body == "Your total investment is $30,100 net." and result.figure_warnings == []
    assert result.campaign_name == "Holiday Push" and "{{PROPOSAL_LINE}}" in result.internal_email_body      # nothing else moved
    assert len(repair_calls) == 1


def test_enrich_proposal_surfaces_what_it_could_not_repair(monkeypatch):
    result, _, repair_calls = _enrich(monkeypatch, _model_reply("Your total investment is $36,000 net."), RuntimeError("down"))
    assert result.client_email_body == "Your total investment is $36,000 net."                    # never blanked
    assert len(result.figure_warnings) == 1 and "$36,000" in result.figure_warnings[0]
    assert result.error is None and len(repair_calls) == 1


def test_enrich_proposal_without_emails_or_totals_checks_nothing(monkeypatch):
    result, _, repair_calls = _enrich(monkeypatch, {"campaign_name": "X", "product_blurbs": []}, {})
    assert result.figure_warnings == [] and repair_calls == []


# ===========================================================================================================
# The endpoints: every generation / revision is fed the CURRENT state
# ===========================================================================================================

def _tier_models():
    return [
        main.TierModel(label="A", name="Independent", line_items=[main.LineItemModel(id="a1", product_name=EMAIL, monthly_budget=2000.0, months=3)]),
        main.TierModel(label="B", name="Democrat", line_items=[
            main.LineItemModel(id="b1", product_name=EMAIL, monthly_budget=4000.0, months=3,
                               monthly_allocations={"2026-10": 6000.0, "2026-11": 3000.0, "2026-12": 3000.0}),
            main.LineItemModel(id="b2", product_name=IAB, monthly_budget=1000.0, months=3)]),
    ]


def test_the_revision_endpoint_prices_every_option_the_addons_and_the_allocations_from_the_request(monkeypatch):
    seen = {}
    monkeypatch.setattr(main.ai_enricher, "reprompt_emails", lambda *a, **k: seen.update(args=a, **k) or {
        "internal_email_subject": "s", "internal_email_body": "b", "client_email_subject": "s", "client_email_body": "b",
        "error": None, "figure_warnings": []})
    monkeypatch.setattr(main, "_get_proposal_metadata", lambda pid: None)
    tiers = _tier_models()
    body = main.RepromptEmailsRequest(
        request={"client_name": "Acme", "start_date": "2026-10-01", "end_date": "2026-12-31", "total_months": 3, "agency_fee": 0.15},
        line_items=tiers[0].line_items,                       # what the old client sent: ONLY the active option
        tiers=tiers, addons=[main.AddonItemModel(product_name="Call Tracking", amount=100.0)],
        monthly_distribution_mode="even", reprompt="shorter", time_unit="month")
    result = asyncio.run(main.reprompt_emails("pid", body))
    totals = seen["totals"]
    assert [o.label for o in totals.options] == ["A", "B"] and totals.multi_option
    a, b = totals.options
    assert (a.total_net, b.total_net) == (6100, 15100)         # 3 x 2,000 + 100 ; 3 x (4,000 + 1,000) + 100
    assert b.custom_allocation and [p.net for p in b.periods] == [7000.0, 4000.0, 4000.0]
    assert a.addons_net == 100 and totals.agency_fee == 0.15 and a.total_gross > a.total_net
    # the legacy positional list is still passed (the old signature), but the figures no longer come from it
    assert [li.product_name for li in seen["args"][1]] == [EMAIL]
    # the response tells the SPA what the figures were, and passes the check's warnings through
    assert result["email_figures"]["options"][1]["total_net"] == 15100 and result["figure_warnings"] == []


def test_the_revision_endpoint_still_accepts_the_legacy_flat_payload_and_survives_odd_state(monkeypatch):
    seen = {}
    monkeypatch.setattr(main.ai_enricher, "reprompt_emails", lambda *a, **k: seen.update(k) or {
        "internal_email_subject": "s", "internal_email_body": "b", "client_email_subject": "s", "client_email_body": "b", "error": None})
    monkeypatch.setattr(main, "_get_proposal_metadata", lambda pid: None)
    body = main.RepromptEmailsRequest(
        request={"client_name": "Acme"}, reprompt="shorter",
        line_items=[main.LineItemModel(product_name=EMAIL, monthly_budget=1000.0, months=2)])
    asyncio.run(main.reprompt_emails("pid", body))
    assert [o.total_net for o in seen["totals"].options] == [2000]
    # a Full Flight revision folds a stale month line and ignores its allocations, like Generate
    ff = main.RepromptEmailsRequest(
        request={"client_name": "Acme"}, reprompt="shorter", time_unit="full_flight",
        line_items=[main.LineItemModel(product_name=EMAIL, monthly_budget=340.0, months=2, monthly_allocations={"2026-11": 1.0})])
    asyncio.run(main.reprompt_emails("pid", ff))
    assert seen["totals"].options[0].total_net == 680 and seen["totals"].options[0].n_periods == 1


def test_the_revision_endpoint_falls_back_when_the_figures_cannot_be_built(monkeypatch):
    seen = {}
    monkeypatch.setattr(main.ai_enricher, "reprompt_emails", lambda *a, **k: seen.update(k) or {
        "internal_email_subject": "s", "internal_email_body": "b", "client_email_subject": "s", "client_email_body": "b", "error": None})
    monkeypatch.setattr(main, "_get_proposal_metadata", lambda pid: None)
    monkeypatch.setattr(main.email_totals_svc, "compute_email_totals", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("bug")))
    result = asyncio.run(main.reprompt_emails("pid", main.RepromptEmailsRequest(request={"client_name": "Acme"}, line_items=[], reprompt="x")))
    assert seen["totals"] is None and "email_figures" not in result              # a bug in the figures never fails the revision


def test_generate_feeds_the_enricher_the_full_current_plan(monkeypatch, tmp_path):
    seen = {}
    saved = {}
    monkeypatch.setattr(main, "PROPOSALS_DIR", tmp_path)
    monkeypatch.setattr(main, "_get_next_short_id", lambda: "EVC-1")
    monkeypatch.setattr(main, "_resolve_draftable_proposal_id", lambda requested: "pid-1")
    monkeypatch.setattr(main, "_save_proposal_metadata", lambda **kw: saved.update(kw))
    monkeypatch.setattr(main.proposal_files_svc, "save_file", lambda *a, **k: None)
    monkeypatch.setattr(main.ai_enricher, "enrich_proposal", lambda *a, **k: seen.update(args=a, **k) or ai_enricher.ProposalEnrichment(
        campaign_name="Holiday", figure_warnings=["$31,000 (client email): “total”"]))
    monkeypatch.setattr(main.disclaimers_svc, "resolve_for_tiers", lambda req, tiers: {})
    monkeypatch.setattr(main.disclaimers_svc, "union_for_ppt", lambda d: [])
    monkeypatch.setattr(main.pptx_builder, "build_signature_deck", lambda *a, **k: None)
    monkeypatch.setattr(main.docx_builder, "build_client_email_docx", lambda **kw: None)
    request = SimpleNamespace(state=SimpleNamespace(user={"email": "planner@entravision.com", "is_admin": False}),
                              client=SimpleNamespace(host="127.0.0.1"), headers={"user-agent": "pytest"})
    body = main.GenerateRequest(
        request={"client_name": "Acme", "start_date": "2026-10-01", "end_date": "2026-12-31", "total_months": 3, "agency_fee": 0.15,
                 "request_type": "New Business"},
        line_items=[], tiers=_tier_models(), addons=[main.AddonItemModel(product_name="Call Tracking", amount=100.0)],
        time_unit="month")
    result = asyncio.run(main.generate(body, request))
    totals = seen["totals"]
    assert [o.total_net for o in totals.options] == [6100, 15100]
    assert totals.options[1].periods is not None and totals.options[0].addons_net == 100
    # the enricher still gets the product-deduplicated union for the blurbs ...
    assert [li.product_name for li in seen["args"][1]] == [EMAIL, IAB]
    # ... while the figures come from the full options
    assert result["email_figures"]["options"][0]["total_net"] == 6100
    assert result["enrichment"]["figure_warnings"] == ["$31,000 (client email): “total”"]
    assert saved["reopen_state"]["enrichment"]["figure_warnings"]


# ===========================================================================================================
# The SPA
# ===========================================================================================================

def _extract_fn(name: str) -> str:
    m = re.search(rf"^(?:async )?function {re.escape(name)}\(", APP_JS, re.M)
    assert m, f"function {name} not found"
    return APP_JS[m.start(): APP_JS.index("\n}\n", m.start()) + 2]


def test_the_spa_sends_the_whole_current_plan_with_every_revision():
    revise = _extract_fn("_submitEmailReprompt")
    assert "..._emailFiguresPayload()" in revise
    assert "line_items: state.lineItems" not in revise                 # the ACTIVE option alone is what used to be sent
    payload = _extract_fn("_emailFiguresPayload")
    for field in ("request: p.request", "tiers:", "addons: p.addons", "monthly_distribution_mode: p.monthly_distribution_mode",
                  "time_unit: p.time_unit", "line_items: p.line_items"):
        assert field in payload, field
    # the payload is built from the same function /api/generate and the drafts use (resync + Full Flight normalization included)
    assert "_buildWizardStatePayload()" in payload and "_buildWizardStatePayload()" in _extract_fn("_emailPlanKey")
    assert "const sentPlanKey = _emailPlanKey();" in _extract_fn("onGenerate")
    assert "state.emailsPlanKey = sentPlanKey;" in _extract_fn("onGenerate")
    assert "_refreshEmailStaleNote();" in _extract_fn("goToStep") and "state.emailsPlanKey = state.enrichment ? _emailPlanKey() : null;" in APP_JS
    for needed in ("email-stale-note", "email-refresh-figures-btn", "email-figures-warning"):
        assert f'id="{needed}"' in INDEX_HTML
    assert 'getElementById("email-refresh-figures-btn").addEventListener("click", onEmailRefreshFigures)' in APP_JS


def _run_node(script: str) -> str:
    node = shutil.which("node")
    if not node:
        pytest.skip("node is not installed")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "scenario.js"
        path.write_text(script, encoding="utf-8")
        proc = subprocess.run([node, str(path)], capture_output=True, text=True, encoding="utf-8", timeout=60)
    assert proc.returncode == 0, proc.stderr + proc.stdout
    return proc.stdout.strip()


_JS_PRELUDE = """
var __els = {};
function makeEl(id) {
  var cls = { hidden: true }, text = "";
  var el = { id: id, style: {}, children: [], appendChild: function (c) { this.children.push(c); },
    classList: { add: function (c) { cls[c] = true; }, remove: function (c) { delete cls[c]; },
      toggle: function (c, f) { var on = f === undefined ? !cls[c] : !!f; if (on) cls[c] = true; else delete cls[c]; return on; },
      contains: function (c) { return !!cls[c]; } } };
  // like the DOM: assigning textContent replaces every child
  Object.defineProperty(el, "textContent", { get: function () { return text; }, set: function (v) { text = v; this.children = []; } });
  return el;
}
var document = { getElementById: function (id) { return __els[id] || (__els[id] = makeEl(id)); },
                 createElement: function (t) { return makeEl(t); } };
var _resetEpoch = 0, _lastSavedDraftSnapshot = null, __alerts = [], __fetches = [], __reply = {};
var state = { generatedProposalId: "p1", proposalId: "p1", timeUnit: "month", parsed: { client_name: "Acme" },
  enrichment: { campaign_name: "Holiday", internal_email_subject: "is", internal_email_body: "ib",
                client_email_subject: "cs", client_email_body: "cb" }, emailsPlanKey: null, emailFigures: null };
var __plan;
function makePlan(budget, note) {
  var lineA = { product_name: "Display", monthly_budget: budget, months: 3, is_added_value: false, monthly_allocations: null, notes_override: note || null };
  var lineB = { product_name: "Search", monthly_budget: 900, months: 3, is_added_value: false, monthly_allocations: { "2026-10": 1, "2026-11": 2, "2026-12": 3 } };
  return { request: { agency_fee: 0.15, start_date: "2026-10-01", end_date: "2026-12-31", total_months: 3, client_name: "Acme" },
    line_items: [lineA],
    tiers: [ { label: "A", name: "Independent", start_date: null, end_date: null, period_merge_groups: [], avails_data: { big: [1, 2, 3] }, line_items: [lineA] },
             { label: "B", name: "", start_date: "2026-10-01", end_date: "2027-01-31", period_merge_groups: [["2026-10", "2026-11"]], avails_data: {}, line_items: [lineB] } ],
    addons: [{ product_name: "Landing Page", amount: 100 }], monthly_distribution_mode: "even", time_unit: state.timeUnit };
}
__plan = makePlan(2000);
function _buildWizardStatePayload() { return JSON.parse(JSON.stringify(__plan)); }
function alert(m) { __alerts.push(String(m)); }
function fetch(url, opts) { __fetches.push({ url: url, body: JSON.parse(opts.body) }); return Promise.resolve({ ok: true, statusText: "OK", json: function () { return Promise.resolve(__reply); } }); }
function setText(id, t) { document.getElementById(id).textContent = t; }
function _flushAndSnapshotDraftState() { return "snap"; }
function _resetEmailRepromptScope() {}
function buildInternalEmailMailtoLink() { return Promise.resolve(); }
"""


def test_the_spa_revision_payload_the_stale_note_and_the_warnings_behave():
    fns = "\n".join(_extract_fn(n) for n in ("_emailFiguresPayload", "_emailPlanKey", "_refreshEmailStaleNote",
                                              "_showEmailFigureWarnings", "_submitEmailReprompt"))
    scenario = _JS_PRELUDE + fns + r"""
(async function () {
  var out = {};
  var note = function () { return document.getElementById("email-stale-note").classList.contains("hidden"); };
  var okReply = function () { return { internal_email_subject: "is2", internal_email_body: "ib2", client_email_subject: "cs2",
                                        client_email_body: "cb2", error: null, figure_warnings: [], email_figures: { options: [] } }; };
  var btn = { disabled: false, textContent: "" };

  // 1. a revision sends EVERY option, the add-ons, the allocations, the unit and the distribution mode (not avails)
  __reply = okReply();
  await _submitEmailReprompt({ text: "shorter", scope: "both", btn: btn, idleLabel: "Idle" });
  var sent = __fetches[0].body;
  out.url = __fetches[0].url;
  out.tiers = sent.tiers.map(function (t) { return [t.label, t.line_items.length, "avails_data" in t]; });
  out.addons = sent.addons; out.unit = sent.time_unit; out.mode = sent.monthly_distribution_mode;
  out.alloc = sent.tiers[1].line_items[0].monthly_allocations;
  out.merge = sent.tiers[1].period_merge_groups; out.tierBDates = [sent.tiers[1].start_date, sent.tiers[1].end_date];
  out.legacy = sent.line_items.length; out.scope = sent.scope; out.reprompt = sent.reprompt; out.current = sent.current_client_body;
  out.btn = [btn.disabled, btn.textContent];
  out.keySet = state.emailsPlanKey === _emailPlanKey();
  out.hiddenAfterFresh = note();
  out.enrichmentUpdated = state.enrichment.client_email_body;

  // 2. editing a figure marks the emails stale; editing only text that has no figure in it does not
  __plan = makePlan(3000); _refreshEmailStaleNote(); out.staleAfterBudgetEdit = !note();
  __plan = makePlan(2000, "only a note changed"); _refreshEmailStaleNote(); out.staleAfterNoteEdit = !note();
  __plan = makePlan(2000); _refreshEmailStaleNote(); out.hiddenWhenRestored = note();
  __plan.addons = []; _refreshEmailStaleNote(); out.staleAfterAddonRemoved = !note();
  __plan = makePlan(2000); __plan.time_unit = "week"; _refreshEmailStaleNote(); out.staleAfterUnitChange = !note();
  __plan = makePlan(2000); __plan.tiers[1].line_items[0].monthly_allocations["2026-12"] = 5; _refreshEmailStaleNote(); out.staleAfterAllocationEdit = !note();

  // 3. a one-sided revision against a changed plan leaves the other email stale; a refresh of both clears it
  __plan = makePlan(3500); _refreshEmailStaleNote();
  __reply = okReply();
  await _submitEmailReprompt({ text: "x", scope: "internal", btn: btn, idleLabel: "Idle" });
  out.staleAfterInternalOnly = !note();
  __reply = okReply();
  await _submitEmailReprompt({ text: "x", scope: "both", btn: btn, idleLabel: "Idle" });
  out.freshAfterBoth = note();

  // 4. a failed revision (error) never marks the emails fresh
  __plan = makePlan(4000); _refreshEmailStaleNote();
  _showEmailFigureWarnings(["$31,000 (client email): still wrong"]);
  __reply = okReply(); __reply.error = "OPENAI_API_KEY not set";
  await _submitEmailReprompt({ text: "x", scope: "both", btn: btn, idleLabel: "Idle" });
  out.staleAfterError = !note(); out.alerts = __alerts.length;
  out.warnKeptOnError = !document.getElementById("email-figures-warning").classList.contains("hidden");

  // 5. an edit while the revision is in flight still reads stale afterwards
  __plan = makePlan(4100); _refreshEmailStaleNote();
  __reply = okReply();
  var pending = _submitEmailReprompt({ text: "x", scope: "both", btn: btn, idleLabel: "Idle" });
  __plan = makePlan(4200);
  await pending;
  out.staleAfterEditDuringRevision = !note();

  // 6. the warnings from the server's figure check are listed (as text), and an empty list hides them
  var w = document.getElementById("email-figures-warning");
  __reply = okReply(); __reply.figure_warnings = ["$31,000 (client email): <b>total</b>", "$27,000 (internal email): x"];
  await _submitEmailReprompt({ text: "x", scope: "both", btn: btn, idleLabel: "Idle" });
  out.warnShown = !w.classList.contains("hidden");
  out.warnItems = w.children[1].children.map(function (li) { return li.textContent; });
  __reply = okReply();
  await _submitEmailReprompt({ text: "x", scope: "both", btn: btn, idleLabel: "Idle" });
  out.warnHidden = w.classList.contains("hidden");

  // 7. a reset while a revision runs: the late answer is dropped
  var before = state.enrichment.client_email_body;
  __reply = okReply(); __reply.client_email_body = "LATE";
  var p2 = _submitEmailReprompt({ text: "x", scope: "both", btn: btn, idleLabel: "Idle" });
  _resetEpoch++;
  await p2;
  out.lateDropped = state.enrichment.client_email_body === before;

  // 8. without a generated proposal nothing is sent
  var n = __fetches.length; state.enrichment = null;
  await _submitEmailReprompt({ text: "x", scope: "both", btn: btn, idleLabel: "Idle" });
  out.noSend = __fetches.length === n;
  console.log(JSON.stringify(out));
})();
"""
    out = json.loads(_run_node(scenario))
    assert out["url"] == "/api/proposal/p1/reprompt-emails"
    assert out["tiers"] == [["A", 1, False], ["B", 1, False]]                              # both options, avails stripped
    assert out["addons"] == [{"product_name": "Landing Page", "amount": 100}] and out["unit"] == "month" and out["mode"] == "even"
    assert out["alloc"] == {"2026-10": 1, "2026-11": 2, "2026-12": 3}
    assert out["merge"] == [["2026-10", "2026-11"]] and out["tierBDates"] == ["2026-10-01", "2027-01-31"]
    assert (out["legacy"], out["scope"], out["reprompt"], out["current"]) == (1, "both", "shorter", "cb")
    assert out["btn"] == [False, "Idle"] and out["keySet"] and out["hiddenAfterFresh"] and out["enrichmentUpdated"] == "cb2"
    assert out["staleAfterBudgetEdit"] and not out["staleAfterNoteEdit"] and out["hiddenWhenRestored"]
    assert out["staleAfterAddonRemoved"] and out["staleAfterUnitChange"] and out["staleAfterAllocationEdit"]
    assert out["staleAfterInternalOnly"] and out["freshAfterBoth"]
    assert out["staleAfterError"] and out["alerts"] == 1 and out["warnKeptOnError"]
    assert out["staleAfterEditDuringRevision"]
    assert out["warnShown"] and out["warnItems"] == ["$31,000 (client email): <b>total</b>", "$27,000 (internal email): x"]
    assert out["warnHidden"] and out["lateDropped"] and out["noSend"]


def test_the_gamma_outline_states_weekly_and_quarterly_amounts_in_their_own_unit():
    body = _extract_fn("buildGammaOutline")
    assert 'const unitAbbr = { week: "wk", quarter: "qtr" }[state.timeUnit] || "mo";' in body
    assert "/${unitAbbr} × ${li.months}${unitAbbr}" in body and "/mo × ${li.months}mo" not in body
    # right after a Generate the option totals come from the server's figures (add-ons and per-period amounts included)
    assert "state.emailFigures" in body and "fig.total_net" in body and "fig.addons_net" in body
