"""
Keyword-triggered disclaimer banners (app/disclaimers.py) and their two
renderers: the Excel footer (below the totals, ABOVE the legal lines) and the
PPT closing slide (which also now carries the Terms of Sales line).

Same DB-mocking approach as test_time_unit_export.py: by_name() and the meta
block's market lookup would otherwise need a real DATABASE_URL.
"""
from __future__ import annotations

import asyncio
import os
from datetime import date

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import pytest
from fastapi import HTTPException

from app import disclaimers as dz
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import LineItem, generate_proposal
from app.catalog import Product


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


def _rule(**kw) -> dz.Disclaimer:
    base = dict(id="r1", name="Rule", keywords=["live sports"], banner_text="Heads up.", color="red")
    base.update(kw)
    return dz.Disclaimer(**base)


# --- normalization / validation ---------------------------------------------

def test_normalize_keywords_splits_lowercases_and_dedupes():
    assert dz.normalize_keywords("Live Sports, tentpole\nLIVE  SPORTS;  nbc sports ,") == [
        "live sports", "tentpole", "nbc sports"]
    assert dz.normalize_keywords(None) == []


def test_normalize_keywords_rejects_oversized_input():
    with pytest.raises(ValueError):
        dz.normalize_keywords("x" * (dz.MAX_KEYWORD_CHARS + 1))
    with pytest.raises(ValueError):
        dz.normalize_keywords([f"kw{i}" for i in range(dz.MAX_KEYWORDS + 1)])


def test_normalize_mmdd_accepts_loose_forms_and_rejects_nonsense():
    assert dz.normalize_mmdd("3/7") == "03-07"
    assert dz.normalize_mmdd("11-15") == "11-15"
    assert dz.normalize_mmdd("02-29") == "02-29"
    assert dz.normalize_mmdd("  ") is None
    for bad in ("13-01", "02-30", "June 1", "6"):
        with pytest.raises(ValueError):
            dz.normalize_mmdd(bad)


def _fields(**kw):
    base = dict(name="N", keywords="a", banner_text="T", color="red", applies_to="both",
                season_start=None, season_end=None)
    base.update(kw)
    return base


def test_validate_fields_requires_a_trigger_and_both_season_dates():
    assert dz.validate_fields(**_fields())["keywords"] == ["a"]
    with pytest.raises(ValueError, match="never trigger"):
        dz.validate_fields(**_fields(keywords=""))
    with pytest.raises(ValueError, match="both season dates"):
        dz.validate_fields(**_fields(season_start="11-01"))
    # season-only (no keywords) is a valid trigger
    ok = dz.validate_fields(**_fields(keywords="", season_start="11-01", season_end="12-31"))
    assert ok["season_start"] == "11-01" and ok["keywords"] == []


@pytest.mark.parametrize("field,value", [("name", " "), ("banner_text", ""), ("color", "pink"), ("applies_to", "web")])
def test_validate_fields_rejects_bad_values(field, value):
    with pytest.raises(ValueError):
        dz.validate_fields(**_fields(**{field: value}))


# --- matching ----------------------------------------------------------------

def test_keywords_match_whole_words_and_phrases_only():
    assert dz.text_matches(["live sports"], "NBC Live Sports Stream") == ["live sports"]
    assert dz.text_matches(["sport"], "sports") == []          # not a substring match
    assert dz.text_matches(["tentpole"], "Tentpole events!") == ["tentpole"]
    assert dz.text_matches(["cbd"], "abcbd") == []


@pytest.mark.parametrize("start,end,flight,expected", [
    ("11-01", "12-31", (date(2026, 11, 15), date(2026, 12, 1)), True),
    ("11-01", "12-31", (date(2026, 9, 1), date(2026, 10, 31)), False),
    ("11-01", "12-31", (date(2026, 10, 20), date(2026, 11, 2)), True),           # overlaps the edge
    ("11-15", "01-05", (date(2026, 12, 20), date(2027, 1, 2)), True),             # wraps New Year
    ("11-15", "01-05", (date(2026, 2, 1), date(2026, 3, 1)), False),
    ("11-15", "01-05", (date(2026, 1, 1), date(2026, 1, 3)), True),               # wrap, early-January flight
    ("02-29", "03-02", (date(2027, 3, 1), date(2027, 3, 1)), True),               # Feb 29 in a non-leap year
])
def test_season_overlap(start, end, flight, expected):
    assert dz.season_overlaps(start, end, *flight) is expected


def test_season_never_matches_without_flight_dates():
    assert dz.season_overlaps("11-01", "12-31", None, date(2026, 11, 2)) is False


def test_matches_requires_keyword_and_season_when_both_are_set():
    rule = _rule(keywords=["political"], season_start="09-01", season_end="11-05")
    fall = (date(2026, 10, 1), date(2026, 10, 31))
    spring = (date(2026, 4, 1), date(2026, 4, 30))
    assert dz.matches(rule, "political campaign", *fall)
    assert not dz.matches(rule, "political campaign", *spring)   # right word, wrong season
    assert not dz.matches(rule, "auto dealer", *fall)            # right season, no word
    season_only = _rule(keywords=[], season_start="09-01", season_end="11-05")
    assert dz.matches(season_only, "anything", *fall)


def _req(**kw):
    base = dict(client_name="Acme Co", request_type="New Business",
                start_date="2026-10-01", end_date="2026-12-31", total_months=3)
    base.update(kw)
    return ProposalRequest(**base)


def test_resolve_is_per_option_so_a_sports_line_only_flags_its_own_option():
    tiers = [
        {"label": "A", "line_items": [LineItem(product_name="Search - SEM", monthly_budget=1000)]},
        {"label": "B", "line_items": [LineItem(product_name="NBC Sports Stream Sponsorship", monthly_budget=1000)]},
    ]
    out = dz.resolve_for_tiers(_req(), tiers, rules=dz.DEFAULT_DISCLAIMERS)
    assert out["A"] == []
    assert [b["name"] for b in out["B"]] == ["Live Sports Inventory"]
    # banner dicts carry their own colors so the exporters need no DB import
    assert out["B"][0]["fill"] == "F8D7DA" and out["B"][0]["font_color"] == "58151C"


def test_resolve_reads_request_level_and_line_level_text():
    rule = _rule(keywords=["election"])
    tiers = [{"label": "A", "line_items": [LineItem(product_name="Search - SEM", monthly_budget=1, target_override="Election-night viewers")]}]
    assert dz.resolve_for_tiers(_req(), tiers, rules=[rule])["A"]
    assert dz.resolve_for_tiers(_req(campaign_goal="Win the election"), tiers[:0] or [{"label": "A", "line_items": []}], rules=[rule])["A"]
    assert not dz.resolve_for_tiers(_req(), [{"label": "A", "line_items": []}], rules=[rule])["A"]


def test_resolve_uses_each_options_own_flight_for_seasonal_rules():
    rule = _rule(keywords=[], season_start="11-01", season_end="12-31")
    tiers = [
        {"label": "A", "line_items": [], "start_date": "2026-11-10", "end_date": "2026-12-10"},
        {"label": "B", "line_items": [], "start_date": "2026-03-01", "end_date": "2026-03-31"},
    ]
    out = dz.resolve_for_tiers(_req(), tiers, rules=[rule])
    assert out["A"] and not out["B"]


def test_applies_to_filters_and_ppt_union_dedupes_across_options():
    both = {"id": "1", "applies_to": "both", "name": "x", "text": "t", "fill": "F8D7DA", "font_color": "58151C"}
    excel_only = {**both, "id": "2", "applies_to": "excel"}
    ppt_only = {**both, "id": "3", "applies_to": "ppt"}
    by_tier = {"A": [both, excel_only], "B": [both, ppt_only]}
    assert [b["id"] for b in dz.union_for_ppt(by_tier)] == ["1", "3"]
    assert [b["id"] for b in dz.for_target([both, excel_only, ppt_only], "excel")] == ["1", "2"]


def test_rules_resolve_in_sort_order():
    r1, r2 = _rule(id="a", name="Zed", sort_order=5), _rule(id="b", name="Alpha", sort_order=50)
    out = dz.resolve_for_tiers(_req(client_name="live sports"), [{"label": "A", "line_items": []}], rules=[r2, r1])
    assert [b["id"] for b in out["A"]] == ["a", "b"]


# --- fail-safe loading --------------------------------------------------------

def test_load_active_falls_back_to_the_builtin_default_when_the_table_is_unreachable(monkeypatch):
    def boom():
        raise RuntimeError("relation \"disclaimers\" does not exist")
    monkeypatch.setattr(dz, "list_all", boom)
    rules = dz.load_active()
    assert [r.id for r in rules] == [dz.SEED_ID]


def test_load_active_excludes_inactive_rules(monkeypatch):
    monkeypatch.setattr(dz, "list_all", lambda: [_rule(id="on"), _rule(id="off", is_active=False)])
    assert [r.id for r in dz.load_active()] == ["on"]


def test_the_seeded_rule_cannot_be_deleted():
    with pytest.raises(ValueError):
        dz.delete(dz.SEED_ID)


def test_schema_declares_the_table_and_seed():
    from pathlib import Path
    sql = (Path(__file__).resolve().parents[1] / "schema.sql").read_text(encoding="utf-8")
    assert "CREATE TABLE IF NOT EXISTS disclaimers" in sql
    assert "'seed-live-sports'" in sql and "ON CONFLICT (id) DO NOTHING" in sql
    assert "CHECK (color IN ('red', 'amber', 'green', 'blue', 'gray'))" in sql


# --- admin endpoints ----------------------------------------------------------

def test_admin_create_validates_before_touching_the_db(monkeypatch):
    import app.main as m
    called = []
    monkeypatch.setattr(dz, "create", lambda f: called.append(f))
    body = m.DisclaimerRequest(name="Bad", keywords="", banner_text="x")
    with pytest.raises(HTTPException) as exc:
        asyncio.run(m.admin_create_disclaimer(body))
    assert exc.value.status_code == 400 and not called


def test_admin_create_passes_normalized_fields(monkeypatch):
    import app.main as m
    captured = {}

    def fake_create(fields):
        captured.update(fields)
        return _rule(id="new", keywords=fields["keywords"])
    monkeypatch.setattr(dz, "create", fake_create)
    body = m.DisclaimerRequest(name=" Holiday ", keywords="Black Friday, CYBER monday", banner_text=" Book early. ", color="amber")
    out = asyncio.run(m.admin_create_disclaimer(body))
    assert out["created"] is True
    assert captured["name"] == "Holiday" and captured["banner_text"] == "Book early."
    assert captured["keywords"] == ["black friday", "cyber monday"]


def test_admin_update_toggle_and_delete_map_missing_ids_to_404(monkeypatch):
    import app.main as m
    monkeypatch.setattr(dz, "update", lambda i, f: None)
    monkeypatch.setattr(dz, "set_active", lambda i, a: None)
    monkeypatch.setattr(dz, "delete", lambda i: False)
    with pytest.raises(HTTPException) as e1:
        asyncio.run(m.admin_update_disclaimer("nope", m.DisclaimerRequest(name="n", keywords="k", banner_text="t")))
    with pytest.raises(HTTPException) as e2:
        asyncio.run(m.admin_toggle_disclaimer("nope", m.DisclaimerActiveRequest(is_active=False)))
    with pytest.raises(HTTPException) as e3:
        asyncio.run(m.admin_delete_disclaimer("nope"))
    assert (e1.value.status_code, e2.value.status_code, e3.value.status_code) == (404, 404, 404)


def test_admin_delete_of_the_seed_rule_is_a_400(monkeypatch):
    import app.main as m
    with pytest.raises(HTTPException) as exc:
        asyncio.run(m.admin_delete_disclaimer(dz.SEED_ID))
    assert exc.value.status_code == 400


def test_preview_endpoint_uses_the_same_resolver(monkeypatch):
    import app.main as m
    monkeypatch.setattr(dz, "load_active", lambda: list(dz.DEFAULT_DISCLAIMERS))
    body = m.DisclaimerPreviewRequest(
        request={"client_name": "Acme", "start_date": "2026-10-01", "end_date": "2026-12-31"},
        tiers=[m.TierModel(label="A", line_items=[m.LineItemModel(product_name="Fox Sports Go Video Sponsorship", monthly_budget=1000)])],
    )
    out = asyncio.run(m.preview_disclaimers(body))
    assert [d["name"] for d in out["tiers"]["A"]] == ["Live Sports Inventory"]


# --- Excel rendering ---------------------------------------------------------

def _banner(text, fill="F8D7DA", font="58151C", bid="b1"):
    return {"id": bid, "name": "n", "text": text, "fill": fill, "font_color": font, "applies_to": "both", "color": "red"}


def _text_rows(ws):
    """{row: first non-empty text} for rows with a string in column C."""
    return {r: ws.cell(row=r, column=3).value for r in range(1, ws.max_row + 1)
            if isinstance(ws.cell(row=r, column=3).value, str)}


def _generate(tmp_path, disclaimers_by_tier, name="x.xlsx", **kw):
    req = _req()
    li = LineItem(product_name="Search - SEM", monthly_budget=1200.0, months=3, id="li1")
    out = tmp_path / name
    generate_proposal(req, [li], out, disclaimers_by_tier=disclaimers_by_tier, **kw)
    import openpyxl
    return openpyxl.load_workbook(out)


def test_excel_banner_sits_below_the_total_and_above_the_legal_lines(tmp_path):
    wb = _generate(tmp_path, {"A": [_banner("Live Sports Inventory: re-avail required.")]})
    ws = wb["Proposal A"]
    rows = _text_rows(ws)
    banner_row = next(r for r, t in rows.items() if t.startswith("Live Sports Inventory"))
    grand_row = next(r for r, t in rows.items() if "TOTAL DIGITAL" in str(t) and r > 17 and ws.cell(row=r, column=12).value
                     and str(ws.cell(row=r, column=12).value).startswith("=ROUNDDOWN"))
    gtm_row = next(r for r, t in rows.items() if t.startswith("To maximize the efficiency"))
    validity_row = next(r for r, t in rows.items() if t.startswith("This proposal will be valid"))
    assert grand_row < banner_row < gtm_row < validity_row
    cell = ws.cell(row=banner_row, column=3)
    assert cell.fill.start_color.rgb == "FFF8D7DA"
    assert cell.font.color.rgb == "FF58151C"
    assert ws.row_dimensions[banner_row].height >= 26
    # merged across the box, wrapping
    assert any(rng.min_row == banner_row and rng.min_col == 3 for rng in ws.merged_cells.ranges)
    assert cell.alignment.wrap_text


def test_excel_without_banners_has_no_banner_and_unchanged_footer(tmp_path):
    wb = _generate(tmp_path, {"A": []})
    ws = wb["Proposal A"]
    rows = _text_rows(ws)
    assert not any("Live Sports" in t for t in rows.values())
    # GTM note still lands two rows below the grand total, as before
    gtm = next(r for r, t in rows.items() if t.startswith("To maximize the efficiency"))
    grand = max(r for r in range(1, gtm) if str(ws.cell(row=r, column=3).value or "").startswith("=\"TOTAL DIGITAL"))
    assert gtm == grand + 2


def test_excel_multiple_banners_and_excel_only_filtering(tmp_path):
    ppt_only = {**_banner("PPT ONLY", bid="p"), "applies_to": "ppt"}
    wb = _generate(tmp_path, {"A": [_banner("First banner", bid="1"), _banner("Second banner", fill="FFF3CD", font="664D03", bid="2"), ppt_only]})
    rows = _text_rows(wb["Proposal A"])
    texts = list(rows.values())
    assert "First banner" in texts and "Second banner" in texts
    assert "PPT ONLY" not in texts
    first = next(r for r, t in rows.items() if t == "First banner")
    second = next(r for r, t in rows.items() if t == "Second banner")
    assert second > first
    assert wb["Proposal A"].cell(row=second, column=3).fill.start_color.rgb == "FFFFF3CD"


def test_excel_each_option_gets_only_its_own_banners(tmp_path):
    req = _req()
    li_a = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="a")
    li_b = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="b")
    tiers = [{"label": "A", "line_items": [li_a], "avails_data": {}}, {"label": "B", "line_items": [li_b], "avails_data": {}}]
    out = tmp_path / "tiers.xlsx"
    generate_proposal(req, [], out, tiers=tiers, disclaimers_by_tier={"B": [_banner("Only on B")]})
    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws_a = next(w for w in wb.worksheets if "Option A" in w.title)
    ws_b = next(w for w in wb.worksheets if "Option B" in w.title)
    assert not any(t == "Only on B" for t in _text_rows(ws_a).values())
    assert any(t == "Only on B" for t in _text_rows(ws_b).values())


def test_excel_gross_and_avails_only_sheets_also_carry_the_banner(tmp_path):
    req = _req(agency_fee=0.15)
    li = LineItem(product_name="Search - SEM", monthly_budget=1200.0, months=3, id="li1")
    out = tmp_path / "gross.xlsx"
    generate_proposal(req, [li], out, force_tabs={"net": True, "gross": True, "avails_only": True},
                      avails_data={"li1": {"max_imps": 100000, "max_spend": 5000}},
                      disclaimers_by_tier={"A": [_banner("Everywhere banner")]})
    import openpyxl
    wb = openpyxl.load_workbook(out)
    carrying = [ws.title for ws in wb.worksheets if any(t == "Everywhere banner" for t in _text_rows(ws).values())]
    assert any("Gross" in t for t in carrying)
    assert any("Avails" in t for t in carrying)
    assert any(t == "Proposal A" for t in carrying)


def test_signature_block_still_follows_the_footer_with_banners(tmp_path):
    wb = _generate(tmp_path, {"A": [_banner("b1", bid="1"), _banner("b2", bid="2")]})
    rows = _text_rows(wb["Proposal A"])
    sig = next(r for r, t in rows.items() if t == "Customer Signature")
    terms = next(r for r, t in rows.items() if t.startswith("Client accepts Entravision"))
    assert sig > terms


# --- PPT rendering -----------------------------------------------------------

pptx = pytest.importorskip("pptx")


def _product(name="Search - SEM"):
    return Product(
        family="Search", name=name, short_label=name, proposal_description="",
        sizes="", buying_model="CPM", base_rate=10.0, estimated_impressions=False,
        discloses_impressions=True, minimum_spend=0.0, minimum_flight_days=(0, 0),
        sla_data_days=None, sla_creative_days=None, sla_activate_days=None, sla_total_days=None,
        media_allocation_pct=1.0, margin_upper=0.5, margin_lower=0.5,
        tech_platform="", wide_orbit_code="",
    )


def _deck(tmp_path, banners=None, products=1, gross=False):
    from app.services import pptx_builder
    prods = [_product(f"Product {i}") for i in range(products)]
    lis = [LineItem(product_name=p.name, monthly_budget=1000.0, months=3, id=f"l{i}") for i, p in enumerate(prods)]
    out = tmp_path / ("g.pptx" if gross else "n.pptx")
    ok = pptx_builder.build_signature_deck(
        _req(agency_fee=0.15 if gross else None),
        [{"name": "Option A", "products": prods, "line_items": lis}],
        out, gross=gross, proposal_title="Title", banners=banners,
    )
    assert ok
    return pptx.Presentation(str(out))


def _all_text(prs):
    return [sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame]


def test_ppt_carries_the_validity_and_terms_of_sales_text_with_a_real_link(tmp_path):
    prs = _deck(tmp_path)
    text = "\n".join(_all_text(prs))
    assert "This proposal will be valid for a period of 1 month after being presented." in text
    assert "Please notify your Account Executive if you require the presented media to remain booked after that time." in text
    assert "Client accepts Entravision's Terms of Sales (https://entravision.com/termsofsales/)" in text
    links = [r.hyperlink.address for sl in prs.slides for sh in sl.shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
    assert links == ["https://entravision.com/termsofsales/"]


def test_ppt_rate_wording_follows_net_vs_gross(tmp_path):
    assert "All rates are NET" in "\n".join(_all_text(_deck(tmp_path)))
    assert "All rates are GROSS" in "\n".join(_all_text(_deck(tmp_path, gross=True)))


def test_ppt_banners_render_above_the_legal_copy_with_their_colors(tmp_path):
    prs = _deck(tmp_path, banners=[{"id": "1", "text": "Live Sports Inventory: re-avail needed.", "fill": "F8D7DA", "font_color": "58151C"}])
    banner_shape = next(sh for sl in prs.slides for sh in sl.shapes
                        if sh.has_text_frame and sh.text_frame.text.startswith("Live Sports Inventory"))
    legal_shape = next(sh for sl in prs.slides for sh in sl.shapes
                       if sh.has_text_frame and sh.text_frame.text.startswith("This proposal will be valid"))
    assert str(banner_shape.fill.fore_color.rgb) == "F8D7DA"
    assert str(banner_shape.text_frame.paragraphs[0].runs[0].font.color.rgb) == "58151C"
    assert banner_shape.top + banner_shape.height <= legal_shape.top


def test_ppt_nothing_overflows_the_slide_even_with_many_long_banners(tmp_path):
    from app.services import pptx_builder as pb
    long_text = "Live Sports Inventory: " + ("geo-based estimates, mandatory re-avail buffer. " * 8)
    banners = [{"id": str(i), "text": long_text, "fill": "F8D7DA", "font_color": "58151C"} for i in range(5)]
    for products in (1, 12):
        prs = _deck(tmp_path, banners=banners, products=products)
        for slide in prs.slides:
            for sh in slide.shapes:
                if sh.has_text_frame and (sh.text_frame.text.startswith("Live Sports") or sh.text_frame.text.startswith("This proposal will")
                                          or sh.text_frame.text in ("Client", "Signature", "Date")):
                    assert sh.top + sh.height <= pb.FOOTER_Y + 1, (products, sh.text_frame.text[:30])
        # the full block always lands on exactly one slide
        holders = [sl for sl in prs.slides if any(sh.has_text_frame and sh.text_frame.text.startswith("This proposal will") for sh in sl.shapes)]
        assert len(holders) == 1


def test_ppt_banners_push_the_closing_block_to_its_own_slide_when_it_no_longer_fits(tmp_path):
    base = len(_deck(tmp_path, products=3).slides)
    long_text = "Warning. " * 60
    banners = [{"id": str(i), "text": long_text, "fill": "F8D7DA", "font_color": "58151C"} for i in range(4)]
    assert len(_deck(tmp_path, banners=banners, products=3).slides) >= base
