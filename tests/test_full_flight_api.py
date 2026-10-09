"""
Full Flight billing period (time_unit == "full_flight") — the pieces below the export: the single-period
builder in monthly_allocation.py, strict/lenient request validation in main.py, the /api/generate
handler's server-side normalization, reopen coercion, and the recommender.

(The workbook itself is covered by test_full_flight_export.py.) No Postgres: the handler is called
directly with a fake Request and monkeypatched DB/AI/file helpers, like test_draft_save.py.
"""
from __future__ import annotations

import asyncio
import os
from datetime import date
from types import SimpleNamespace

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import openpyxl
import pydantic
import pytest

from app import main
from app.services import monthly_allocation as mo
from app.services.ai_enricher import ProposalEnrichment
from app.services.notion_parser import ProposalRequest
from app.services.recommender import recommend_line_items

EMAIL = "Email Campaigns - Display Re-targeting"
IAB = "eDigital Network Display - Standard IAB"


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.catalog.resolve_product_alias", lambda name: None)
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


# ---------------------------------------------------------------------------
# monthly_allocation: the single period
# ---------------------------------------------------------------------------

def _flight(start, end):
    return mo.periods_between(date.fromisoformat(start), date.fromisoformat(end), "full_flight")


def test_the_flight_is_exactly_one_period_with_a_constant_key():
    periods = _flight("2026-11-24", "2026-12-31")
    assert len(periods) == 1
    p = periods[0]
    assert p["key"] == mo.FULL_FLIGHT_KEY == "full_flight" and "+" not in p["key"]
    assert p["label"] == "November–December 2026"
    assert p["date_range_label"] == "Nov 24 – Dec 31, 2026"
    assert (p["start"], p["end"]) == (date(2026, 11, 24), date(2026, 12, 31))
    # the flight spans two calendar months but is ONE period: nothing (minimum spend, SOV) is scaled by a month count
    assert (p["active_days"], p["days_in_month"], p["period_count"]) == (38, 61, 1)


@pytest.mark.parametrize("start,end,label", [
    ("2026-11-03", "2026-11-20", "November 2026"),
    ("2026-12-15", "2027-01-20", "December 2026–January 2027"),
    ("2026-11-01", "2027-01-31", "November 2026–January 2027"),
    ("2026-12-31", "2026-11-24", "November–December 2026"),               # swapped dates are repaired, like the other units
])
def test_flight_labels_and_the_period_count_is_always_one(start, end, label):
    [p] = _flight(start, end)
    assert (p["label"], p["period_count"], p["key"]) == (label, 1, "full_flight")


def test_the_key_does_not_change_when_the_dates_do():
    assert _flight("2026-11-24", "2026-12-31")[0]["key"] == _flight("2027-02-01", "2027-06-30")[0]["key"]


def test_allocation_helpers_degenerate_to_one_hundred_percent():
    periods = _flight("2026-11-24", "2026-12-31")
    for mode in ("even", "prorated"):
        assert mo.compute_default_allocation(680.0, periods, mode) == {"full_flight": 680.0}
    assert mo.reconcile_allocation(680.0, {"full_flight": 680.0})["balanced"]


def test_unit_constants_and_lenient_normalizer():
    assert mo.TIME_UNITS == ("week", "month", "quarter", "full_flight")
    assert mo.is_full_flight("full_flight") and not mo.is_full_flight("month")
    for raw, expected in [(None, "month"), ("", "month"), (" Full_Flight ", "full_flight"), ("banana", "month"), (3, "month"),
                          ("WEEK", "week")]:
        assert mo.normalize_time_unit(raw) == expected
    # the pre-existing fallback is intentionally still there for DIRECT callers of periods_between ...
    assert [p["key"] for p in mo.periods_between(date(2026, 11, 24), date(2026, 12, 31), "bogus")] == ["2026-11", "2026-12"]


# ---------------------------------------------------------------------------
# main.py request validation
# ---------------------------------------------------------------------------

def _generate_body(**overrides) -> "main.GenerateRequest":
    fields = dict(request={"client_name": "Acme", "start_date": "2026-11-24", "end_date": "2026-12-31", "total_months": 2,
                           "agency_fee": 0.15, "request_type": "New Business"},
                  line_items=[main.LineItemModel(product_name=EMAIL, monthly_budget=680.0, months=1, id="a")],
                  time_unit="full_flight")
    fields.update(overrides)
    return main.GenerateRequest(**fields)


def test_generate_accepts_full_flight_and_rejects_unknown_units():
    assert _generate_body().time_unit == "full_flight"
    assert _generate_body(time_unit=" Full_Flight ").time_unit == "full_flight"
    assert _generate_body(time_unit="month").time_unit == "month"
    assert main.GenerateRequest(request={}, line_items=[]).time_unit == "month"        # default
    assert _generate_body(time_unit=None).time_unit == "month"
    for bad in ("banana", "full-flight", "flight", "quarterly", 7):
        with pytest.raises(pydantic.ValidationError):
            _generate_body(time_unit=bad)


def test_drafts_and_the_other_requests_coerce_instead_of_failing():
    # a draft must ALWAYS be saveable
    assert main.DraftSaveRequest(time_unit="full_flight").time_unit == "full_flight"
    assert main.DraftSaveRequest(time_unit="banana").time_unit == "month"
    assert main.DraftSaveRequest(time_unit=None).time_unit == "month"
    assert main.StrategyRequest(request={}, time_unit="full_flight").time_unit == "full_flight"
    assert main.StrategyRequest(request={}, time_unit="x").time_unit == "month"
    assert main.StrategyDocRebuildRequest(request={}, time_unit="full_flight").time_unit == "full_flight"
    assert main.RepromptEmailsRequest(request={}, line_items=[], reprompt="r", time_unit="full_flight").time_unit == "full_flight"
    # Suggest Mix is not the export: an unknown unit degrades to month exactly as it always did
    assert main.RecommendRequest(request={}, monthly_budget=1000, time_unit="full_flight").time_unit == "full_flight"
    assert main.RecommendRequest(request={}, monthly_budget=1000, time_unit="banana").time_unit == "month"
    assert main.RecommendRequest(request={}, monthly_budget=1000, time_unit=" Full_Flight ").time_unit == "full_flight"
    # a client still sending the retired `flight_months` hint is simply ignored (Full Flight is one period: no month count)
    assert not hasattr(main.RecommendRequest(request={}, monthly_budget=1000, time_unit="full_flight", flight_months=2), "flight_months")


def test_reopen_state_round_trips_the_unit_and_reopen_coerces_unknown_ones(monkeypatch):
    state = main._build_reopen_state(
        request_dict={}, line_items=[], tiers=None, avails_data=None, strategy_brief=None, roadblocks=None,
        force_tabs=None, addons=[], raw_notion_text=None, enrichment=None, time_unit="full_flight",
        monthly_distribution_mode="even", campaign_name_override=None, wizard_step=8)
    assert state["time_unit"] == "full_flight"
    for stored, expected in [("full_flight", "full_flight"), (None, "month"), ("banana", "month"), ("week", "week")]:
        rs = dict(state, time_unit=stored)
        monkeypatch.setattr(main, "_get_proposal_metadata", lambda pid, rs=rs: {"reopen_state": rs, "status": "generated"})
        assert asyncio.run(main.reopen_proposal("p1"))["time_unit"] == expected


# ---------------------------------------------------------------------------
# The /api/generate handler
# ---------------------------------------------------------------------------

def _run_generate(monkeypatch, tmp_path, body, deck_spy=None):
    saved = {}
    monkeypatch.setattr(main, "PROPOSALS_DIR", tmp_path)
    monkeypatch.setattr(main, "_get_next_short_id", lambda: "EVC-1")
    monkeypatch.setattr(main, "_resolve_draftable_proposal_id", lambda requested: "pid-1")
    monkeypatch.setattr(main, "_save_proposal_metadata", lambda **kw: saved.update(kw))
    monkeypatch.setattr(main.ai_enricher, "enrich_proposal", lambda *a, **k: ProposalEnrichment(campaign_name="Holiday"))
    monkeypatch.setattr(main.disclaimers_svc, "resolve_for_tiers", lambda req, tiers: {})
    monkeypatch.setattr(main.disclaimers_svc, "union_for_ppt", lambda d: [])
    monkeypatch.setattr(main.pptx_builder, "build_signature_deck", deck_spy or (lambda *a, **k: None))
    monkeypatch.setattr(main.docx_builder, "build_client_email_docx", lambda **kw: None)
    request = SimpleNamespace(state=SimpleNamespace(user={"email": "planner@entravision.com", "is_admin": False}),
                              client=SimpleNamespace(host="127.0.0.1"), headers={"user-agent": "pytest"})
    return asyncio.run(main.generate(body, request)), saved


def test_generate_exports_a_stale_month_payload_as_one_full_flight_period(monkeypatch, tmp_path):
    """A line saved in month mode (340 x 2, month-keyed allocations) arriving as full_flight must come out as one
    period with its total (680) intact — never 2x'd, never silently replaced by the stale allocations."""
    stale = main.LineItemModel(product_name=EMAIL, monthly_budget=340.0, months=2, id="a",
                               monthly_allocations={"2026-11": 100.0, "2026-12": 580.0})
    result, saved = _run_generate(monkeypatch, tmp_path, _generate_body(line_items=[stale]))

    assert result["summary"]["total_net"] == pytest.approx(680.0)
    assert result["summary"]["total_gross"] == pytest.approx(800.0)
    assert any("folded into one period" in w for w in result["summary"]["warnings"])
    ws = openpyxl.load_workbook(tmp_path / result["filename"])["Proposal A"]
    assert ws["I10"].value == 1 and ws["L19"].value == 680.0
    assert ws["S17"].value == "NOVEMBER–DECEMBER 2026"
    # no false below-minimum warnings from stale month keys, and none for the folded line either: the flight is ONE
    # period, so $680 only has to clear ONE $375 monthly minimum (not 2 x)
    assert result["monthly_breakdown_warnings"] == []
    # the reopen state keeps what the CLIENT sent and records the unit
    assert saved["reopen_state"]["time_unit"] == "full_flight"
    assert saved["reopen_state"]["line_items"][0]["months"] == 2


def test_generate_a_clean_full_flight_payload_has_no_warnings_and_exact_totals(monkeypatch, tmp_path):
    lines = [main.LineItemModel(product_name=EMAIL, monthly_budget=1500.0, months=1, id="a"),
             main.LineItemModel(product_name=IAB, monthly_budget=4000.0, months=1, id="b")]
    result, _ = _run_generate(monkeypatch, tmp_path, _generate_body(line_items=lines))
    assert result["summary"]["total_net"] == pytest.approx(5500.0)
    assert result["summary"]["warnings"] == []
    assert result["monthly_breakdown_warnings"] == [] and result["monthly_breakdown_balance_notes"] == []


def test_generate_month_mode_is_unchanged(monkeypatch, tmp_path):
    lines = [main.LineItemModel(product_name=EMAIL, monthly_budget=1000.0, months=2, id="a")]
    result, _ = _run_generate(monkeypatch, tmp_path, _generate_body(line_items=lines, time_unit="month"))
    assert result["summary"]["total_net"] == pytest.approx(2000.0)
    ws = openpyxl.load_workbook(tmp_path / result["filename"])["Proposal A"]
    assert ws["I10"].value == 2 and ws["H10"].value == "Months:"


# ---------------------------------------------------------------------------
# Recommender (Suggest Mix)
# ---------------------------------------------------------------------------

def _recommend_request(**overrides) -> ProposalRequest:
    fields = dict(client_name="Acme", request_type="New Business", start_date="2026-11-24", end_date="2026-12-31",
                  total_months=2, campaign_goal="Awareness", products_selected=[])
    fields.update(overrides)
    return ProposalRequest(**fields)


def test_recommender_gives_every_full_flight_line_one_period_and_one_monthly_minimum():
    req = _recommend_request()
    items = recommend_line_items(req, 15000.0, time_unit="full_flight")
    assert items and all(li.months == 1 for li in items)
    assert sum(li.monthly_budget for li in items) <= 15000.0 + 50
    from app.catalog import by_name
    for li in items:
        assert li.monthly_budget >= (by_name(li.product_name).minimum_spend or 0) - 1e-6    # ONE monthly minimum
    # month mode still stamps the calendar span
    assert all(li.months == 2 for li in recommend_line_items(req, 15000.0, time_unit="month"))


def test_the_flight_length_does_not_change_a_full_flight_suggestion():
    short = recommend_line_items(_recommend_request(), 15000.0, time_unit="full_flight")
    long_ = recommend_line_items(_recommend_request(start_date="2026-09-01", end_date="2027-02-28", total_months=6),
                                 15000.0, time_unit="full_flight")
    assert [(li.product_name, li.monthly_budget) for li in short] == [(li.product_name, li.monthly_budget) for li in long_]
    assert all(li.months == 1 for li in long_)


def test_generate_a_full_flight_line_without_months_is_one_period_not_the_models_default_of_three(monkeypatch, tmp_path):
    """LineItemModel.months defaults to 3: a non-SPA client that omits it in full flight must not get its budget folded
    x3 (the SPA always sends months = 1). An EXPLICIT months > 1 is still folded, total preserved."""
    omitted = main.LineItemModel(product_name=EMAIL, monthly_budget=680.0, id="a")
    assert "months" not in omitted.model_fields_set
    result, _ = _run_generate(monkeypatch, tmp_path, _generate_body(line_items=[omitted]))
    assert result["summary"]["total_net"] == pytest.approx(680.0)
    assert result["summary"]["warnings"] == []
    # month mode still takes the model default exactly as before
    result_m, _ = _run_generate(monkeypatch, tmp_path, _generate_body(line_items=[omitted], time_unit="month"))
    assert result_m["summary"]["total_net"] == pytest.approx(680.0 * 3)
    explicit = main.LineItemModel(product_name=EMAIL, monthly_budget=340.0, months=2, id="b")
    result_e, _ = _run_generate(monkeypatch, tmp_path, _generate_body(line_items=[explicit]))
    assert result_e["summary"]["total_net"] == pytest.approx(680.0)
    assert any("folded into one period" in w for w in result_e["summary"]["warnings"])


def test_the_pptx_tiers_carry_each_options_own_flight_window(monkeypatch, tmp_path):
    seen = {}
    body = _generate_body(
        line_items=[],
        tiers=[main.TierModel(label="A", line_items=[main.LineItemModel(product_name=EMAIL, monthly_budget=680.0, months=1, id="a")]),
               main.TierModel(label="B", start_date="2027-01-04", end_date="2027-03-31",
                              line_items=[main.LineItemModel(product_name=EMAIL, monthly_budget=500.0, months=1, id="b")])])
    _run_generate(monkeypatch, tmp_path, body, deck_spy=lambda req, tiers, *a, **k: seen.setdefault("tiers", tiers))
    windows = [(t["start_date"], t["end_date"]) for t in seen["tiers"]]
    assert windows[1] == ("2027-01-04", "2027-03-31")
    assert not (windows[0][0] or windows[0][1])            # an option with no override carries none (the deck uses the campaign dates)


def test_the_brief_path_is_one_period_with_one_monthly_minimum_too():
    req = _recommend_request()
    brief = {"recommended_tactics": [{"product_family": "Email", "suggested_budget_pct": 30},
                                     {"product_family": "Display", "suggested_budget_pct": 40},
                                     {"product_family": "Social", "suggested_budget_pct": 30}]}
    from app.catalog import by_name
    items = recommend_line_items(req, 3000.0, strategy_brief=brief, time_unit="full_flight")
    assert items
    for li in items:
        assert li.months == 1
        assert li.monthly_budget >= (by_name(li.product_name).minimum_spend or 0) - 1e-6, (li.product_name, li.monthly_budget)


def test_the_reprompt_handler_passes_the_billing_period_to_the_revision(monkeypatch):
    seen = {}
    monkeypatch.setattr(main.ai_enricher, "reprompt_emails", lambda *a, **k: seen.update(k) or {
        "internal_email_subject": "s", "internal_email_body": "b", "client_email_subject": "s", "client_email_body": "b", "error": None})
    monkeypatch.setattr(main, "_get_proposal_metadata", lambda pid: None)
    body = main.RepromptEmailsRequest(request={"client_name": "Acme"}, line_items=[], reprompt="shorter", time_unit="full_flight")
    asyncio.run(main.reprompt_emails("pid", body))
    assert seen["time_unit"] == "full_flight" and seen["scope"] == "both"
    seen.clear()
    asyncio.run(main.reprompt_emails("pid", main.RepromptEmailsRequest(request={"client_name": "Acme"}, line_items=[], reprompt="shorter")))
    assert seen["time_unit"] == "month"
