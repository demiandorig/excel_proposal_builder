"""
Regression tests for the adversarial-review fixes on the Disclaimers / Restrictions / blurb / export work:
restrictions parsing + renamed products, Roadblocks contradiction filtering, per-option disclaimer scoping,
banner sanitizing, PPT overflow, Excel edge cases, abbreviation-aware blurb tightening.
"""
from __future__ import annotations

import io
import os
from datetime import date

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import pytest

from app import disclaimers as dz
from app.catalog import Product
from app.services import restrictions as rs
from app.services import roadblocks as rb
from app.services.ai_enricher import _tighten_blurb
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import LineItem, generate_proposal


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


@pytest.fixture
def builtin_state(monkeypatch):
    monkeypatch.setattr(rs, "load_state", lambda: {"categories": rs.builtin_categories(), "using_builtin": True,
                                                   "sync": None, "error": False})


def _items(*names):
    return [LineItem(product_name=n, monthly_budget=1000, months=3, id=f"i{i}") for i, n in enumerate(names)]


def _result(*roadblocks, summary="s"):
    return {"overall_summary": summary, "product_roadblocks": list(roadblocks), "categories": [],
            "used_web_search": True, "error": None}


def _workbook(builder) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    builder(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Roadblocks: conditional caveats survive, flat denials don't ------------------------------------

POLITICAL = [c for c in rs.builtin_categories() if c.name == "Political"]


@pytest.mark.parametrize("sentence", [
    "Political ads cannot run on Meta without authorization and a Paid for by disclaimer.",
    "Ads will not be served until the candidate's disclosure is approved.",
    "Political advertisers cannot run ads on Google's display network until verification is complete.",
    "Political ads will not be served without a 'Paid for by' disclosure.",
    "Candidate creative cannot be served without sponsor ID.",
    "Pac-12 Network inventory is not available.",
])
def test_conditional_caveats_are_not_treated_as_denials(sentence):
    assert not rb._denies_category(sentence, POLITICAL)


@pytest.mark.parametrize("sentence", [
    "Political ads are not allowed on Display.",
    "Political ads aren't allowed on this product.",
    "This product doesn't accept political ads.",
    "Political ads are unavailable on this product.",
    "Political creative was rejected by the platform.",
    "Political ads are not an option.",
    "Political ads should not be run on Display.",
    "Political advertisers are ineligible.",
    "Political ads are prohibited.",
])
def test_flat_denials_of_an_allowed_vertical_are_caught(sentence):
    assert rb._denies_category(sentence, POLITICAL)


def test_denials_are_scrubbed_from_the_mitigation_and_summary_and_the_risk_level_is_recomputed(builtin_state):
    items = _items("Video - Pre-roll (OLV)", "YouTube Ads")
    matrix = rs.build_matrix([li.product_name for li in items], ["Political"])
    result = _result(
        {"product_name": "Video - Pre-roll (OLV)", "risk_level": "high",
         "risks": [{"issue": "Political ads not feasible", "detail": "Political ads are not allowed on OLV.", "source": "x"},
                   {"issue": "Additional: disclosure", "detail": "Political ads cannot run without a Paid for by disclaimer.", "source": "y"}],
         "recommended_mitigation": "Do not run political ads on OLV. Add the Paid for by disclaimer to every creative."},
        summary="Political ads are not allowed on OLV. Plan for creative approval lead time.",
    )
    rb._apply_matrix(result, matrix, items)
    entry = next(r for r in result["product_roadblocks"] if r["product_name"] == "Video - Pre-roll (OLV)")
    assert [r["issue"] for r in entry["risks"]] == ["Additional: disclosure"]     # the caveat stayed, the denial went
    assert "Do not run" not in entry["recommended_mitigation"] and "disclaimer" in entry["recommended_mitigation"]
    assert entry["risk_level"] == "medium"                                          # no longer "high" off a claim that was removed
    assert "not allowed on OLV" not in result["overall_summary"] and "lead time" in result["overall_summary"]


def test_a_summary_that_was_only_a_denial_gets_a_neutral_replacement(builtin_state):
    items = _items("Video - Pre-roll (OLV)")
    matrix = rs.build_matrix(["Video - Pre-roll (OLV)"], ["Political"])
    result = _result({"product_name": "Video - Pre-roll (OLV)", "risk_level": "low", "risks": [], "recommended_mitigation": ""},
                     summary="Political ads cannot be run on OLV.")
    rb._apply_matrix(result, matrix, items)
    assert "Restricted verticals confirmed: Political" in result["overall_summary"]


def test_a_summary_sentence_about_a_refused_product_is_left_alone(builtin_state):
    items = _items("Video - Pre-roll (OLV)", "YouTube Ads")
    matrix = rs.build_matrix([li.product_name for li in items], ["Political"])
    summary = "Political ads are not allowed on YouTube."
    result = _result(summary=summary)
    rb._apply_matrix(result, matrix, items)
    assert result["overall_summary"] == summary


def test_the_models_retyped_product_name_still_matches_and_repeats_dont_duplicate_the_injected_risk(builtin_state):
    items = _items("YouTube Ads", "YouTube Ads")
    matrix = rs.build_matrix(["YouTube Ads"], ["Political"])
    result = _result({"product_name": "  youtube ads ", "risk_level": "low", "risks": [], "recommended_mitigation": ""})
    rb._apply_matrix(result, matrix, items)
    assert len(result["product_roadblocks"]) == 1
    assert [r["issue"] for r in result["product_roadblocks"][0]["risks"]] == ["Not accepted for Political"]


def test_confirmed_verticals_missing_from_the_sheet_are_reported(builtin_state):
    items = _items("YouTube Ads")
    matrix = rs.build_matrix(["YouTube Ads"], ["Political", "Gone Vertical"])
    result = _result()
    rb._apply_matrix(result, matrix, items)
    assert result["unresolved_categories"] == ["Gone Vertical"]


# --- Restrictions: renamed / custom products ---------------------------------------------------------

def test_a_renamed_product_keeps_its_sheet_verdict(monkeypatch, builtin_state):
    monkeypatch.setattr("app.catalog.load_rate_overrides",
                        lambda: {"Video - Pre-roll (OLV)": {"name": "Pre-Roll Video (OLV)"}})
    matrix = rs.build_matrix(["Pre-Roll Video (OLV)", "Video - Pre-roll (OLV)"], ["Political"])
    assert matrix["verdicts"]["Pre-Roll Video (OLV)"]["verdict"] == "allowed"
    assert matrix["verdicts"]["Video - Pre-roll (OLV)"]["verdict"] == "allowed"


def test_a_custom_product_the_sheet_doesnt_list_reads_as_guidance_not_a_hard_refusal(monkeypatch, builtin_state):
    custom = Product(
        family="Display", name="Custom Display Package", short_label="x", proposal_description="", sizes="",
        buying_model="CPM", base_rate=5.0, estimated_impressions=False, discloses_impressions=True, minimum_spend=0.0,
        minimum_flight_days=(0, 0), sla_data_days=None, sla_creative_days=None, sla_activate_days=None,
        sla_total_days=None, media_allocation_pct=1.0, margin_upper=0.5, margin_lower=0.5, tech_platform="", wide_orbit_code="")
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [custom])
    matrix = rs.build_matrix(["Custom Display Package"], ["Political"])
    verdict = matrix["verdicts"]["Custom Display Package"]
    assert verdict["verdict"] == "guidance"
    assert "Sales Planning" in verdict["checks"][0]["note"]


def test_services_and_measurement_are_exempt_from_guidance_tabs_too():
    guidance = rs.Category(name="Casino Gambling", kind="guidance", summary="Platform guidance", body="TTD: ...")
    assert rs.evaluate_product("Anything", "Measurement", [guidance])["verdict"] == "n/a"


# --- Restrictions: the sheet parser --------------------------------------------------------------------

def _tab(name, rows):
    def build(wb):
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r if isinstance(r, list) else [r])
    return _workbook(build)


@pytest.mark.parametrize("header", [
    "Political Ads are only accepted for the below products",
    "These ads can only be run on the below products",
    "We only accept political advertising on the following products",
    "Political is limited to the following products",
    "Political ads are allowed on the below products",
])
def test_allowlist_headers_in_different_words_are_still_recognised(header):
    cats = rs.parse_workbook(_tab("Political", [header, "eDigital Display", "eDigital OLV"]))
    assert cats[0].kind == "allowlist" and cats[0].items == ["eDigital Display", "eDigital OLV"]


def test_a_blanket_hold_is_recognised_without_the_word_all():
    cats = rs.parse_workbook(_tab("Cannabis", ["Entravision Digital Products are on hold and will not accept cannabis ads"]))
    assert cats[0].kind == "blanket_hold"


def test_a_stray_all_other_products_line_does_not_turn_an_allowlist_into_a_blanket_hold():
    cats = rs.parse_workbook(_tab("Political", [
        "Political Ads are only accepted for the below products", "eDigital Display",
        "All other products are not accepting political ads"]))
    assert cats[0].kind == "allowlist"


def test_numbered_lists_and_a_single_spacer_row_keep_every_item():
    cats = rs.parse_workbook(_tab("Political", [
        "Political Ads are only accepted for the below products",
        ["1", "eDigital Display"], ["2", "eDigital OLV"], [None], ["3", "Digital Out of Home"],
        [None], [None], "Trailing note that is not a product"]))
    assert cats[0].items == ["eDigital Display", "eDigital OLV", "Digital Out of Home"]


def test_an_allowlist_with_no_entries_degrades_to_guidance_with_a_warning():
    cats, warnings = rs.parse_workbook_ex(_tab("Political", ["Political Ads are only accepted for the below products"]))
    assert cats[0].kind == "guidance"
    assert any("no entries" in w for w in warnings)


def test_a_grid_keeps_the_text_above_it_and_notes_under_it_and_finds_a_platform_column_b():
    cats = rs.parse_workbook(_tab("Casino Gambling", [
        ["Read this first"], [None, "Platforms", "Casinos"], [None, "TTD", "We accept brick+mortar."],
        ["Note: TTD and Amazon never allow online casino ads"]]))
    body = cats[0].body
    assert "Read this first" in body and "TTD: We accept brick+mortar." in body and "never allow online casino" in body


def test_sheet_names_that_collapse_to_the_same_name_are_rejected():
    def build(wb):
        wb.create_sheet("Casino_Gambling").append(["x"])
        wb.create_sheet("Casino Gambling").append(["y"])
    with pytest.raises(ValueError, match="Two tabs"):
        rs.parse_workbook(_workbook(build))


def test_hidden_tabs_are_skipped_with_a_warning():
    def build(wb):
        wb.create_sheet("Political").append(["Political Ads are only accepted for the below products"])
        hidden = wb.create_sheet("Old Notes")
        hidden.append(["whatever"])
        hidden.sheet_state = "hidden"
    cats, warnings = rs.parse_workbook_ex(_workbook(build))
    assert [c.name for c in cats] == ["Political"]
    assert any("hidden tab" in w for w in warnings)


def test_hyperlink_targets_are_kept_next_to_their_text():
    def build(wb):
        ws = wb.create_sheet("Casino Gambling")
        ws.append(["Platform", "Guidance", "Link"])
        ws.append(["Meta", "See policy", "Meta's Gambling Policy"])
        ws["C2"].hyperlink = "https://example.com/meta-gambling"
    cats = rs.parse_workbook(_workbook(build))
    assert "https://example.com/meta-gambling" in cats[0].body


def test_entries_that_matched_no_catalog_product_are_called_out():
    cats, warnings = rs.parse_workbook_ex(_tab("Political", [
        "Political Ads are only accepted for the below products", "eDigital Display", "Netflix"]))
    assert any("matched no catalog product" in w and "Netflix" in w for w in warnings)


def test_rows_past_the_cap_are_not_read_and_the_admin_is_told():
    def build(wb):
        ws = wb.create_sheet("Casino Gambling")
        ws.append(["Platform", "Guidance"])
        for i in range(rs.MAX_ROWS + 60):
            ws.append([f"P{i}", f"line {i}"])
    cats, warnings = rs.parse_workbook_ex(_workbook(build))
    assert len(cats[0].body.splitlines()) == rs.MAX_ROWS - 1
    assert any("has content past" in w for w in warnings)


# --- Restrictions: keywords, carry-over, load_state ---------------------------------------------------------

def test_generic_words_dont_suggest_a_vertical():
    cats = rs.builtin_categories() + [rs.Category(name="Blood Donation", kind="guidance",
                                                  keywords=rs.default_keywords("Blood Donation"))]
    for text in ("Pac-12 Network live sports", "holiday donation drive", "plasma TV", "find the best candidate for the job"):
        assert rs.suggest_categories(text, cats) == [], text
    assert "Political" in rs.suggest_categories("Friends of Maria for the ballot measure", cats)
    assert "Blood Donation" in rs.suggest_categories("Spring blood drive", cats)


def test_a_resync_keeps_admin_keyword_and_mapping_edits():
    old = rs.Category(name="Political", kind="allowlist", keywords=["my custom keyword"], items=["eDigital Display", "Gone"],
                      product_map={"eDigital Display": ["My Custom Display"], "Gone": ["X"]})
    new = rs.Category(name="political", kind="allowlist", keywords=["political"], items=["eDigital Display", "New Item"],
                      product_map={"eDigital Display": ["eDigital Network Display - Standard IAB"], "New Item": []})
    rs.carry_over_admin_edits([new], [old])
    assert new.keywords == ["my custom keyword"]
    assert new.product_map["eDigital Display"] == ["My Custom Display"]
    assert new.product_map["New Item"] == []


def test_load_state_flags_a_failed_read_apart_from_an_empty_table(monkeypatch):
    def boom():
        raise RuntimeError("db down")
    monkeypatch.setattr(rs, "get_connection", boom)
    state = rs.load_state()
    assert state["using_builtin"] is True and state["error"] is True


def test_editing_the_builtin_snapshot_is_refused_while_the_table_cant_be_read(monkeypatch):
    import asyncio
    import app.main as m
    from fastapi import HTTPException
    monkeypatch.setattr(rs, "load_state", lambda: {"categories": rs.builtin_categories(), "using_builtin": True,
                                                   "sync": None, "error": True})
    seeded = []
    monkeypatch.setattr(rs, "seed_builtin_into_db", lambda email: seeded.append(email))

    class _Req:
        class state:
            user = {"email": "a@x"}
    with pytest.raises(HTTPException) as exc:
        asyncio.run(m.admin_update_restriction("Political", m.RestrictionUpdateRequest(keywords="political", product_map={}), _Req()))
    assert exc.value.status_code == 503 and not seeded


def test_a_missing_table_becomes_a_503_with_a_plain_message():
    import asyncio
    import psycopg
    import app.main as m

    class _Req:
        class url:
            path = "/api/admin/disclaimers"
    resp = asyncio.run(m._schema_not_applied(_Req(), psycopg.errors.UndefinedTable("relation does not exist")))
    assert resp.status_code == 503 and b"schema.sql" in resp.body


# --- Disclaimers: scoping + sanitizing --------------------------------------------------------------------------

def _rule(**kw) -> dz.Disclaimer:
    base = dict(id="r1", name="Rule", keywords=["live sports"], banner_text="Heads up.", color="red")
    base.update(kw)
    return dz.Disclaimer(**base)


def test_the_pasted_product_list_never_triggers_a_banner_but_still_feeds_the_vertical_suggestion():
    req = ProposalRequest(client_name="Acme", products_selected=["NBC Sports Stream Sponsorship"],
                          products_selected_raw="NBC Sports Stream Sponsorship, live sports")
    assert "live sports" not in dz.request_text(req).lower()
    assert "live sports" in dz.request_text(req, include_products=True).lower()


def test_a_product_dropped_in_step_04_stops_triggering_the_banner_on_that_option():
    req = ProposalRequest(client_name="Acme", products_selected=["Fox Sports Go Video Sponsorship"],
                          products_selected_raw="Fox Sports Go Video Sponsorship")
    tiers = [
        {"label": "A", "name": "A", "line_items": _items("Search - SEM")},
        {"label": "B", "name": "B", "line_items": _items("Fox Sports Go Video Sponsorship")},
    ]
    out = dz.resolve_for_tiers(req, tiers, [_rule(keywords=["live sports", "fox sports"])])
    assert "A" not in out or not out["A"]
    assert out["B"]


def test_validate_fields_strips_control_characters_and_caps_the_name():
    clean = dz.validate_fields(name="Name\x07", keywords="a", banner_text="Line\x0b one\x0c", color="red",
                               applies_to="both", season_start=None, season_end=None)
    assert clean["name"] == "Name" and clean["banner_text"] == "Line one"
    with pytest.raises(ValueError, match="too long"):
        dz.validate_fields(name="x" * 121, keywords="a", banner_text="t", color="red", applies_to="both",
                           season_start=None, season_end=None)


def test_banner_text_starting_with_an_equals_sign_stays_text_in_the_workbook(tmp_path):
    req = ProposalRequest(client_name="Acme", request_type="New Business", start_date="2026-09-01",
                          end_date="2026-11-30", total_months=3)
    li = LineItem(product_name="Search - SEM", monthly_budget=1200.0, months=3, id="li1")
    banner = {"id": "1", "name": "n", "text": "=== NOTICE === read this\x07", "fill": "F8D7DA", "font_color": "58151C",
              "applies_to": "both", "color": "red"}
    out = tmp_path / "eq.xlsx"
    generate_proposal(req, [li], out, disclaimers_by_tier={"A": [banner]})
    import openpyxl
    ws = openpyxl.load_workbook(out)["Proposal A"]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "NOTICE" in c.value]
    assert cells and cells[0].data_type == "s" and cells[0].value == "=== NOTICE === read this"


# --- PPT: overflow + labels -----------------------------------------------------------------------------------

pptx = pytest.importorskip("pptx")


def _product(name):
    return Product(
        family="Search", name=name, short_label=name, proposal_description="", sizes="", buying_model="CPM",
        base_rate=10.0, estimated_impressions=False, discloses_impressions=True, minimum_spend=0.0,
        minimum_flight_days=(0, 0), sla_data_days=None, sla_creative_days=None, sla_activate_days=None,
        sla_total_days=None, media_allocation_pct=1.0, margin_upper=0.5, margin_lower=0.5, tech_platform="", wide_orbit_code="")


def _deck(tmp_path, *, banners=None, products=2, target=None, time_unit="month"):
    from app.services import pptx_builder
    prods = [_product(f"Product {i}") for i in range(products)]
    lis = [LineItem(product_name=p.name, monthly_budget=2010.0, months=3, id=f"l{i}",
                    target_override=target) for i, p in enumerate(prods)]
    out = tmp_path / "d.pptx"
    assert pptx_builder.build_signature_deck(
        ProposalRequest(client_name="Acme", start_date="2026-09-01", end_date="2026-11-30"),
        [{"name": "Option A", "products": prods, "line_items": lis}], out, gross=False, proposal_title="T",
        banners=banners, time_unit=time_unit)
    return pptx.Presentation(str(out))


def _banner(text="Banner text.", i=0):
    return {"id": str(i), "text": text, "fill": "F8D7DA", "font_color": "58151C"}


def test_a_long_target_no_longer_lets_the_banner_cover_the_table(tmp_path):
    long_target = "Hispanic adults 25-54 in Los Angeles, Orange and Riverside counties, " * 3
    prs = _deck(tmp_path, banners=[_banner()], target=long_target)
    slide = next(sl for sl in prs.slides if any(sh.has_table for sh in sl.shapes))
    table = next(sh for sh in slide.shapes if sh.has_table)
    banner = next(sh for sl in prs.slides for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text == "Banner text.")
    assert table.top + table.height <= banner.top + 1          # the banner starts below where the table is declared to end


def test_many_banners_flow_across_slides_instead_of_running_off_the_footer(tmp_path):
    from app.services import pptx_builder as pb
    banners = [_banner("Heads up. " * 60, i) for i in range(8)]
    prs = _deck(tmp_path, banners=banners)
    seen = 0
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith("Heads up."):
                seen += 1
                assert sh.top + sh.height <= pb.FOOTER_Y + 1
            if sh.has_text_frame and sh.text_frame.text in ("Client", "Signature", "Date"):
                assert sh.top + sh.height <= pb.FOOTER_Y + 1
    assert seen == 8
    legal = [sl for sl in prs.slides if any(sh.has_text_frame and sh.text_frame.text.startswith("This proposal will be valid") for sh in sl.shapes)]
    assert len(legal) == 1


def test_the_investment_line_is_a_flight_total_and_unit_labels_follow_the_time_unit(tmp_path):
    prs = _deck(tmp_path)
    text = "\n".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    assert "Total Investment: $12,060" in text and "per month" not in text
    weekly = _deck(tmp_path, time_unit="week")
    table = next(sh for sl in weekly.slides for sh in sl.shapes if sh.has_table).table
    assert [table.cell(0, c).text for c in (2, 3)] == ["Weeks", "Weekly Budget"]
    assert table.cell(1, 3).text.endswith("/wk")


# --- Excel edge cases ---------------------------------------------------------------------------------------------

def _req(**kw):
    base = dict(client_name="Acme", request_type="New Business", start_date="2026-09-01", end_date="2026-11-30", total_months=3)
    base.update(kw)
    return ProposalRequest(**base)


def test_forced_dooh_tabs_are_not_discarded_by_the_all_off_guard(tmp_path):
    li = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="a")
    out = tmp_path / "dooh.xlsx"
    summary = generate_proposal(_req(), [li], out, force_tabs={"net": False, "wsections": False, "gross": False,
                                                              "avails_only": False, "dooh_summary": True, "dooh_screenlist": True})
    assert not any("switched off" in w for w in summary.get("warnings", []))


def test_added_value_lines_leave_the_sov_cell_blank_not_a_green_zero(tmp_path):
    paid = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="p")
    gift = LineItem(product_name="eDigital Network Display - Standard IAB", monthly_budget=0.0, months=3, id="g", is_added_value=True)
    out = tmp_path / "av.xlsx"
    generate_proposal(_req(), [paid, gift], out, avails_data={"p": {"max_imps": 100000, "max_spend": 5000},
                                                              "g": {"max_imps": 200000, "max_spend": 9000}})
    import openpyxl
    ws = openpyxl.load_workbook(out)["Proposal A"]
    def row_of(name):
        return next(r for r in range(1, ws.max_row + 1) if str(ws.cell(row=r, column=3).value or "").startswith(name))
    assert str(ws.cell(row=row_of("Search - SEM"), column=17).value).startswith("=IFERROR(L")     # paid line: live SOV
    assert ws.cell(row=row_of("eDigital Network Display"), column=17).value is None             # Added Value: nothing, not 0.0%


def test_an_avails_only_line_with_just_uniques_keeps_its_live_formulas_and_the_uniques(tmp_path):
    li = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="u")
    out = tmp_path / "uniq.xlsx"
    generate_proposal(_req(), [li], out, force_tabs={"net": False, "avails_only": True},
                      avails_data={"u": {"est_uniques": 12345}})
    import openpyxl
    ws = next(w for w in openpyxl.load_workbook(out).worksheets if "Avails" in w.title)
    row = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(row=r, column=3).value or "").startswith("Search - SEM"))
    assert ws.cell(row=row, column=12).value == 12345                       # uniques kept
    assert str(ws.cell(row=row, column=11).value).startswith("=IFERROR(")   # K still a live formula for a typed J


# --- Blurbs ---------------------------------------------------------------------------------------------------------

def test_abbreviations_dont_split_a_blurb_sentence():
    # The sentence after "U.S." starts with a capital, so a naive splitter breaks there; once the second half also
    # trips the language rule it would leave the stub "Ads reach U.S." — the whole sentence must go or stay together.
    text = "Ads reach shoppers across the U.S. Hispanic market on the sites they visit most. Plain sentence."
    assert _tighten_blurb(text, "Display") == text
    dropped = "Ads reach shoppers across the U.S. Spanish-language radio stations carry them. Plain sentence."
    assert _tighten_blurb(dropped, "Display") == "Plain sentence."


def test_a_dropped_fact_takes_its_orphaned_citation_with_it():
    out = _tighten_blurb("Audio ads play inside radio streams. Listeners can skip some ads in skippable formats. (eMarketer, 2025)", "Audio")
    assert out == "Audio ads play inside radio streams."


def test_language_of_the_inventory_claims_in_more_phrasings_are_dropped_unless_the_name_says_it():
    text = "Display ads run on Spanish-language radio stations and podcasts. Banner ads on a network of sites."
    assert _tighten_blurb(text, "Display") == "Banner ads on a network of sites."
    assert _tighten_blurb("Runs on Spanish-language sites. Plain sentence.", "Spanish Display") == "Runs on Spanish-language sites. Plain sentence."


def test_the_prompt_no_longer_tells_the_model_to_reuse_stats_in_blurbs(monkeypatch):
    from app.services import ai_enricher
    monkeypatch.setattr("app.catalog.resolve_product_alias", lambda name: None)
    prompt = ai_enricher._build_prompt(_req(), _items("Search - SEM", "Not In The Catalog"), strategy_brief={
        "recommended_tactics": [{"product_family": "Search", "rationale": "r", "data_point": "d", "citation": "c"}],
        "strategy_summary": "s"})
    assert "product blurbs below MUST stay consistent" not in prompt
    assert "never on the Target Audience section" in prompt
    assert "no rate-card description on file" in prompt            # a product without a catalog description
    assert "EMAIL AND OUTPUT RULES" in prompt
