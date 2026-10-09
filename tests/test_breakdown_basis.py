"""
The breakdown block says WHICH dollars it shows, in every billing period.

  * The Gross sheet's breakdown is in GROSS dollars — each net amount marked up by 1 / (1 - agency fee), exactly like its
    GROSS BUDGET column — whatever the billing period (it used to be net dollars next to a gross budget everywhere but
    Full Flight); the Net sheets stay NET.
  * With no agency fee selected there is nothing to mark up, so it is NET on every sheet.
  * Each block's title names its basis: "MONTHLY BREAKDOWN (GROSS)" / "(NET)" ("WEEKLY ..." / "QUARTERLY ..." for those
    periods; Full Flight keeps the planner's own "MONTHLY BREAKDOWN" over its single column).
  * The standalone breakdown tab follows the same rule: gross with an agency fee ("Monthly Breakdown (Gross)"), net without.

Same DB-mocking pattern as the other export tests (no Postgres needed).
"""
from __future__ import annotations

import os

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import openpyxl
import pytest

from app.services import monthly_allocation as mo
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import LineItem, generate_proposal

EMAIL = "Email Campaigns - Display Re-targeting"
IAB = "eDigital Network Display - Standard IAB"
PLAN = {EMAIL: 680.0, IAB: 5470.0}          # net flight totals
FEE = 0.15
START, END = "2026-11-24", "2026-12-31"      # two calendar months, six weeks, one calendar quarter


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})
    monkeypatch.setattr("app.catalog.resolve_product_alias", lambda name: None)


def _req(**overrides) -> ProposalRequest:
    fields = dict(client_name="Acme Corp", request_type="New Business", start_date=START, end_date=END,
                  total_months=2, geo="Los Angeles", agency_fee=FEE)
    fields.update(overrides)
    return ProposalRequest(**fields)


def _periods(unit):
    return mo.periods_between(mo.parse_flexible_date(START), mo.parse_flexible_date(END), unit)


def _lines(unit):
    """The PLAN spread evenly over the unit's periods, with per-period allocations (what Step 06 stores)."""
    periods = _periods(unit)
    if unit == "full_flight":
        return [LineItem(product_name=n, monthly_budget=t, months=1, id=f"l{i}") for i, (n, t) in enumerate(PLAN.items())]
    return [LineItem(product_name=n, monthly_budget=t / len(periods), months=len(periods), id=f"l{i}",
                     monthly_allocations=mo.compute_default_allocation(t, periods, "even"))
            for i, (n, t) in enumerate(PLAN.items())]


def _export(tmp_path, unit, *, req=None, force_tabs=None):
    out = tmp_path / f"{unit}.xlsx"
    # the sectioned Net tab is only built by default for some mixes; force it so every sheet under test exists
    generate_proposal(req or _req(), _lines(unit), out, time_unit=unit, force_tabs={"wsections": True, **(force_tabs or {})}, addons=[])
    return openpyxl.load_workbook(out)


def _find(ws, text, col="C"):
    return next(r for r in range(1, ws.max_row + 1) if ws[f"{col}{r}"].value == text)


UNIT_WORD = {"week": "WEEKLY", "month": "MONTHLY", "quarter": "QUARTERLY", "full_flight": "MONTHLY"}


@pytest.mark.parametrize("unit", ["week", "month", "quarter", "full_flight"])
def test_every_billing_period_titles_its_breakdown_with_the_basis(tmp_path, unit):
    wb = _export(tmp_path, unit)
    word = UNIT_WORD[unit]
    assert wb["Proposal A"]["S16"].value == f"{word} BREAKDOWN (NET)"
    assert wb["Proposal A (wsections)"]["S16"].value == f"{word} BREAKDOWN (NET)"
    assert wb["Proposal A (Gross)"]["U16"].value == f"{word} BREAKDOWN (GROSS)"


@pytest.mark.parametrize("unit", ["week", "month", "quarter"])
def test_the_gross_sheet_breakdown_is_in_gross_dollars_and_the_net_sheet_stays_net(tmp_path, unit):
    wb = _export(tmp_path, unit)
    net_ws, gross_ws = wb["Proposal A"], wb["Proposal A (Gross)"]
    keys = [p["key"] for p in _periods(unit)]
    lines = _lines(unit)
    for row, li in zip((19, 20), lines):
        total = li.monthly_budget * li.months
        for i, key in enumerate(keys):
            net = li.monthly_allocations[key]
            pct = net / total * 100
            assert net_ws.cell(row=row, column=19 + i).value == f"${net:,.0f} ({pct:.0f}%)", (unit, row, key)
            assert gross_ws.cell(row=row, column=21 + i).value == f"${net / (1 - FEE):,.0f} ({pct:.0f}%)", (unit, row, key)
    # the plan-total row: per-period sums, gross on the Gross sheet (the same markup its GROSS BUDGET column applies)
    net_total_row = _find(net_ws, f"TOTAL DIGITAL {UNIT_WORD[unit]}")
    gross_total_row = _find(gross_ws, f"TOTAL DIGITAL {UNIT_WORD[unit]}")
    for i, key in enumerate(keys):
        net_sum = sum(li.monthly_allocations[key] for li in lines)
        assert net_ws.cell(row=net_total_row, column=19 + i).value == pytest.approx(net_sum)
        assert gross_ws.cell(row=gross_total_row, column=21 + i).value == pytest.approx(net_sum / (1 - FEE))
    # the first row's gross cell agrees with the sheet's own GROSS BUDGET for that line when it is one period
    if len(keys) == 1:
        gross_budget = lines[0].monthly_budget * lines[0].months / (1 - FEE)
        assert gross_ws["U19"].value.startswith(f"${gross_budget:,.0f} ")


@pytest.mark.parametrize("unit", ["week", "month", "quarter", "full_flight"])
def test_without_an_agency_fee_every_breakdown_is_net(tmp_path, unit):
    # a Gross tab can still be forced on; with no fee selected its breakdown is NET — and says so
    for fee in (None, 0.0):
        wb = _export(tmp_path, unit, req=_req(agency_fee=fee), force_tabs={"gross": True})
        word = UNIT_WORD[unit]
        assert wb["Proposal A"]["S16"].value == f"{word} BREAKDOWN (NET)"
        assert wb["Proposal A (Gross)"]["U16"].value == f"{word} BREAKDOWN (NET)", (unit, fee)
        if unit == "full_flight":
            # live formulas over the NET BUDGET column, on the Gross sheet too
            assert wb["Proposal A (Gross)"]["U19"].value == '=IFERROR("$"&TEXT(L19,"#,##0")&IF(L19>0," (100%)"," (0%)"),"—")'
        else:
            key = _periods(unit)[0]["key"]
            li = _lines(unit)[0]
            net = li.monthly_allocations[key]
            pct = net / (li.monthly_budget * li.months) * 100
            assert wb["Proposal A (Gross)"]["U19"].value == f"${net:,.0f} ({pct:.0f}%)"


@pytest.mark.parametrize("unit", ["week", "month", "quarter"])
def test_the_standalone_breakdown_tab_is_gross_with_an_agency_fee(tmp_path, unit):
    wb = _export(tmp_path, unit)
    sheet = f"{UNIT_WORD[unit].title()} Breakdown"
    ws = wb[sheet]
    assert ws["B2"].value == f"{sheet} (Gross)"
    keys = [p["key"] for p in _periods(unit)]
    for row, li in zip((5, 6), _lines(unit)):
        total = li.monthly_budget * li.months
        for i, key in enumerate(keys):
            net = li.monthly_allocations[key]
            assert ws.cell(row=row, column=3 + i).value == f"${net / (1 - FEE):,.0f} ({net / total * 100:.0f}%)", (unit, row, key)
    total_row = _find(ws, f"TOTAL PLAN SPEND BY {unit.upper()}", col="B")
    for i, key in enumerate(keys):
        expected = sum(li.monthly_allocations[key] for li in _lines(unit)) / (1 - FEE)
        assert ws.cell(row=total_row, column=3 + i).value == pytest.approx(expected)


@pytest.mark.parametrize("unit", ["week", "month", "quarter"])
def test_the_standalone_breakdown_tab_is_net_without_an_agency_fee(tmp_path, unit):
    wb = _export(tmp_path, unit, req=_req(agency_fee=None))
    sheet = f"{UNIT_WORD[unit].title()} Breakdown"
    ws = wb[sheet]
    assert ws["B2"].value == f"{sheet} (Net)"
    li = _lines(unit)[0]
    key = _periods(unit)[0]["key"]
    net = li.monthly_allocations[key]
    assert ws.cell(row=5, column=3).value == f"${net:,.0f} ({net / (li.monthly_budget * li.months) * 100:.0f}%)"


def test_a_single_column_block_is_widened_for_its_banner(tmp_path):
    # a quarterly plan on a one-quarter flight is a ONE-column block (not merged): its banner now carries the basis
    wb = _export(tmp_path, "quarter")
    assert len(_periods("quarter")) == 1
    for sheet, col, banner in (("Proposal A", "S", "QUARTERLY BREAKDOWN (NET)"), ("Proposal A (Gross)", "U", "QUARTERLY BREAKDOWN (GROSS)")):
        ws = wb[sheet]
        assert ws[f"{col}16"].value == banner
        assert ws.column_dimensions[col].width >= len(banner) + 2
    # a multi-column block is merged across its columns, as before
    wb = _export(tmp_path, "month")
    assert any(str(r) == "S16:T16" for r in wb["Proposal A"].merged_cells.ranges)
    assert any(str(r) == "U16:V16" for r in wb["Proposal A (Gross)"].merged_cells.ranges)
