"""
Round-3 regression tests: what the final adversarial pass found in the round-2 code (sentence splitting at bullets and
camelCase brands, the approved-list terminator, hyperlinked entries, header clauses, schema-missing vs flaky reads,
fuzzy product-name matching, request-value stop-words, and the merged-periods grand total).
"""
from __future__ import annotations

import io
import json
import os

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import pytest

from app.services import restrictions as rs
from app.services import roadblocks as rb
from app.services.ai_enricher import _compile_avoid, _parse_response, _targeting_phrases, _tighten_blurb
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import LineItem, generate_proposal
from app.services.text_utils import split_sentences

POLITICAL = [c for c in rs.builtin_categories() if c.name == "Political"]


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.catalog.resolve_product_alias", lambda name: None)
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


@pytest.fixture
def builtin_state(monkeypatch):
    monkeypatch.setattr(rs, "load_state", lambda: {"categories": rs.builtin_categories(), "using_builtin": True,
                                                   "sync": None, "error": False})


def _items(*names):
    return [LineItem(product_name=n, monthly_budget=1000, months=3, id=f"i{i}") for i, n in enumerate(names)]


def _book(rows_by_tab, mutate=None) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in rows_by_tab.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r if isinstance(r, list) else [r])
    if mutate:
        mutate(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse(name, rows):
    cats, warnings = rs.parse_workbook_ex(_book({name: rows}))
    return cats[0], warnings


# --- sentence splitting ---------------------------------------------------------------------------------------------

@pytest.mark.parametrize("bullet", ["-", "*", "•"])
def test_a_denial_bullet_removes_only_that_bullet(bullet):
    text = f"{bullet} Submit creative 5 days early.\n{bullet} Political ads are not allowed on Display.\n{bullet} Use the Meta Ad Library."
    cleaned, gone = rb._strip_denials(text, POLITICAL)
    assert len(gone) == 1 and "not allowed" in gone[0]
    assert "Submit creative" in cleaned and "Meta Ad Library" in cleaned and "not allowed" not in cleaned


def test_a_camelcase_brand_starts_a_new_sentence_and_survives_a_neighbouring_denial():
    cleaned, gone = rb._strip_denials("Political ads are not allowed on this product. eDigital OLV is the approved option. Plan early.", POLITICAL)
    assert cleaned == "eDigital OLV is the approved option. Plan early." and len(gone) == 1
    out = _tighten_blurb("Ads run in Spanish-language sites and apps. eDigital placements reach shoppers. Text stays short.", "Display")
    assert out == "eDigital placements reach shoppers. Text stays short."


def test_section_and_month_abbreviations_dont_split_and_a_paragraph_break_survives_a_removed_denial():
    assert split_sentences("Approvals close before Oct. 21. Submit early. See Sec. 4. Plan ahead.") == [
        "Approvals close before Oct. 21. ", "Submit early. ", "See Sec. 4. ", "Plan ahead."]
    cleaned, _ = rb._strip_denials("Approvals take a week.\nPolitical ads are not allowed on Display.\n\nSubmit early.", POLITICAL)
    assert cleaned == "Approvals take a week.\n\nSubmit early."


def test_split_sentences_stays_fast_on_long_abbreviation_runs():
    import time
    start = time.time()
    split_sentences("U.S. " * 6000)
    split_sentences("Acme Inc. " * 8000)
    assert time.time() - start < 2


# --- the sheet parser -----------------------------------------------------------------------------------------------

def test_a_long_approved_list_survives_trailing_periods_and_see_below_remarks():
    entries = [e + "." if i == 4 else e for i, e in enumerate(rs._POLITICAL_ITEMS)]
    entries[2] = "eDigital OLV Geo Fence (see restrictions below)"
    cat, _ = _parse("Political", ["Political Ads are only accepted for the below products", *entries])
    assert len(cat.items) == 11


def test_a_hyperlinked_entry_keeps_its_name_clean_and_still_matches():
    def link_them(wb):
        ws = wb["Political"]
        for i in range(len(rs._POLITICAL_ITEMS)):
            ws.cell(row=i + 2, column=1).hyperlink = "https://support.google.com/" + "x" * 140
    data = _book({"Political": ["Political Ads are only accepted for the below products", *rs._POLITICAL_ITEMS]}, link_them)
    cats, warnings = rs.parse_workbook_ex(data)
    assert cats[0].items == rs._POLITICAL_ITEMS and not any("too long" in w for w in warnings)
    assert "https://support.google.com/" in cats[0].body                      # still in what the AI reads
    assert all(cats[0].product_map[item] for item in cats[0].items if "Geo Fence" not in item or "Display" in item)


def test_a_header_with_a_second_negated_sentence_is_still_an_approved_list():
    for header in ("Political Ads are only accepted for the below products. All other products are not accepting political ads.",
                   "Political Ads are only accepted for the below products (all others are NOT accepted)"):
        cat, _ = _parse("Political", [header, "eDigital Display", "eDigital OLV"])
        assert cat.kind == "allowlist", header
    for header in ("Political ads are NOT accepted for the below products", "Previously approved for the following (now on hold)"):
        cat, _ = _parse("Political", [header, "eDigital Display"])
        assert cat.kind == "guidance", header


def test_an_allowlist_whose_entries_match_no_product_degrades_to_guidance_with_the_grid_kept():
    cat, warnings = _parse("Casino Gambling", ["Gambling ads are only accepted on the following platforms:", "See links in column C",
                                                [None], ["Platform", "Guidance", "Link"], ["TTD", "Brick and mortar only", "http://x"]])
    assert cat.kind == "guidance" and "TTD: Brick and mortar only" in cat.body


def test_a_list_that_stops_early_says_where():
    _, warnings = _parse("Political", ["Political Ads are only accepted for the below products", "eDigital Display",
                                      "Note: this list was last reviewed by Legal in 2024 and is subject to change."])
    assert any("anything after it was not treated" in w for w in warnings)


def test_hold_headlines_in_other_wordings_and_below_a_title_row_are_recognised():
    for rows in ([["Cannabis ads are on hold across all Entravision digital products."]],
                 [["Cannabis", "Updated 9/1", "Owner: Legal"], ["** All Entravision Digital Products are currently on hold **"]]):
        cat, _ = _parse("Cannabis", rows)
        assert cat.kind == "blanket_hold", rows


def test_a_bare_column_title_on_the_header_row_is_not_an_entry():
    cat, warnings = _parse("Political", [["Political Ads are only accepted for the below products", "Notes"], "eDigital Display"])
    assert cat.items == ["eDigital Display"]


def test_geofencing_spellings_all_map_to_the_geo_fence_product():
    for entry in ("eDigital Display Geo-Fencing", "Geo Fencing Display", "Geofencing Display", "Display Geo Fence"):
        assert rs.match_catalog_products(entry) == ["Display - Geo Fence"], entry


def test_formatting_only_cells_past_the_cap_do_not_raise_the_truncation_warning():
    from openpyxl.styles import Alignment

    def style_down(wb):
        ws = wb["Political"]
        for r in range(4, 1000):
            ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True)
    data = _book({"Political": ["Political Ads are only accepted for the below products", "eDigital Display"]}, style_down)
    _, warnings = rs.parse_workbook_ex(data)
    assert not any("past" in w for w in warnings)


def test_ordinary_words_no_longer_suggest_a_vertical():
    cats = rs.builtin_categories() + [rs.Category(name="Sexual Enhancement Procedures", kind="guidance",
                                                  keywords=rs.default_keywords("Sexual Enhancement Procedures"))]
    assert rs.suggest_categories("Valley Image Enhancement Studio - photo editing", cats) == []
    assert rs.suggest_categories("Please vote no later than Friday", cats) == []
    assert "Political" in rs.suggest_categories("Vote no on Measure B", cats)


def test_a_missing_table_is_told_apart_from_a_flaky_read(monkeypatch):
    import psycopg
    import app.main as m

    def boom():
        raise psycopg.errors.UndefinedTable("relation does not exist")
    monkeypatch.setattr(rs, "get_connection", boom)
    state = rs.load_state()
    assert state["error"] is True and state["missing_table"] is True
    assert "schema.sql" in m._restrictions_read_failed(state).detail
    assert "try again" in m._restrictions_read_failed({"error": True}).detail


# --- Roadblocks -----------------------------------------------------------------------------------------------------

def test_the_model_shortened_name_prefers_the_closest_product_and_never_merges_a_product_not_in_the_plan(builtin_state):
    std, geo = "eDigital Network Display - Standard IAB", "Display - Geo Fence"
    items = _items(std, geo)
    matrix = rs.build_matrix([std, geo], ["Political"])
    denial = {"issue": "Political ads not allowed", "detail": "Political ads are not allowed on Display.", "source": "s"}
    result = {"overall_summary": "", "used_web_search": True, "error": None, "categories": [], "product_roadblocks": [
        {"product_name": "eDigital Display", "risk_level": "high", "risks": [denial], "recommended_mitigation": "Do not run political ads on Display."}]}
    rb._apply_matrix(result, matrix, items)
    cards = {c["product_name"]: c for c in result["product_roadblocks"]}
    assert cards[std]["risks"] == [] and cards[std]["recommended_mitigation"] == ""      # merged into Standard, then filtered

    only_std = _items(std)
    unrelated = {"overall_summary": "", "used_web_search": True, "error": None, "categories": [], "product_roadblocks": [
        {"product_name": "Display - Geo Fence", "risk_level": "high",
         "risks": [{"issue": "Carrier limits", "detail": "Geo fences need a 50m radius.", "source": "s"}], "recommended_mitigation": ""}]}
    rb._apply_matrix(unrelated, rs.build_matrix([std], ["Political"]), only_std)
    names = [c["product_name"] for c in unrelated["product_roadblocks"]]
    assert "Display - Geo Fence" in names                                    # kept as its own card, not folded into Standard Display
    assert next(c for c in unrelated["product_roadblocks"] if c["product_name"] == std)["risk_level"] == "low"


def test_an_unmatched_model_entry_cannot_deny_a_vertical_allowed_on_every_plan_product(builtin_state):
    items = _items("Video - Pre-roll (OLV)", "eDigital Network Display - Standard IAB")
    matrix = rs.build_matrix([li.product_name for li in items], ["Political"])
    result = {"overall_summary": "", "used_web_search": True, "error": None, "categories": [], "product_roadblocks": [
        {"product_name": "Programmatic video", "risk_level": "high", "recommended_mitigation": "Do not run political ads here.",
         "risks": [{"issue": "Political ads not allowed", "detail": "Political ads are not allowed.", "source": "s"}]}]}
    rb._apply_matrix(result, matrix, items)
    orphan = next(c for c in result["product_roadblocks"] if c["product_name"] == "Programmatic video")
    assert orphan["risks"] == [] and orphan["recommended_mitigation"] == ""
    assert result["removed_claims"]


def test_a_malformed_risk_level_cannot_break_the_merge(builtin_state):
    items = _items("YouTube Ads")
    matrix = rs.build_matrix(["YouTube Ads"], ["Political"])
    raw = json.dumps({"product_roadblocks": [
        {"product_name": "YouTube Ads", "risk_level": ["high"], "risks": []},
        {"product_name": "youtube ads", "risk_level": "High", "risks": []}]})
    result = rb._parse(raw, True, matrix=matrix, line_items=items)
    assert len(result["product_roadblocks"]) == 1 and result["product_roadblocks"][0]["risk_level"] == "high"


@pytest.mark.parametrize("sentence", [
    "Google rejects political ads from unverified advertisers.",
    "Meta rejects political ads from advertisers who have not completed the ID check.",
    "Political ads are prohibited from targeting users under 18.",
])
def test_more_real_caveats_are_kept(sentence):
    assert not rb._denies_category(sentence, POLITICAL)


@pytest.mark.parametrize("sentence", [
    "Political ads are not currently accepted on the Display network.",
    "Political ads are no longer accepted on Display.",
    "Entravision has paused political ads on this product.",
])
def test_more_leaks_are_caught(sentence):
    assert rb._denies_category(sentence, POLITICAL)


# --- blurbs + exports -----------------------------------------------------------------------------------------------

def test_single_generic_words_in_geo_or_demo_do_not_drop_ordinary_sentences():
    assert _targeting_phrases(ProposalRequest(client_name="Acme", geo="Local", demo="Hispanic")) == []
    phrases = _targeting_phrases(ProposalRequest(client_name="Acme", geo="Phoenix, Orange County"))
    assert "phoenix" in phrases and "orange county" in phrases and "orange" not in phrases
    out = _tighten_blurb("Reach shoppers in Phoenix. Ads follow them across sites.", "Display", _compile_avoid(phrases))
    assert out == "Ads follow them across sites."


def test_a_huge_geo_list_does_not_slow_blurb_parsing_down():
    import time
    req = ProposalRequest(client_name="Acme", geo=", ".join(f"9{i:04d}" for i in range(1500)))
    raw = json.dumps({"campaign_name": "C", "internal_email_body": "b", "client_email_body": "b", "product_blurbs": [
        {"product_name": f"P{i}", "blurb": "Plain sentence one. Plain sentence two. Plain sentence three."} for i in range(30)]})
    start = time.time()
    _parse_response(raw, req, [])
    assert time.time() - start < 2


def test_the_grand_total_fallback_counts_base_periods_not_merged_buckets(tmp_path):
    lis = [LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="a"),
           LineItem(product_name="eDigital Network Display - Standard IAB", monthly_budget=1000.0, months=2, id="b")]
    tiers = [{"label": "A", "line_items": lis, "avails_data": {}, "start_date": "2026-09-01", "end_date": "2026-12-31",
              "period_merge_groups": [["2026-09", "2026-10"]]}]
    out = tmp_path / "m.xlsx"
    generate_proposal(ProposalRequest(client_name="Acme", request_type="New Business", start_date="2026-09-01",
                                      end_date="2026-12-31", total_months=4), [], out, tiers=tiers)
    import openpyxl
    assert openpyxl.load_workbook(out)["Proposal A"]["I10"].value == 4         # four calendar months bought; a merge only regroups
