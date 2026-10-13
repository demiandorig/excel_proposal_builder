"""
AI research quality in the Excel export (specific, sourced, no unfounded audience claims) and the "never cite
Wikipedia" rule (use the primary source it references instead).

Everything is offline: no OpenAI call is possible here, so the prompts are checked by rendering their text, the code
layers (claim_rules.py, source_rules.py and their wiring into ai_enricher / strategy_brief / roadblocks) by feeding
them the kind of text a model returns, and the model call itself by stubbing llm_utils.
"""
from __future__ import annotations

import json
import os

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import pytest

from app.catalog import CATALOG
from app.services import ai_enricher, claim_rules, llm_utils, roadblocks as rb, source_rules
from app.services import strategy_brief as sb
from app.services.ai_enricher import _blurb_sentence_ok, _compile_avoid, _parse_response, _tighten_blurb
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import LineItem, generate_proposal
from app.services.text_utils import split_sentences, strip_variant_pricing


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


# The user's own example of what must never reach the Details cell.
USER_BAD_EXAMPLE = "Hispanics are always researching for news on YouTube."


# ============================================================================================================
# 1. The claim validator (claim_rules.is_unsupported_claim)
# ============================================================================================================

# Sweeping audience / behavior / market / performance statements with no (number AND named source): all dropped.
UNSUPPORTED = [
    USER_BAD_EXAMPLE,
    "Hispanics are always researching for news on YouTube (YouTube).",              # a "source" with no year/publisher
    "Many Latino families increasingly rely on streaming audio.",
    "Most listeners trust podcast host recommendations.",
    "Podcast listeners trust host recommendations over display ads.",
    "Latinos often turn to YouTube for how-to videos.",
    "Spanish-speaking consumers prefer video content in their own language.",
    "Hispanic consumers over-index on mobile video.",
    "YouTube is where Hispanics go to research news.",
    "Audio is increasingly popular with Hispanic listeners.",
    "Streaming audio is the fastest-growing channel.",
    "CTV is the most popular way to watch TV.",
    "Podcasts are a trusted, highly engaging channel.",
    "Reach an engaged, loyal audience.",
    "Audiences are always online.",
    "Latino households spend more time streaming than the general market.",
    "Millions of Hispanics stream audio every day.",
    "This format outperforms standard display.",
    "Video ads drive higher completion rates than display.",
    "Podcast listeners are more likely to buy after hearing an ad.",
    "Hispanic audiences increasingly turn to streaming.",
    "Nearly all Latinos watch YouTube.",
    "YouTube is strong among Hispanics.",
    "Spanish-language content resonates with Hispanic audiences.",
    "Users tend to trust sponsored podcast reads.",
    "Parents search for answers on YouTube before they ever visit a store.",
    "Latino families love YouTube.",
    "Hispanics watch YouTube on their phones.",
    "Shoppers research online before they buy, and search captures them.",
    "Hispanic viewers are heavy YouTube users (Nielsen, 2025).",                    # sourced, but no number
    "Hispanic viewers are heavy YouTube users (Industry Report, 2025).",            # a vague "source" is no source
    "62% of Hispanics watch YouTube every week.",                                   # a number, but no named source
    "Over 160MM people listen to digital audio.",
    "62% of Hispanics watch YouTube every week (Source, Year).",                    # a template placeholder
    "62% of Hispanics watch YouTube every week (Wikipedia, 2024).",                 # Wikipedia is not a source
]

# What a product is / does, and sourced statistics: all kept.
SUPPORTED = [
    "Text ads that appear on Google when people search for what you offer, so your business shows up at the moment they're ready to call or visit.",
    "Banner ads shown across a network of websites and apps, retargeting local shoppers as they browse so your offer stays in front of them after they leave your site.",
    "Video ads delivered to connected TVs and streaming apps.",
    "Skippable in-stream ads that viewers can skip after five seconds.",
    "Viewers watch the full 15-second spot before their video starts.",
    "Audio ads across internet radio, podcasts and streaming.",
    "Your ad plays between songs and episodes while listeners are tuned in.",
    "Reach shoppers within a virtual perimeter around the places that matter to you.",
    "Customers who search for your service see your ad first.",
    "People searching for a plumber see your ad at the top of the results.",
    "Users click your ad and land on a page built for that offer.",
    "Pays only when people click.",
    "Shoppers ready to buy see your offer as they browse.",
    "Listeners tune in to radio stations online through apps and websites, and your audio ad plays during the stream.",
    "Hispanic households watching Amazon Prime Video see your ad on the platform's ad-supported tier.",
    "Short vertical videos appear in people's TikTok feeds, so shoppers discover your offer while scrolling.",
    "Parents scrolling Facebook in the evening see your ad.",
    "Email lands in the inbox of people who already opted in to hear from you.",
    "Households see your ad on the biggest screen in the home.",
    "300x250 and 728x90 banners served to the people you choose.",
    "Ads run on 15- and 30-second spots.",
    "62% of U.S. Hispanic adults watch YouTube weekly (Nielsen, 2025).",
    "YouTube reaches 2.5 billion logged-in users each month (Google, 2025).",
    "Podcast listeners are 1.5x more likely to act on an ad (Edison Research, 2025).",
    "Hispanic adults spend 3.5 hours a day on streaming video (Nielsen, 2025).",
    "Audio reaches 160 million listeners a month ([edisonresearch.com](https://www.edisonresearch.com/x)).",
]


@pytest.mark.parametrize("sentence", UNSUPPORTED)
def test_unfounded_or_unsourced_claims_are_flagged(sentence):
    assert claim_rules.is_unsupported_claim(sentence), sentence


@pytest.mark.parametrize("sentence", SUPPORTED)
def test_product_descriptions_and_sourced_statistics_pass(sentence):
    assert not claim_rules.is_unsupported_claim(sentence), sentence


def test_a_statistic_needs_both_the_number_and_a_named_source():
    claim = "Hispanic viewers are heavy YouTube users"
    assert claim_rules.is_unsupported_claim(claim + ".")                                   # neither
    assert claim_rules.is_unsupported_claim(claim + " (Nielsen, 2025).")                   # source, no number
    assert claim_rules.is_unsupported_claim("62% of " + claim.lower() + ".")               # number, no source
    assert not claim_rules.is_unsupported_claim("62% of Hispanic viewers are heavy YouTube users (Nielsen, 2025).")


def test_vague_and_template_sources_do_not_count_as_a_source():
    for source in ("Industry Report, 2025", "Market Research, 2024", "Company Data, 2025", "Source, Year", "Studies, 2025"):
        assert not source_rules.has_named_citation(f"62% of adults stream video ({source}).")
    for source in ("Nielsen, 2025", "Edison Research, 2025", "Pew Research Center, 2024", "Think with Google, 2025", "IAB, 2025",
                   "nielsen.com"):
        assert source_rules.has_named_citation(f"62% of adults stream video ({source}).")
    assert source_rules.has_named_citation("According to Edison Research, 62% of adults stream video.")
    assert not source_rules.has_named_citation("62% of adults stream video (25-54).")


def test_an_empty_sentence_is_not_a_claim():
    assert not claim_rules.is_unsupported_claim("")
    assert not claim_rules.is_unsupported_claim("   ")


def test_no_catalog_description_sentence_trips_the_validator():
    # The catalog text prints in the same cell as the blurb; if the validator flagged ordinary rate-card prose it
    # would be hostile to every product description it is meant to leave alone.
    flagged = []
    for product in CATALOG:
        for seg in split_sentences(strip_variant_pricing(product.proposal_description or "")):
            if seg.strip() and claim_rules.is_unsupported_claim(seg):
                flagged.append((product.name, seg.strip()))
    assert flagged == []


# ============================================================================================================
# 2. _tighten_blurb: the validator in the blurb pipeline
# ============================================================================================================

DESCRIPTION = "Video ads that play before and during YouTube videos, so you pay for people who actually watch."


def test_the_users_example_is_dropped_and_the_plain_description_stays():
    assert _tighten_blurb(f"{DESCRIPTION} {USER_BAD_EXAMPLE}", "YouTube Ads - Skippable") == DESCRIPTION
    assert _tighten_blurb(f"{USER_BAD_EXAMPLE} {DESCRIPTION}", "YouTube Ads - Skippable") == DESCRIPTION


def test_a_stat_sentence_with_its_citation_survives_untouched():
    text = f"{DESCRIPTION} 62% of U.S. Hispanic adults watch YouTube weekly (Nielsen, 2025)."
    assert _tighten_blurb(text, "YouTube Ads - Skippable") == text


def test_an_uncited_or_vaguely_cited_stat_is_dropped():
    assert _tighten_blurb(f"{DESCRIPTION} 62% of Hispanic adults watch YouTube weekly.", "Any") == DESCRIPTION
    assert _tighten_blurb(f"{DESCRIPTION} 62% of Hispanic adults watch YouTube weekly (Industry Report, 2025).", "Any") == DESCRIPTION


def test_a_blurb_made_only_of_unfounded_claims_comes_back_empty_so_the_cell_keeps_the_catalog_text():
    assert _tighten_blurb(USER_BAD_EXAMPLE, "Any") == ""
    assert _tighten_blurb(f"{USER_BAD_EXAMPLE} Many Latinos increasingly rely on streaming.", "Any") == ""


def test_an_empty_blurb_never_blanks_the_export_cell(tmp_path):
    from app.services.ai_enricher import ProposalEnrichment
    import openpyxl
    raw = json.dumps({
        "campaign_name": "Fall Push",
        "product_blurbs": [{"product_name": "AudioEngage", "blurb": f"{USER_BAD_EXAMPLE} Most listeners trust hosts."}],
        "internal_email_body": "b", "client_email_body": "b",
    })
    req = ProposalRequest(client_name="Acme", request_type="New Business", start_date="2026-10-01",
                          end_date="2026-12-31", total_months=3)
    enrichment: ProposalEnrichment = _parse_response(raw, req, [])
    assert enrichment.product_blurbs[0].blurb == ""
    out = tmp_path / "x.xlsx"
    generate_proposal(req, [LineItem(product_name="AudioEngage", monthly_budget=2000.0, months=3, id="a")], out,
                      enrichment=enrichment)
    ws = openpyxl.load_workbook(out)["Proposal A"]
    cells = [str(ws.cell(row=r, column=5).value) for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=5).value]
    audio = next(t for t in cells if "Internet radio, podcasts, and streaming channels" in t)
    assert "Hispanics" not in audio and "trust" not in audio and audio.strip()


LONG_CITATION = ("(Edison Research, The Infinite Dial 2025 — U.S. Hispanic audio report with a very long title that runs "
                 "past a hundred characters easily, 2025)")


def test_a_dropped_claim_takes_its_citation_with_it_even_when_the_citation_is_long():
    text = f"Audio ads play inside radio streams. Hispanics always listen to podcasts. {LONG_CITATION}"
    assert _tighten_blurb(text, "Any") == "Audio ads play inside radio streams."


def test_a_kept_stat_keeps_its_long_citation_as_one_unit():
    text = f"Audio ads play inside radio streams. 62% of Hispanic adults listen to podcasts monthly {LONG_CITATION}."
    assert _tighten_blurb(text, "Any") == text
    # ...and the pair is one unit for the word cap too
    big = "Audio ads play inside radio streams. " + ("Word " * 60).strip() + " 62% " + LONG_CITATION + "."
    assert _tighten_blurb(big, "Any").endswith("streams.")


@pytest.mark.parametrize("citation", ["(Nielsen, 2025).", "[Nielsen, 2025].", "(Edison Research, 2025 Infinite Dial)."])
def test_a_trailing_citation_of_any_shape_follows_its_sentence(citation):
    kept = f"Audio ads play inside radio streams. 62% of adults listen to podcasts monthly {citation}"
    assert _tighten_blurb(kept, "Any") == kept
    dropped = f"Audio ads play inside radio streams. Listeners always trust podcast hosts {citation}"
    assert _tighten_blurb(dropped, "Any") == "Audio ads play inside radio streams."


def test_template_placeholder_citations_are_removed_but_ordinary_parentheticals_stay():
    assert _tighten_blurb("Audio ads play inside radio streams (Source, Year).", "Any") == "Audio ads play inside radio streams."
    assert _tighten_blurb("Audio ads play inside radio streams (Industry Report, 2025).", "Any") == "Audio ads play inside radio streams."
    assert _tighten_blurb("Audio ads play inside radio streams (since 2020).", "Any") == "Audio ads play inside radio streams (since 2020)."


def test_markdown_links_from_web_search_become_plain_text_citations():
    text = ("Audio ads play inside radio streams. 62% of adults listen to podcasts monthly "
            "([edisonresearch.com](https://www.edisonresearch.com/the-infinite-dial-2025/?utm_source=openai)).")
    assert _tighten_blurb(text, "Any") == (
        "Audio ads play inside radio streams. 62% of adults listen to podcasts monthly (edisonresearch.com).")


def test_a_sourced_stat_may_name_the_audience_even_when_the_request_does_too():
    avoid = _compile_avoid(["hispanic adults 25-54"])
    stat = "62% of Hispanic adults 25-54 watch YouTube weekly (Nielsen, 2025)."
    assert _tighten_blurb(f"{DESCRIPTION} {stat}", "Any", avoid) == f"{DESCRIPTION} {stat}"
    # an unsourced restatement of the targeting is still dropped, as before
    assert _tighten_blurb(f"{DESCRIPTION} Built for Hispanic adults 25-54.", "Any", avoid) == DESCRIPTION


def test_the_existing_blurb_rules_still_apply_to_a_sourced_stat_sentence():
    assert not _blurb_sentence_ok("Skippable ads reach 62% of adults (Nielsen, 2025).", "Search - SEM")
    assert not _blurb_sentence_ok("62% of viewers like the Hispanic version of this product (Nielsen, 2025).", "Any")


# ============================================================================================================
# 3. Wikipedia: detection and stripping, across the citation formats the code handles
# ============================================================================================================

@pytest.mark.parametrize("url", [
    "https://en.wikipedia.org/wiki/Hispanic_and_Latino_Americans",
    "https://es.wikipedia.org/wiki/Latinoamericano",
    "https://en.m.wikipedia.org/wiki/Podcast?utm_source=openai",
    "en.wikipedia.org/wiki/Podcast",
    "https://commons.wikimedia.org/wiki/File:X.png",
    "https://www.wikidata.org/wiki/Q1",
    "https://www.wikiwand.com/en/articles/Podcast",
    "https://dbpedia.org/page/Podcast",
    "https://web.archive.org/web/2024/https://en.wikipedia.org/wiki/Podcast",
])
def test_wikipedia_urls_and_mirrors_are_recognised(url):
    assert source_rules.is_wikipedia_url(url)


@pytest.mark.parametrize("url", [
    "https://www.nielsen.com/insights/2025/hispanic-audio/",
    "https://www.nielsen.com/wikipedia-traffic-study",
    "https://notwikipedia.org/x",
    "https://www.pewresearch.org/race-and-ethnicity/",
    "",
])
def test_other_urls_are_not_wikipedia(url):
    assert not source_rules.is_wikipedia_url(url)


STAT_AFTER_DESCRIPTION = "Podcast ads reach adults while they listen."


@pytest.mark.parametrize("citation", [
    "(Wikipedia, 2024)",
    "(Wikipedia)",
    "(Source: Wikipedia)",
    "(en.wikipedia.org)",
    "(Wikipedia: Podcast, 2024)",
    "([en.wikipedia.org](https://en.wikipedia.org/wiki/Podcast?utm_source=openai))",
    "([Wikipedia](https://en.wikipedia.org/wiki/Hispanic_(ethnicity)))",
    "[Wikipedia, 2024]",
    "https://en.wikipedia.org/wiki/Podcast",
])
def test_a_wikipedia_citation_is_removed_and_a_stat_that_rested_on_it_goes_with_it(citation):
    # Descriptive text with a Wikipedia citation keeps its words and loses the citation...
    assert source_rules.strip_wikipedia(f"{STAT_AFTER_DESCRIPTION} {citation}") == STAT_AFTER_DESCRIPTION
    assert source_rules.strip_wikipedia(f"Podcast ads reach adults while they listen {citation}.") == STAT_AFTER_DESCRIPTION
    # ...but a sentence that carried a statistic on it is dropped, not left as an unsourced number.
    assert source_rules.strip_wikipedia(f"{STAT_AFTER_DESCRIPTION} 62% of adults listen monthly {citation}.") == STAT_AFTER_DESCRIPTION


def test_a_wikipedia_part_of_a_multi_source_citation_is_removed_and_the_real_source_kept():
    assert source_rules.strip_wikipedia("62% of adults listen monthly (Nielsen, 2025; Wikipedia, 2024).") == \
        "62% of adults listen monthly (Nielsen, 2025)."
    # "via Wikipedia" means the real source is the one Wikipedia relays: keep it
    assert source_rules.strip_wikipedia("62% of adults listen monthly (Nielsen, 2025, via Wikipedia).") == \
        "62% of adults listen monthly (Nielsen, 2025)."
    assert source_rules.strip_wikipedia("62% of adults listen monthly (Pew Research Center, 2024 (via Wikipedia)).") == \
        "62% of adults listen monthly (Pew Research Center, 2024)."


def test_prose_attribution_to_wikipedia_is_removed():
    assert source_rules.strip_wikipedia("According to Wikipedia, the company sells furniture. Streaming is growing.") == \
        "The company sells furniture. Streaming is growing."
    assert source_rules.strip_wikipedia("The company sells furniture, according to Wikipedia.") == "The company sells furniture."
    # a figure that only Wikipedia vouched for is dropped with its sentence
    assert source_rules.strip_wikipedia("According to Wikipedia, 19% of the U.S. population is Hispanic. Streaming is growing.") == \
        "Streaming is growing."
    assert source_rules.strip_wikipedia("Wikipedia lists 40 locations. The company sells furniture.") == "The company sells furniture."
    assert source_rules.strip_wikipedia("See https://en.wikipedia.org/wiki/Foo for details. The company sells furniture.") == \
        "The company sells furniture."


def test_text_without_wikipedia_is_returned_byte_for_byte():
    text = "Line one.\n\nLine two with a link [Nielsen](https://www.nielsen.com/wikipedia-study)  and  odd   spacing (Nielsen, 2025).\n"
    assert source_rules.strip_wikipedia(text) is text
    assert source_rules.strip_wikipedia("") == "" and source_rules.strip_wikipedia(None) is None
    assert source_rules.strip_wikipedia(42) == 42


def test_paragraph_breaks_survive_a_dropped_sentence():
    text = "Line one.\nLine two says Wikipedia reports 5 things.\nLine three."
    assert source_rules.strip_wikipedia(text) == "Line one.\nLine three."


def test_a_figure_vouched_for_by_a_separate_citation_field_is_kept():
    assert source_rules.strip_wikipedia("62% of adults listen monthly (Wikipedia).", supported_elsewhere=True) == \
        "62% of adults listen monthly."


@pytest.mark.parametrize("value,expected", [
    ("Wikipedia, 2024", ""),
    ("Wikipedia", ""),
    ("https://en.wikipedia.org/wiki/Podcast", ""),
    ("[Wikipedia](https://en.wikipedia.org/wiki/Podcast)", ""),
    ("Nielsen, 2025", "Nielsen, 2025"),
    ("Nielsen, 2025; Wikipedia, 2024", "Nielsen, 2025"),
    ("Nielsen, 2025 and Wikipedia, 2024", "Nielsen, 2025"),
    ("Nielsen, 2025 (via Wikipedia)", "Nielsen, 2025"),
    ("Pew Research Center, 2024, from Wikipedia", "Pew Research Center, 2024"),
])
def test_a_citation_field_loses_its_wikipedia_part(value, expected):
    assert source_rules.strip_wikipedia_source(value) == expected


def test_scrub_model_output_walks_nested_dicts_and_lists_in_place():
    data = {"a": "Fine (Wikipedia, 2024).", "b": ["x", {"c": "See https://en.wikipedia.org/wiki/X for more. Kept."}], "n": 3}
    out = source_rules.scrub_model_output(data)
    assert out is data
    assert data == {"a": "Fine.", "b": ["x", {"c": "Kept."}], "n": 3}


# ============================================================================================================
# 4. The rule is in every research prompt
# ============================================================================================================

def _request() -> ProposalRequest:
    return ProposalRequest(
        client_name="Acme", demo="Hispanic adults 25-54", language="Spanish", geo="Los Angeles",
        start_date="2026-10-01", end_date="2026-12-31", total_months=3, request_type="New Business",
        campaign_goal="Awareness", products_selected=["AudioEngage"],
    )


def _items():
    return [
        LineItem(product_name="AudioEngage", monthly_budget=2000, months=3, id="a"),
        LineItem(product_name="Search - SEM", monthly_budget=1500, months=3, id="b"),
    ]


def _assert_has_no_wikipedia_rule(prompt: str):
    assert source_rules.NO_WIKIPEDIA_RULE in prompt
    assert "NEVER CITE WIKIPEDIA" in prompt
    assert "primary" in source_rules.NO_WIKIPEDIA_RULE and "references" in source_rules.NO_WIKIPEDIA_RULE


def test_the_rule_text_names_every_wikimedia_host_and_the_primary_source_fallback():
    rule = source_rules.NO_WIKIPEDIA_RULE
    for word in ("wikipedia.org", "wikimedia.org", "wikidata.org", "any language", "study", "report", "filing", "news outlet",
                 "leave the number out"):
        assert word in rule, word


def test_the_blurb_and_email_prompt_carries_the_no_wikipedia_rule_and_the_new_claim_rules():
    prompt = ai_enricher._build_prompt(_request(), _items())
    _assert_has_no_wikipedia_rule(prompt)
    assert "Never cite Wikipedia in a blurb or an email" in prompt
    # item 4: specific and sourced, no unfounded audience claims
    assert "NO UNFOUNDED CLAIMS" in prompt
    assert USER_BAD_EXAMPLE in prompt                                  # spelled out as a BAD example
    assert "BOTH a specific number AND a named source" in prompt
    assert "ONE concrete statistic" in prompt
    for banned in ("always", "often", "typically", "many", "most", "increasingly"):
        assert f'"{banned}"' in prompt, banned
    assert "Plain text only: no URLs, links" in prompt
    assert prompt.count("GOOD") >= 3 and prompt.count("BAD") >= 6        # several good / bad examples, not one


def test_the_prompt_examples_use_slots_so_a_model_cannot_copy_a_made_up_stat():
    prompt = ai_enricher._build_prompt(_request(), _items())
    assert "[N]% of [specific group]" in prompt
    assert "never copy a figure or a source from these examples" in prompt


def test_the_blurb_prompt_still_keeps_the_established_short_format():
    prompt = ai_enricher._build_prompt(_request(), _items())
    assert "ONE short paragraph, 25–55 words" in prompt
    assert "(Publisher, Year)" in prompt
    assert "Do NOT repeat the audience, geography, language" in prompt
    assert "Never mention variants, tiers, versions or format options" in prompt


def test_the_prompts_good_examples_pass_the_validator():
    for good in (
        "Banner ads shown across a network of websites and apps, retargeting local shoppers as they browse so your offer stays in front of them after they leave your site.",
        "Text ads that appear on Google when people search for what you offer, so your business shows up at the moment they're ready to call or visit.",
    ):
        assert good in " ".join(ai_enricher._build_prompt(_request(), _items()).split())
        assert not claim_rules.is_unsupported_claim(good)
    # ...and its BAD claim examples are exactly what the validator drops
    for bad in (USER_BAD_EXAMPLE, "Many Latino families increasingly rely on streaming audio.",
                "Podcasts are a trusted, highly engaging channel."):
        assert claim_rules.is_unsupported_claim(bad)


def test_the_strategy_brief_prompt_carries_the_no_wikipedia_rule():
    prompt = sb._build_prompt(_request(), None)
    _assert_has_no_wikipedia_rule(prompt)
    assert "never Wikipedia" in prompt                                 # also on the citation-field instruction


def test_the_roadblocks_prompt_carries_the_no_wikipedia_rule():
    prompt = rb._build_prompt(_request(), _items(), None)
    _assert_has_no_wikipedia_rule(prompt)
    assert "never Wikipedia" in prompt                                 # also on the "source" RULE
    assert "platform's or network's own published policy page" in prompt


def _stub_openai(monkeypatch, module):
    monkeypatch.setenv("OPENAI_API_KEY", "test-key-not-real")
    monkeypatch.setattr(module, "_HAS_OPENAI", True)
    monkeypatch.setattr(module, "_OpenAI", lambda api_key=None, **kw: object())


def test_the_gamma_outline_refiner_carries_the_rule_and_strips_wikipedia_from_its_answer(monkeypatch):
    _stub_openai(monkeypatch, ai_enricher)
    seen = {}

    def fake_chat(client, *, model, prompt, max_completion_tokens):
        seen["prompt"] = prompt
        return json.dumps({"outline": "Title\nStreaming is growing (Wikipedia, 2024).\nBudget: $5,000 for the flight."})

    monkeypatch.setattr(llm_utils, "chat_text", fake_chat)
    result = ai_enricher.refine_gamma_outline("Title\nBudget: $5,000 for the flight.", "make it shorter")
    assert result["error"] is None
    _assert_has_no_wikipedia_rule(seen["prompt"])
    assert "Wikipedia" not in result["outline"]
    assert "$5,000" in result["outline"]                              # a budget figure is never collateral damage


def test_every_web_search_call_uses_the_plain_supported_tool_config(monkeypatch):
    """The pinned SDK's web_search tool filter only offers `allowed_domains` (an allow-list); there is no block-list
    to put Wikipedia on, so the callers send exactly the plain tool and the prompt + code layers do the work."""
    calls = []

    def fake_responses(client, *, model, prompt, max_output_tokens, tools=None):
        calls.append(tools)
        if "ad operations compliance" in prompt:
            return json.dumps({"overall_summary": "s", "product_roadblocks": []})
        if "senior digital media sales strategist" in prompt:
            return json.dumps({"client_summary": "c", "recommended_tactics": [], "key_insights": []})
        return json.dumps({"campaign_name": "X", "product_blurbs": [], "internal_email_body": "b", "client_email_body": "b"})

    monkeypatch.setattr(llm_utils, "responses_text", fake_responses)
    _stub_openai(monkeypatch, ai_enricher)
    _stub_openai(monkeypatch, sb)
    _stub_openai(monkeypatch, rb)
    ai_enricher.enrich_proposal(_request(), _items(), "EVC-1")
    sb._run_brief_call(object(), sb._build_prompt(_request(), None))
    rb.generate_roadblocks(_request(), _items())
    assert len(calls) == 3
    assert all(tools == [{"type": "web_search"}] for tools in calls)


# ============================================================================================================
# 5. The code layer on what each service gets back
# ============================================================================================================

def test_enrichment_strips_wikipedia_from_blurbs_and_emails_and_drops_a_stat_that_rested_on_it():
    raw = json.dumps({
        "campaign_name": "Fall Push",
        "product_blurbs": [
            {"product_name": "AudioEngage",
             "blurb": "Audio ads across internet radio, podcasts and streaming. 62% of adults listen to podcasts monthly (Wikipedia, 2024)."},
            {"product_name": "Search - SEM",
             "blurb": ("Text ads on Google that appear when people search for what you offer. "
                       "62% of consumers research a local business online first "
                       "([en.wikipedia.org](https://en.wikipedia.org/wiki/Search_engine?utm_source=openai)).")},
        ],
        "internal_email_subject": "s",
        "internal_email_body": "Hi Sam,\nPodcasts are growing (Wikipedia, 2024).\n{{PROPOSAL_LINE}}\n{{PLANNER_SIGNATURE}}",
        "client_email_subject": "s",
        "client_email_body": "Digital matters. According to Wikipedia, 19% of the U.S. is Hispanic. See the plan below.",
    })
    out = _parse_response(raw, _request(), [])
    by = {b.product_name: b.blurb for b in out.product_blurbs}
    assert by["AudioEngage"] == "Audio ads across internet radio, podcasts and streaming."
    assert by["Search - SEM"] == "Text ads on Google that appear when people search for what you offer."
    assert "Wikipedia" not in out.internal_email_body and "{{PROPOSAL_LINE}}" in out.internal_email_body
    assert "{{PLANNER_SIGNATURE}}" in out.internal_email_body
    assert "Wikipedia" not in out.client_email_body and "See the plan below." in out.client_email_body
    assert "19%" not in out.client_email_body


def test_enrich_proposal_end_to_end_with_a_stubbed_model(monkeypatch):
    raw = json.dumps({
        "campaign_name": "Fall Push",
        "product_blurbs": [
            {"product_name": "AudioEngage",
             "blurb": f"Audio ads across internet radio, podcasts and streaming. {USER_BAD_EXAMPLE}"},
            {"product_name": "Search - SEM",
             "blurb": ("Text ads on Google that appear when people search for what you offer. "
                       "46% of Hispanic adults research local businesses on their phones (Google, 2025).")},
        ],
        "internal_email_subject": "s", "internal_email_body": "b", "client_email_subject": "s", "client_email_body": "b",
    })
    _stub_openai(monkeypatch, ai_enricher)
    monkeypatch.setattr(llm_utils, "responses_text", lambda client, **kw: raw)
    result = ai_enricher.enrich_proposal(_request(), _items(), "EVC-1")
    by = {b.product_name: b.blurb for b in result.product_blurbs}
    assert by["AudioEngage"] == "Audio ads across internet radio, podcasts and streaming."
    assert by["Search - SEM"].endswith("(Google, 2025).") and "46%" in by["Search - SEM"]
    assert result.used_web_search is True


def _brief_json(**tactic):
    base = {"product_family": "Audio", "rationale": "Reaches the audience.", "data_point": "62% of adults listen monthly",
            "citation": "Nielsen, 2025", "entravision_advantage": "Network.", "suggested_budget_pct": 50}
    base.update(tactic)
    return json.dumps({
        "client_summary": "Acme sells furniture (Wikipedia). It has 40 stores according to Wikipedia.",
        "market_context": "The market is steady.",
        "objectives_analysis": "o",
        "strategy_summary": "s",
        "recommended_tactics": [base],
        "key_insights": ["Podcasts grew 12% (Wikipedia, 2024).", "Keep creative short (Nielsen, 2025).", "Plain insight."],
    })


def test_strategy_brief_drops_a_wikipedia_only_citation_together_with_its_data_point():
    brief = sb._parse(_brief_json(citation="Wikipedia, 2024"))
    t = brief["recommended_tactics"][0]
    assert t["citation"] == "" and t["data_point"] == ""
    assert t["rationale"] == "Reaches the audience."                    # the rest of the tactic is untouched
    assert brief["error"] is None


def test_strategy_brief_keeps_the_real_source_when_wikipedia_was_one_of_several():
    t = sb._parse(_brief_json(citation="Nielsen, 2025; Wikipedia, 2024"))["recommended_tactics"][0]
    assert t["citation"] == "Nielsen, 2025" and t["data_point"] == "62% of adults listen monthly"


def test_strategy_brief_leaves_an_ordinary_tactic_alone_and_cleans_the_prose_fields():
    brief = sb._parse(_brief_json())
    t = brief["recommended_tactics"][0]
    assert t["citation"] == "Nielsen, 2025" and t["data_point"] == "62% of adults listen monthly"
    assert brief["client_summary"] == "Acme sells furniture."          # the 40-store figure rested on Wikipedia alone
    assert brief["key_insights"] == ["Keep creative short (Nielsen, 2025).", "Plain insight."]


def test_strategy_brief_data_point_that_names_wikipedia_keeps_its_figure_when_the_citation_field_is_real():
    t = sb._parse(_brief_json(data_point="62% of adults listen monthly (Wikipedia, 2024)"))["recommended_tactics"][0]
    assert t["data_point"] == "62% of adults listen monthly" and t["citation"] == "Nielsen, 2025"


def _roadblocks_json(source):
    return json.dumps({
        "overall_summary": "Low risk overall. Policy 7 changed 3 times (Wikipedia, 2024).",
        "product_roadblocks": [{
            "product_name": "Search - SEM", "risk_level": "low",
            "risks": [{"issue": "Ad review", "detail": "Ads go through review (Wikipedia).", "source": source}],
            "recommended_mitigation": "Submit early. See https://en.wikipedia.org/wiki/Google_Ads for more.",
        }],
    })


def test_roadblocks_replace_a_wikipedia_only_source_with_a_visible_no_primary_source_note():
    result = rb._parse(_roadblocks_json("https://en.wikipedia.org/wiki/Google_Ads"), used_web_search=True)
    risk = result["product_roadblocks"][0]["risks"][0]
    assert risk["source"] == rb._NO_PRIMARY_SOURCE and "Wikipedia" not in risk["source"]
    assert risk["detail"] == "Ads go through review."
    assert result["product_roadblocks"][0]["recommended_mitigation"] == "Submit early."
    assert result["overall_summary"] == "Low risk overall."
    assert result["error"] is None


def test_roadblocks_keep_the_real_source_when_wikipedia_was_one_of_several():
    result = rb._parse(_roadblocks_json("Google Ads policy, 2025; Wikipedia"), used_web_search=True)
    assert result["product_roadblocks"][0]["risks"][0]["source"] == "Google Ads policy, 2025"


def test_roadblocks_leave_a_clean_result_alone():
    raw = json.dumps({"overall_summary": "Fine.", "product_roadblocks": [{
        "product_name": "Search - SEM", "risk_level": "medium",
        "risks": [{"issue": "Review", "detail": "Takes 3 business days.", "source": "Google Ads policy, 2025"}],
        "recommended_mitigation": "Submit early."}]})
    result = rb._parse(raw, used_web_search=True)
    risk = result["product_roadblocks"][0]["risks"][0]
    assert risk == {"issue": "Review", "detail": "Takes 3 business days.", "source": "Google Ads policy, 2025"}


# --- template slots and footnote markers (a model copying the prompt's examples) ---------------------------------------

def test_unfilled_template_slots_are_dropped_and_footnote_numbers_stripped():
    slot = "[N]% of [specific group] watch YouTube weekly ([Publisher], [Year])."
    assert _tighten_blurb(f"{DESCRIPTION} {slot}", "Any") == DESCRIPTION
    assert _tighten_blurb(f"{DESCRIPTION}[1] It plays on phones and TVs.[2]", "Any") == f"{DESCRIPTION} It plays on phones and TVs."
    # a real bracketed citation is not a slot
    kept = f"{DESCRIPTION} 62% of adults watch YouTube weekly [Nielsen, 2025]."
    assert _tighten_blurb(kept, "Any") == kept
