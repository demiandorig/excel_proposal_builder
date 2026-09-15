"""
Regression test for a real bug found while verifying that Curate-step rate
overrides consistently flow into every rate-dependent Excel cell: the
IMPRESSIONS (I) column on Proposal A / Proposal A (Gross) — for Fixed /
estimated-impressions products (Meta, YouTube, TikTok, LinkedIn, Spotify,
Branded Content, ...) — was hardcoding the CATALOG's estimated_cpm_for_imps
into the formula, silently ignoring LineItem.estimated_cpm_override even
though that field exists specifically to feed this calculation. The Avails
step's own "Max Recommended Monthly Imps/Spend" figures were never
affected (they already resolved the override correctly) — only this one
column, on the main proposal sheets.
"""
from unittest.mock import patch

import openpyxl
import pytest

from app import excel_template as et
from app.services.proposal_generator import LineItem, _populate_line_items


@pytest.fixture(autouse=True)
def _mock_catalog_db():
    with patch("app.catalog.load_rate_overrides", return_value={}), \
         patch("app.catalog.load_custom_products", return_value=[]), \
         patch("app.catalog.load_deleted_builtin_names", return_value=set()):
        yield


def _first_fixed_estimate_product():
    from app.catalog import CATALOG
    for p in CATALOG:
        if (p.buying_model == "Fixed" or p.estimated_impressions) and p.estimated_cpm_for_imps:
            return p
    pytest.skip("no Fixed/estimated-impressions catalog product with estimated_cpm_for_imps found")


def test_impressions_formula_uses_catalog_default_when_no_override():
    p = _first_fixed_estimate_product()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = et.build_proposal_a(wb, [p], with_sections=False, total_months=1)
    li = LineItem(product_name=p.name, monthly_budget=1000, months=1, id="1")
    _populate_line_items(ws, [p], [li], gross=False)
    assert str(p.estimated_cpm_for_imps) in ws["I19"].value


def test_impressions_formula_uses_estimated_cpm_override_when_set():
    p = _first_fixed_estimate_product()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = et.build_proposal_a(wb, [p], with_sections=False, total_months=1)
    li = LineItem(product_name=p.name, monthly_budget=1000, months=1, id="1", estimated_cpm_override=99.99)
    _populate_line_items(ws, [p], [li], gross=False)
    formula = ws["I19"].value
    assert "99.99" in formula
    assert str(p.estimated_cpm_for_imps) not in formula


def test_impressions_formula_override_applies_on_gross_sheet_too():
    p = _first_fixed_estimate_product()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = et.build_proposal_a_gross(wb, [p], total_months=1)
    li = LineItem(product_name=p.name, monthly_budget=1000, months=1, id="1", estimated_cpm_override=42.5)
    _populate_line_items(ws, [p], [li], gross=True)
    assert "42.5" in ws["I19"].value


def test_added_value_line_impressions_cell_does_not_error_on_text_budget():
    """L{row} becomes non-numeric ("Estimated $X value") for an Added
    Value line — the IMPRESSIONS formula must degrade to "NA" via
    IFERROR, never a raw #VALUE!, and must NOT apply estimated_cpm_override
    (an AV line has no real budget to estimate impressions from)."""
    p = _first_fixed_estimate_product()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = et.build_proposal_a(wb, [p], with_sections=False, total_months=1)
    li = LineItem(product_name=p.name, monthly_budget=0, months=1, id="1",
                  estimated_cpm_override=99.99, is_added_value=True, added_value_pct=5)
    _populate_line_items(ws, [p], [li], gross=False)
    formula = ws["I19"].value
    assert formula.startswith("=IFERROR(")
    assert "99.99" not in formula  # override never applied to an AV line
