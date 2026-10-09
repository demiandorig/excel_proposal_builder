"""
Full Flight billing period (time_unit == "full_flight") through the REAL pipeline:
generate_proposal() -> openpyxl cells. The whole flight is ONE period: "Months: 1",
one breakdown column spanning the flight, and NOTHING multiplied — totals are the
plain sum of the line items.

Reference scenario = the manually edited export a planner supplied: an email/display
flight 2026-11-24 -> 2026-12-31 at a 15% agency fee whose lines sum to $12,750 net /
$15,000 gross, the email line reading 48,571 impressions, $16.47 gross rate, $800 gross
budget and a "NOVEMBER–DECEMBER 2026" column holding "$800 (100%)".

Same DB-mocking pattern as test_time_unit_export.py (no Postgres needed).
"""
from __future__ import annotations

import os

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import openpyxl
import pytest

from app import excel_template as et
from app.services import monthly_allocation as mo
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import AddonItem, LineItem, generate_proposal, normalize_full_flight_lines

EMAIL = "Email Campaigns - Display Re-targeting"
REDROP = "Email Campaigns and/or Email Campaigns - Re-Drop"
IAB = "eDigital Network Display - Standard IAB"
GEOFENCE = "Display - Geo Fence"
FLIGHT_NET = {EMAIL: 680.0, REDROP: 3600.0, IAB: 5470.0, GEOFENCE: 3000.0}   # sums to 12,750 net = 15,000 gross at 15%


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})
    monkeypatch.setattr("app.catalog.resolve_product_alias", lambda name: None)   # no alias table without Postgres


def _req(**overrides) -> ProposalRequest:
    fields = dict(client_name="Acme Corp", request_type="New Business", start_date="2026-11-24",
                  end_date="2026-12-31", total_months=2, geo="Los Angeles", agency_fee=0.15)
    fields.update(overrides)
    return ProposalRequest(**fields)


def _lines(**overrides) -> list[LineItem]:
    return [LineItem(product_name=name, monthly_budget=total, months=1, id=f"li{i}", **overrides)
            for i, (name, total) in enumerate(FLIGHT_NET.items())]


def _export(tmp_path, *, lines=None, req=None, addons=None, tiers=None, force_tabs=None, time_unit="full_flight"):
    out = tmp_path / "ff.xlsx"
    lines = lines if lines is not None else _lines()
    summary = generate_proposal(req or _req(), lines, out, tiers=tiers, time_unit=time_unit, force_tabs=force_tabs,
                                addons=addons if addons is not None else [AddonItem("Call Tracking", 100.0), AddonItem("Non Media Offering - Brand Lift", 0.0)])
    return summary, openpyxl.load_workbook(out)


def _find(ws, text, col="C"):
    return next(r for r in range(1, ws.max_row + 1) if ws[f"{col}{r}"].value == text)


# ---------------------------------------------------------------------------
# The reference scenario
# ---------------------------------------------------------------------------

def test_net_sheet_matches_the_reference_layout(tmp_path):
    summary, wb = _export(tmp_path)
    ws = wb["Proposal A"]
    assert ws["H10"].value == "Months:" and ws["I10"].value == 1
    assert isinstance(ws["I10"].value, int)

    total_row = _find(ws, "TOTAL DIGITAL")                     # plain label — not "TOTAL DIGITAL MONTHLY"
    assert ws[f"L{total_row}"].value == "=ROUNDDOWN(SUM(L19:L22),0)"
    assert [ws[f"L{r}"].value for r in range(19, 23)] == list(FLIGHT_NET.values())

    grand_row = _find(ws, "TOTAL DIGITAL — CAMPAIGN")           # static label, no "N-MONTH" suffix
    assert not str(ws[f"C{grand_row}"].value).startswith("=")
    for col in "ILN":
        formula = ws[f"{col}{grand_row}"].value
        if formula is None:
            continue
        assert "*" not in formula and "I10" not in formula, (col, formula)   # NOTHING is multiplied
    assert ws[f"L{grand_row}"].value == f'=ROUNDDOWN(L{total_row} + SUMIF(L27:L28,">0"),0)'
    assert ws[f"I{grand_row}"].value == f"=ROUNDDOWN(I{total_row},0)"

    # exactly one breakdown column, headed with the flight's date span
    assert ws["S16"].value == "MONTHLY BREAKDOWN (NET)"            # the planner's own title for the block, with its basis
    assert ws["S17"].value == "NOVEMBER–DECEMBER 2026"
    assert ws["T17"].value in (None, "") and ws["T19"].value is None
    assert ws["S19"].value == '=IFERROR("$"&TEXT(L19,"#,##0")&IF(L19>0," (100%)"," (0%)"),"—")'
    assert ws[f"S{total_row}"].value == f"=L{total_row}"          # == the TOTAL DIGITAL cell, no second rounding

    # the summary the API reports is the plain sum too
    assert summary["total_net"] == pytest.approx(sum(FLIGHT_NET.values()))
    assert summary["total_gross"] == pytest.approx(15000.0)
    assert summary["warnings"] == []


def test_gross_sheet_shows_gross_dollars_in_the_single_column(tmp_path):
    _, wb = _export(tmp_path)
    ws = wb["Proposal A (Gross)"]
    assert ws["I14"].value == 0.15
    assert ws["H10"].value == "Months:" and ws["I10"].value == 1
    assert ws["U16"].value == "MONTHLY BREAKDOWN (GROSS)"
    assert ws["U17"].value == "NOVEMBER–DECEMBER 2026"
    # live formula over GROSS BUDGET (N) — the reference reads "$800 (100%)" for the email line
    assert ws["U19"].value == '=IFERROR("$"&TEXT(N19,"#,##0")&IF(N19>0," (100%)"," (0%)"),"—")'
    assert ws["N19"].value == "=L19/(1-$I$14)"
    total_row = _find(ws, "TOTAL DIGITAL")
    assert ws[f"U{total_row}"].value == f"=N{total_row}"
    grand_row = _find(ws, "TOTAL DIGITAL — CAMPAIGN")
    assert ws[f"N{grand_row}"].value == f'=ROUNDDOWN(N{total_row} + SUMIF(N27:N28,">0"),0)'
    assert "*" not in ws[f"L{grand_row}"].value and "*" not in ws[f"N{grand_row}"].value


def test_no_standalone_breakdown_tab_and_no_single_cell_merges(tmp_path):
    summary, wb = _export(tmp_path)
    assert not [n for n in wb.sheetnames if "Breakdown" in n]            # it would just repeat the inline column
    assert not [t for t in summary["tabs_built"] if "Breakdown" in t]
    for ws in wb.worksheets:
        for rng in ws.merged_cells.ranges:                              # Excel flags one-cell merges for "repair"
            assert (rng.min_row, rng.min_col) != (rng.max_row, rng.max_col), (ws.title, str(rng))


def test_wsections_sheet_aligns_breakdown_cells_with_the_product_rows(tmp_path):
    _, wb = _export(tmp_path)
    ws = wb["Proposal A (wsections)"]
    product_rows = [r for r in range(19, ws.max_row) if str(ws[f"C{r}"].value or "").split("\n")[0] in FLIGHT_NET]
    assert len(product_rows) == 4
    for r in product_rows:
        assert ws[f"S{r}"].value == f'=IFERROR("$"&TEXT(L{r},"#,##0")&IF(L{r}>0," (100%)"," (0%)"),"—")'
    banner_rows = [r for r in range(19, ws.max_row) if str(ws[f"C{r}"].value or "").startswith("SECTION")]
    assert banner_rows and all(ws[f"S{r}"].value is None for r in banner_rows)


def test_month_mode_is_untouched_by_the_new_unit(tmp_path):
    """Same plan in monthly mode keeps every month label/multiplier — Full Flight must not leak into it (the one visible
    change in this mode is that the breakdown title now names its basis: "(NET)" here, "(GROSS)" on the Gross sheet)."""
    s, e = mo.parse_flexible_date("2026-11-24"), mo.parse_flexible_date("2026-12-31")
    periods = mo.periods_between(s, e, "month")
    lines = [LineItem(product_name=n, monthly_budget=t / 2, months=2, id=f"m{i}",
                      monthly_allocations=mo.compute_default_allocation(t, periods, "even"))
             for i, (n, t) in enumerate(FLIGHT_NET.items())]
    _, wb = _export(tmp_path, lines=lines, time_unit="month")
    ws = wb["Proposal A"]
    assert ws["I10"].value == 2
    assert ws[f"C{_find(ws, 'TOTAL DIGITAL MONTHLY')}"].value == "TOTAL DIGITAL MONTHLY"
    assert any(str(ws[f"C{r}"].value or "").startswith('="TOTAL DIGITAL — "&IF(ISNUMBER(I10),I10,3)&"-MONTH CAMPAIGN"')
               for r in range(20, ws.max_row))
    assert ws["S16"].value == "MONTHLY BREAKDOWN (NET)"
    assert (ws["S17"].value, ws["T17"].value) == ("NOVEMBER 2026", "DECEMBER 2026")
    assert "Monthly Breakdown" in wb.sheetnames


# ---------------------------------------------------------------------------
# Labels and edge flights
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("start,end,label", [
    ("2026-11-03", "2026-11-20", "NOVEMBER 2026"),                     # inside one month: NOT "NOVEMBER–NOVEMBER 2026"
    ("2026-11-24", "2026-12-31", "NOVEMBER–DECEMBER 2026"),
    ("2026-12-15", "2027-01-20", "DECEMBER 2026–JANUARY 2027"),       # crosses a year
    ("2026-11-01", "2027-01-31", "NOVEMBER 2026–JANUARY 2027"),
])
def test_single_column_header_follows_the_flight_span(tmp_path, start, end, label):
    _, wb = _export(tmp_path, req=_req(start_date=start, end_date=end))
    assert wb["Proposal A"]["S17"].value == label
    assert wb["Proposal A"]["I10"].value == 1                            # the calendar span never leaks into Months


def test_flight_without_parseable_dates_still_reads_months_1_and_does_not_crash(tmp_path):
    _, wb = _export(tmp_path, req=_req(start_date="", end_date="", total_months=3))
    ws = wb["Proposal A"]
    assert ws["I10"].value == 1
    assert ws["S17"].value is None and ws["S16"].value is None            # no dates -> no breakdown column (and no merge)
    assert any(ws[f"C{r}"].value == "TOTAL DIGITAL — CAMPAIGN" for r in range(1, ws.max_row))


def test_a_per_option_date_override_gets_its_own_single_period(tmp_path):
    lines_a = _lines()
    lines_b = [LineItem(product_name=EMAIL, monthly_budget=500.0, months=1, id="b0")]
    tiers = [
        {"label": "A", "line_items": lines_a, "avails_data": {}},
        {"label": "B", "name": "Short", "line_items": lines_b, "avails_data": {},
         "start_date": "2026-12-01", "end_date": "2026-12-20"},
    ]
    _, wb = _export(tmp_path, lines=lines_a, tiers=tiers)
    headers = {ws.title: ws["S17"].value for ws in wb.worksheets if ws["S17"].value}
    assert "NOVEMBER–DECEMBER 2026" in headers.values() and "DECEMBER 2026" in headers.values()
    assert all(ws["I10"].value == 1 for ws in wb.worksheets if ws["H10"].value == "Months:")


# ---------------------------------------------------------------------------
# Normalization (stale / legacy payloads can never multiply a total)
# ---------------------------------------------------------------------------

def test_a_legacy_month_payload_is_folded_into_one_period_with_the_total_preserved(tmp_path):
    """A month-mode line (340 x 2, month-keyed allocations) sent as full_flight is folded: 680 x 1."""
    stale = [LineItem(product_name=EMAIL, monthly_budget=340.0, months=2, id="old",
                      monthly_allocations={"2026-11": 100.0, "2026-12": 580.0},
                      period_merge_groups=[["2026-11", "2026-12"]])]
    summary, wb = _export(tmp_path, lines=stale, addons=[])
    ws = wb["Proposal A"]
    assert ws["I10"].value == 1 and ws["L19"].value == 680.0
    assert summary["total_net"] == pytest.approx(680.0)
    assert ws["S19"].value.startswith('=IFERROR("$"&TEXT(L19')           # stale month keys are not what's printed
    assert any("folded into one period" in w for w in summary["warnings"])


def test_normalize_full_flight_lines_contract():
    paid = LineItem(product_name=EMAIL, monthly_budget=340.0, months=2, id="p", monthly_allocations={"2026-11": 1.0})
    av = LineItem(product_name=GEOFENCE, monthly_budget=0.0, months=3, id="av", is_added_value=True,
                  monthly_allocations={"2026-11": 5.0})
    clean = LineItem(product_name=IAB, monthly_budget=1000.0, months=1, id="ok")
    warnings = normalize_full_flight_lines([paid, av, clean])
    assert (paid.months, paid.monthly_budget) == (1, 680.0)
    assert paid.monthly_allocations == {mo.FULL_FLIGHT_KEY: 680.0}
    assert (av.months, av.monthly_allocations) == (1, None)                # Added Value: no allocation, $0 stays $0
    assert av.monthly_budget == 0.0
    assert (clean.months, clean.monthly_budget) == (1, 1000.0)
    assert clean.monthly_allocations == {mo.FULL_FLIGHT_KEY: 1000.0}
    assert len(warnings) == 1 and "'" + EMAIL + "'" in warnings[0]
    assert normalize_full_flight_lines([paid, av, clean]) == []            # idempotent: second pass changes nothing
    assert paid.monthly_budget == 680.0


def test_added_value_and_addons_are_never_multiplied_or_given_a_period(tmp_path):
    lines = _lines() + [LineItem(product_name=GEOFENCE, monthly_budget=0.0, months=1, id="av", is_added_value=True)]
    _, wb = _export(tmp_path, lines=lines)
    ws = wb["Proposal A"]
    # paid lines occupy rows 19-22; the Added Value line sorts to the bottom (row 23) and gets no breakdown cell
    assert all(ws[f"S{r}"].value for r in range(19, 23))
    assert str(ws["C23"].value).startswith(GEOFENCE) and ws["S23"].value is None
    total_row = _find(ws, "TOTAL DIGITAL")
    assert ws[f"L{total_row}"].value == "=ROUNDDOWN(SUM(L19:L23),0)"
    grand_row = _find(ws, "TOTAL DIGITAL — CAMPAIGN")
    assert "SUMIF(" in ws[f"L{grand_row}"].value and "*" not in ws[f"L{grand_row}"].value   # add-ons added once, not x months


# ---------------------------------------------------------------------------
# SOV and minimums — a flight-long budget against MONTHLY avails ceilings
# ---------------------------------------------------------------------------

def test_sov_for_a_flight_is_the_whole_flight_budget_over_the_monthly_ceiling():
    """Full Flight is ONE period: its (whole-flight) budget is compared with the monthly ceiling exactly as it stands —
    no division by the calendar months the flight touches."""
    p = next(x for x in __import__("app.catalog", fromlist=["CATALOG"]).CATALOG if x.name == EMAIL)
    avail = {"max_spend": 1400.0}
    flight = et.compute_sov_pct(p, 680.0, avail, time_unit="full_flight")
    assert flight == pytest.approx(680.0 / 1400.0 * 100)
    # a monthly plan of the same dollars (680 a month) reads the same; the per-unit conversions are unchanged
    assert et.compute_sov_pct(p, 680.0, avail, time_unit="month") == pytest.approx(flight)
    assert et.compute_sov_pct(p, 340.0, avail, time_unit="week") == pytest.approx(340.0 * 4 / 1400.0 * 100)
    assert et.compute_sov_pct(p, 340.0, avail, time_unit="quarter") == pytest.approx(340.0 / 3 / 1400.0 * 100)


def test_live_sov_formula_compares_the_whole_flight_budget_with_the_monthly_ceiling(tmp_path):
    avails = {"li0": {"max_spend": 1400.0}}
    tiers = [{"label": "A", "line_items": _lines(), "avails_data": avails}]
    for start, end in (("2026-11-24", "2026-12-31"), ("2026-11-03", "2026-11-20"), ("2026-09-01", "2026-12-31")):
        _, wb = _export(tmp_path, req=_req(start_date=start, end_date=end), tiers=tiers)
        # never "(L19/N)": one period, however many calendar months it touches
        assert wb["Proposal A"]["Q19"].value == '=IFERROR(L19/O19,"")', (start, end)
        assert wb["Proposal A (Gross)"]["S19"].value == '=IFERROR(L19/Q19,"")', (start, end)


def test_sov_formula_term_shapes():
    assert et._sov_budget_formula_term("L", 5, "full_flight") == "L5"
    assert et._sov_budget_formula_term("L", 5, "month") == "L5"
    assert et._sov_budget_formula_term("L", 5, "week") == "(L5*4)"                      # existing units untouched
    assert et._sov_budget_formula_term("L", 5, "quarter") == "(L5/3)"
    assert et.SOV_MONTHLY_EQUIVALENT_SCALE["full_flight"] == 1.0


def test_avails_only_sheet_relabels_the_budget_column_and_compares_the_flight_budget(tmp_path):
    tiers = [{"label": "A", "line_items": _lines(), "avails_data": {}}]
    force = {"avails_only": True, "net": False, "wsections": False, "gross": False}
    _, wb = _export(tmp_path, tiers=tiers, force_tabs=force)
    ws = wb["Avails-Only"]
    assert ws["I10"].value == "Full-Flight\nBudget"
    assert ws["M10"].value == "% of Avails Reached with\nFull-Flight Budget"
    assert ws["M12"].value == '=IFERROR(I12/K12,"")'      # the whole-flight budget over the MONTHLY ceiling, as it stands
    # the same whether or not the dates can be read
    _, wb0 = _export(tmp_path, req=_req(start_date="", end_date=""), tiers=tiers, force_tabs=force)
    assert wb0["Avails-Only"]["M10"].value == "% of Avails Reached with\nFull-Flight Budget"
    assert wb0["Avails-Only"]["M12"].value == '=IFERROR(I12/K12,"")'


def test_minimum_spend_for_a_flight_is_one_monthly_minimum():
    [period] = mo.periods_between(mo.parse_flexible_date("2026-11-24"), mo.parse_flexible_date("2026-12-31"), "full_flight")
    assert period["period_count"] == 1 and mo.granularity_scale("full_flight") == 1.0
    # Email Re-targeting's minimum is $375 a month; a flight is ONE period however long it runs, so it owes $375 — once:
    # $374 is flagged, $375 (and, of course, $680) is not
    short = mo.check_minimum_violations("l", EMAIL, 375.0, False, [period], {period["key"]: 374.0}, "full_flight")
    assert [v["minimum_required"] for v in short] == [375.0]
    for amount in (375.0, 680.0):
        assert mo.check_minimum_violations("l", EMAIL, 375.0, False, [period], {period["key"]: amount}, "full_flight") == []
    assert mo.check_minimum_violations("l", EMAIL, 375.0, True, [period], {period["key"]: 0.0}, "full_flight") == []   # AV exempt
    # a longer flight changes nothing
    [long_period] = mo.periods_between(mo.parse_flexible_date("2026-09-01"), mo.parse_flexible_date("2027-02-28"), "full_flight")
    assert mo.check_minimum_violations("l", EMAIL, 375.0, False, [long_period], {long_period["key"]: 375.0}, "full_flight") == []


# ---------------------------------------------------------------------------
# Review fixes
# ---------------------------------------------------------------------------

def test_a_one_column_banner_is_wide_enough_for_its_own_title(tmp_path):
    """The one-column block isn't merged, so its banner (now "MONTHLY BREAKDOWN (GROSS)") must fit the column's width."""
    _, wb = _export(tmp_path)
    for sheet, col, banner in (("Proposal A", "S", "MONTHLY BREAKDOWN (NET)"), ("Proposal A (Gross)", "U", "MONTHLY BREAKDOWN (GROSS)")):
        ws = wb[sheet]
        assert ws[f"{col}16"].value == banner
        assert ws.column_dimensions[col].width >= len(banner) + 2, (sheet, ws.column_dimensions[col].width)
        assert ws.column_dimensions[col].width >= len("NOVEMBER–DECEMBER 2026")      # and the date-span header


def test_an_added_value_line_with_a_percent_does_not_poison_the_gross_totals(tmp_path):
    """PRE-EXISTING bug that Full Flight would inherit (its breakdown total cell is =N{total}): on the Gross sheet the AV
    row's GROSS BUDGET was =L/(1-$I$14) over the TEXT 'Estimated $X value' -> #VALUE!, which propagated to TOTAL DIGITAL
    and the campaign total. It is an estimated gift value, so it reads the same text on both sheets."""
    lines = _lines() + [LineItem(product_name=GEOFENCE, monthly_budget=0.0, months=1, id="av", is_added_value=True,
                                 added_value_pct=10.0)]
    for unit in ("full_flight", "month"):
        if unit == "month":
            s, e = mo.parse_flexible_date("2026-11-24"), mo.parse_flexible_date("2026-12-31")
            periods = mo.periods_between(s, e, "month")
            lines = [LineItem(product_name=n, monthly_budget=t / 2, months=2, id=f"m{i}",
                              monthly_allocations=mo.compute_default_allocation(t, periods, "even"))
                     for i, (n, t) in enumerate(FLIGHT_NET.items())] + [
                LineItem(product_name=GEOFENCE, monthly_budget=0.0, months=2, id="av", is_added_value=True, added_value_pct=10.0)]
        _, wb = _export(tmp_path, lines=lines, time_unit=unit)
        ws = wb["Proposal A (Gross)"]
        av_row = next(r for r in range(19, 30) if str(ws[f"C{r}"].value or "").startswith(GEOFENCE) and ws[f"L{r}"].data_type == "s")
        assert str(ws[f"L{av_row}"].value).startswith("Estimated $")
        assert ws[f"N{av_row}"].value == ws[f"L{av_row}"].value, (unit, ws[f"N{av_row}"].value)      # text, not a formula over text
        assert not str(ws[f"N{av_row}"].value).startswith("=")
        total_row = _find(ws, "TOTAL DIGITAL" if unit == "full_flight" else "TOTAL DIGITAL MONTHLY")
        assert ws[f"N{total_row}"].value == f"=ROUNDDOWN(SUM(N19:N{av_row}),0)"


def test_a_full_flight_tier_with_unreadable_dates_warns_instead_of_silently_dropping_the_column(tmp_path):
    summary, wb = _export(tmp_path, req=_req(start_date="Nov 24, 2026", end_date="TBD"))
    assert wb["Proposal A"]["S17"].value is None                                      # no period to label ...
    assert any("couldn't be read" in w and "Full Flight" in w for w in summary["warnings"]), summary["warnings"]   # ... and we say so
    clean, _ = _export(tmp_path)
    assert clean["warnings"] == []
    # week/month/quarter never had this warning (their breakdown is optional) — unchanged
    month_summary, _ = _export(tmp_path, req=_req(start_date="TBD", end_date="TBD"), time_unit="month",
                               lines=[LineItem(product_name=EMAIL, monthly_budget=340.0, months=2, id="m")])
    assert not any("couldn't be read" in w for w in month_summary["warnings"])


# ---------------------------------------------------------------------------
# "Minimum 3 month Commitment." only when this option's dates cover 3+ months
# ---------------------------------------------------------------------------

NET_BASIS = "All rates are NET."
GROSS_BASIS = "Rates shown are GROSS (inclusive of agency commission)."
COMMITMENT = " Minimum 3 month Commitment."


def test_the_three_month_commitment_prints_only_when_the_dates_cover_three_months(tmp_path):
    # Nov 24 - Dec 31 touches two calendar months: no sentence — even with "Total months: 3" typed on the request
    _, wb = _export(tmp_path, req=_req(total_months=3))
    assert wb["Proposal A"]["C14"].value == NET_BASIS
    assert wb["Proposal A (wsections)"]["C14"].value == NET_BASIS
    assert wb["Proposal A (Gross)"]["C14"].value == GROSS_BASIS
    # Oct 1 - Dec 31 covers three: it prints on the Net and the Gross sheets (and the Months cell still reads 1)
    _, wb = _export(tmp_path, req=_req(start_date="2026-10-01", end_date="2026-12-31", total_months=3))
    assert wb["Proposal A"]["C14"].value == NET_BASIS + COMMITMENT
    assert wb["Proposal A (Gross)"]["C14"].value == GROSS_BASIS + COMMITMENT
    assert wb["Proposal A"]["I10"].value == 1
    # a partial first/last month counts whole, like Step 02's Months field: Nov 24 - Jan 31 touches three
    _, wb = _export(tmp_path, req=_req(start_date="2026-11-24", end_date="2027-01-31", total_months=3))
    assert wb["Proposal A"]["C14"].value == NET_BASIS + COMMITMENT


def test_the_three_month_commitment_needs_dates_that_can_be_read(tmp_path):
    for start, end in (("", ""), ("TBD", "TBD"), ("2026-10-01", "")):
        _, wb = _export(tmp_path, req=_req(start_date=start, end_date=end, total_months=3))
        assert wb["Proposal A"]["C14"].value == NET_BASIS, (start, end)        # nothing covers the 3 months it would promise


def test_the_three_month_commitment_follows_each_options_own_dates(tmp_path):
    lines_a = _lines()
    lines_b = [LineItem(product_name=EMAIL, monthly_budget=500.0, months=1, id="b0")]
    tiers = [
        {"label": "A", "line_items": lines_a, "avails_data": {}},
        {"label": "B", "name": "Short", "line_items": lines_b, "avails_data": {}, "start_date": "2026-12-01", "end_date": "2026-12-20"},
    ]
    # the campaign covers three months but Option B's override window is 20 days: only Option A promises the commitment
    _, wb = _export(tmp_path, lines=lines_a, tiers=tiers, req=_req(start_date="2026-10-01", end_date="2026-12-31", total_months=3))
    assert wb["Option A"]["C14"].value == NET_BASIS + COMMITMENT
    assert wb["Option A (Gross)"]["C14"].value == GROSS_BASIS + COMMITMENT
    assert wb["Short"]["C14"].value == NET_BASIS
    assert wb["Short (Gross)"]["C14"].value == GROSS_BASIS
    # ... and the reverse: a short campaign whose Option B overrides to a 3-month window
    tiers[1].update(start_date="2027-01-01", end_date="2027-03-15")
    _, wb = _export(tmp_path, lines=lines_a, tiers=tiers, req=_req(total_months=2))
    assert wb["Option A"]["C14"].value == NET_BASIS
    assert wb["Short"]["C14"].value == NET_BASIS + COMMITMENT


def test_the_other_billing_periods_keep_the_long_standing_commitment_rule(tmp_path):
    """Weekly / Monthly / Quarterly are deliberately unchanged: total_months (which follows the dates when there are any,
    and is the typed figure when there are none) decides."""
    line = [LineItem(product_name=EMAIL, monthly_budget=340.0, months=3, id="m")]
    for unit in ("month", "week", "quarter"):
        _, wb = _export(tmp_path, lines=line, time_unit=unit, req=_req(start_date="", end_date="", total_months=3))
        assert wb["Proposal A"]["C14"].value == NET_BASIS + COMMITMENT, unit
        _, wb = _export(tmp_path, lines=line, time_unit=unit, req=_req(start_date="", end_date="", total_months=2))
        assert wb["Proposal A"]["C14"].value == NET_BASIS, unit


# ---------------------------------------------------------------------------
# Finance policy: billing is always invoiced on a monthly (30-day) cycle — the Full Flight comment says so too
# ---------------------------------------------------------------------------

def test_the_full_flight_comment_keeps_the_monthly_invoicing_policy(tmp_path):
    _, wb = _export(tmp_path)
    for sheet, cell in (("Proposal A", "S16"), ("Proposal A (Gross)", "U16")):
        note = wb[sheet][cell].comment.text
        assert "the entire flight is treated as a single period" in note, sheet
        assert "no monthly split" in note, sheet
        assert "Billing is always invoiced on a monthly (30-day) cycle, regardless of this plan's breakdown granularity." in note, sheet
    # week/month/quarter comments are untouched by this
    assert et._billing_cadence_note("month") == et._BILLING_CADENCE_NOTE_BASE
