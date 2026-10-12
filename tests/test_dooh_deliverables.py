"""
DOOH Deliverables: the DOOH Screenlist table, the pre-made DOOH Summary pivot table, the single "dooh" export-tab key
(old dooh_summary / dooh_screenlist keys still accepted), and DOOH detection by product family.

The pivot's real behaviour (it opens in Excel without a repair prompt, refreshes, sums per-row correctly) was checked
in Excel itself through COM automation; these tests pin the parts a Python test can see: the package structure of the
injected pivot (well-formed XML, rels / content types / ids consistent) and the logic around it.
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import openpyxl
import pytest
from openpyxl import Workbook

from app import excel_template as et
from app.services import dooh_pivot as dp
from app.services.notion_parser import ProposalRequest, classify_output_tabs, normalize_tab_keys
from app.services.proposal_generator import LineItem, generate_proposal

ROOT = Path(__file__).resolve().parents[1]
APP_JS = ROOT / "app" / "static" / "app.js"
INDEX_HTML = ROOT / "app" / "templates" / "index.html"

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

# The user's 40 columns and sample row, verbatim (the source of truth the module is checked against).
USER_COLUMNS = ["id", "adunit_name", "asset_name", "adunit_id", "status", "integration_type", "publisher", "network",
                "media_type", "venue_type", "venue_name", "city", "postal_code", "county", "state", "dma", "country",
                "latitude", "longitude", "address", "cpm", "currency", "four_week_impressions", "impression_multiplier",
                "bidstream_impressions", "bidstream_inferred_impressions", "screen_count", "min_duration", "max_duration",
                "supports_video", "supports_banner", "supports_audio", "supports_dynamic_creative", "slot_dimensions",
                "preferred_ad_format", "ad_formats", "restrictions", "languages", "open_ooh_venue_category_id",
                "asset_image_url"]
USER_SAMPLE = ["1348310", "61e86cc70b573d05e4183fe0", "DVG288", "550eceb3-b6fa-4daf-a9e7-456bfb47846f", "live", "API",
               "GSTV", "GSTV Fast", "Point of Sale", "Retail|Gas Stations|Pump", "61e86cc70b573d05e4183fe0", "DENVER",
               "80216", "Denver County - CO", "CO", "Denver, CO", "USA", "39.774694", "-104.924262",
               "4140 Grape Street, Denver, CO", "32.5", "USD", "2353", "1.483757764", "1675", "1486", "8", "3", "61",
               "TRUE", "TRUE", "TRUE", "FALSE", "1920x1080", "1920x1080", "1920x1080", "", "", "20101",
               "https://static1.squarespace.com/static/64d54fd8cd03f240928c1cf0/t/657372a51e23d159d4ce9520/"
               "1702064835985/Programmatic_Inventory"]


@pytest.fixture(autouse=True)
def _offline(monkeypatch):
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


def _dooh_workbook(path: Path, client_name: str = "") -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    et.build_dooh_summary(wb, client_name=client_name)
    et.build_dooh_screenlist(wb)
    wb.save(path)
    return path


def _parts(path: Path) -> dict:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


# --- the Screenlist: columns, sample row, table --------------------------------------------------------------------

def test_the_forty_screenlist_columns_are_exactly_the_requested_ones_in_order():
    assert dp.DOOH_SCREENLIST_COLUMNS == USER_COLUMNS
    assert len(dp.DOOH_SCREENLIST_COLUMNS) == 40
    assert dp.DOOH_ALL_COLUMNS == USER_COLUMNS + ["max_avails", "max_spend"]


def test_the_sample_row_matches_the_request_with_numbers_as_numbers_and_codes_as_text():
    sample = dict(zip(dp.DOOH_SCREENLIST_COLUMNS, dp.DOOH_SAMPLE_ROW))
    for name, requested in zip(USER_COLUMNS, USER_SAMPLE):
        value = sample[name]
        if requested in ("TRUE", "FALSE"):
            assert value is (requested == "TRUE"), name            # real booleans
        elif name in ("postal_code", "id", "open_ooh_venue_category_id", "adunit_name", "adunit_id", "asset_name", "venue_name"):
            assert value == requested and isinstance(value, str), name   # codes stay text
        elif re.fullmatch(r"-?\d+(\.\d+)?", requested):
            assert isinstance(value, (int, float)) and not isinstance(value, bool) and float(value) == float(requested), name
        else:
            assert value == requested, name


def test_the_screenlist_sheet_has_the_header_one_sample_row_and_the_calculated_columns(tmp_path):
    wb = openpyxl.load_workbook(_dooh_workbook(tmp_path / "d.xlsx"))
    ws = wb[dp.DOOH_SCREENLIST_SHEET]
    assert [c.value for c in ws[1]] == dp.DOOH_ALL_COLUMNS
    row2 = [c.value for c in ws[2]]
    assert row2[:40] == [None if v == "" else v for v in dp.DOOH_SAMPLE_ROW]      # "" is an empty cell
    assert ws.max_row == 2                                                          # one sample row, nothing else
    assert ws["AO2"].value == "=IFERROR(DOOHScreenlist[[#This Row],[four_week_impressions]]*0.4,0)"
    assert ws["AP2"].value == ("=IFERROR((DOOHScreenlist[[#This Row],[four_week_impressions]]*0.4)"
                               "*DOOHScreenlist[[#This Row],[cpm]]/1000,0)")
    assert ws.freeze_panes == "A2"
    assert ws["V2"].number_format == "General" and ws["W2"].number_format == "#,##0"
    assert ws.column_dimensions["W"].width >= 22                                    # the header text fits


def test_the_screenlist_is_an_excel_table_covering_header_and_sample_with_calculated_columns(tmp_path):
    path = _dooh_workbook(tmp_path / "d.xlsx")
    ws = openpyxl.load_workbook(path)[dp.DOOH_SCREENLIST_SHEET]
    assert list(ws.tables) == ["DOOHScreenlist"]
    assert ws.tables["DOOHScreenlist"].ref == "A1:AP2"
    xml = _parts(path)["xl/tables/table1.xml"].decode("utf-8")
    root = ET.fromstring(xml)
    cols = root.findall(".//m:tableColumn", NS)
    assert [c.get("name") for c in cols] == dp.DOOH_ALL_COLUMNS
    calc = {c.get("name"): c.find("m:calculatedColumnFormula", NS).text for c in cols if c.find("m:calculatedColumnFormula", NS) is not None}
    assert set(calc) == {"max_avails", "max_spend"}                                 # the 40 inventory columns are plain
    assert calc["max_avails"] == dp.table_formula("max_avails")
    assert root.find("m:autoFilter", NS) is not None                                # filter buttons


def test_max_avails_and_max_spend_follow_the_requested_formulas_and_one_named_divisor(monkeypatch):
    assert dp.max_avails_for(2353) == pytest.approx(941.2)
    assert dp.max_spend_for(2353, 32.5) == pytest.approx(2353 * 0.4 * 32.5 / 1000) == pytest.approx(30.589)    # CPM is per thousand
    assert dp.SAMPLE_MAX_AVAILS == pytest.approx(941.2) and dp.SAMPLE_MAX_SPEND == pytest.approx(30.589)
    assert dp.DOOH_SPEND_DIVISOR == 1000 and dp.DOOH_MAX_AVAILS_SHARE == 0.4
    assert "/1000" in dp.table_formula("max_spend") and "*0.4" in dp.table_formula("max_avails")
    # the divisor is the single knob: change it and the formula text follows
    monkeypatch.setattr(dp, "DOOH_SPEND_DIVISOR", 100)
    assert "/100," in dp.table_formula("max_spend") and "/1000" not in dp.table_formula("max_spend")
    assert dp.max_spend_for(2353, 32.5) == pytest.approx(305.89)


# --- the Summary sheet: no header block, no CPM assumption --------------------------------------------------------

def _all_text(ws) -> str:
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def test_the_summary_has_no_proposal_header_comms_details_client_address_or_cpm_assumption(tmp_path):
    ws = openpyxl.load_workbook(_dooh_workbook(tmp_path / "d.xlsx"))[dp.DOOH_SUMMARY_SHEET]
    text = _all_text(ws)
    for forbidden in ("ENTRAVISION COMMUNICATIONS", "Estrella", "Burbank", "Contact", "CUSTOMER BILLING", "Billing",
                      "Attn", "Address", "Media Proposal", "Order Description", "Geo:", "NET.", "CPM Assumption", "Assumption"):
        assert forbidden not in text, forbidden
    # and none of the old placeholder machinery: no formulas at all, no logo, no merged contact block
    assert not [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert not ws._images
    assert ws["B2"].value == "DOOH Summary"
    assert "DOOH Screenlist" in ws["B4"].value
    assert str(ws.merged_cells) == "B4:F4"


def test_the_summary_title_can_carry_the_client_name_but_nothing_else(tmp_path):
    ws = openpyxl.load_workbook(_dooh_workbook(tmp_path / "d.xlsx", client_name="Acme Co"))[dp.DOOH_SUMMARY_SHEET]
    assert ws["B2"].value == "DOOH Summary — Acme Co"


def test_the_summary_cells_hold_the_pivot_rendering_of_the_sample_row(tmp_path):
    ws = openpyxl.load_workbook(_dooh_workbook(tmp_path / "d.xlsx"))[dp.DOOH_SUMMARY_SHEET]
    r = dp.PIVOT_FIRST_ROW
    assert [c.value for c in ws[r]][1:6] == ["Venue Type", "Network ", "Number of Screens", "Max. Avails", "Max. Spend"]
    assert [c.value for c in ws[r + 1]][1:6] == ["Retail|Gas Stations|Pump", "GSTV", 1, pytest.approx(941.2), pytest.approx(30.589)]
    assert ws.cell(row=r + 1, column=6).number_format == '"$"#,##0.00'
    assert ws.cell(row=r + len(dp.pivot_preview_rows()) - 1, column=2).value == "Grand Total"
    # one row height for the whole sheet — evenly spaced rows however far the pivot grows
    assert ws.sheet_format.defaultRowHeight == dp.SUMMARY_ROW_HEIGHT
    assert {ws.row_dimensions[i].height for i in range(r, r + 40)} == {dp.SUMMARY_ROW_HEIGHT}
    assert ws.sheet_view.showGridLines is False


# --- the injected pivot: package structure ---------------------------------------------------------------------------

def _pivot_package(tmp_path) -> tuple[Path, dict]:
    path = _dooh_workbook(tmp_path / "d.xlsx")
    assert dp.add_dooh_pivot(path) is True
    return path, _parts(path)


def test_every_part_of_the_injected_workbook_is_well_formed_xml(tmp_path):
    _, parts = _pivot_package(tmp_path)
    for name, data in parts.items():
        if name.endswith((".xml", ".rels", ".vml")) or name == "[Content_Types].xml":
            ET.fromstring(data)                                                    # raises on bad XML


def test_the_pivot_parts_are_registered_in_content_types_and_rels_and_all_targets_exist(tmp_path):
    _, parts = _pivot_package(tmp_path)
    ct = ET.fromstring(parts["[Content_Types].xml"])
    override_list = [o for o in ct if o.tag.endswith("Override")]
    overrides = {o.get("PartName"): o.get("ContentType") for o in override_list}
    for part, ctype in (("/xl/pivotCache/pivotCacheDefinition1.xml", "pivotCacheDefinition"),
                        ("/xl/pivotCache/pivotCacheRecords1.xml", "pivotCacheRecords"),
                        ("/xl/pivotTables/pivotTable1.xml", "pivotTable")):
        assert part.lstrip("/") in parts and overrides[part].endswith(f"{ctype}+xml"), part
    # every relationship target in the package resolves to a part that exists
    for name, data in parts.items():
        if not name.endswith(".rels"):
            continue
        base = os.path.dirname(os.path.dirname(name))                              # the folder the .rels describes
        for rel in ET.fromstring(data):
            target = rel.get("Target")
            resolved = target.lstrip("/") if target.startswith("/") else os.path.normpath(os.path.join(base, target)).replace("\\", "/")
            assert resolved in parts, (name, target, resolved)
    # no duplicated override, no duplicated rel id
    assert len(overrides) == len(override_list)
    for name, data in parts.items():
        if name.endswith(".rels"):
            ids = [r.get("Id") for r in ET.fromstring(data)]
            assert len(ids) == len(set(ids)), name


def test_the_workbook_points_at_the_cache_and_the_pivot_table_points_at_the_same_cache_id(tmp_path):
    _, parts = _pivot_package(tmp_path)
    wb_root = ET.fromstring(parts["xl/workbook.xml"])
    caches = wb_root.findall("m:pivotCaches/m:pivotCache", NS)
    assert len(caches) == 1
    cache_id = caches[0].get("cacheId")
    rid = caches[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    rels = {r.get("Id"): r for r in ET.fromstring(parts["xl/_rels/workbook.xml.rels"])}
    assert rels[rid].get("Type").endswith("/pivotCacheDefinition")
    assert rels[rid].get("Target") == "pivotCache/pivotCacheDefinition1.xml"
    pivot = ET.fromstring(parts["xl/pivotTables/pivotTable1.xml"])
    assert pivot.get("cacheId") == cache_id
    # CT_Workbook element order: pivotCaches comes after calcPr (Excel is strict about it)
    order = [c.tag.replace(M, "") for c in wb_root]
    assert order.index("pivotCaches") > order.index("calcPr")


def test_the_summary_sheet_has_exactly_one_pivot_table_and_it_belongs_to_the_summary(tmp_path):
    _, parts = _pivot_package(tmp_path)
    wb_root = ET.fromstring(parts["xl/workbook.xml"])
    rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(parts["xl/_rels/workbook.xml.rels"])}
    sheet_part = {s.get("name"): rels[s.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")].lstrip("/")
                  for s in wb_root.findall("m:sheets/m:sheet", NS)}
    summary_rels = ET.fromstring(parts["xl/worksheets/_rels/" + os.path.basename(sheet_part[dp.DOOH_SUMMARY_SHEET]) + ".rels"])
    pivots = [r for r in summary_rels if r.get("Type").endswith("/pivotTable")]
    assert len(pivots) == 1
    pivot_parts = [n for n in parts if n.startswith("xl/pivotTables/pivotTable") and n.endswith(".xml")]
    assert len(pivot_parts) == 1


def test_the_pivot_definition_is_venue_type_then_network_with_the_three_requested_values(tmp_path):
    _, parts = _pivot_package(tmp_path)
    pivot = ET.fromstring(parts["xl/pivotTables/pivotTable1.xml"])
    index = {name: i for i, name in enumerate(dp.DOOH_ALL_COLUMNS)}
    assert pivot.find("m:location", NS).get("ref") == dp.pivot_ref() == "B6:F9"
    assert [f.get("x") for f in pivot.findall("m:rowFields/m:field", NS)] == [str(index["venue_type"]), str(index["publisher"])]
    fields = pivot.findall("m:pivotFields/m:pivotField", NS)
    assert len(fields) == 42
    assert fields[index["venue_type"]].get("name") == "Venue Type" and fields[index["venue_type"]].get("axis") == "axisRow"
    network = fields[index["publisher"]]
    assert network.get("axis") == "axisRow"
    assert network.get("name") == "Network " and network.get("name").strip() == "Network"   # see PIVOT_ROW_FIELDS
    data = pivot.findall("m:dataFields/m:dataField", NS)
    assert [(d.get("name"), d.get("fld"), d.get("subtotal")) for d in data] == [
        ("Number of Screens", str(index["id"]), "count"),                    # a COUNT of id
        ("Max. Avails", str(index["max_avails"]), None),                     # sums of the per-row helper columns
        ("Max. Spend", str(index["max_spend"]), None),
    ]
    assert [fields[index[c]].get("dataField") for c in ("id", "max_avails", "max_spend")] == ["1", "1", "1"]
    # tabular layout (compact off), grand totals on (no rowGrandTotals/colGrandTotals="0")
    assert pivot.get("compact") == "0" and pivot.get("rowGrandTotals") is None and pivot.get("colGrandTotals") is None
    assert pivot.find("m:pivotTableStyleInfo", NS).get("name") == dp.PIVOT_STYLE_NAME
    # row / column items match the rendered shape: one data line, its venue-type subtotal, the grand total
    assert [i.get("t") for i in pivot.findall("m:rowItems/m:i", NS)] == [None, "default", "grand"]
    assert len(pivot.findall("m:colItems/m:i", NS)) == 3


def test_turning_the_venue_type_subtotals_off_changes_the_pivot_shape_consistently(tmp_path, monkeypatch):
    """PIVOT_SHOW_SUBTOTALS is a one-line switch; the definition, the rendered cells and the range must follow it
    (checked in Excel too: both shapes refresh to the right totals)."""
    monkeypatch.setattr(dp, "PIVOT_SHOW_SUBTOTALS", False)
    assert [kind for kind, _ in dp.pivot_preview_rows()] == ["header", "data", "grand"]
    assert dp.pivot_ref() == "B6:F8"
    path, parts = _pivot_package(tmp_path)
    pivot = ET.fromstring(parts["xl/pivotTables/pivotTable1.xml"])
    assert pivot.find("m:location", NS).get("ref") == "B6:F8"
    assert [i.get("t") for i in pivot.findall("m:rowItems/m:i", NS)] == [None, "grand"]
    venue = pivot.findall("m:pivotFields/m:pivotField", NS)[dp.DOOH_ALL_COLUMNS.index("venue_type")]
    assert venue.get("defaultSubtotal") == "0" and len(venue.findall("m:items/m:item", NS)) == 1
    ws = openpyxl.load_workbook(path)[dp.DOOH_SUMMARY_SHEET]
    assert [ws.cell(row=r, column=2).value for r in (6, 7, 8)] == ["Venue Type", "Retail|Gas Stations|Pump", "Grand Total"]


def test_a_network_caption_never_equals_a_source_column_name():
    """Excel refuses a pivot caption equal to a source field name (case-insensitive) — 'network' IS a column."""
    source = {c.lower() for c in dp.DOOH_ALL_COLUMNS}
    captions = [c for _, c in dp.PIVOT_ROW_FIELDS] + [f[2] for f in dp.PIVOT_DATA_FIELDS]
    for caption in captions:
        assert caption.lower() not in source, caption


def test_the_cache_reads_the_screenlist_table_by_name_refreshes_on_load_and_describes_all_42_columns(tmp_path):
    _, parts = _pivot_package(tmp_path)
    cache = ET.fromstring(parts["xl/pivotCache/pivotCacheDefinition1.xml"])
    assert cache.get("refreshOnLoad") == "1"
    source = cache.find("m:cacheSource", NS)
    assert source.get("type") == "worksheet" and source.find("m:worksheetSource", NS).get("name") == "DOOHScreenlist"
    names = [f.get("name") for f in cache.findall("m:cacheFields/m:cacheField", NS)]
    assert names == dp.DOOH_ALL_COLUMNS
    # the cache's records part is wired through its own rels
    rels = ET.fromstring(parts["xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels"])
    assert [r.get("Target") for r in rels] == ["pivotCacheRecords1.xml"]
    assert cache.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id") == rels[0].get("Id")
    records = ET.fromstring(parts["xl/pivotCache/pivotCacheRecords1.xml"])
    rows = records.findall("m:r", NS)
    assert len(rows) == 1 and len(rows[0]) == 42 and records.get("count") == "1"
    row = list(rows[0])
    assert row[dp.DOOH_ALL_COLUMNS.index("max_avails")].get("v") == "941.2"
    assert row[dp.DOOH_ALL_COLUMNS.index("max_spend")].get("v") == "30.589"
    assert row[dp.DOOH_ALL_COLUMNS.index("supports_dynamic_creative")].tag == M + "b"
    assert row[dp.DOOH_ALL_COLUMNS.index("restrictions")].tag == M + "m"            # blank cell


def test_the_pivot_number_formats_exist_in_styles_and_the_custom_style_is_consistent(tmp_path):
    _, parts = _pivot_package(tmp_path)
    styles = ET.fromstring(parts["xl/styles.xml"])
    formats = {n.get("numFmtId"): n.get("formatCode") for n in styles.findall("m:numFmts/m:numFmt", NS)}
    pivot = ET.fromstring(parts["xl/pivotTables/pivotTable1.xml"])
    for d in pivot.findall("m:dataFields/m:dataField", NS):
        fid = d.get("numFmtId")
        assert fid == "3" or formats.get(fid) in ('"$"#,##0.00',), (d.get("name"), fid)
    dxfs = styles.find("m:dxfs", NS)
    assert int(dxfs.get("count")) == len(dxfs.findall("m:dxf", NS))
    style = next(s for s in styles.findall("m:tableStyles/m:tableStyle", NS) if s.get("name") == dp.PIVOT_STYLE_NAME)
    assert style.get("table") == "0" and int(style.get("count")) == len(style)
    assert {e.get("type") for e in style} == {"wholeTable", "headerRow", "firstSubtotalRow", "totalRow"}
    for e in style:
        assert 0 <= int(e.get("dxfId")) < int(dxfs.get("count"))
    ts = styles.find("m:tableStyles", NS)
    assert int(ts.get("count")) == len(ts.findall("m:tableStyle", NS))


def test_existing_differential_formats_are_kept_and_the_new_ones_are_offset_past_them():
    styles = ('<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><numFmts count="1">'
              '<numFmt numFmtId="164" formatCode="0.0"/></numFmts><dxfs count="2"><dxf><font><b/></font></dxf>'
              '<dxf><font><i/></font></dxf></dxfs><tableStyles count="0" defaultTableStyle="TableStyleMedium9" '
              'defaultPivotStyle="PivotStyleLight16"/></styleSheet>')
    styles, money_id = dp._money_format_id(styles)
    styles = dp._add_pivot_style(styles)
    root = ET.fromstring(styles)
    assert money_id == 165 and root.find("m:numFmts", NS).get("count") == "2"
    dxfs = root.find("m:dxfs", NS)
    assert dxfs.get("count") == "6" and len(dxfs) == 6
    ids = sorted(int(e.get("dxfId")) for e in root.findall("m:tableStyles/m:tableStyle/m:tableStyleElement", NS))
    assert ids == [2, 3, 4, 5]                                                       # the two existing ones untouched
    assert dp._add_pivot_style(styles) == styles                                     # idempotent


def test_the_helper_formulas_carry_their_cached_values_and_number_stored_as_text_is_not_flagged(tmp_path):
    _, parts = _pivot_package(tmp_path)
    sheet = parts["xl/worksheets/sheet2.xml"].decode("utf-8")
    assert re.search(r'<c r="AO2"[^>]*><f>[^<]*</f><v>941\.2</v></c>', sheet)
    assert re.search(r'<c r="AP2"[^>]*><f>[^<]*</f><v>30\.589</v></c>', sheet)
    assert '<ignoredError sqref="A2:AP1048576" numberStoredAsText="1"/>' in sheet
    root = ET.fromstring(sheet)
    order = [c.tag.replace(M, "") for c in root]
    assert order.index("ignoredErrors") < order.index("tableParts")                  # schema order
    cached = openpyxl.load_workbook(tmp_path / "d.xlsx", data_only=True)[dp.DOOH_SCREENLIST_SHEET]
    assert cached["AO2"].value == pytest.approx(941.2) and cached["AP2"].value == pytest.approx(30.589)


def test_openpyxl_can_still_read_the_injected_workbook_including_its_pivot(tmp_path):
    path, _ = _pivot_package(tmp_path)
    ws = openpyxl.load_workbook(path)[dp.DOOH_SUMMARY_SHEET]
    assert len(ws._pivots) == 1 and ws._pivots[0].location.ref == "B6:F9"
    assert [f.name for f in ws._pivots[0].dataFields] == ["Number of Screens", "Max. Avails", "Max. Spend"]


# --- add_dooh_pivot: idempotent, safe, atomic ----------------------------------------------------------------------

def test_adding_the_pivot_twice_changes_nothing_the_second_time(tmp_path):
    path = _dooh_workbook(tmp_path / "d.xlsx")
    assert dp.add_dooh_pivot(path) is True
    once = path.read_bytes()
    assert dp.add_dooh_pivot(path) is False
    assert path.read_bytes() == once


def test_a_workbook_without_the_dooh_sheets_is_left_byte_for_byte_alone(tmp_path):
    wb = Workbook()
    wb.active.title = "Proposal A"
    path = tmp_path / "plain.xlsx"
    wb.save(path)
    before = path.read_bytes()
    assert dp.add_dooh_pivot(path) is False
    assert path.read_bytes() == before
    # a Screenlist alone (no Summary) has nowhere to put a pivot either
    wb = Workbook()
    wb.remove(wb.active)
    et.build_dooh_screenlist(wb)
    only = tmp_path / "only.xlsx"
    wb.save(only)
    before = only.read_bytes()
    assert dp.add_dooh_pivot(only) is False and only.read_bytes() == before


def test_a_screenlist_that_is_not_the_expected_table_raises_and_leaves_the_file_untouched(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    et.build_dooh_summary(wb)
    ws = wb.create_sheet(dp.DOOH_SCREENLIST_SHEET)
    ws.append(["just", "some", "cells"])                                             # no table at all
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    before = path.read_bytes()
    with pytest.raises(dp.DoohPivotError):
        dp.add_dooh_pivot(path)
    assert path.read_bytes() == before
    assert [p.name for p in tmp_path.iterdir()] == ["bad.xlsx"]                      # no temp file left behind


def test_a_screenlist_table_with_other_columns_is_refused(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    et.build_dooh_summary(wb)
    from openpyxl.worksheet.table import Table
    ws = wb.create_sheet(dp.DOOH_SCREENLIST_SHEET)
    ws.append(["a", "b"])
    ws.append([1, 2])
    ws.add_table(Table(displayName="DOOHScreenlist", ref="A1:B2"))
    path = tmp_path / "cols.xlsx"
    wb.save(path)
    with pytest.raises(dp.DoohPivotError):
        dp.add_dooh_pivot(path)


def test_the_pivot_survives_alongside_other_sheets_with_images_and_conditional_formats(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    products = et.effective_catalog()[:3]
    et.build_proposal_a(wb, products, with_sections=False)         # carries the logo image
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font
    wb["Proposal A"].conditional_formatting.add("Z1:Z5", CellIsRule(operator="greaterThan", formula=["1"], font=Font(bold=True)))
    et.build_dooh_summary(wb)
    et.build_dooh_screenlist(wb)
    et.build_process_faqs(wb)
    path = tmp_path / "full.xlsx"
    wb.save(path)
    assert dp.add_dooh_pivot(path) is True
    parts = _parts(path)
    for name, data in parts.items():
        if name.endswith((".xml", ".rels")):
            ET.fromstring(data)
    styles = ET.fromstring(parts["xl/styles.xml"])
    dxfs = styles.find("m:dxfs", NS)
    assert int(dxfs.get("count")) == len(dxfs.findall("m:dxf", NS)) == 5          # the sheet's own one + the 4 of the pivot style
    style = next(s for s in styles.findall("m:tableStyles/m:tableStyle", NS) if s.get("name") == dp.PIVOT_STYLE_NAME)
    assert sorted(int(e.get("dxfId")) for e in style) == [1, 2, 3, 4]              # offset past the existing one
    loaded = openpyxl.load_workbook(path)
    assert loaded.sheetnames == ["Proposal A", "DOOH Summary", "DOOH Screenlist", "Process FAQs"]
    assert any(n.startswith("xl/media/") for n in parts)           # the other sheet's logo is still there


def test_the_template_builder_also_ships_the_pivot(tmp_path):
    out = tmp_path / "template.xlsx"
    et.build_workbook(out, families_to_include=[et.families()[0]])
    parts = _parts(out)
    assert "xl/pivotTables/pivotTable1.xml" in parts
    assert openpyxl.load_workbook(out)["DOOH Summary"]._pivots


# --- tab keys ------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("given, expected", [
    ({"dooh_summary": True}, {"dooh": True}),
    ({"dooh_screenlist": True}, {"dooh": True}),
    ({"dooh_summary": True, "dooh_screenlist": False}, {"dooh": True}),               # either true -> both sheets
    ({"dooh_summary": False, "dooh_screenlist": False}, {"dooh": False}),
    ({"dooh_summary": False}, {"dooh": False}),
    ({"dooh": False, "dooh_summary": True}, {"dooh": False}),                          # an explicit new key wins
    ({"dooh": True, "dooh_screenlist": False}, {"dooh": True}),
    ({"net": True, "gross": False, "_v": 2, "dooh_summary": True}, {"net": True, "gross": False, "_v": 2, "dooh": True}),
    ({"net": False}, {"net": False}),
    ({}, {}),
    (None, {}),
])
def test_normalize_tab_keys(given, expected):
    original = dict(given) if given else given
    assert normalize_tab_keys(given) == expected
    assert given == original                                                            # the input is not mutated


def test_legacy_dooh_keys_in_force_tabs_still_build_both_sheets(tmp_path):
    li = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="a")
    for forced in ({"dooh_summary": True}, {"dooh_screenlist": True}, {"dooh_summary": True, "dooh_screenlist": True},
                   {"dooh": True}):
        out = tmp_path / f"f{len(list(tmp_path.iterdir()))}.xlsx"
        summary = generate_proposal(_req(), [li], out, force_tabs=forced)
        assert "DOOH Summary" in summary["tabs_built"] and "DOOH Screenlist" in summary["tabs_built"], forced
        assert openpyxl.load_workbook(out)["DOOH Summary"]._pivots


def test_dooh_forced_off_builds_neither_sheet(tmp_path):
    li = LineItem(product_name="Digital Out of Home", monthly_budget=1000.0, months=3, id="a")
    out = tmp_path / "off.xlsx"
    summary = generate_proposal(_req(), [li], out, force_tabs={"dooh": False})
    assert not [t for t in summary["tabs_built"] if "DOOH" in t]
    summary = generate_proposal(_req(), [li], out, force_tabs={"dooh_summary": False, "dooh_screenlist": False})
    assert not [t for t in summary["tabs_built"] if "DOOH" in t]


def test_only_dooh_forced_on_is_a_valid_override_not_an_all_off_one(tmp_path):
    li = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="a")
    out = tmp_path / "only.xlsx"
    summary = generate_proposal(_req(), [li], out, force_tabs={"net": False, "wsections": False, "gross": False,
                                                              "avails_only": False, "dooh": True})
    assert not any("switched off" in w for w in summary["warnings"])
    assert summary["tabs_built"] == ["DOOH Summary", "DOOH Screenlist", "Process FAQs"]


# --- classification by product family -----------------------------------------------------------------------------

def test_the_dooh_product_is_what_the_catalog_says_it_is():
    from app.catalog import by_name
    assert by_name("Digital Out of Home").family == "DOOH"


def test_a_dooh_product_switches_the_dooh_deliverables_on_for_every_request_type():
    for request_type in ("New Business", "Renewal Proposal", "Full Presentation",
                         "Avails / Estimates Only (I don't need a proposal right now)"):
        tabs = classify_output_tabs(request_type, ["Digital Out of Home", "Search - SEM"], False)
        assert tabs["dooh"] is True, request_type
    avails = classify_output_tabs("Avails / Estimates Only", ["Digital Out of Home"], False)
    assert avails == {"net": False, "wsections": False, "gross": False, "avails_only": True, "dooh": True}


def test_no_dooh_product_no_dooh_deliverables_and_the_old_keys_are_gone_from_the_classification():
    tabs = classify_output_tabs("New Business", ["Search - SEM", "Meta Ads"], True)
    assert tabs == {"net": True, "wsections": False, "gross": True, "avails_only": False, "dooh": False}
    assert "dooh_summary" not in tabs and "dooh_screenlist" not in tabs


def test_detection_resolves_names_through_the_catalog_and_judges_by_family(monkeypatch):
    class Fake:
        def __init__(self, name, family):
            self.name, self.family = name, family

    products = {"My Screens": Fake("My Screens", "DOOH"), "DOOH Lookalike Display": Fake("DOOH Lookalike Display", "Display")}
    monkeypatch.setattr("app.catalog.by_name", lambda n: products.get(n))
    # resolves to a DOOH-family product whose NAME says nothing about DOOH -> on
    assert classify_output_tabs("New Business", ["My Screens"], False)["dooh"] is True
    # the name says DOOH but the product resolves to another family -> the family wins
    assert classify_output_tabs("New Business", ["DOOH Lookalike Display"], False)["dooh"] is False


def test_names_that_do_not_resolve_fall_back_to_the_old_name_check(monkeypatch):
    monkeypatch.setattr("app.catalog.by_name", lambda n: None)
    assert classify_output_tabs("New Business", ["Mystery DOOH Network"], False)["dooh"] is True
    assert classify_output_tabs("New Business", ["Some Out-of-Home thing"], False)["dooh"] is True
    assert classify_output_tabs("New Business", ["Digital Out of Home"], False)["dooh"] is True
    assert classify_output_tabs("New Business", ["Search - SEM"], False)["dooh"] is False
    assert classify_output_tabs("New Business", [], False)["dooh"] is False


def test_a_catalog_failure_cannot_break_tab_classification(monkeypatch):
    def boom(_name):
        raise RuntimeError("database is down")
    monkeypatch.setattr("app.catalog.by_name", boom)
    assert classify_output_tabs("New Business", ["Search - SEM"], False)["net"] is True
    assert classify_output_tabs("New Business", ["Mystery DOOH Network"], False)["dooh"] is True


def test_the_suggest_tabs_endpoint_and_parse_suggestions_carry_the_single_dooh_key():
    import app.main as m
    out = asyncio.run(m.suggest_tabs(m.SuggestTabsRequest(request={"request_type": "New Business"},
                                                          product_names=["Digital Out of Home"])))["tabs"]
    assert out["dooh"] is True and "dooh_summary" not in out
    plain = asyncio.run(m.suggest_tabs(m.SuggestTabsRequest(request={"request_type": "New Business"},
                                                            product_names=["Search - SEM"])))["tabs"]
    assert plain["dooh"] is False


# --- end to end: a DOOH line through the real generator ------------------------------------------------------------

def _req(**kw):
    base = dict(client_name="Acme", request_type="New Business", start_date="2026-10-01", end_date="2026-12-31",
                total_months=3)
    base.update(kw)
    return ProposalRequest(**base)


def test_a_dooh_product_in_the_plan_produces_the_deliverables_next_to_the_other_tabs(tmp_path):
    names = ["Digital Out of Home", "eDigital Network Display - Standard IAB", "Search - SEM"]
    lines = [LineItem(product_name=n, monthly_budget=2000.0, months=3, id=f"l{i}") for i, n in enumerate(names)]
    out = tmp_path / "e2e.xlsx"
    summary = generate_proposal(_req(agency_fee=0.15, client_name="Fronteras del Norte"), lines, out)
    assert summary["tabs_built"] == ["Proposal A", "Proposal A (Gross)", "DOOH Summary", "DOOH Screenlist", "Process FAQs"]
    assert not summary["warnings"]
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Proposal A", "Proposal A (Gross)", "DOOH Summary", "DOOH Screenlist", "Process FAQs"]
    assert wb["DOOH Summary"]["B2"].value == "DOOH Summary — Fronteras del Norte"
    assert wb["DOOH Summary"]["A1"].value is None                     # the old stray A1 title stamp is gone
    assert len(wb["DOOH Summary"]._pivots) == 1
    assert wb["DOOH Screenlist"].tables["DOOHScreenlist"].ref == "A1:AP2"


def test_a_dooh_line_in_an_option_other_than_the_first_still_gets_the_deliverables(tmp_path):
    a = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="a")
    c = LineItem(product_name="Digital Out of Home", monthly_budget=1000.0, months=3, id="c")
    tiers = [{"label": "A", "line_items": [a], "avails_data": {}}, {"label": "B", "line_items": [c], "avails_data": {}}]
    summary = generate_proposal(_req(), [], tmp_path / "t.xlsx", tiers=tiers)
    assert "DOOH Summary" in summary["tabs_built"]


def test_no_tier_can_take_a_dooh_sheet_name(tmp_path):
    a = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="a")
    c = LineItem(product_name="Digital Out of Home", monthly_budget=1000.0, months=3, id="c")
    tiers = [{"label": "A", "name": "DOOH Summary", "line_items": [a], "avails_data": {}},
             {"label": "B", "name": "DOOH Screenlist", "line_items": [c], "avails_data": {}}]
    out = tmp_path / "names.xlsx"
    generate_proposal(_req(), [], out, tiers=tiers)
    wb = openpyxl.load_workbook(out)
    assert wb["DOOH Summary"]["B6"].value == "Venue Type" and len(wb["DOOH Summary"]._pivots) == 1
    assert wb["DOOH Screenlist"].tables["DOOHScreenlist"].ref == "A1:AP2"


def test_a_failing_pivot_injection_keeps_the_workbook_and_warns(tmp_path, monkeypatch):
    def boom(_path):
        raise dp.DoohPivotError("simulated")
    monkeypatch.setattr("app.services.dooh_pivot.add_dooh_pivot", boom)
    li = LineItem(product_name="Digital Out of Home", monthly_budget=1000.0, months=3, id="a")
    out = tmp_path / "warn.xlsx"
    summary = generate_proposal(_req(), [li], out)
    assert any("pivot table could not be added" in w for w in summary["warnings"])
    wb = openpyxl.load_workbook(out)
    assert "DOOH Screenlist" in wb.sheetnames and not wb["DOOH Summary"]._pivots


# --- reopening an older draft --------------------------------------------------------------------------------------

def test_reopening_a_saved_draft_with_the_old_dooh_keys_returns_the_new_key(monkeypatch):
    import app.main as m
    saved = {"proposal_title": "T", "status": "draft",
             "reopen_state": {"request": {}, "force_tabs": {"dooh_screenlist": True, "net": False, "_v": 2}}}
    monkeypatch.setattr(m, "_get_proposal_metadata", lambda _id: saved)
    out = asyncio.run(m.reopen_proposal("x"))
    assert out["force_tabs"] == {"net": False, "_v": 2, "dooh": True}


def test_saving_a_draft_stores_the_single_key():
    import app.main as m
    state = m._build_reopen_state(request_dict={}, line_items=[], tiers=None, avails_data=None, strategy_brief=None,
                                  roadblocks=None, force_tabs={"dooh_summary": True, "_v": 2}, addons=[],
                                  raw_notion_text=None, enrichment=None, time_unit="month",
                                  monthly_distribution_mode="even", campaign_name_override=None, wizard_step=3)
    assert state["force_tabs"] == {"_v": 2, "dooh": True}
    empty = m._build_reopen_state(request_dict={}, line_items=[], tiers=None, avails_data=None, strategy_brief=None,
                                  roadblocks=None, force_tabs={}, addons=[], raw_notion_text=None, enrichment=None,
                                  time_unit="month", monthly_distribution_mode="even", campaign_name_override=None,
                                  wizard_step=3)
    assert empty["force_tabs"] == {}


# --- the UI ----------------------------------------------------------------------------------------------------------

def test_step_08_offers_one_dooh_deliverables_box_instead_of_two():
    html = INDEX_HTML.read_text(encoding="utf-8")
    boxes = re.findall(r'data-tab="([a-z_]+)"', html)
    assert boxes == ["net", "wsections", "gross", "avails_only", "dooh"]
    assert "DOOH Deliverables" in html
    assert "dooh_summary" not in html and "dooh_screenlist" not in html


def _node_run(script: str):
    node = shutil.which("node")
    if not node:
        pytest.skip("node not installed")
    return subprocess.run([node, "-e", script], capture_output=True, text=True, check=True).stdout


def _js_function(name: str) -> str:
    src = APP_JS.read_text(encoding="utf-8")
    start = src.index(f"function {name}(")
    depth, i = 0, src.index("{", start)
    for j in range(i, len(src)):
        depth += src[j] == "{"
        depth -= src[j] == "}"
        if depth == 0:
            return src[start:j + 1]
    raise AssertionError(name)


def test_the_client_folds_the_old_dooh_keys_the_same_way_the_server_does():
    cases = [{"dooh_summary": True}, {"dooh_screenlist": True}, {"dooh_summary": False, "dooh_screenlist": False},
             {"dooh": False, "dooh_summary": True}, {"net": True, "dooh_summary": True}, {}, None]
    script = _js_function("_normalizeTabKeys") + f"\nconsole.log(JSON.stringify({json.dumps(cases)}.map(_normalizeTabKeys)));"
    client = json.loads(_node_run(script))
    assert client == [normalize_tab_keys(c) for c in cases]


def test_the_client_names_the_single_dooh_tab_and_no_longer_the_two_old_ones():
    src = APP_JS.read_text(encoding="utf-8")
    pretty = _js_function("prettyTabName")
    assert 'dooh: "DOOH Deliverables"' in pretty
    assert "dooh_summary:" not in pretty and "dooh_screenlist:" not in pretty
    assert 'const allTabKeys = ["net", "wsections", "gross", "avails_only", "dooh"];' in src
