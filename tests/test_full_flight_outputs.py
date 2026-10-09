"""
Full Flight billing period (time_unit == "full_flight") in the NON-Excel outputs: the PowerPoint deck, the AI
enrichment prompt, the Strategy Brief prompt and its Word document.

Full Flight = the whole flight is billed as ONE period. A line's dollars are the flat whole-flight figure (months == 1,
monthly_budget == the flight total), so nothing in these outputs may read as a monthly rate or be multiplied by a month
count. Week / month / quarter must keep rendering EXACTLY as before — pinned here by a semantic dump hash captured from
the pre-feature builder (see _GOLDEN) plus explicit layout assertions.
"""
from __future__ import annotations

import asyncio
import difflib
import hashlib
import json
import logging
import os
from types import SimpleNamespace

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import pytest

pptx = pytest.importorskip("pptx")
docx = pytest.importorskip("docx")

from pptx.enum.text import PP_ALIGN

from app.catalog import Product
from app.services import ai_enricher, docx_builder, period_copy, pptx_builder
from app.services import strategy_brief as sb
from app.services.monthly_allocation import FULL_FLIGHT, TIME_UNITS
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import LineItem

NBSP = " "


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.catalog.resolve_product_alias", lambda name: None)
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


def _product(name, *, family="Display", description="Banner ads shown across a network of websites."):
    return Product(
        family=family, name=name, short_label=name, proposal_description=description, sizes="", buying_model="CPM",
        base_rate=10.0, estimated_impressions=False, discloses_impressions=True, minimum_spend=0.0,
        minimum_flight_days=(0, 0), sla_data_days=None, sla_creative_days=None, sla_activate_days=None,
        sla_total_days=None, media_allocation_pct=1.0, margin_upper=0.5, margin_lower=0.5, tech_platform="",
        wide_orbit_code="")


# ===========================================================================================================
# period_copy — the one shared wording helper
# ===========================================================================================================

def test_period_copy_resolves_every_known_unit_to_itself_and_forgives_case_and_whitespace(caplog):
    with caplog.at_level(logging.WARNING):
        for unit in TIME_UNITS:
            assert period_copy.resolve_time_unit(unit) == unit
        assert period_copy.resolve_time_unit(" Full_Flight ") == FULL_FLIGHT
        assert period_copy.resolve_time_unit(None) == "month" and period_copy.resolve_time_unit("") == "month"
    assert not caplog.records                      # absent / well-formed values never warn


def test_period_copy_warns_on_an_unknown_unit_and_falls_back_to_month_without_raising(caplog):
    with caplog.at_level(logging.WARNING, logger="app.services.period_copy"):
        assert period_copy.resolve_time_unit("fortnight", where="unit test") == "month"
    assert len(caplog.records) == 1
    message = caplog.records[0].getMessage()
    assert "fortnight" in message and "unit test" in message and FULL_FLIGHT in message


def test_period_copy_pluralizes_month_counts():
    assert period_copy.months_phrase(1) == "1 month"
    assert period_copy.months_phrase(2) == "2 months" and period_copy.months_phrase(12) == "12 months"


def test_budget_sentence_per_unit():
    # month / week / quarter: the Step 02 intake figure has always been a MONTHLY one
    for unit in ("month", "week", "quarter"):
        assert period_copy.budget_sentence(5000, 3, unit) == "Budget: $5,000/month × 3 months = $15,000 total flight"
    assert period_copy.budget_sentence(5000, 1, "month") == "Budget: $5,000/month × 1 month = $5,000 total flight"
    # full flight: ONE period and nothing multiplied — the Step 02 budget is the flight's budget exactly as typed (Step 04
    # seeds it unchanged), so the sentence states that flat total whatever the flight's length or month count
    assert period_copy.budget_sentence(6375, 2, "full_flight") == "Budget: $6,375 total flight (single billing period)"
    assert period_copy.budget_sentence(5000, 3, "full_flight") == "Budget: $5,000 total flight (single billing period)"
    assert period_copy.budget_sentence(5000, 1, "full_flight") == "Budget: $5,000 total flight (single billing period)"
    for ignored in (0, None, "", -2, "x"):          # the month count plays no part
        assert period_copy.budget_sentence(5000, ignored, "full_flight") == "Budget: $5,000 total flight (single billing period)", ignored
    assert not hasattr(period_copy, "flight_months")


def test_budget_sentence_is_only_stated_when_the_figures_exist():
    assert period_copy.has_budget_sentence(5000, 3, "month") and not period_copy.has_budget_sentence(5000, 0, "month")
    assert not period_copy.has_budget_sentence(0, 3, "month") and not period_copy.has_budget_sentence(None, None, "month")
    assert period_copy.has_budget_sentence(15000, 0, "full_flight")        # a missing month count just means x1
    assert not period_copy.has_budget_sentence(0, 2, "full_flight")


# ===========================================================================================================
# PowerPoint
# ===========================================================================================================

def _req(**kw):
    base = dict(client_name="Acme", start_date="2026-11-24", end_date="2026-12-31")
    base.update(kw)
    return ProposalRequest(**base)


def _lines(budgets, *, months=1, target=None, added_value=(), prefix="Product"):
    prods = [_product(f"{prefix} {i}") for i in range(len(budgets))]
    lis = [LineItem(product_name=p.name, monthly_budget=0.0 if i in added_value else b, months=months, id=f"l{i}",
                    target_override=target, is_added_value=i in added_value)
           for i, (p, b) in enumerate(zip(prods, budgets))]
    return prods, lis


def _build(tmp_path, tiers, *, time_unit="month", gross=False, banners=None, request=None, name="d.pptx"):
    out = tmp_path / name
    ok = pptx_builder.build_signature_deck(
        request or _req(agency_fee=0.15 if gross else None), tiers, out, gross=gross, proposal_title="T",
        banners=banners, time_unit=time_unit)
    assert ok
    return pptx.Presentation(str(out))


def _tier(budgets, name="Option A", **kw):
    prods, lis = _lines(budgets, **kw)
    return {"name": name, "products": prods, "line_items": lis}


def _tables(prs):
    return [sh.table for sl in prs.slides for sh in sl.shapes if sh.has_table]


def _text(prs):
    return "\n".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)


def _cells(table):
    return [[table.cell(r, c).text for c in range(len(table.columns))] for r in range(len(table.rows))]


def _align(table, r, c):
    return table.cell(r, c).text_frame.paragraphs[0].alignment


def test_pptx_full_flight_table_is_product_target_flight_budget(tmp_path):
    prs = _build(tmp_path, [_tier([5000.0, 7750.0])], time_unit="full_flight")
    (table,) = _tables(prs)
    assert len(table.columns) == 3
    assert _cells(table) == [
        ["Product", "Target", "Flight Budget"],
        ["Product 0\nEst. 500,000 impressions", "TBD", "$5,000"],
        ["Product 1\nEst. 775,000 impressions", "TBD", "$7,750"],
        ["TOTAL", NBSP, "$12,750"],
    ]


def test_pptx_full_flight_has_no_per_period_wording_anywhere(tmp_path):
    prs = _build(tmp_path, [_tier([5000.0, 7750.0])], time_unit="full_flight")
    text = _text(prs) + "\n" + "\n".join(c for row in _cells(_tables(prs)[0]) for c in row)
    for banned in ("/mo", "/wk", "/qtr", "Monthly Budget", "Weekly Budget", "Quarterly Budget", "Months", "per month"):
        assert banned not in text, banned
    # the impressions sub-line is the whole-flight estimate, so it carries no unit suffix at all
    sub_lines = [line for cell in (c for row in _cells(_tables(prs)[0]) for c in row) for line in cell.split("\n") if line.startswith("Est. ")]
    assert sub_lines and all(line.endswith(" impressions") for line in sub_lines)


def test_pptx_full_flight_total_investment_equals_the_table_subtotal(tmp_path):
    prs = _build(tmp_path, [_tier([5000.0, 7750.0])], time_unit="full_flight")
    (table,) = _tables(prs)
    assert table.cell(3, 2).text == "$12,750"
    assert "Total Investment: $12,750" in _text(prs)


def test_pptx_full_flight_gross_deck_divides_by_the_agency_fee_without_multiplying_by_months(tmp_path):
    # net 12,750 / (1 - 0.15) = 15,000 — the reference proposal's gross total
    prs = _build(tmp_path, [_tier([5000.0, 7750.0])], time_unit="full_flight", gross=True)
    (table,) = _tables(prs)
    assert [table.cell(r, 2).text for r in (1, 2, 3)] == ["$5,882", "$9,118", "$15,000"]
    assert "Total Investment: $15,000" in _text(prs)


def test_pptx_full_flight_added_value_lines_show_the_tag_not_zero_dollars(tmp_path):
    prs = _build(tmp_path, [_tier([5000.0, 0.0, 7750.0], added_value={1})], time_unit="full_flight")
    (table,) = _tables(prs)
    assert [table.cell(r, 2).text for r in (1, 2, 3, 4)] == ["$5,000", "Added Value", "$7,750", "$12,750"]
    assert table.cell(2, 0).text == "Product 1"                      # no impressions sub-line for a $0 line
    assert "$0" not in "\n".join(c for row in _cells(table) for c in row)
    assert _align(table, 2, 2) == PP_ALIGN.RIGHT
    assert table.cell(2, 2).text_frame.paragraphs[0].runs[0].font.bold is False   # same plain tag as the 5-column layout


def test_pptx_full_flight_alignment_follows_the_column_role(tmp_path):
    prs = _build(tmp_path, [_tier([5000.0])], time_unit="full_flight")
    (table,) = _tables(prs)
    for r in range(len(table.rows)):
        assert _align(table, r, 0) == PP_ALIGN.LEFT
    assert [_align(table, 0, c) for c in range(3)] == [PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.RIGHT]
    assert _align(table, 1, 1) == PP_ALIGN.LEFT and _align(table, 1, 2) == PP_ALIGN.RIGHT
    assert _align(table, 2, 2) == PP_ALIGN.RIGHT                      # the subtotal's figure sits in the budget column


def test_pptx_full_flight_column_widths_fill_the_content_width_and_keep_product_and_target_wide(tmp_path):
    prs = _build(tmp_path, [_tier([5000.0])], time_unit="full_flight")
    (table,) = _tables(prs)
    widths = [col.width for col in table.columns]
    assert abs(sum(widths) - pptx_builder.CONTENT_W) <= 3             # EMU: nothing left over for another column to absorb
    assert min(widths[:2]) >= pptx.util.Inches(3.6)                   # the row-height estimate is calibrated to 3.6in
    assert abs(sum(pptx_builder._FULL_FLIGHT_COL_WIDTHS_IN) - 12.333) < 0.01


def test_pptx_full_flight_multi_option_deck_has_bars_subtotals_and_one_total_per_option(tmp_path):
    tiers = [_tier([5000.0, 7750.0], name="Independent"), _tier([3000.0], name="Democrat", prefix="Other")]
    prs = _build(tmp_path, tiers, time_unit="full_flight")
    (table,) = _tables(prs)
    cells = _cells(table)
    assert cells[1][0] == "Independent" and table.cell(1, 0).is_merge_origin and table.cell(1, 0).span_width == 3
    subtotals = [row for row in cells if row[0].endswith(" — Total")]
    assert subtotals == [["Independent — Total", NBSP, "$12,750"], ["Democrat — Total", NBSP, "$3,000"]]
    assert "Independent: $12,750" in _text(prs) and "Democrat: $3,000" in _text(prs)


def test_pptx_closing_investment_line_always_matches_the_table_subtotals(tmp_path):
    # Two options that SHARE a display name (the old recompute-by-name hazard) plus a gross fee: still one total each.
    tiers = [_tier([1000.0, 2000.0], name="Same", months=3), _tier([500.0], name="Same", months=2, prefix="Other")]
    for unit in ("month", "week", "quarter", "full_flight"):
        prs = _build(tmp_path, [dict(t, line_items=list(t["line_items"])) for t in tiers], time_unit=unit, gross=True,
                     name=f"{unit}.pptx")
        (table,) = _tables(prs)
        total_col = len(table.columns) - 1
        shown = [row[total_col] for row in _cells(table) if row[0].endswith(" — Total")]
        assert f"Same: {shown[0]}  ·  Same: {shown[1]}" in _text(prs)


def test_pptx_full_flight_long_rows_paginate_without_overrunning_the_footer(tmp_path):
    long_target = "Hispanic adults 25-54 in Los Angeles, Orange and Riverside counties, " * 5
    prods, lis = _lines([1000.0 + i for i in range(14)], target=long_target)
    banners = [{"id": str(i), "text": "Heads up. " * 60, "fill": "F8D7DA", "font_color": "58151C"} for i in range(3)]
    prs = _build(tmp_path, [{"name": "Option A", "products": prods, "line_items": lis}], time_unit="full_flight",
                 banners=banners)
    assert len(prs.slides) > 2
    footer = pptx_builder.FOOTER_Y
    names_seen, totals_seen = [], 0
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_table:
                t = sh.table
                assert sh.top + sh.height <= footer + 1
                assert abs(sh.height - sum(row.height for row in t.rows)) <= len(t.rows)      # declared == real row heights
                names_seen += [t.cell(r, 0).text.split("\n")[0] for r in range(1, len(t.rows)) if t.cell(r, 0).text.startswith("Product ")]
                totals_seen += sum(1 for r in range(len(t.rows)) if t.cell(r, 0).text == "TOTAL")
            elif sh.has_text_frame and sh.text_frame.text in ("Client", "Signature", "Date"):
                assert sh.top + sh.height <= footer + 1
            elif sh.has_text_frame and sh.text_frame.text.startswith("Total Investment"):
                assert sh.top + sh.height <= footer + 1
    assert names_seen == [p.name for p in prods]                          # every row exactly once, in order
    assert totals_seen == 1                                               # one TOTAL row, never dropped or duplicated


def test_pptx_full_flight_rows_use_the_conservative_calibrated_height_estimate(tmp_path):
    # The estimate is calibrated to 3.6in columns; the wider full-flight columns can only wrap LESS than estimated,
    # so a row is booked at least as tall as it renders (never shorter, which is what would overrun the footer).
    long_name = "Programmatic Display retargeting " * 4
    prods = [_product(long_name.strip())]
    lis = [LineItem(product_name=prods[0].name, monthly_budget=5000.0, months=1, id="a", target_override="T " * 90)]
    prs = _build(tmp_path, [{"name": "A", "products": prods, "line_items": lis}], time_unit="full_flight")
    (table,) = _tables(prs)
    expected = pptx_builder._row_height_in(prods[0].name, "T " * 90, has_impressions=True)
    assert abs(table.rows[1].height - pptx.util.Inches(expected)) <= 2


def test_pptx_unknown_unit_falls_back_to_the_month_layout_but_warns(tmp_path, caplog):
    with caplog.at_level(logging.WARNING, logger="app.services.period_copy"):
        prs = _build(tmp_path, [_tier([2000.0], months=3)], time_unit="fortnight")
    (table,) = _tables(prs)
    assert [table.cell(0, c).text for c in range(5)] == ["Product", "Target", "Months", "Monthly Budget", "Total"]
    assert any("fortnight" in rec.getMessage() for rec in caplog.records)


def test_every_billing_period_has_a_powerpoint_table_layout():
    assert set(pptx_builder._TABLE_SPECS) == set(TIME_UNITS)
    for spec in pptx_builder._TABLE_SPECS.values():
        assert len(spec.columns) == len(spec.widths_in)
        assert abs(sum(spec.widths_in) - 12.333) < 0.01
        roles = [role for role, _ in spec.columns]
        assert roles[0] == pptx_builder._ROLE_PRODUCT and roles[-1] == pptx_builder._ROLE_TOTAL


# --- week / month / quarter stay exactly as they were -----------------------------------------------------

_PER_PERIOD = {"week": ("Weeks", "Weekly Budget", "/wk"), "month": ("Months", "Monthly Budget", "/mo"),
               "quarter": ("Quarters", "Quarterly Budget", "/qtr")}


@pytest.mark.parametrize("unit", sorted(_PER_PERIOD))
def test_pptx_per_period_layout_is_the_unchanged_five_column_table(tmp_path, unit):
    periods, budget, suffix = _PER_PERIOD[unit]
    prs = _build(tmp_path, [_tier([2010.0, 0.0], months=3, added_value={1})], time_unit=unit)
    (table,) = _tables(prs)
    assert [col.width for col in table.columns] == [pptx.util.Inches(w) for w in pptx_builder._COL_WIDTHS_IN]
    assert pptx_builder._COL_WIDTHS_IN == [3.6, 3.6, 0.9, 2.1, 2.133]
    assert _cells(table) == [
        ["Product", "Target", periods, budget, "Total"],
        ["Product 0\nEst. 201,000 impressions" + suffix, "TBD", "3", "$2,010" + suffix, "$6,030"],
        ["Product 1", "TBD", "3", "Added Value", "$0"],
        ["TOTAL", NBSP, NBSP, NBSP, "$6,030"],
    ]
    assert [_align(table, 0, c) for c in range(5)] == [PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT]
    assert [_align(table, 1, c) for c in range(5)] == [PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT]
    assert _align(table, 3, 4) == PP_ALIGN.RIGHT
    assert "Total Investment: $6,030" in _text(prs)


def _dump_presentation(prs) -> str:
    """Everything a viewer can see about a deck — text, fonts, colors, fills, alignment, geometry, table structure —
    read back through python-pptx (so it doesn't depend on how the XML happens to be serialized)."""
    def rgb(color):
        try:
            return str(color.rgb)
        except (AttributeError, TypeError):
            return None

    def paragraphs(tf):
        return [[p.alignment and int(p.alignment), [(r.text, r.font.name, r.font.size and int(r.font.size), r.font.bold,
                                                    rgb(r.font.color), r.hyperlink.address) for r in p.runs]]
                for p in tf.paragraphs]

    def fill(f):
        try:
            return rgb(f.fore_color) if f.type == 1 else None
        except (AttributeError, TypeError):
            return None

    slides = []
    for slide in prs.slides:
        shapes = []
        for sh in slide.shapes:
            entry = {"type": int(sh.shape_type) if sh.shape_type else None, "box": [sh.left, sh.top, sh.width, sh.height]}
            if sh.has_text_frame:
                entry["text"] = paragraphs(sh.text_frame)
            if sh.has_table:
                t = sh.table
                entry["cols"] = [c.width for c in t.columns]
                entry["rows"] = [r.height for r in t.rows]
                entry["cells"] = [[{"merge": [c.is_merge_origin, c.is_spanned, c.span_width, c.span_height],
                                    "fill": fill(c.fill), "anchor": int(c.vertical_anchor) if c.vertical_anchor else None,
                                    "margins": [c.margin_left, c.margin_right, c.margin_top, c.margin_bottom],
                                    "text": paragraphs(c.text_frame)}
                                   for c in (t.cell(r, k) for k in range(len(t.columns)))] for r in range(len(t.rows))]
            elif sh.shape_type is not None and hasattr(sh, "fill"):
                entry["fill"] = fill(sh.fill)
            shapes.append(entry)
        slides.append(shapes)
    return json.dumps(slides, sort_keys=True, default=str)


def _golden_deck(tmp_path, scenario, unit):
    long_target = "Hispanic adults 25-54 in Los Angeles, Orange and Riverside counties. " * 2
    if scenario == "single":
        prods, lis = _lines([2010.0, 1500.5, 0.0, 900.0], months=3, target=long_target, added_value={2})
        tiers = [{"name": "Option A", "products": prods, "line_items": lis}]
        return _build(tmp_path, tiers, time_unit=unit, name=f"{scenario}-{unit}.pptx")
    tiers = [_tier([2010.0, 1500.0], name="Independent", months=3), _tier([900.0, 400.0], name="Democrat", months=3, prefix="Other")]
    banners = [{"id": "1", "text": "Live Sports Inventory: re-avail needed.", "fill": "F8D7DA", "font_color": "58151C"}]
    return _build(tmp_path, tiers, time_unit=unit, gross=True, banners=banners, name=f"{scenario}-{unit}.pptx")


# sha256 of _dump_presentation() for each deck below, captured from the PRE-feature pptx_builder. If one of these
# fails, a week/month/quarter deck no longer renders exactly as it used to — that is a regression unless the change
# to the deck's look was deliberate (then regenerate the hashes and say so in the commit).
_GOLDEN = {
    ("single", "week"):
        "fabf90a6fcbe481cf14f67302cba7776b4898fbfd7c13cf7c6403c870bc36712",
    ("single", "month"):
        "34f4701b7dd4259718407162d503f39679a52f6420d947749287968bc69f84e7",
    ("single", "quarter"):
        "c5eefc8d072a0c1709131aa42b95b7043d6d1f4f42efe1cd8610c473152599bd",
    ("tiers_gross_banner", "week"):
        "ae83c372d99067df86c231fd79cb031dc477caaee23afedde01ffa5373d40547",
    ("tiers_gross_banner", "month"):
        "d41ef2799f27c4cfdf19d716a75ab5cfcc5ed9764302919931a8fafb685b2045",
    ("tiers_gross_banner", "quarter"):
        "a1a423355643250ddfb30f8bd6c12a656160abef6c4e3ccbeb2aaa64003f3eb6",
}


@pytest.mark.parametrize("scenario,unit", sorted(_GOLDEN))
def test_pptx_week_month_quarter_decks_render_exactly_as_before(tmp_path, scenario, unit):
    dump = _dump_presentation(_golden_deck(tmp_path, scenario, unit))
    assert hashlib.sha256(dump.encode()).hexdigest() == _GOLDEN[(scenario, unit)], dump[:600]


# ===========================================================================================================
# AI enrichment prompt
# ===========================================================================================================

@pytest.fixture(autouse=True)
def stub_catalog(monkeypatch):
    """Neutral, synthetic products so the wording assertions can't be tripped by a rate-card description that
    happens to say 'per month'."""
    catalog = {n: _product(n, family=f, description="Banner ads shown across a network of websites and apps.")
               for n, f in (("Display Alpha", "Display"), ("Search Beta", "Search"))}
    monkeypatch.setattr(ai_enricher, "_catalog_by_name", lambda name: catalog.get(name))
    return catalog


def _ai_request(**kw):
    base = dict(client_name="Acme", demo="Hispanic adults 25-54", language="Spanish", geo="Los Angeles",
                start_date="2026-11-24", end_date="2026-12-31", total_months=2, request_type="New Business",
                requested_by="Pat Planner", salesperson_email="pat@example.com")
    base.update(kw)
    return ProposalRequest(**base)


def _ai_items(months=1, display=5000.0, search=7750.0):
    return [LineItem(product_name="Display Alpha", monthly_budget=display, months=months, id="a"),
            LineItem(product_name="Search Beta", monthly_budget=search, months=months, id="b")]


def _tier_context(items):
    return [{"label": "A", "name": "Independent", "line_items": items[:1]}, {"label": "B", "name": "", "line_items": items}]


def _exemplar(prompt):
    start = prompt.index('{\n  "campaign_name"')
    exemplar, _ = json.JSONDecoder().raw_decode(prompt[start:])
    return exemplar


def test_ai_prompt_full_flight_states_every_figure_as_a_flat_flight_total():
    prompt = ai_enricher._build_prompt(_ai_request(), _ai_items(), time_unit="full_flight")
    assert "  - Display Alpha: $5,000 flat for the whole flight (single billing period — not a monthly rate)" in prompt
    assert "  - Search Beta: $7,750 flat for the whole flight (single billing period — not a monthly rate)" in prompt
    assert "- Flight: 2026-11-24 → 2026-12-31 (single flight, billed as one period)" in prompt
    assert "(2 months)" not in prompt
    # 5,000 / $10 CPM + 7,750 / $10 CPM = 1,275,000 impressions for the WHOLE flight
    assert "- Total Net Investment: $12,750 (~1.3M total flight impressions — pre-computed from real rates" in prompt
    for banned in ("/month", "monthly impressions", "× 1 month", " = $5,000 total"):
        assert banned not in prompt, banned


def test_ai_prompt_full_flight_option_lines_drop_the_monthly_clause_and_label_impressions():
    items = _ai_items()
    prompt = ai_enricher._build_prompt(_ai_request(), items, tiers=_tier_context(items), time_unit="full_flight")
    assert "  - Independent: $5,000 total flight (single billing period, ~500K total flight impressions) — Display Alpha" in prompt
    assert "  - Option B: $12,750 total flight (single billing period, ~1.3M total flight impressions) — Display Alpha, Search Beta" in prompt
    assert "The total flight impressions figures below (where" in prompt
    assert "/month" not in prompt and "monthly impressions" not in prompt


def test_ai_prompt_full_flight_adds_the_billing_rule_to_the_rules_block_only():
    prompt = ai_enricher._build_prompt(_ai_request(), _ai_items(), time_unit="full_flight")
    rule_lines = [line for line in prompt.splitlines() if line.startswith("- Billing period:")]
    assert len(rule_lines) == 1
    rule = rule_lines[0]
    assert "flat total for the ENTIRE flight" in rule and "Never write one as per month, monthly, each month" in rule
    assert "never multiply or divide one by a number of months" in rule
    assert prompt.index("EMAIL AND OUTPUT RULES:") < prompt.index(rule)
    assert prompt.rstrip().endswith("- Respond ONLY with the JSON object, starting with { and ending with }")
    assert "Billing period" not in json.dumps(_exemplar(prompt))               # never inside the JSON exemplar's strings


def test_ai_prompt_full_flight_keeps_the_json_exemplar_valid_and_the_placeholders():
    for tiers in (None, _tier_context(_ai_items())):
        prompt = ai_enricher._build_prompt(_ai_request(), _ai_items(), tiers=tiers, time_unit="full_flight")
        exemplar = _exemplar(prompt)
        assert set(exemplar) >= {"campaign_name", "product_blurbs", "internal_email_body", "client_email_body"}
        assert "{{PROPOSAL_LINE}}" in exemplar["internal_email_body"] and "{{PLANNER_SIGNATURE}}" in exemplar["internal_email_body"]


def test_ai_prompt_full_flight_only_changes_the_period_sensitive_lines():
    items = _ai_items()
    for tiers in (None, _tier_context(items)):
        month = ai_enricher._build_prompt(_ai_request(), items, tiers=tiers).splitlines()
        flat = ai_enricher._build_prompt(_ai_request(), items, tiers=tiers, time_unit="full_flight").splitlines()
        changed = [line for line in flat if line not in set(month)]
        allowed = ("- Flight:", "- Total Net Investment:", "  - Display Alpha:", "  - Search Beta:", "  - Independent:",
                   "  - Option B:", "given (no planner name set). The total flight impressions", "- Billing period:")
        assert changed and all(line.startswith(allowed) for line in changed), changed
        assert len(flat) == len(month) + 1                                      # exactly one new line: the billing rule


def test_ai_prompt_month_wording_is_unchanged():
    items = _ai_items(months=3, display=2000.0, search=1500.0)
    prompt = ai_enricher._build_prompt(_ai_request(total_months=3, start_date="2026-10-01", end_date="2026-12-31"), items)
    assert "- Flight: 2026-10-01 → 2026-12-31 (3 months)" in prompt
    assert "  - Display Alpha: $2,000/month × 3 months = $6,000 total" in prompt
    assert "  - Search Beta: $1,500/month × 3 months = $4,500 total" in prompt
    assert "- Total Net Investment: $10,500 (~350K monthly impressions — pre-computed from real rates, use verbatim, never recompute)" in prompt
    assert "- Billing period:" not in prompt and "whole flight" not in prompt
    tiered = ai_enricher._build_prompt(_ai_request(), _ai_items(months=3, display=2000.0, search=1500.0),
                                       tiers=_tier_context(_ai_items(months=3, display=2000.0, search=1500.0)))
    assert "  - Independent: $2,000/month ($6,000 total flight, ~200K monthly impressions) — Display Alpha" in tiered
    assert "The monthly impressions figures below (where" in tiered


def test_ai_prompt_no_longer_prints_a_double_tilde_or_1_months():
    prompt = ai_enricher._build_prompt(_ai_request(total_months=1), _ai_items(months=1))
    assert "~~" not in prompt
    assert "(~1.3M monthly impressions" in prompt                                # the single '~' from the formatter
    assert "- Flight: 2026-11-24 → 2026-12-31 (1 month)" in prompt
    assert "$5,000/month × 1 month = $5,000 total" in prompt
    assert "1 months" not in prompt


def test_ai_prompt_week_and_quarter_keep_the_original_wording():
    # Pre-existing behavior, pinned so any change to it is deliberate: only Full Flight has its own period wording.
    items = _ai_items(months=3)
    month = ai_enricher._build_prompt(_ai_request(), items, tiers=_tier_context(items))
    for unit in ("week", "quarter", None, ""):
        assert ai_enricher._build_prompt(_ai_request(), items, tiers=_tier_context(items), time_unit=unit) == month


def test_ai_prompt_unknown_unit_warns_and_uses_the_month_wording(caplog):
    month = ai_enricher._build_prompt(_ai_request(), _ai_items())
    with caplog.at_level(logging.WARNING, logger="app.services.period_copy"):
        assert ai_enricher._build_prompt(_ai_request(), _ai_items(), time_unit="fortnight") == month
    assert any("fortnight" in rec.getMessage() for rec in caplog.records)


def test_ai_prompt_default_call_signature_is_unchanged(stub_catalog):
    # existing callers (tests/test_blurbs.py, main.py before wiring) pass no time_unit
    assert ai_enricher._build_prompt(_ai_request(), _ai_items()) == ai_enricher._build_prompt(_ai_request(), _ai_items(), time_unit="month")


class _FakeResponses:
    def __init__(self, text):
        self._text, self.calls = text, []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(output_text=self._text, status="completed", incomplete_details=None, usage=None)


def test_enrich_proposal_forwards_time_unit_to_the_prompt(monkeypatch, stub_catalog):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    monkeypatch.setattr(ai_enricher, "_HAS_OPENAI", True)
    reply = json.dumps({"campaign_name": "Holiday Push", "product_blurbs": [], "internal_email_subject": "s",
                        "internal_email_body": "b", "client_email_subject": "s", "client_email_body": "b"})
    prompts = {}
    for unit in ("full_flight", None):
        fake = SimpleNamespace(responses=_FakeResponses(reply))
        monkeypatch.setattr(ai_enricher, "_OpenAI", lambda api_key, fake=fake: fake)
        kwargs = {} if unit is None else {"time_unit": unit}
        result = ai_enricher.enrich_proposal(_ai_request(), _ai_items(), "short", **kwargs)
        assert result.error is None and result.campaign_name == "Holiday Push"
        prompts[unit] = fake.responses.calls[0]["input"]
    assert "flat for the whole flight" in prompts["full_flight"] and "- Billing period:" in prompts["full_flight"]
    assert "$5,000/month × 1 month" in prompts[None] and "whole flight" not in prompts[None]


# ===========================================================================================================
# Strategy Brief — prompt + Word document share ONE budget sentence
# ===========================================================================================================

def _brief_request(**kw):
    base = dict(client_name="Acme", monthly_budget=15000, total_months=2)
    base.update(kw)
    return ProposalRequest(**base)


def test_strategy_prompt_budget_line_for_every_unit():
    flat = sb._build_prompt(_brief_request(), None, time_unit="full_flight")
    # the figure is the flight's budget as typed (the same figure Step 04 seeds), stated as ONE flat total
    assert "\n- Budget: $15,000 total flight (single billing period)\n" in flat
    assert "/month" not in flat and "$30,000" not in flat and "×" not in flat   # no "$15,000/month × 2 months" arithmetic
    monthly = "\n- Budget: $15,000/month × 2 months = $30,000 total flight\n"
    assert monthly in sb._build_prompt(_brief_request(), None)
    for unit in ("month", "week", "quarter"):
        assert monthly in sb._build_prompt(_brief_request(), None, time_unit=unit)


def test_strategy_prompt_pluralizes_one_month_and_keeps_the_three_month_fallback():
    assert "\n- Budget: $5,000/month × 1 month = $5,000 total flight\n" in sb._build_prompt(_brief_request(monthly_budget=5000, total_months=1), None)
    assert "\n- Budget: $5,000/month × 3 months = $15,000 total flight\n" in sb._build_prompt(_brief_request(monthly_budget=5000, total_months=0), None)
    # full flight never uses a month count at all: the same flat total whatever it is (and never the fallback of 3)
    for months in (0, 1, 2, 6):
        assert "\n- Budget: $5,000 total flight (single billing period)\n" in sb._build_prompt(
            _brief_request(monthly_budget=5000, total_months=months), None, time_unit="full_flight"), months


def test_strategy_prompt_differs_between_units_only_on_the_budget_line():
    month = sb._build_prompt(_brief_request(), None).splitlines()
    flat = sb._build_prompt(_brief_request(), None, time_unit="full_flight").splitlines()
    diff = [line for line in difflib.ndiff(month, flat) if line[:2] in ("- ", "+ ")]
    assert diff == ["- - Budget: $15,000/month × 2 months = $30,000 total flight",
                    "+ - Budget: $15,000 total flight (single billing period)"]


class _FakeBriefResponses:
    def __init__(self):
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        text = json.dumps({"client_summary": "c", "market_context": "m", "objectives_analysis": "o",
                           "strategy_summary": "s", "recommended_tactics": [], "key_insights": ["i"]})
        return SimpleNamespace(output_text=text, status="completed", incomplete_details=None, usage=None)


def test_generate_brief_forwards_time_unit(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    monkeypatch.setattr(sb, "_HAS_OPENAI", True)
    fake = SimpleNamespace(responses=_FakeBriefResponses())
    monkeypatch.setattr(sb, "_OpenAI", lambda api_key: fake)
    brief = asyncio.run(sb.generate_brief(_brief_request(), time_unit="full_flight"))
    assert brief["error"] is None
    assert "Budget: $15,000 total flight (single billing period)" in fake.responses.calls[0]["input"]
    asyncio.run(sb.generate_brief(_brief_request()))                              # default keeps the original wording
    assert "Budget: $15,000/month × 2 months = $30,000 total flight" in fake.responses.calls[1]["input"]


def _strategy_docx(tmp_path, **kw):
    out = tmp_path / "brief.docx"
    assert docx_builder.build_strategy_brief_docx(
        output_path=out, title="Acme", client_summary="Acme sells widgets.", market_context="", objectives_analysis="",
        strategy_summary="Plan.", recommended_tactics=[], key_insights=[], **kw)
    return docx.Document(str(out))


def _budget_paragraphs(doc):
    return [p for p in doc.paragraphs if p.text.startswith("Budget:")]


@pytest.mark.parametrize("unit,expected", [
    ("full_flight", "Budget: $15,000 total flight (single billing period)"),
    ("month", "Budget: $15,000/month × 2 months = $30,000 total flight"),
    ("week", "Budget: $15,000/month × 2 months = $30,000 total flight"),
    ("quarter", "Budget: $15,000/month × 2 months = $30,000 total flight"),
])
def test_strategy_docx_budget_paragraph_per_unit(tmp_path, unit, expected):
    paragraphs = _budget_paragraphs(_strategy_docx(tmp_path, monthly_budget=15000, total_months=2, time_unit=unit))
    assert [p.text for p in paragraphs] == [expected]
    assert paragraphs[0].runs[0].italic is True


def test_strategy_docx_default_unit_is_month_and_one_month_is_singular(tmp_path):
    assert [p.text for p in _budget_paragraphs(_strategy_docx(tmp_path, monthly_budget=5000, total_months=3))] == \
        ["Budget: $5,000/month × 3 months = $15,000 total flight"]
    assert [p.text for p in _budget_paragraphs(_strategy_docx(tmp_path, monthly_budget=5000, total_months=1))] == \
        ["Budget: $5,000/month × 1 month = $5,000 total flight"]


def test_strategy_docx_full_flight_states_the_budget_as_typed_and_is_omitted_without_a_budget(tmp_path):
    # the figure is the flight's budget — nothing is multiplied, whatever the month count (or its absence)
    assert [p.text for p in _budget_paragraphs(_strategy_docx(tmp_path, monthly_budget=15000, total_months=0, time_unit="full_flight"))] == \
        ["Budget: $15,000 total flight (single billing period)"]
    assert [p.text for p in _budget_paragraphs(_strategy_docx(tmp_path, monthly_budget=6375, total_months=2, time_unit="full_flight"))] == \
        ["Budget: $6,375 total flight (single billing period)"]
    assert not _budget_paragraphs(_strategy_docx(tmp_path, monthly_budget=0, total_months=2, time_unit="full_flight"))
    # unchanged for the monthly units: both figures are still required
    assert not _budget_paragraphs(_strategy_docx(tmp_path, monthly_budget=15000, total_months=0))
    assert not _budget_paragraphs(_strategy_docx(tmp_path))


def test_strategy_docx_unknown_unit_warns_and_uses_month_wording(tmp_path, caplog):
    with caplog.at_level(logging.WARNING, logger="app.services.period_copy"):
        paragraphs = _budget_paragraphs(_strategy_docx(tmp_path, monthly_budget=15000, total_months=2, time_unit="fortnight"))
    assert [p.text for p in paragraphs] == ["Budget: $15,000/month × 2 months = $30,000 total flight"]
    assert any("fortnight" in rec.getMessage() for rec in caplog.records)


@pytest.mark.parametrize("unit", sorted(TIME_UNITS))
def test_strategy_prompt_and_docx_always_state_the_same_budget(tmp_path, unit):
    prompt = sb._build_prompt(_brief_request(), None, time_unit=unit)
    (para,) = _budget_paragraphs(_strategy_docx(tmp_path, monthly_budget=15000, total_months=2, time_unit=unit))
    assert f"\n- {para.text}\n" in prompt


def test_strategy_docx_builds_with_none_of_the_optional_budget_arguments(tmp_path):
    # tests/test_strategy_doc_rebuild.py-style call: no monthly_budget / total_months / time_unit at all
    assert not _budget_paragraphs(_strategy_docx(tmp_path))


# ===========================================================================================================
# Email revision prompt (the planner's "make it shorter / mention the budget" pass) carries the same billing rule
# ===========================================================================================================

def _revision_prompt(monkeypatch, **kwargs):
    captured = []
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    monkeypatch.setattr(ai_enricher, "_HAS_OPENAI", True)
    monkeypatch.setattr(ai_enricher, "_OpenAI", lambda api_key: object())
    monkeypatch.setattr(ai_enricher.llm_utils, "chat_text", lambda client, model, prompt, **k: captured.append(prompt) or json.dumps(
        {"internal_email_subject": "s", "internal_email_body": "b", "client_email_subject": "s", "client_email_body": "b"}))
    result = ai_enricher.reprompt_emails(
        SimpleNamespace(client_name="Acme"), [SimpleNamespace(monthly_budget=680.0, months=1)], "Holiday",
        "int subject", "int body", "cli subject", "cli body", "show the cost per month", **kwargs)
    assert result["error"] is None
    return captured[0]


def test_the_revision_prompt_carries_the_full_flight_rule_and_nothing_else_changes(monkeypatch):
    flat = _revision_prompt(monkeypatch, time_unit="full_flight")
    month = _revision_prompt(monkeypatch)
    assert "CRITICAL — BILLING PERIOD: this is a FULL-FLIGHT proposal" in flat
    assert "ONE period" in flat and "never multiply or divide one by a number of months" in flat
    assert "give the flat flight total instead" in flat
    # the rule sits in the instructions, before (never inside) the JSON structure the model must return
    assert flat.index("CRITICAL — BILLING PERIOD") < flat.index("Respond with this exact JSON structure")
    # every other unit — and the default call — is the original prompt, byte for byte (verified against the pre-feature module)
    for unit in ("month", "week", "quarter"):
        assert _revision_prompt(monkeypatch, time_unit=unit) == month
    assert "BILLING PERIOD" not in month
    assert flat.replace(flat[flat.index("CRITICAL — BILLING PERIOD"):flat.index("Respond with this exact JSON structure")], "") == month


def test_the_email_prompt_and_the_revision_prompt_share_one_billing_rule(monkeypatch):
    rule = ai_enricher._FULL_FLIGHT_RULE_TEXT
    assert f"\n- Billing period: {rule}\n" in ai_enricher._build_prompt(_ai_request(), _ai_items(), time_unit="full_flight")
    assert rule in _revision_prompt(monkeypatch, time_unit="full_flight")


# ===========================================================================================================
# Full Flight decks name each option's own flight (they have no Months column to hint at different windows)
# ===========================================================================================================

def _with_window(tier, start=None, end=None):
    return dict(tier, start_date=start, end_date=end)


def test_pptx_full_flight_option_bars_carry_an_overridden_flight_window(tmp_path):
    tiers = [_tier([680.0], name="Holiday"),
             _with_window(_tier([2040.0], name="Q1", prefix="Other"), "2027-01-04", "2027-03-31")]
    prs = _build(tmp_path, tiers, time_unit="full_flight")
    (table,) = _tables(prs)
    bars = [row[0] for row in _cells(table) if row[0] and row[0].startswith(("Holiday", "Q1")) and not row[0].endswith("Total")]
    assert bars == ["Holiday", "Q1  ·  2027-01-04 – 2027-03-31"]          # the option on the campaign dates says nothing extra


def test_pptx_full_flight_a_partial_override_falls_back_per_bound_and_a_no_op_override_adds_nothing(tmp_path):
    tiers = [_with_window(_tier([100.0], name="Late start"), "2026-12-01", None),            # only the start moves
             _with_window(_tier([200.0], name="Same dates", prefix="Other"), "2026-11-24", "2026-12-31"),   # == campaign
             _with_window(_tier([300.0], name="Blank", prefix="Third"), "", "  ")]
    (table,) = _tables(_build(tmp_path, tiers, time_unit="full_flight"))
    bars = [row[0] for row in _cells(table) if row[0] and not row[0].endswith("Total") and row[0] in
            ("Late start  ·  2026-12-01 – 2026-12-31", "Same dates", "Blank", "Late start")]
    assert bars == ["Late start  ·  2026-12-01 – 2026-12-31", "Same dates", "Blank"]


def test_pptx_full_flight_override_still_shows_when_the_campaign_has_no_dates_and_dateless_options_say_nothing(tmp_path):
    tiers = [_with_window(_tier([100.0], name="A"), "2027-01-04", "2027-03-31"), _tier([200.0], name="B", prefix="Other")]
    (table,) = _tables(_build(tmp_path, tiers, time_unit="full_flight", request=_req(start_date="", end_date="")))
    bars = [row[0] for row in _cells(table) if row[0] in ("B",) or row[0].startswith("A  ·")]
    assert bars == ["A  ·  2027-01-04 – 2027-03-31", "B"]     # A's window is information the (date-less) header lacks; B has none to give
    # a half-specified window (end only) is never printed as a window
    half = [_with_window(_tier([100.0], name="A"), None, "2027-03-31"), _tier([200.0], name="B", prefix="Other")]
    (table2,) = _tables(_build(tmp_path, half, time_unit="full_flight", request=_req(start_date="", end_date=""), name="h.pptx"))
    assert [row[0] for row in _cells(table2) if row[0] in ("A", "B")] == ["A", "B"]


def test_pptx_full_flight_single_option_header_uses_its_own_flight(tmp_path):
    overridden = _with_window(_tier([680.0]), "2027-01-04", "2027-03-31")
    text = _text(_build(tmp_path, [overridden], time_unit="full_flight"))
    assert "2027-01-04 – 2027-03-31" in text and "2026-11-24 – 2026-12-31" not in text
    plain = _text(_build(tmp_path, [_tier([680.0])], time_unit="full_flight", name="p.pptx"))
    assert "2026-11-24 – 2026-12-31" in plain


@pytest.mark.parametrize("unit", ["month", "week", "quarter"])
def test_pptx_per_period_decks_ignore_the_option_flight_keys(tmp_path, unit):
    """week/month/quarter must render exactly as before: the new start_date/end_date keys on a tier are never read."""
    base = [_tier([5000.0], name="One", months=2), _tier([3000.0], name="Two", months=3, prefix="Other")]
    keyed = [_with_window(base[0], "2027-01-04", "2027-03-31"), _with_window(base[1], "2027-02-01", "2027-02-28")]
    a = _build(tmp_path, base, time_unit=unit, name="a.pptx")
    b = _build(tmp_path, keyed, time_unit=unit, name="b.pptx")
    assert [_cells(t) for t in _tables(a)] == [_cells(t) for t in _tables(b)]
    assert _text(a) == _text(b) and "2027" not in _text(b)
    single = _build(tmp_path, [_with_window(_tier([5000.0], months=2), "2027-01-04", "2027-03-31")], time_unit=unit, name="c.pptx")
    assert "2026-11-24 – 2026-12-31" in _text(single) and "2027-01-04" not in _text(single)
