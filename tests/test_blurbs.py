"""
Product blurbs (Excel column E / "DETAILS"): short, plain, no targeting
repeats, no variants, no language-of-inventory claims — enforced both in the
prompt AND deterministically on whatever the model returns. Plus the catalog's
internal "Variants (contact planning for exact tier): ..." rate-tier tail,
which no longer prints into the client-facing cell.
"""
from __future__ import annotations

import os

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import pytest

from app.catalog import CATALOG
from app.services import ai_enricher
from app.services.ai_enricher import _tighten_blurb, _blurb_sentence_ok
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import LineItem, generate_proposal
from app.services.text_utils import strip_variant_pricing


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


# --- the catalog's internal variants tail ----------------------------------------

def test_strip_variant_pricing_removes_only_the_trailing_rate_tier_list():
    desc = ("Use our mix of Internet radio, podcasts, and streaming channels. "
            "Variants (contact planning for exact tier): Standard: $13 CPM; Custom Audience: $15 CPM.")
    assert strip_variant_pricing(desc) == "Use our mix of Internet radio, podcasts, and streaming channels."


def test_strip_variant_pricing_leaves_other_descriptions_alone_and_blanks_a_variants_only_one():
    assert strip_variant_pricing("Serve your banner ads within a virtual perimeter.") == "Serve your banner ads within a virtual perimeter."
    assert strip_variant_pricing("Variants (contact planning for exact tier): Search - AdWords - SEM: $1000 CPP.") == ""
    assert strip_variant_pricing(None) == "" and strip_variant_pricing("") == ""


def test_every_catalog_description_loses_the_tail_and_nothing_else():
    for p in CATALOG:
        cleaned = strip_variant_pricing(p.proposal_description)
        assert "contact planning for exact tier" not in cleaned, p.name
        # Only the trailing variants chunk may go: everything before it is kept verbatim.
        original = p.proposal_description or ""
        marker = original.find("Variants (contact planning")
        assert cleaned == (original[:marker] if marker >= 0 else original).rstrip(), p.name


# --- deterministic blurb enforcement ------------------------------------------------

def test_a_clean_short_blurb_passes_through_untouched():
    text = "Banner ads shown across a network of websites and apps, retargeting shoppers as they browse."
    assert _tighten_blurb(text, "eDigital Network Display - Standard IAB") == text


@pytest.mark.parametrize("bad", [
    "This is the Hispanic-focused version of Entravision's programmatic CTV/OTT video.",
    "Several variants are available depending on your budget.",
    "Runs on Spanish-language sites and apps.",
    "Delivers across English-only inventory.",
    "Skippable and non-skippable formats are available, though skippable is recommended.",
])
def test_forbidden_sentences_are_dropped(bad):
    good = "Video ads delivered to connected TVs and streaming apps."
    assert _tighten_blurb(f"{good} {bad}", "Some Product") == good
    assert _tighten_blurb(f"{bad} {good}", "Some Product") == good


def test_a_blurb_made_only_of_forbidden_sentences_becomes_empty_so_the_cell_shows_catalog_text_only():
    assert _tighten_blurb("This is the Hispanic-focused version of our product.", "X") == ""
    assert _tighten_blurb("", "X") == ""


def test_skippable_is_allowed_when_the_product_itself_is_named_that():
    sentence = "Skippable in-stream ads that viewers can skip after five seconds."
    assert _blurb_sentence_ok(sentence, "YouTube Skippable In-Stream")
    assert not _blurb_sentence_ok(sentence, "YouTube Ads")


def test_over_long_blurbs_drop_trailing_sentences_but_never_cut_mid_sentence():
    sentences = [f"Sentence number {i} says something plain about the channel and how it works for buyers." for i in range(12)]
    out = _tighten_blurb(" ".join(sentences), "Any")
    assert len(out.split()) <= 70 and ai_enricher._BLURB_MAX_WORDS == 70
    assert out.endswith(".") and out.split(". ")[0].startswith("Sentence number 0")
    # a single oversized sentence is kept whole rather than truncated
    long_one = "word " * 200 + "end."
    assert _tighten_blurb(long_one, "Any") == long_one.strip()


# --- the prompt itself ---------------------------------------------------------------

def _prompt():
    req = ProposalRequest(
        client_name="Acme", demo="Hispanic adults 25-54", language="Spanish", geo="Los Angeles",
        start_date="2026-10-01", end_date="2026-12-31", total_months=3, request_type="New Business",
    )
    items = [
        LineItem(product_name="AudioEngage", monthly_budget=2000, months=3, id="a"),
        LineItem(product_name="Search - SEM", monthly_budget=1500, months=3, id="b"),
    ]
    return ai_enricher._build_prompt(req, items)


def test_prompt_demands_short_plain_blurbs_without_targeting_or_variants():
    prompt = _prompt()
    assert "ONE short paragraph, 25–55 words" in prompt
    assert "Do NOT repeat the audience, geography, language" in prompt
    assert "Never mention variants, tiers, versions or format options" in prompt
    assert "Never make claims about the language of the content, sites, apps, stations or publishers" in prompt
    assert "never the client's audience or market" in prompt


def test_prompt_no_longer_forces_geo_demo_stats_or_hispanic_stats_into_every_blurb():
    prompt = _prompt()
    for old in (
        "use these SPECIFIC values by name in every blurb",
        "must name at least one specific targeting value",
        "this matters because [specific recent stat tied to THAT audience",
        "For Hispanic/Spanish targets: use U.S. Hispanic-specific stats",
        "50–80 words",
    ):
        assert old not in prompt, old


def test_prompt_still_gives_the_emails_their_audience_context_and_the_voice_guide():
    prompt = _prompt()
    assert "Hispanic adults 25-54" in prompt          # target block still present for the emails
    assert "background for the EMAILS" in prompt
    assert "governs the EMAILS" in prompt
    assert ai_enricher.HOUSE_VOICE_GUIDE.strip()[:40] in prompt


def test_prompt_shows_the_model_each_product_without_its_internal_rate_tiers():
    prompt = _prompt()
    assert "AudioEngage" in prompt and "Use our mix of Internet radio" in prompt
    assert "contact planning for exact tier" not in prompt
    assert "Custom Audience: $15 CPM" not in prompt


def test_prompt_json_exemplar_is_still_valid_json_shape():
    import json
    prompt = _prompt()
    start = prompt.index('{\n  "campaign_name"')
    exemplar, _ = json.JSONDecoder().raw_decode(prompt[start:])       # the exemplar really is valid JSON
    assert set(exemplar) >= {"campaign_name", "product_blurbs", "internal_email_body", "client_email_body"}
    assert exemplar["product_blurbs"][0]["blurb"].startswith("[One or two plain sentences")


# --- parsed responses are tightened -------------------------------------------------

def test_parse_response_applies_the_tightening_to_every_blurb():
    import json
    raw = json.dumps({
        "campaign_name": "Fall Push",
        "product_blurbs": [
            {"product_name": "eDigital Network Display - Standard IAB",
             "blurb": "Banner ads across a network of sites. This is the Hispanic-focused version of the product."},
            {"product_name": "AudioEngage", "blurb": "Runs on Spanish-language sites."},
        ],
        "internal_email_subject": "s", "internal_email_body": "b",
        "client_email_subject": "s", "client_email_body": "b",
    })
    req = ProposalRequest(client_name="Acme")
    out = ai_enricher._parse_response(raw, req, [])
    by = {b.product_name: b.blurb for b in out.product_blurbs}
    assert by["eDigital Network Display - Standard IAB"] == "Banner ads across a network of sites."
    assert by["AudioEngage"] == ""


# --- the export cell ----------------------------------------------------------------

def _cell_e_texts(ws):
    return [str(ws.cell(row=r, column=5).value) for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=5).value]


def test_export_cell_shows_catalog_text_without_the_variants_tail_then_the_blurb(tmp_path):
    from app.services.ai_enricher import ProposalEnrichment, ProductBlurb
    req = ProposalRequest(client_name="Acme", request_type="New Business", start_date="2026-10-01",
                          end_date="2026-12-31", total_months=3)
    items = [
        LineItem(product_name="AudioEngage", monthly_budget=2000.0, months=3, id="a"),
        LineItem(product_name="Search - SEM", monthly_budget=1500.0, months=3, id="b"),
    ]
    enrichment = ProposalEnrichment(
        campaign_name="Fall Push",
        product_blurbs=[ProductBlurb("AudioEngage", "Audio ads across internet radio, podcasts and streaming."),
                        ProductBlurb("Search - SEM", "Text ads that appear when people search for what you sell.")],
    )
    out = tmp_path / "x.xlsx"
    generate_proposal(req, items, out, enrichment=enrichment)
    import openpyxl
    ws = openpyxl.load_workbook(out)["Proposal A"]
    texts = _cell_e_texts(ws)
    audio = next(t for t in texts if "Internet radio, podcasts, and streaming channels" in t)
    assert "Variants" not in audio and "$15 CPM" not in audio
    assert audio.index("Use our mix of Internet radio") < audio.index("Audio ads across internet radio")
    # SEM's catalog sentence, then the blurb (its variants list is stripped like everyone else's)
    sem = next(t for t in texts if "Text ads that appear when people search" in t)
    assert sem.startswith("Text ads on Google Search that appear when people look for what you offer")
    assert sem.endswith("Text ads that appear when people search for what you sell.")
    assert not any("contact planning for exact tier" in t for t in texts)


def test_export_without_any_blurb_still_hides_the_variants_tail(tmp_path):
    req = ProposalRequest(client_name="Acme", request_type="New Business", start_date="2026-10-01",
                          end_date="2026-12-31", total_months=3)
    out = tmp_path / "y.xlsx"
    generate_proposal(req, [LineItem(product_name="AudioEngage", monthly_budget=2000.0, months=3, id="a")], out)
    import openpyxl
    ws = openpyxl.load_workbook(out)["Proposal A"]
    texts = _cell_e_texts(ws)
    assert any("Use our mix of Internet radio" in t for t in texts)
    assert not any("Variants" in t for t in texts)
