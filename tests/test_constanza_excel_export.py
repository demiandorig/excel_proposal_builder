"""
Verifies a Constanza-derived avails estimate actually survives into the
real generated .xlsx — not just the UI's own form fields, which is all
the other Constanza tests check. Generating an actual file caught a real
bug the UI-level tests couldn't: applyConstanzaEstimate() (app.js) sets
max_imps_estimated=True together with basis="imps" — a combination that
never arose before Constanza (Step 05's own UI only ever sets
max_imps_estimated=False when basis="imps"; a directly-typed value isn't
"estimated" by definition). app/excel_template.py's write_avails_cells()
wrote that combination as literal "Est. 12,081,928" TEXT into the imps
cell while ALSO writing a live spend formula that multiplies that same
cell as a NUMBER (every branch of _spend_formula_from_imps does) — text
times a number is a formula error in real Excel, silently caught by
IFERROR into a blank spend cell. Fixed in write_avails_cells() by writing
a real number with a custom "Est. "#,##0 display format instead of literal
text, whenever a formula depends on the cell. This test locks that in.

Needs no real Postgres — mocks get_connection() in every module that
touches it during proposal generation (catalog overrides, market config,
auth, Drive tokens), matching how the app itself degrades when
DATABASE_URL is unset for anything that doesn't actually need a row.
"""
import sys
from contextlib import ExitStack, contextmanager
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

print("=" * 80)
print("Constanza -> real .xlsx export")
print("=" * 80)


@contextmanager
def _fake_connection():
    conn = MagicMock()
    conn.execute.return_value.fetchall.return_value = []
    conn.execute.return_value.fetchone.return_value = None
    yield conn


import app.catalog as catalog_mod
import app.market_config as market_config_mod
import app.auth as auth_mod
import app.services.drive_uploader as drive_mod

with ExitStack() as stack:
    for mod in (catalog_mod, market_config_mod, auth_mod, drive_mod):
        stack.enter_context(patch.object(mod, "get_connection", _fake_connection))

    from app.catalog import CATALOG
    from app.services.notion_parser import parse_notion
    from app.services.proposal_generator import LineItem, generate_proposal
    from app.services.constanza_geo import resolve_geo
    from app.services.constanza_engine import estimate_avails

    import openpyxl

    catalog_names = [p.name for p in CATALOG]
    NOTION_TEXT = """Requested by: Test Planner
Salesperson market: Los Angeles (Tier 1)
Salesperson email: test.planner@entravision.com
CCs:
Request type: Proposal to Sign
────────────
>>> Avails/Presentation/Proposal Page details <<<
Client name: Test Client Inc
Client website: https://example.com
Agency name:
Agency Fee:
────────────
Start date: 2026-11-01
End date: 2027-01-31
Total months: 3
Monthly budget: 5000
Tiered budget?: false
────────────
Chosen campaign goal: Awareness
────────────
Target details:
Language of campaign: English
(Other languages: )
Geo: Sacramento, CA
"""
    req = parse_notion(NOTION_TEXT, catalog_names)
    assert req.request_type == "Proposal to Sign"

    ROKU = "Entravision Plus - Roku Ads (Includes The Roku Channel and the popular Espacio Latino Hub)"
    roku_product = next(p for p in CATALOG if p.name == ROKU)
    geo = resolve_geo("Sacramento, CA", geo_type_hint="city")
    est = estimate_avails(ROKU, geo, months=1)
    max_spend = round(est.imps_mid * roku_product.base_rate / 1000)

    li = LineItem(product_name=ROKU, monthly_budget=5000, months=3, id="li-roku")
    avails_data = {
        "li-roku": {
            "max_imps": round(est.imps_mid), "max_imps_estimated": True,
            "max_spend": max_spend, "max_spend_estimated": True,
            "est_uniques": round(est.est_uniques), "basis": "imps",
        },
    }

    out_path = Path("/tmp/test_constanza_excel_export.xlsx")
    result = generate_proposal(req, [li], out_path, avails_data=avails_data)
    assert out_path.exists()
    print(f"Generated real .xlsx via generate_proposal(): {result['tabs_built']}: OK")

    wb = openpyxl.load_workbook(out_path, data_only=False)
    ws = wb[result["tabs_built"][0]]

    # Find the Roku line's row (row 19 in the standard template layout,
    # but check by content instead of a hardcoded row in case the
    # template shifts).
    imps_cell = spend_cell = None
    for row in ws.iter_rows(min_row=15, max_row=25):
        for cell in row:
            if isinstance(cell.value, str) and "Roku Ads" in cell.value:
                r = cell.row
                imps_cell, spend_cell = ws[f"N{r}"], ws[f"O{r}"]
                break
        if imps_cell:
            break
    assert imps_cell is not None, "Could not find the Roku line item's row in the generated sheet"

    # THE regression check: the imps cell must be a NUMBER (not text),
    # even though it displays as "Est. ..." via number formatting — a
    # live formula (in the spend cell) depends on it arithmetically.
    assert imps_cell.data_type == "n", f"imps cell should be numeric, got data_type={imps_cell.data_type!r} value={imps_cell.value!r}"
    assert imps_cell.value == round(est.imps_mid)
    assert "Est." in imps_cell.number_format
    print(f"Imps cell is a real NUMBER ({imps_cell.value:,}) displayed as \"Est. ...\" via number_format, not literal text: OK")

    assert isinstance(spend_cell.value, str) and spend_cell.value.startswith("=")
    assert f"{imps_cell.coordinate}*" in spend_cell.value or f"*{imps_cell.coordinate}" in spend_cell.value or imps_cell.coordinate in spend_cell.value
    print(f"Spend cell is a live formula referencing the (numeric) imps cell: {spend_cell.value}: OK")

    # Manually evaluate what the formula would compute in real Excel —
    # confirms it's not just syntactically plausible but numerically
    # correct, since no recalculation engine (LibreOffice) is available
    # in this environment to actually open and recalc the file.
    rate_cell_coord = spend_cell.value.split("*")[1].split("/")[0].strip()
    rate = ws[rate_cell_coord].value
    expected_spend = imps_cell.value * rate / 1000
    assert abs(expected_spend - max_spend) < 1, (expected_spend, max_spend)
    print(f"Formula would evaluate to ${expected_spend:,.2f}, matching the Constanza-computed spend (${max_spend:,}): OK")

print()
print("All Constanza Excel-export checks passed.")
