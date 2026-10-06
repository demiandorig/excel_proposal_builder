"""
Round-2 regression tests: the second adversarial pass over the Disclaimers / Restrictions / blurb / export work
(sheet parser edge cases, the Roadblocks denial filter, model-shortened product names, PPT line estimates, blurb
tightening corners, request-size limits).
"""
from __future__ import annotations

import io
import json
import os
from types import SimpleNamespace

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import pytest

from app.services import restrictions as rs
from app.services import roadblocks as rb
from app.services.ai_enricher import _parse_response, _tighten_blurb
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import LineItem, generate_proposal
from app.services.text_utils import split_sentences


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.catalog.resolve_product_alias", lambda name: None)
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


@pytest.fixture
def builtin_state(monkeypatch):
    monkeypatch.setattr(rs, "load_state", lambda: {"categories": rs.builtin_categories(), "using_builtin": True,
                                                   "sync": None, "error": False})


def _items(*names):
    return [LineItem(product_name=n, monthly_budget=1000, months=3, id=f"i{i}") for i, n in enumerate(names)]


def _tab(name, rows):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(name)
    for r in rows:
        ws.append(r if isinstance(r, list) else [r])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse(name, rows):
    cats, warnings = rs.parse_workbook_ex(_tab(name, rows))
    return cats[0], warnings


# --- the sheet parser ------------------------------------------------------------------------------------------

def test_a_header_that_says_not_accepted_is_not_an_approved_list():
    cat, _ = _parse("Political", ["Political ads are not accepted on the following products:", "YouTube Ads"])
    assert cat.kind == "guidance"


def test_a_hold_headline_followed_by_a_previously_approved_list_stays_a_blanket_hold():
    cat, _ = _parse("Cannabis", [
        None, "** All Entravision Digital Products are currently on hold and will not accept cannabis ads **",
        "Previously approved for the following products (currently on hold as well):", "CTV/OTT", "eDigital OLV"])
    assert cat.kind == "blanket_hold"
    assert rs.evaluate_product("Video - Pre-roll (OLV)", "Online Video", [cat])["verdict"] == "not_allowed"


def test_a_second_list_after_a_spacer_row_is_not_absorbed_as_approved_products():
    cat, _ = _parse("Political", [
        "Political Ads are only accepted for the below products", "eDigital Display", "eDigital OLV", [None],
        "Political ads are NOT accepted for the below products", "YouTube Ads", "Search"])
    assert cat.items == ["eDigital Display", "eDigital OLV"]


def test_notes_after_a_spacer_row_are_not_products():
    cat, _ = _parse("Political", [
        "Political Ads are only accepted for the below products", "eDigital Display", [None],
        "Notes:", "All creative must carry a Paid for by disclaimer."])
    assert cat.items == ["eDigital Display"]


def test_a_headline_above_a_platform_grid_still_decides_the_tabs_kind():
    allow, _ = _parse("Political", ["Political Ads are only accepted for the below products", "eDigital Display",
                                    [None], [None], ["Platform", "Notes"], ["TTD", "ok"]])
    assert allow.kind == "allowlist" and allow.items == ["eDigital Display"]
    hold, _ = _parse("Cannabis", [None, "** All Entravision Digital Products are currently on hold **", [None],
                                  ["Platform", "Notes"], ["TTD", "ok"]])
    assert hold.kind == "blanket_hold"
    grid, _ = _parse("Casino Gambling", [["Platform", "Casinos"],
                                         ["TTD", "We only accept brick+mortar ads on the following platforms"]])
    assert grid.kind == "guidance" and "TTD:" in grid.body       # a grid's own cells never turn it into a list


def test_a_note_about_one_product_family_is_not_a_blanket_hold():
    for line in ("Entravision's audio products are currently on hold pending legal review.",
                 "* Plasma centers: all platforms will not accept ads that promise payment."):
        cat, _ = _parse("Blood Donation", ["* Interest targeting", line])
        assert cat.kind == "guidance", line


def test_an_entry_on_the_headers_own_row_is_kept_and_an_overlong_entry_is_reported():
    cat, _ = _parse("Political", [["Political Ads are only accepted for the below products", "eDigital Display"], "eDigital OLV"])
    assert cat.items == ["eDigital Display", "eDigital OLV"]
    cat, warnings = _parse("Political", ["Political Ads are only accepted for the below products", "eDigital Display",
                                         " ".join(["word"] * 30)])
    assert cat.items == ["eDigital Display"] and any("too long" in w for w in warnings)


def test_the_fallback_matcher_no_longer_guesses():
    assert rs.match_catalog_products("Audio Streaming (Spotify)") == []
    # a suffixed entry maps to ITS OWN product, never to the plain Re-Drop one
    assert rs.match_catalog_products("Email Campaigns - Hashed Email File") == ["Email Campaigns - Hashed Email File"]
    assert not any("Re-Drop" in n for n in rs.match_catalog_products("Email Campaigns - Postal Matching"))
    assert rs.match_catalog_products("Geofencing Display and Geofencing OLV") == ["Display - Geo Fence"]


def test_political_suggestions_recognise_campaign_phrasing():
    cats = rs.builtin_categories()
    for text in ("Smith for Congress", "Committee to elect John Garcia", "Yes on Prop 50", "Vote yes on Measure A"):
        assert "Political" in rs.suggest_categories(text, cats), text


# --- request sizes ---------------------------------------------------------------------------------------------------

def test_admin_keyword_strings_are_limited_by_their_content_not_by_a_200_character_cap():
    import app.main as m
    from pydantic import ValidationError
    many = ", ".join(f"keyword number {i}" for i in range(40))           # ~600 characters, 40 keywords
    assert m.RestrictionUpdateRequest(keywords=many, product_map={}).keywords == many
    assert m.DisclaimerRequest(name="n", keywords=many, banner_text="t").keywords == many
    with pytest.raises(ValidationError):
        m.RestrictionUpdateRequest(keywords=[f"k{i}" for i in range(201)], product_map={})


def test_a_failed_pre_sync_read_refuses_to_overwrite_the_stored_rows(monkeypatch):
    import app.main as m
    from fastapi import HTTPException
    monkeypatch.setattr(rs, "load_state", lambda: {"categories": [], "using_builtin": True, "sync": None, "error": True})
    monkeypatch.setattr(rs, "replace_all", lambda *a, **k: pytest.fail("must not overwrite"))
    with pytest.raises(HTTPException) as exc:
        m._ingest_restrictions(_tab("Political", ["Political Ads are only accepted for the below products", "eDigital Display"]),
                               rs.DEFAULT_SHEET_URL, "t", "a@x")
    assert exc.value.status_code == 503


def test_a_mapping_typed_under_an_intermediate_display_name_still_matches(monkeypatch, builtin_state):
    monkeypatch.setattr("app.catalog.load_rate_overrides",
                        lambda: {"Video - Pre-roll (OLV)": {"name": "Pre-roll C"}})
    monkeypatch.setattr("app.catalog.all_product_aliases",
                        lambda: {"Pre-roll B": "Pre-roll C", "Pre-roll A": "Pre-roll B"})
    from app.catalog import name_variants
    assert {"Video - Pre-roll (OLV)", "Pre-roll C", "Pre-roll B"} <= name_variants("Pre-roll C")


# --- Roadblocks: what the sheet may remove ----------------------------------------------------------------------------

POLITICAL = [c for c in rs.builtin_categories() if c.name == "Political"]


@pytest.mark.parametrize("sentence", [
    "Political advertisers are not permitted to use lookalike audiences on Meta.",
    "Political ads are not allowed to use detailed targeting on Meta.",
    "Political ads do not support interest-based targeting.",
    "Political spots are not allowed to exceed 30 seconds on some CTV apps.",
    "Political ads are frequently rejected during Meta's ad review for missing disclaimers.",
    "Political ads on Meta are blocked in the EU.",
    "Political ads aren't served in Washington state due to state law.",
])
def test_restrictions_on_how_a_vertical_may_run_are_kept_as_caveats(sentence):
    assert not rb._denies_category(sentence, POLITICAL)


@pytest.mark.parametrize("sentence", [
    "This product can't carry political ads.",
    "Political ads are off-limits on Display.",
    "Display is closed to political advertisers.",
    "We can't sell political ads on this product.",
    "Political ads are a no-go on Display.",
])
def test_more_ways_of_saying_no_are_caught(sentence):
    assert rb._denies_category(sentence, POLITICAL)


def test_a_mixed_risk_keeps_its_real_sentence_and_what_was_removed_is_recorded(builtin_state):
    items = _items("Video - Pre-roll (OLV)")
    matrix = rs.build_matrix(["Video - Pre-roll (OLV)"], ["Political"])
    result = {"overall_summary": "", "used_web_search": True, "error": None, "categories": [], "product_roadblocks": [{
        "product_name": "Video - Pre-roll (OLV)", "risk_level": "high",
        "risks": [{"issue": "Additional: approvals", "source": "s",
                   "detail": "Political ads are not allowed on OLV. Meta needs ID authorization before launch."}],
        "recommended_mitigation": "Submit creative early."}]}
    rb._apply_matrix(result, matrix, items)
    entry = result["product_roadblocks"][0]
    assert entry["risks"][0]["detail"] == "Meta needs ID authorization before launch."
    assert entry["recommended_mitigation"] == "Submit creative early."
    assert [c["text"] for c in result["removed_claims"]] == ["Political ads are not allowed on OLV."]
    assert result["removed_claims"][0]["product"] == "Video - Pre-roll (OLV)"


def test_untouched_text_keeps_its_line_breaks_and_does_not_lower_the_risk(builtin_state):
    items = _items("Video - Pre-roll (OLV)")
    matrix = rs.build_matrix(["Video - Pre-roll (OLV)"], ["Political"])
    mitigation = "1. Submit early.\n2. Add the Paid for by disclosure."
    result = {"overall_summary": "", "used_web_search": True, "error": None, "categories": [], "product_roadblocks": [{
        "product_name": "Video - Pre-roll (OLV)", "risk_level": "high",
        "risks": [{"issue": "Additional: lead time", "detail": "Approvals take a week.", "source": "s"}],
        "recommended_mitigation": mitigation}]}
    rb._apply_matrix(result, matrix, items)
    entry = result["product_roadblocks"][0]
    assert entry["recommended_mitigation"] == mitigation          # not flattened onto one line
    assert entry["risk_level"] == "high" and result["removed_claims"] == []


def test_denial_removal_does_not_cut_sentences_at_abbreviations():
    text = ("Political ads are not allowed on Display. Plan DMA targeting, e.g. Fresno. "
            "Use state-level targeting in the U.S. Hispanic market.")
    cleaned, gone = rb._strip_denials(text, POLITICAL)
    assert gone == ["Political ads are not allowed on Display."]
    assert cleaned == "Plan DMA targeting, e.g. Fresno. Use state-level targeting in the U.S. Hispanic market."


def test_a_model_shortened_product_name_is_filtered_like_the_real_one(builtin_state):
    name = "eDigital Network Display - Standard IAB"
    items = _items(name)
    matrix = rs.build_matrix([name], ["Political"])
    result = {"overall_summary": "", "used_web_search": True, "error": None, "categories": [], "product_roadblocks": [{
        "product_name": "eDigital Display", "risk_level": "high",
        "risks": [{"issue": "Political ads not allowed", "detail": "Political ads are not allowed on Display.", "source": "s"}],
        "recommended_mitigation": "Do not run political ads on Display."}]}
    rb._apply_matrix(result, matrix, items)
    assert len(result["product_roadblocks"]) == 1
    entry = result["product_roadblocks"][0]
    assert entry["product_name"] == name and entry["matrix"]["verdict"] == "allowed"
    assert entry["risks"] == [] and entry["recommended_mitigation"] == "" and entry["risk_level"] == "low"


def test_a_product_the_model_repeats_is_merged_so_neither_copy_escapes_the_filter(builtin_state):
    items = _items("Video - Pre-roll (OLV)")
    matrix = rs.build_matrix(["Video - Pre-roll (OLV)"], ["Political"])
    denial = {"issue": "Political ads not allowed", "detail": "Political ads are not allowed on OLV.", "source": "s"}
    caveat = {"issue": "Additional: lead time", "detail": "Approvals take a week.", "source": "s"}
    result = {"overall_summary": "", "used_web_search": True, "error": None, "categories": [], "product_roadblocks": [
        {"product_name": "Video - Pre-roll (OLV)", "risk_level": "low", "risks": [caveat], "recommended_mitigation": ""},
        {"product_name": "video - pre-roll (olv)", "risk_level": "high", "risks": [denial], "recommended_mitigation": ""}]}
    rb._apply_matrix(result, matrix, items)
    assert len(result["product_roadblocks"]) == 1
    assert [r["issue"] for r in result["product_roadblocks"][0]["risks"]] == ["Additional: lead time"]


def test_a_summary_denial_is_removed_when_the_vertical_is_allowed_on_every_product_even_unnamed(builtin_state):
    items = _items("Video - Pre-roll (OLV)", "eDigital Network Display - Standard IAB")
    matrix = rs.build_matrix([li.product_name for li in items], ["Political"])
    result = {"overall_summary": "Political ads are not accepted on this campaign's digital products. Budget is tight.",
              "used_web_search": True, "error": None, "categories": [], "product_roadblocks": []}
    rb._apply_matrix(result, matrix, items)
    assert result["overall_summary"] == "Budget is tight."


# --- exports -----------------------------------------------------------------------------------------------------------

def test_the_pptx_row_estimate_counts_hard_line_breaks():
    from app.services import pptx_builder as pb
    assert pb._estimate_wrapped_lines("a\nb\nc\nd\ne") == 5
    assert pb._estimate_wrapped_lines("x" * 100) == 3
    assert pb._estimate_wrapped_lines("") == 1


def test_the_pptx_impressions_sub_line_follows_the_time_unit():
    from app.services import pptx_builder as pb
    from app.catalog import by_name
    product = by_name("Video - Pre-roll (OLV)")
    li = LineItem(product_name=product.name, monthly_budget=2010.0, months=3, id="x")
    assert pb._product_cell_text(product, li).splitlines()[1].endswith("impressions/mo")
    assert pb._product_cell_text(product, li, "/wk").splitlines()[1].endswith("impressions/wk")


def test_the_pptx_closing_copy_keeps_the_not_a_guarantee_wording():
    from app.services import pptx_builder as pb
    assert "not a guarantee of delivery" in pb._signature_authorization_text("Acme", False)


def test_the_pptx_total_row_has_no_empty_run_that_falls_back_to_the_18pt_default(tmp_path):
    pptx = pytest.importorskip("pptx")
    from app.services import pptx_builder
    from app.catalog import by_name
    items = _items("Video - Pre-roll (OLV)")
    out = tmp_path / "t.pptx"
    pptx_builder.build_signature_deck(ProposalRequest(client_name="Acme"), [{"name": "A", "products": [by_name(items[0].product_name)], "line_items": items}],
                                      out, gross=False, proposal_title="T")
    table = next(sh for sl in pptx.Presentation(str(out)).slides for sh in sl.shapes if sh.has_table).table
    assert table.cell(2, 0).text == "TOTAL" and table.cell(2, 1).text == " "


def test_avails_only_added_value_lines_leave_the_sov_cell_blank(tmp_path):
    paid = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="p")
    gift = LineItem(product_name="eDigital Network Display - Standard IAB", monthly_budget=0.0, months=3, id="g", is_added_value=True)
    out = tmp_path / "ao.xlsx"
    generate_proposal(ProposalRequest(client_name="Acme", request_type="New Business", start_date="2026-09-01", end_date="2026-11-30",
                                      total_months=3), [paid, gift], out, force_tabs={"net": False, "avails_only": True},
                      avails_data={"p": {"max_imps": 100000, "max_spend": 5000}, "g": {"max_imps": 200000, "max_spend": 9000}})
    import openpyxl
    ws = next(w for w in openpyxl.load_workbook(out).worksheets if "Avails" in w.title)
    row = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(row=r, column=3).value or "").startswith("eDigital Network Display"))
    assert ws.cell(row=row, column=13).value is None


def test_a_tier_whose_lines_disagree_on_months_is_warned_about(tmp_path):
    lis = [LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="a"),
           LineItem(product_name="eDigital Network Display - Standard IAB", monthly_budget=1000.0, months=2, id="b")]
    summary = generate_proposal(ProposalRequest(client_name="Acme", request_type="New Business", start_date="2026-09-01",
                                                end_date="2026-11-30", total_months=3), lis, tmp_path / "w.xlsx")
    assert any("different month counts" in w for w in summary.get("warnings", []))


def test_add_on_descriptions_lose_the_internal_variants_tail(tmp_path):
    out = tmp_path / "addon.xlsx"
    li = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="a")
    addon = SimpleNamespace(product_name="Video - Pre-roll (OLV)", amount=100.0, notes_override=None)
    generate_proposal(ProposalRequest(client_name="Acme", request_type="New Business", start_date="2026-09-01",
                                      end_date="2026-11-30", total_months=3), [li], out, addons=[addon])
    import openpyxl
    wb = openpyxl.load_workbook(out)
    texts = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value]
    assert not any("contact planning for exact tier" in t for t in texts)


# --- text + blurbs -------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("text", [
    "Ads reach the U.S. market. Plain sentence.",
    "1. Submit early.\n2. Add disclosure.",
    "Audio plays in streams. (Edison Research, 2025 Infinite Dial)",
    "It grew 3.5 million. See https://example.com/x. Done.",
    "",
    "no terminal punctuation",
])
def test_split_sentences_is_lossless(text):
    assert "".join(split_sentences(text)) == text


def test_numbered_steps_and_abbreviations_stay_whole_and_citations_stay_attached():
    assert split_sentences("1. Submit early.\n2. Add disclosure.") == ["1. Submit early.\n", "2. Add disclosure."]
    assert len(split_sentences("Ads reach the U.S. market. Plain sentence.")) == 2
    assert len(split_sentences("Audio plays in streams. (IAB 2025 Outlook)")) == 1
    assert len(split_sentences("Audio plays in streams. [Nielsen, 2025]")) == 1


def test_a_citation_of_any_shape_goes_with_the_fact_it_supported():
    for citation in ("(Edison Research, 2025 Infinite Dial)", "(IAB 2025 Outlook)", "[Nielsen, 2025]"):
        out = _tighten_blurb(f"Audio ads play inside radio streams. Listeners can skip some ads in skippable formats. {citation}", "Audio")
        assert out == "Audio ads play inside radio streams.", citation


def test_a_citation_survives_with_its_fact_even_at_the_word_cap():
    long_fact = "Word " * 66 + "end."
    out = _tighten_blurb(f"Intro sentence. {long_fact} (eMarketer, 2025)", "Any")
    assert out == "Intro sentence."          # the fact and its source are one unit: both fit or neither does


def test_version_of_only_trips_on_a_product_variant_not_on_ordinary_prose():
    prose = "Streaming audio is the online version of radio, reaching listeners through station websites and apps."
    assert _tighten_blurb(prose, "Audio - EVC Audio Streaming") == prose
    assert _tighten_blurb("This is the Hispanic version of our CTV product. Plain sentence.", "CTV") == "Plain sentence."


def test_only_a_stated_language_exempts_a_product_from_the_language_rule():
    claim = "Ads run on Spanish-language sites and apps."
    assert _tighten_blurb(claim, "Entravision Plus - Hispanics CTV/OTT") == ""
    assert _tighten_blurb(claim, "Netflix - Spanish Content :30s Ads") == claim


def test_a_malformed_blurb_entry_costs_that_blurb_not_the_whole_enrichment():
    raw = json.dumps({"campaign_name": "C", "internal_email_body": "b", "client_email_body": "b", "product_blurbs": [
        {"product_name": "AudioEngage", "blurb": ["a", "b"]}, {"product_name": 5, "blurb": "x"},
        {"product_name": "Search - SEM", "blurb": "Text ads that appear on Google."}]})
    out = _parse_response(raw, ProposalRequest(client_name="Acme"), [])
    assert [(b.product_name, b.blurb) for b in out.product_blurbs] == [("Search - SEM", "Text ads that appear on Google.")]
    assert out.campaign_name == "C"


def test_the_prompt_tells_the_model_how_to_use_the_knowledge_base_and_not_to_restate_the_rate_card():
    from app.services import ai_enricher
    prompt = ai_enricher._build_prompt(ProposalRequest(client_name="Acme"), _items("Search - SEM"))
    assert "never copy its numbers, audience, language or geography wording into a blurb" in prompt
    assert "do not restate it" in prompt
    assert "leave the number out" in prompt and "fold it into the blurb as directional context" not in prompt
    assert "(for the product_blurbs field only" in prompt


def test_a_blurb_sentence_that_repeats_the_requests_geo_or_demo_is_dropped():
    req = ProposalRequest(client_name="Acme", geo="Los Angeles DMA, Orange County", demo="Hispanic adults 25-54")
    raw = json.dumps({"campaign_name": "C", "internal_email_body": "b", "client_email_body": "b", "product_blurbs": [
        {"product_name": "Display - Standard", "blurb": "Banner ads shown across a network of sites. Reach shoppers in Los Angeles every day. "
                                                        "Built for Hispanic adults 25-54 on mobile."},
        {"product_name": "Search - SEM", "blurb": "Text ads on Google when people search. Entravision's hispanic display network helps too."}]})
    out = _parse_response(raw, req, [])
    blurbs = {b.product_name: b.blurb for b in out.product_blurbs}
    assert blurbs["Display - Standard"] == "Banner ads shown across a network of sites."
    # generic wording that merely contains a word of the demo is not a repeat of it
    assert "hispanic display network" in blurbs["Search - SEM"]
