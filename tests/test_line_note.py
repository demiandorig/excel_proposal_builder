"""
The per-line "line note" (field `line_modifier`): a short plain-text label a planner types on a Step 04 line ("Hispanic,
Phoenix only") so two lines of the SAME product can be told apart.

  * the export shows it in the brand purple, in a small font, directly under the product name + objective, on every tab
    that lists the lines (Net, with-sections, Gross, Avails-Only, the Monthly Breakdown tab) — and a line WITHOUT a note
    exports exactly as it did before the field existed;
  * the PowerPoint decks, the Gamma outline and the email prompts carry it as a label (never a number);
  * the other steps show it read-only for reference; it survives option tabs, duplication, Suggest Mix, draft saves and
    reopen (an older save without the field opens with empty notes);
  * it is sanitized in ONE place per language (trimmed, control characters out, capped at 100, plain text) and always
    HTML-escaped where it is rendered.

It is NOT the "Planner note" (`notes_override`), which goes to the export's Planner Notes column — the two are independent.

No OpenAI / Google / Gamma call is made anywhere. The JS helpers are run in Node (skipped when Node is missing); the full
click-through of the real SPA was driven separately with Playwright (see the session notes), not from pytest.
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from types import SimpleNamespace

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText, TextBlock

from app import excel_template as et
from app import main
from app.services import ai_enricher, email_totals, pptx_builder
from app.services.ai_enricher import ProposalEnrichment
from app.services.line_modifier import (
    LINE_MODIFIER_MAX_LEN, clean_line_modifier, clean_line_modifier_or_none, line_modifier_of,
)
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import LineItem, generate_proposal
from app.services.recommender import recommend_line_items

ROOT = Path(__file__).resolve().parent.parent
APP_JS = (ROOT / "app" / "static" / "app.js").read_text(encoding="utf-8")
STYLES_CSS = (ROOT / "app" / "static" / "styles.css").read_text(encoding="utf-8")

PURPLE = "FF910E95"
NOTE_A = "Hispanic, Phoenix only"
NOTE_B = "Brand terms"
NOTE_AV = "Bonus: weekend geofence"

SEM = "Search - SEM"
VIDEO = "Video - Pre-roll (OLV)"
GEOFENCE = "Display - Geo Fence"
EMAIL = "Email Campaigns - Display Re-targeting"


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.catalog.resolve_product_alias", lambda name: None)
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


# =============================================================================================================
# The sanitizer (one rule set, used by the API model, the dataclass, the exporters and the email prompts)
# =============================================================================================================

@pytest.mark.parametrize("raw,expected", [
    (None, ""),
    ("", ""),
    ("   ", ""),
    ("Hispanic, Phoenix only", "Hispanic, Phoenix only"),
    ("   Hispanic,   Phoenix\tonly \n", "Hispanic, Phoenix only"),                 # trimmed, whitespace collapsed
    ("line one\nline two\r\nline three", "line one line two line three"),          # never multi-line (it is one cell line)
    ("nul\x00byte", "nul byte"),                                                   # C0 control characters
    ("del\x7fand\x9fc1", "del and c1"),                                            # DEL / C1
    ("zero​width‍marks", "zero width marks"),
    ("bidi‮flip", "bidi flip"),
    ("sep para end", "sep para end"),
    ("bom﻿here", "bom here"),
    ("lone\ud800surrogate", "lone surrogate"),                                     # would break UTF-8 when the workbook is saved
    ("non breaking　space", "non breaking space"),
    ("<b>bold</b> & \"quoted\" 'x'", "<b>bold</b> & \"quoted\" 'x'"),              # plain text: markup is escaped on RENDER, not stripped
    ("=SUM(A1:A9)", "=SUM(A1:A9)"),                                                # never a formula: it is only ever a rich-text run
    (12345, "12345"),
    (True, ""),
    (3.5, "3.5"),
])
def test_clean_line_modifier(raw, expected):
    assert clean_line_modifier(raw) == expected


def test_the_cap_is_a_hundred_characters_and_never_leaves_a_trailing_space():
    assert LINE_MODIFIER_MAX_LEN == 100
    assert clean_line_modifier("x" * 150) == "x" * 100
    assert len(clean_line_modifier("é" * 300)) == 100
    assert len(clean_line_modifier("😀" * 300)) == 100                              # counted in characters, not UTF-16 units
    cut_on_a_space = "y" * 99 + " " + "z" * 20
    assert clean_line_modifier(cut_on_a_space) == "y" * 99                          # the cut lands on the space: it is trimmed too


def test_cleaning_is_idempotent():
    for raw in ["  a\tb  ", "x" * 150, "😀" * 150, "a​b", None, "  "]:
        once = clean_line_modifier(raw)
        assert clean_line_modifier(once) == once


def test_line_modifier_of_reads_a_dataclass_a_dict_or_nothing():
    assert line_modifier_of(LineItem(product_name=SEM, monthly_budget=100, line_modifier="  Phoenix  ")) == "Phoenix"
    assert line_modifier_of({"product_name": SEM, "line_modifier": " Phoenix\n"}) == "Phoenix"
    assert line_modifier_of({"product_name": SEM}) == ""                            # a line from before the field existed
    assert line_modifier_of(SimpleNamespace(product_name=SEM)) == ""                # an object that never had the attribute
    assert line_modifier_of(None) == ""
    assert clean_line_modifier_or_none("   ") is None and clean_line_modifier_or_none(" x ") == "x"


# =============================================================================================================
# The model, the dataclass and everything that persists them
# =============================================================================================================

def test_line_item_model_field_defaults_to_none_and_is_in_the_dump():
    m = main.LineItemModel(product_name=SEM, monthly_budget=100)
    assert m.line_modifier is None
    assert "line_modifier" in m.model_dump() and m.model_dump()["line_modifier"] is None
    assert "line_modifier" not in m.model_fields_set                                # an old payload simply never sent it
    assert "line_modifier" in main.LineItemModel.model_json_schema()["properties"]


def test_line_item_model_sanitizes_whatever_arrives():
    def clean(v):
        return main.LineItemModel(product_name=SEM, monthly_budget=1, line_modifier=v).line_modifier
    assert clean("  Hispanic,\tPhoenix only \n") == NOTE_A
    assert clean("x" * 500) == "x" * 100
    assert clean("   ") is None and clean("") is None and clean(None) is None       # blank = no note
    assert clean("nul\x00 byte​") == "nul byte"
    assert clean(7) == "7"
    assert clean("<img src=x onerror=alert(1)>") == "<img src=x onerror=alert(1)>"  # kept as plain text; escaped on render


def test_an_older_payload_without_the_field_is_accepted_everywhere():
    old_line = {"id": "a", "product_name": SEM, "monthly_budget": 100.0, "months": 3, "notes_override": "keep me"}
    tier = main.TierModel(label="A", line_items=[old_line])
    assert tier.line_items[0].line_modifier is None and tier.line_items[0].notes_override == "keep me"
    body = main.GenerateRequest(request={}, line_items=[old_line], tiers=[{"label": "A", "line_items": [old_line]}])
    assert body.line_items[0].line_modifier is None and body.tiers[0].line_items[0].line_modifier is None
    draft = main.DraftSaveRequest(request={}, line_items=[old_line])
    assert draft.line_items[0].line_modifier is None


def test_line_modifier_and_notes_override_are_two_independent_fields():
    m = main.LineItemModel(product_name=SEM, monthly_budget=1, notes_override="Client list to upload", line_modifier=NOTE_A)
    assert (m.notes_override, m.line_modifier) == ("Client list to upload", NOTE_A)
    only_planner_note = main.LineItemModel(product_name=SEM, monthly_budget=1, notes_override="Client list to upload")
    assert only_planner_note.line_modifier is None


def test_the_line_item_dataclass_holds_a_clean_note_and_exposes_it_to_asdict():
    from dataclasses import asdict
    li = LineItem(product_name=SEM, monthly_budget=100.0, months=3, line_modifier="  Hispanic,\tPhoenix only ")
    assert li.line_modifier == NOTE_A
    assert LineItem(product_name=SEM, monthly_budget=1).line_modifier is None
    assert LineItem(product_name=SEM, monthly_budget=1, line_modifier="  ").line_modifier is None
    assert asdict(li)["line_modifier"] == NOTE_A
    # Suggest Mix's server answer carries the key (None): the SPA decides what a recommended line keeps
    items = recommend_line_items(ProposalRequest(client_name="Acme", request_type="New Business", start_date="2026-11-24",
                                                 end_date="2026-12-31", total_months=2, campaign_goal="Awareness",
                                                 products_selected=[]), 15000.0, time_unit="month")
    assert items and all(asdict(i)["line_modifier"] is None for i in items)


def _tiers_with_notes():
    return [
        main.TierModel(label="A", name="Independent", line_items=[
            main.LineItemModel(id="a1", product_name=SEM, monthly_budget=1200.0, line_modifier=NOTE_A),
            main.LineItemModel(id="a2", product_name=SEM, monthly_budget=800.0, line_modifier=NOTE_B),
            main.LineItemModel(id="a3", product_name=VIDEO, monthly_budget=2000.0)]),
        main.TierModel(label="B", name="Democrat", line_items=[
            main.LineItemModel(id="b1", product_name=SEM, monthly_budget=1500.0, line_modifier="Spanish-language ads")]),
    ]


def _reopen_state(tiers, line_items=None):
    return main._build_reopen_state(
        request_dict={"client_name": "Acme"}, line_items=line_items or [], tiers=tiers, avails_data=None, strategy_brief=None,
        roadblocks=None, force_tabs=None, addons=[], raw_notion_text=None, enrichment=None, time_unit="month",
        monthly_distribution_mode="even", campaign_name_override=None, wizard_step=4)


def test_reopen_state_saves_and_restores_the_notes_for_every_option():
    state = json.loads(json.dumps(_reopen_state(_tiers_with_notes())))              # through JSONB and back
    a, b = state["tiers"]
    assert [l["line_modifier"] for l in a["line_items"]] == [NOTE_A, NOTE_B, None]
    assert [l["line_modifier"] for l in b["line_items"]] == ["Spanish-language ads"]
    # the legacy flat list (the first option's lines) carries them too
    assert [l["line_modifier"] for l in state["line_items"]] == [NOTE_A, NOTE_B, None]
    restored = main._with_current_product_names(state)
    assert [l["line_modifier"] for l in restored["tiers"][0]["line_items"]] == [NOTE_A, NOTE_B, None]


def test_a_draft_goes_through_the_same_state_builder():
    body = main.DraftSaveRequest(request={"client_name": "Acme"}, tiers=_tiers_with_notes())
    state = main._build_reopen_state(
        request_dict=body.request, line_items=body.line_items, tiers=body.tiers, avails_data=body.avails_data,
        strategy_brief=None, roadblocks=None, force_tabs=None, addons=[], raw_notion_text=None, enrichment=None,
        time_unit=body.time_unit, monthly_distribution_mode=body.monthly_distribution_mode, campaign_name_override=None,
        wizard_step=body.wizard_step)
    assert [l["line_modifier"] for l in state["tiers"][0]["line_items"]] == [NOTE_A, NOTE_B, None]


def test_reopen_endpoint_returns_notes_and_tolerates_an_old_save(monkeypatch):
    saved = json.loads(json.dumps(_reopen_state(_tiers_with_notes())))
    monkeypatch.setattr(main, "_get_proposal_metadata", lambda pid: {"reopen_state": saved, "status": "draft"})
    got = asyncio.run(main.reopen_proposal("p1"))
    assert [l["line_modifier"] for l in got["tiers"][0]["line_items"]] == [NOTE_A, NOTE_B, None]
    old = {"request": {}, "line_items": [{"id": "a", "product_name": SEM, "monthly_budget": 100.0, "months": 3}], "tiers": None}
    monkeypatch.setattr(main, "_get_proposal_metadata", lambda pid: {"reopen_state": old, "status": "generated"})
    got = asyncio.run(main.reopen_proposal("p1"))
    assert got["line_items"][0].get("line_modifier") is None                         # absent / null: the SPA turns it into ""


# =============================================================================================================
# The Excel export
# =============================================================================================================

REQ = ProposalRequest(client_name="Acme Corp", request_type="Avails + Proposal", start_date="2026-09-28",
                      end_date="2026-11-30", total_months=3, geo="Phoenix", agency_fee=0.15, demo="A18+")
ALL_TABS = {"net": True, "wsections": True, "gross": True, "avails_only": True, "dooh": False}


def _lines(*, notes=True, blank=False):
    """Two lines of ONE product with different notes, an un-noted line carrying a Step 06 breakdown, and an Added Value
    line with a note. notes=False: no line has the field at all; blank=True: every note is whitespace only."""
    def L(**kw):
        if not notes:
            kw.pop("line_modifier", None)
        elif blank and "line_modifier" in kw:
            kw["line_modifier"] = "  \n\t "
        return LineItem(**kw)
    lines = [
        L(id="l1", product_name=SEM, monthly_budget=1200.0, months=3, objective_override="Lead Generation",
          line_modifier=NOTE_A, notes_override="Client list to upload"),
        L(id="l2", product_name=SEM, monthly_budget=800.0, months=3, objective_override="Awareness", line_modifier=NOTE_B),
        L(id="l3", product_name=VIDEO, monthly_budget=2000.0, months=3, objective_override="Awareness"),
        L(id="l4", product_name=GEOFENCE, monthly_budget=0.0, months=3, is_added_value=True, added_value_pct=5,
          line_modifier=NOTE_AV),
    ]
    lines[2].monthly_allocations = {"2026-09": 500.0, "2026-10": 3000.0, "2026-11": 2500.0}
    return lines


def _generate(tmp_path, name, **kw):
    out = tmp_path / f"{name}.xlsx"
    generate_proposal(REQ, _lines(**kw), out, force_tabs=ALL_TABS, time_unit="month",
                      avails_data={"l1": {"max_imps": 100000, "max_spend": 3000}, "l3": {"max_imps": 400000, "max_spend": 6000}})
    return out


@pytest.fixture(scope="module")
def books(tmp_path_factory):
    # module-scoped: the catalog stub is applied inside the build via monkeypatch-free module state, so patch by hand
    mp = pytest.MonkeyPatch()
    mp.setattr("app.catalog.load_rate_overrides", lambda: {})
    mp.setattr("app.catalog.load_custom_products", lambda: [])
    mp.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    mp.setattr("app.catalog.resolve_product_alias", lambda name: None)
    mp.setattr("app.market_config.load_market_config", lambda: {})
    try:
        tmp = tmp_path_factory.mktemp("line_note_books")
        paths = {"noted": _generate(tmp, "noted"), "plain": _generate(tmp, "plain", notes=False),
                 "blank": _generate(tmp, "blank", blank=True)}
        yield {k: openpyxl.load_workbook(p, rich_text=True) for k, p in paths.items()}
    finally:
        mp.undo()


def _blocks(cell):
    """The rich-text runs of a cell as [(text, font-or-None), ...] — a plain string is one run with no font."""
    v = cell.value
    if isinstance(v, CellRichText):
        return [(p, None) if isinstance(p, str) else (p.text, p.font) for p in v]
    if hasattr(v, "ref") and hasattr(v, "text"):                  # an array formula: compare what it says, not its identity
        return [(f"{{{v.ref}}} {v.text}", None)]
    return [(v, None)]


def _purple_runs(ws):
    """{coordinate: text} for every small purple rich-text run on the sheet (the note's own style)."""
    found = {}
    for row in ws.iter_rows():
        for c in row:
            for text, font in _blocks(c):
                if font is not None and getattr(font.color, "rgb", None) == PURPLE and font.sz == 8:
                    found[c.coordinate] = text
    return found


def _name_cells(ws, col, products=(SEM, VIDEO, GEOFENCE)):
    """[(row, product_name)] for every line row, read off the name column in sheet order."""
    out = []
    for r in range(1, ws.max_row + 1):
        text = str(ws[f"{col}{r}"].value or "")
        for p in products:
            if text.startswith(p):
                out.append((r, p))
                break
    return out


LINE_SHEETS = {"Proposal A": "C", "Proposal A (wsections)": "C", "Proposal A (Gross)": "C", "Avails-Only": "C"}


def test_the_workbook_has_every_tab_that_lists_the_lines(books):
    assert set(books["noted"].sheetnames) >= set(LINE_SHEETS) | {"Monthly Breakdown"}


@pytest.mark.parametrize("sheet", sorted(LINE_SHEETS))
def test_the_note_is_a_small_purple_last_line_under_the_name_and_objective(books, sheet):
    ws = books["noted"][sheet]
    rows = _name_cells(ws, "C")
    names = [p for _, p in rows]
    assert names.count(SEM) == 2 and VIDEO in names and GEOFENCE in names
    purple = _purple_runs(ws)
    assert set(purple.values()) == {"\n" + NOTE_A, "\n" + NOTE_B, "\n" + NOTE_AV}
    for (r, product), want in zip(rows, [NOTE_A, NOTE_B, None, NOTE_AV]):
        cell = ws[f"C{r}"]
        text = str(cell.value)
        if want is None:
            assert not any(coord == cell.coordinate for coord in purple), "an un-noted line must carry no purple run"
            continue
        # the note is the LAST line of the cell, directly after the name and the objective
        assert text.splitlines()[-1] == want and text.splitlines()[0] == product and len(text.splitlines()) == 3
        runs = _blocks(cell)
        assert runs[-1][0] == "\n" + want
        font = runs[-1][1]
        assert font.color.rgb == PURPLE == et.ENTRAVISION_PURPLE and font.sz == 8 and not font.b


def test_the_objective_line_and_the_name_are_untouched_by_the_note(books):
    ws = books["noted"]["Proposal A"]
    (r1, _), (r2, _), (r3, _), _ = _name_cells(ws, "C")
    # name (cell font) + the objective sub-line in the same muted grey the cell always used
    first = _blocks(ws[f"C{r1}"])
    assert first[0] == (SEM, None) and first[1][0] == "\nLead Generation" and first[1][1].color.rgb == et.DARK_GREY
    # the un-noted line is exactly the original two-part value
    third = _blocks(ws[f"C{r3}"])
    assert len(third) == 2 and third[0] == (VIDEO, None) and third[1][0] == "\nAwareness"


def test_the_monthly_breakdown_tab_shows_it_under_the_name_for_the_lines_it_lists(books):
    ws = books["noted"]["Monthly Breakdown"]
    rows = _name_cells(ws, "B")
    assert [p for _, p in rows] == [SEM, SEM, VIDEO]                                  # Added Value lines are not on this tab
    purple = _purple_runs(ws)
    assert sorted(purple.values()) == sorted(["\n" + NOTE_A, "\n" + NOTE_B])
    for (r, product), want in zip(rows, [NOTE_A, NOTE_B, None]):
        text = str(ws[f"B{r}"].value)
        assert (text == f"{product}\n{want}") if want else (text == product)
    # the un-noted row stays a plain string cell
    assert ws[f"B{rows[2][0]}"].value == VIDEO


def test_no_note_means_no_purple_anywhere(books):
    for ws in books["plain"].worksheets:
        assert _purple_runs(ws) == {}, ws.title


def _signature(ws, skip=frozenset()):
    """Everything a reader sees about a sheet except `skip` cells: values (rich-text runs included), fonts, fills,
    borders, alignment, number formats, column widths, merges, print setup, conditional formats."""
    cells = {}
    for row in ws.iter_rows():
        for c in row:
            if (c.row, c.column) in skip:
                continue
            value = [(t, repr(f)) for t, f in _blocks(c)]
            cells[(c.row, c.column)] = (value, repr(c.font), repr(c.fill), repr(c.border), repr(c.alignment), c.number_format)
    return {
        "dims": ws.dimensions, "cells": cells,
        "cols": {k: v.width for k, v in ws.column_dimensions.items()},
        "merged": sorted(str(m) for m in ws.merged_cells.ranges),
        "print_area": ws.print_area, "freeze": ws.freeze_panes, "grid": ws.sheet_view.showGridLines,
        "cf": sorted((str(rng.sqref), [r.type for r in rules]) for rng, rules in ws.conditional_formatting._cf_rules.items()),
    }


def _row_heights(ws):
    return {r: d.height for r, d in ws.row_dimensions.items()}


def test_blank_notes_export_exactly_like_no_notes_at_all(books):
    """A whitespace-only note is no note: the two workbooks are the same cell for cell, style for style, row for row."""
    assert books["blank"].sheetnames == books["plain"].sheetnames
    for name in books["plain"].sheetnames:
        a, b = books["plain"][name], books["blank"][name]
        assert _signature(a) == _signature(b), name
        assert _row_heights(a) == _row_heights(b), name


@pytest.mark.parametrize("sheet,col", [("Proposal A", "C"), ("Proposal A (wsections)", "C"), ("Proposal A (Gross)", "C"),
                                       ("Avails-Only", "C"), ("Monthly Breakdown", "B")])
def test_only_the_noted_name_cells_and_their_row_heights_differ(books, sheet, col):
    """Lines (and everything else on the sheet — layout, merges, widths, formulas, formats, print area) are identical to
    the same workbook without notes; the only differences are the noted lines' name cells, and a noted row may only
    have GROWN."""
    noted, plain = books["noted"][sheet], books["plain"][sheet]
    noted_rows = {r for r, _ in _name_cells(noted, col) if str(noted[f"{col}{r}"].value).count("\n") >= (1 if col == "B" else 2)
                  and any(t.strip() in (NOTE_A, NOTE_B, NOTE_AV) for t, f in _blocks(noted[f"{col}{r}"]))}
    assert noted_rows, "the scenario must have noted rows on this sheet"
    skip = {(r, openpyxl.utils.column_index_from_string(col)) for r in noted_rows}
    sig_n, sig_p = _signature(noted, skip), _signature(plain, skip)
    assert sig_n == sig_p
    hn, hp = _row_heights(noted), _row_heights(plain)
    for r in set(hn) | set(hp):
        if r in noted_rows:
            assert (hn.get(r) or 0) >= (hp.get(r) or 0)
        else:
            assert hn.get(r) == hp.get(r), (sheet, r)


def test_formulas_and_numbers_are_identical_with_and_without_notes(books):
    for name in books["plain"].sheetnames:
        for row in books["plain"][name].iter_rows():
            for c in row:
                other = books["noted"][name][c.coordinate]
                if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("=")):
                    assert other.value == c.value, (name, c.coordinate)


def test_the_planner_note_still_goes_to_its_column_and_not_under_the_name(books):
    ws = books["noted"]["Proposal A"]
    (r1, _), *_ = _name_cells(ws, "C")
    row_values = [str(c.value) for c in ws[r1] if c.value is not None]
    # the Planner Notes COLUMN holds the planner note (never the line note) ...
    assert any("Client list to upload" in v for v in row_values)
    assert not any(NOTE_A in v for v in row_values if v != str(ws[f"C{r1}"].value))
    # ... and the name cell holds the line note (never the planner note)
    assert NOTE_A in str(ws[f"C{r1}"].value) and "Client list to upload" not in str(ws[f"C{r1}"].value)


def test_row_height_helper_only_grows_and_only_for_a_note():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.column_dimensions["C"].width = 38
    ws.row_dimensions[5].height = 45
    et.fit_row_to_line_modifier(ws, 5, "Search - SEM\nLead Generation", None)
    et.fit_row_to_line_modifier(ws, 5, "Search - SEM\nLead Generation", "  ")
    assert ws.row_dimensions[5].height == 45                                          # no note: untouched
    et.fit_row_to_line_modifier(ws, 5, "Search - SEM\nLead Generation", NOTE_A)
    assert ws.row_dimensions[5].height == 45                                          # roomy row: never shrunk
    et.fit_row_to_line_modifier(ws, 6, "Search - SEM\nLead Generation", NOTE_A)       # no height yet: gets what it needs
    assert ws.row_dimensions[6].height == et.line_modifier_row_height("Search - SEM\nLead Generation", NOTE_A, 38) > 30
    long_name = "Entravision Plus - VIX 360 (Includes VIX + UnivisionNow + Univision.com / Univision Apps)\nAwareness"
    tall = et.line_modifier_row_height(long_name, "w" * 100, 38)
    assert tall > et.line_modifier_row_height("Search - SEM\nAwareness", "w", 38)
    ws.row_dimensions[7].height = 30
    et.fit_row_to_line_modifier(ws, 7, long_name, "w" * 100)
    assert ws.row_dimensions[7].height == tall


def test_a_note_cannot_turn_the_cell_into_a_formula_or_inject_xml(tmp_path):
    nasty = "=1+1 </t><x> & \"q\""
    line = LineItem(id="a", product_name=SEM, monthly_budget=1000.0, months=3, line_modifier=nasty)
    out = tmp_path / "nasty.xlsx"
    generate_proposal(REQ, [line], out, force_tabs=ALL_TABS, time_unit="month")
    wb = openpyxl.load_workbook(out, rich_text=True)
    cell = wb["Proposal A"]["C19"]
    assert cell.data_type != "f" and str(cell.value).splitlines()[-1] == nasty.strip()
    assert openpyxl.load_workbook(out)["Avails-Only"]["C12"].data_type != "f"


def test_the_template_builder_without_line_items_is_unchanged(tmp_path):
    out = tmp_path / "template.xlsx"
    et.build_workbook(out)
    wb = openpyxl.load_workbook(out, rich_text=True)
    for ws in wb.worksheets:
        assert _purple_runs(ws) == {}


def test_the_purple_is_the_apps_own_brand_purple():
    assert et.ENTRAVISION_PURPLE == PURPLE
    assert re.search(r"--accent:\s*#910e95", STYLES_CSS)                              # the UI's purple (note text uses it)
    assert re.search(r"^\.line-note \{[^}]*color: var\(--accent\)", STYLES_CSS, re.M | re.S)


# =============================================================================================================
# /api/generate end to end: what the endpoint hands each consumer
# =============================================================================================================

def _run_generate(monkeypatch, tmp_path, body):
    saved, spy = {}, {}
    monkeypatch.setattr(main, "PROPOSALS_DIR", tmp_path)
    monkeypatch.setattr(main, "_get_next_short_id", lambda: "EVC-1")
    monkeypatch.setattr(main, "_resolve_draftable_proposal_id", lambda requested: "pid-1")
    monkeypatch.setattr(main, "_save_proposal_metadata", lambda **kw: saved.update(kw))
    monkeypatch.setattr(main.proposal_files_svc, "save_file", lambda *a, **k: None)

    def fake_enrich(req, lines, short_id, **kw):
        spy["enrich_lines"] = lines
        spy["enrich_tiers"] = kw.get("tiers")
        spy["enrich_totals"] = kw.get("totals")
        return ProposalEnrichment(campaign_name="Holiday")
    monkeypatch.setattr(main.ai_enricher, "enrich_proposal", fake_enrich)
    monkeypatch.setattr(main.disclaimers_svc, "resolve_for_tiers", lambda req, tiers: {})
    monkeypatch.setattr(main.disclaimers_svc, "union_for_ppt", lambda d: [])
    monkeypatch.setattr(main.pptx_builder, "build_signature_deck",
                        lambda req, tiers, path, **kw: spy.setdefault("deck_tiers", tiers) and None)
    monkeypatch.setattr(main.docx_builder, "build_client_email_docx", lambda **kw: None)
    request = SimpleNamespace(state=SimpleNamespace(user={"email": "planner@entravision.com", "is_admin": False}),
                              client=SimpleNamespace(host="127.0.0.1"), headers={"user-agent": "pytest"})
    result = asyncio.run(main.generate(body, request))
    return result, saved, spy


def _generate_body(tiers=None, **kw):
    fields = dict(request={"client_name": "Acme", "start_date": "2026-11-24", "end_date": "2026-12-31", "total_months": 2,
                           "agency_fee": 0.15, "request_type": "New Business"}, time_unit="month")
    if tiers is not None:
        fields["tiers"] = tiers
    else:
        fields["line_items"] = [main.LineItemModel(id="a", product_name=EMAIL, monthly_budget=500.0, months=2, line_modifier=NOTE_A),
                                main.LineItemModel(id="b", product_name=EMAIL, monthly_budget=300.0, months=2, line_modifier=NOTE_B)]
    fields.update(kw)
    return main.GenerateRequest(**fields)


def test_generate_puts_the_notes_in_the_workbook_the_reopen_state_and_what_the_ai_and_the_deck_are_given(monkeypatch, tmp_path):
    result, saved, spy = _run_generate(monkeypatch, tmp_path, _generate_body())
    ws = openpyxl.load_workbook(tmp_path / result["filename"], rich_text=True)["Proposal A"]
    purple = _purple_runs(ws)
    assert sorted(purple.values()) == sorted(["\n" + NOTE_A, "\n" + NOTE_B])
    assert [l["line_modifier"] for l in saved["reopen_state"]["line_items"]] == [NOTE_A, NOTE_B]
    # the dataclass lines reach the deck builder with their notes
    deck_lines = spy["deck_tiers"][0]["line_items"]
    assert [l.line_modifier for l in deck_lines] == [NOTE_A, NOTE_B]
    # the figures the emails are written from are keyed by line and carry the note as a label
    lines = spy["enrich_totals"].options[0].lines
    assert [l.modifier for l in lines] == [NOTE_A, NOTE_B] and [l.total for l in lines] == [1000.0, 600.0]


def test_generate_with_several_options_keeps_each_options_notes(monkeypatch, tmp_path):
    tiers = _tiers_with_notes()
    body = _generate_body(tiers=[tiers[0].model_copy(update={"line_items": [
        main.LineItemModel(id="a1", product_name=EMAIL, monthly_budget=500.0, months=2, line_modifier=NOTE_A)]}),
        tiers[1].model_copy(update={"line_items": [
            main.LineItemModel(id="b1", product_name=EMAIL, monthly_budget=700.0, months=2, line_modifier="Option B only")]})])
    result, saved, spy = _run_generate(monkeypatch, tmp_path, body)
    assert [[l["line_modifier"] for l in t["line_items"]] for t in saved["reopen_state"]["tiers"]] == [[NOTE_A], ["Option B only"]]
    wb = openpyxl.load_workbook(tmp_path / result["filename"], rich_text=True)
    runs = {ws.title: sorted(_purple_runs(ws).values()) for ws in wb.worksheets if _purple_runs(ws)}
    assert sorted(v for vals in runs.values() for v in vals).count("\n" + NOTE_A) >= 1
    assert sorted(v for vals in runs.values() for v in vals).count("\nOption B only") >= 1
    assert [o.lines[0].modifier for o in spy["enrich_totals"].options] == [NOTE_A, "Option B only"]


def test_the_email_revision_endpoint_gets_the_notes_too(monkeypatch):
    seen = {}
    monkeypatch.setattr(main.ai_enricher, "reprompt_emails", lambda *a, **k: seen.update(args=a, kwargs=k) or {
        "internal_email_subject": "s", "internal_email_body": "b", "client_email_subject": "s", "client_email_body": "b"})
    monkeypatch.setattr(main, "_get_proposal_metadata", lambda pid: None)
    body = main.RepromptEmailsRequest(
        request={"client_name": "Acme", "start_date": "2026-11-24", "end_date": "2026-12-31", "agency_fee": 0.15},
        line_items=[main.LineItemModel(id="a", product_name=EMAIL, monthly_budget=500.0, months=2, line_modifier=NOTE_A),
                    main.LineItemModel(id="b", product_name=EMAIL, monthly_budget=300.0, months=2, line_modifier=NOTE_B)],
        tiers=[main.TierModel(label="A", line_items=[
            main.LineItemModel(id="a", product_name=EMAIL, monthly_budget=500.0, months=2, line_modifier=NOTE_A),
            main.LineItemModel(id="b", product_name=EMAIL, monthly_budget=300.0, months=2, line_modifier=NOTE_B)])],
        reprompt="shorter")
    asyncio.run(main.reprompt_emails("pid", body))
    lines = seen["args"][1]
    assert [l.line_modifier for l in lines] == [NOTE_A, NOTE_B]
    assert [l.modifier for l in seen["kwargs"]["totals"].options[0].lines] == [NOTE_A, NOTE_B]


# =============================================================================================================
# PowerPoint decks
# =============================================================================================================

def _product(name):
    from app.catalog import Product
    return Product(
        family="Display", name=name, short_label=name, proposal_description="Banner ads.", sizes="", buying_model="CPM",
        base_rate=10.0, estimated_impressions=False, discloses_impressions=True, minimum_spend=0.0,
        minimum_flight_days=(0, 0), sla_data_days=None, sla_creative_days=None, sla_activate_days=None,
        sla_total_days=None, media_allocation_pct=1.0, margin_upper=0.5, margin_lower=0.5, tech_platform="", wide_orbit_code="")


def _deck(tmp_path, notes, name="deck", unit="month"):
    pptx = pytest.importorskip("pptx")
    products = [_product("Product A"), _product("Product A"), _product("Product B")]
    lines = [LineItem(id="1", product_name="Product A", monthly_budget=2000.0, months=3, line_modifier=notes[0]),
             LineItem(id="2", product_name="Product A", monthly_budget=1000.0, months=3, line_modifier=notes[1]),
             LineItem(id="3", product_name="Product B", monthly_budget=500.0, months=3, line_modifier=notes[2])]
    path = tmp_path / f"{name}.pptx"
    ok = pptx_builder.build_signature_deck(
        ProposalRequest(client_name="Acme", start_date="2026-11-24", end_date="2026-12-31"),
        [{"name": "Option A", "products": products, "line_items": lines}], path, gross=False, proposal_title="T",
        banners=[], time_unit=unit)
    assert ok
    return pptx.Presentation(str(path))


def _product_cells(prs):
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_table:
                t = sh.table
                return [t.cell(r, 0) for r in range(1, len(t.rows)) if t.cell(r, 0).text.startswith("Product")], t
    raise AssertionError("no table")


def test_pptx_shows_the_note_right_under_the_product_name_in_purple(tmp_path):
    prs = _deck(tmp_path, [NOTE_A, NOTE_B, None])
    cells, table = _product_cells(prs)
    first = [p.text for p in cells[0].text_frame.paragraphs]
    assert first[0] == "Product A" and first[1] == NOTE_A and first[2].startswith("Est. ")        # name, note, impressions
    note_run = cells[0].text_frame.paragraphs[1].runs[0]
    assert str(note_run.font.color.rgb) == "910E95" and note_run.font.size.pt == 8.5
    assert str(cells[0].text_frame.paragraphs[2].runs[0].font.color.rgb) != "910E95"             # the impressions line keeps its muted grey
    # two lines of the same product are now different rows; the third has no note line
    assert [p.text for p in cells[1].text_frame.paragraphs][1] == NOTE_B
    assert [p.text for p in cells[2].text_frame.paragraphs][1].startswith("Est. ")


def test_pptx_without_notes_is_exactly_the_old_cell_text(tmp_path):
    prs = _deck(tmp_path, [None, None, None])
    cells, _ = _product_cells(prs)
    assert [c.text for c in cells] == ["Product A\nEst. 200,000 impressions/mo", "Product A\nEst. 100,000 impressions/mo",
                                       "Product B\nEst. 50,000 impressions/mo"]
    for c in cells:
        for p in c.text_frame.paragraphs:
            assert all(str(r.font.color.rgb) != "910E95" for r in p.runs)


def test_pptx_books_extra_height_for_a_noted_row_and_never_less_than_before(tmp_path):
    plain = _deck(tmp_path, [None, None, None], "plain")
    noted = _deck(tmp_path, [NOTE_A, "w" * 100, None], "noted")
    _, tp = _product_cells(plain)
    _, tn = _product_cells(noted)
    assert tn.rows[1].height > tp.rows[1].height and tn.rows[2].height > tn.rows[1].height       # a 100-char note wraps to more lines
    assert tn.rows[3].height == tp.rows[3].height                                                  # the un-noted row is unchanged
    assert pptx_builder._row_height_in("P", "T", has_impressions=True) == pptx_builder._row_height_in("P", "T", has_impressions=True, line_note="")


def test_pptx_product_cell_text_helper_keeps_its_old_contract_for_a_line_without_a_note():
    p = _product("Product A")
    li = LineItem(product_name="Product A", monthly_budget=2000.0, months=3)
    assert pptx_builder._product_cell_text(p, li).splitlines()[1].endswith("impressions/mo")
    noted = LineItem(product_name="Product A", monthly_budget=2000.0, months=3, line_modifier=NOTE_A)
    assert pptx_builder._product_cell_text(p, noted).splitlines()[:2] == ["Product A", NOTE_A]


# =============================================================================================================
# The email prompts (the note is a label so same-product lines can be told apart — never a number)
# =============================================================================================================

def _email_lines(notes=True):
    kw = (lambda n: {"line_modifier": n}) if notes else (lambda n: {})
    return [LineItem(id="e1", product_name=EMAIL, monthly_budget=1000.0, months=3, **kw(NOTE_A)),
            LineItem(id="e2", product_name=EMAIL, monthly_budget=500.0, months=3, **kw(NOTE_B)),
            LineItem(id="e3", product_name=GEOFENCE, monthly_budget=250.0, months=3)]


EMAIL_REQ = ProposalRequest(client_name="Acme", request_type="New Business", start_date="2026-10-01", end_date="2026-12-31",
                            total_months=3, agency_fee=None, requested_by="Pat Planner", salesperson_email="pat@example.com")


def test_the_figures_block_labels_each_line_with_its_note_and_the_numbers_do_not_move():
    with_notes = email_totals.compute_email_totals(EMAIL_REQ, [{"label": "A", "line_items": _email_lines(True)}], None, "month")
    without = email_totals.compute_email_totals(EMAIL_REQ, [{"label": "A", "line_items": _email_lines(False)}], None, "month")
    block = email_totals.render_prompt_block(with_notes)
    assert f"- {EMAIL} [line note: {NOTE_A}]: $1,000 per month" in block
    assert f"- {EMAIL} [line note: {NOTE_B}]: $500 per month" in block
    assert f"- {GEOFENCE}: $250 per month" in block                                    # an un-noted line is just its name
    assert "planner's own label for THAT line" in block
    # the label touches no figure: every number is identical with and without notes
    a, b = with_notes.options[0], without.options[0]
    assert (a.per_period_net, a.media_net, a.total_net) == (b.per_period_net, b.media_net, b.total_net)
    assert [(l.per_period, l.total) for l in a.lines] == [(l.per_period, l.total) for l in b.lines]


def test_the_prompt_is_unchanged_when_no_line_has_a_note():
    plain = email_totals.render_prompt_block(
        email_totals.compute_email_totals(EMAIL_REQ, [{"label": "A", "line_items": _email_lines(False)}], None, "month"))
    assert "line note" not in plain
    prompt = ai_enricher._build_prompt(EMAIL_REQ, _email_lines(False), time_unit="month")
    assert "line note" not in prompt and "planner's own label" not in prompt


def test_the_enrichment_prompt_carries_the_labels_for_one_option_and_for_several():
    single = ai_enricher._build_prompt(EMAIL_REQ, _email_lines(True), time_unit="month")
    assert f"{EMAIL} [line note: {NOTE_A}]" in single and f"{EMAIL} [line note: {NOTE_B}]" in single
    # the product LIST keeps clean product names (the blurbs are keyed by them): the label lives in the figures block only
    products_section = single.split("## PRODUCTS IN THIS PROPOSAL")[1].split("## ENTRAVISION KNOWLEDGE BASE")[0]
    assert "line note" not in products_section
    tiers = [{"label": "A", "name": "Independent", "line_items": _email_lines(True)[:2]},
             {"label": "B", "name": "Democrat", "line_items": [LineItem(product_name=EMAIL, monthly_budget=900.0, months=3)]}]
    multi = ai_enricher._build_prompt(EMAIL_REQ, _email_lines(True)[:2], tiers=tiers, time_unit="month")
    assert f"{EMAIL} [line note: {NOTE_A}], {EMAIL} [line note: {NOTE_B}]" in multi
    assert "- Independent:" in multi


def test_a_dollar_amount_the_planner_typed_into_a_note_is_not_flagged_but_others_still_are():
    lines = [LineItem(id="e1", product_name=EMAIL, monthly_budget=1000.0, months=3, line_modifier="Target $75 per lead"),
             LineItem(id="e2", product_name=EMAIL, monthly_budget=500.0, months=3)]
    totals = email_totals.compute_email_totals(EMAIL_REQ, [{"label": "A", "line_items": lines}], None, "month")
    assert 75.0 in email_totals.allowed_amounts(totals)
    ok = email_totals.check_text("Your investment for the Email Campaigns line is $75 for each lead goal.",
                                 email_totals.allowed_amounts(totals), plan_terms=[EMAIL])
    assert ok == []
    bad = email_totals.check_text("Your total investment is $777 for the flight.", email_totals.allowed_amounts(totals),
                                  plan_terms=[EMAIL])
    assert bad and bad[0].value == 777.0


def test_label_for_and_the_revision_tiers_carry_the_label():
    assert email_totals.label_for("Display", "") == "Display"
    assert email_totals.label_for("Display", "Hispanic") == "Display [line note: Hispanic]"
    tiers = main._email_tiers_from_models(
        [main.TierModel(label="A", line_items=[main.LineItemModel(id="a", product_name=EMAIL, monthly_budget=100.0, months=3,
                                                                  line_modifier=NOTE_A)])], [], "month")
    assert tiers[0]["line_items"][0].line_modifier == NOTE_A


# =============================================================================================================
# The SPA (app.js): the helpers run in Node, the markup contracts are checked statically
# =============================================================================================================

def _extract_fn(name: str) -> str:
    m = re.search(rf"^(?:async )?function {re.escape(name)}\(", APP_JS, re.M)
    assert m, f"function {name} not found"
    return APP_JS[m.start(): APP_JS.index("\n}\n", m.start()) + 2]


def _extract_const(name: str) -> str:
    m = re.search(rf"^const {re.escape(name)} = ", APP_JS, re.M)
    assert m, f"const {name} not found"
    return APP_JS[m.start(): APP_JS.index(";\n", m.end()) + 1]


def _node(body: str) -> str:
    node = shutil.which("node")
    if not node:
        pytest.skip("node is not available")
    with tempfile.TemporaryDirectory() as tmp:
        script = Path(tmp) / "scenario.js"
        script.write_text(f"console.log((function () {{\n{body}\n}})());\n", encoding="utf-8")
        proc = subprocess.run([node, str(script)], capture_output=True, text=True, encoding="utf-8", timeout=120)
    if proc.returncode != 0:
        raise AssertionError(f"node failed:\n{proc.stderr}\n{proc.stdout}")
    return proc.stdout.strip()


def _node_async(body: str) -> str:
    """Like _node, for a scenario that awaits: `body` is the inside of an async function."""
    node = shutil.which("node")
    if not node:
        pytest.skip("node is not available")
    with tempfile.TemporaryDirectory() as tmp:
        script = Path(tmp) / "scenario.js"
        script.write_text("(async function () {\n" + body + "\n})().then(function (r) { console.log(r); }, "
                          "function (e) { console.error(e && e.stack || e); process.exit(1); });\n", encoding="utf-8")
        proc = subprocess.run([node, str(script)], capture_output=True, text=True, encoding="utf-8", timeout=120)
    if proc.returncode != 0:
        raise AssertionError(f"node failed:\n{proc.stderr}\n{proc.stdout}")
    return proc.stdout.strip()


JS_HELPERS = "\n".join([_extract_const("LINE_MODIFIER_MAX")] + [
    _extract_fn(n) for n in ("_LINE_MODIFIER_UNWANTED", "_cleanLineModifier", "_lineNote", "_carryLineNotes", "_migrateReopenedTier",
                              "escapeHtml", "_lineNoteHtml")])

PARITY_CORPUS = [
    None, "", "   ", "Hispanic, Phoenix only", "   Hispanic,   Phoenix\tonly \n", "line one\nline two\r\nline three", "nul\x00byte",
    "del\x7fand\x9fc1", "zero​width‍marks", "bidi‮flip", "sep para end", "bom﻿here",
    "non breaking　space", "<b>bold</b> & \"quoted\" 'x'", "=SUM(A1:A9)", 12345, True, 3.5, "x" * 150, "é" * 300,
    "😀" * 300, "y" * 99 + " " + "z" * 20, "tab\t\t\tend", "\x01\x02\x03", "a b",
]


def test_the_js_cleaner_agrees_with_the_python_one_on_a_corpus():
    script = (JS_HELPERS + "\nreturn JSON.stringify(" + json.dumps(PARITY_CORPUS)
              + ".map(function (v) { return _cleanLineModifier(v); }));")
    assert json.loads(_node(script)) == [clean_line_modifier(v) for v in PARITY_CORPUS]


def test_js_line_note_helpers_escape_and_default():
    out = json.loads(_node(JS_HELPERS + """
return JSON.stringify({
  none: _lineNote({}), nul: _lineNote({ line_modifier: null }), noLine: _lineNote(null),
  clean: _lineNote({ line_modifier: "  Phoenix  " }),
  html: _lineNoteHtml({ line_modifier: '"><img src=x onerror="alert(1)"> & <b>x</b>' }, "extra"),
  empty: _lineNoteHtml({ line_modifier: "   " }),
});"""))
    assert out["none"] == "" and out["nul"] == "" and out["noLine"] == "" and out["clean"] == "Phoenix"
    assert "<img" not in out["html"] and "<b>" not in out["html"] and "&lt;img src=x" in out["html"] and "&quot;" in out["html"]
    assert out["html"].startswith('<span class="line-note extra"') and out["empty"] == ""


def test_js_reopen_migration_defaults_cleans_and_keeps_the_note():
    out = json.loads(_node(JS_HELPERS + """
var _n = 0;
function newLineItemId() { return "new" + (++_n); }
var t = _migrateReopenedTier({ label: "A", line_items: [
  { id: "a", product_name: "P", line_modifier: "  Phoenix  " },      // saved note, to be cleaned like a typed one
  { id: "b", product_name: "P", line_modifier: null },                // server null
  { id: "c", product_name: "P" },                                     // saved before the field existed
  { product_name: "P", line_modifier: "x\\ny" }                       // no id yet: id back-filled AND note kept
]});
return JSON.stringify(t.lineItems.map(function (l) { return [l.id, l.line_modifier]; }));"""))
    assert out == [["a", "Phoenix"], ["b", ""], ["c", ""], ["new1", "x y"]]


def test_js_suggest_mix_carries_each_note_to_the_same_product_and_reports_the_ones_it_drops():
    out = json.loads(_node(JS_HELPERS + """
var old = [
  { product_name: "Search", line_modifier: "Hispanic" }, { product_name: "Search", line_modifier: "Brand" },
  { product_name: "Video", line_modifier: "" },          { product_name: "Display", line_modifier: "Bonus" } ];
var mix = [ { product_name: "Search" }, { product_name: "Video" }, { product_name: "Search" }, { product_name: "Audio" } ];
var all = _carryLineNotes(old, mix);
var same = _carryLineNotes(old, [ { product_name: "Search" }, { product_name: "Search" }, { product_name: "Display" } ]);
return JSON.stringify({ all: all, same: same, empty: _carryLineNotes([], mix), noOld: _carryLineNotes(null, null) });"""))
    assert out["all"]["notes"] == ["Hispanic", "", "Brand", ""]                       # in order, by product; Video had no note
    assert out["all"]["lost"] == ["Display — Bonus"]                                  # no Display line in the new mix
    assert out["same"] == {"notes": ["Hispanic", "Brand", "Bonus"], "lost": []}       # nothing lost -> nothing to ask
    assert out["empty"]["notes"] == ["", "", "", ""] and out["empty"]["lost"] == []
    assert out["noOld"] == {"notes": [], "lost": []}


def test_js_suggest_mix_button_keeps_notes_asks_before_dropping_one_and_declining_changes_nothing():
    """The REAL onRecommend with a faked page: the server's answer is a new mix; notes ride along by product."""
    script = JS_HELPERS + "\n" + _extract_fn("onRecommend") + """
var __confirms = [], __answer = true, __renders = 0, __id = 0, __alerts = [];
function confirm(m) { __confirms.push(m); return __answer; }
function alert(m) { __alerts.push(m); }
var els = {};
function mk(id) { return els[id] || (els[id] = { id: id, value: "3000", disabled: false, innerHTML: "x" }); }
var document = { getElementById: mk };
var state = { lineItems: [], availsData: {}, parsed: {}, timeUnit: "month", rateOverrideOpen: new Set(), objectiveOtherOpen: new Set() };
function _mixEditBlocks() { return false; }
function parseFormattedInput(v) { return Number(v); }
function _mbUnitAdjective() { return "Monthly"; }
function _availsEntryHasData() { return false; }
function _briefWithSelectedTactics() { return null; }
function _curateDerivedMonths() { return 3; }
function newLineItemId() { return "id" + (++__id); }
function renderLineItems() { __renders++; }
var __server = [];
function fetch() { return Promise.resolve({ ok: true, json: function () { return Promise.resolve({ line_items: __server.map(function (l) { return Object.assign({ id: null, months: 3, line_modifier: null }, l); }) }); } }); }
function mkOld() { return [ { id: "a", product_name: "Search", line_modifier: "Hispanic" }, { id: "b", product_name: "Search", line_modifier: "Brand" },
                            { id: "c", product_name: "Video", line_modifier: "" } ]; }
var R = {};
// 1. the new mix drops one noted Search line -> asked first; DECLINED -> every line stays exactly as it was
state.lineItems = mkOld(); __server = [ { product_name: "Search" }, { product_name: "Display" } ]; __answer = false;
await onRecommend();
R.declined = { asked: __confirms.length, mentions: /Search — Brand/.test(__confirms[0] || ""), same: JSON.stringify(state.lineItems) === JSON.stringify(mkOld()), renders: __renders, btn: mk("recommend-btn").disabled };
// 2. accepted -> the first Search keeps its note, the new Display starts empty
__confirms = []; __answer = true;
await onRecommend();
R.accepted = { asked: __confirms.length, lines: state.lineItems.map(function (l) { return [l.product_name, l.line_modifier, !!l.id && l.id !== "a"]; }), btn: mk("recommend-btn").disabled };
// 3. a mix that keeps every noted product (same count) -> no question, both notes kept in order
state.lineItems = mkOld(); __confirms = []; __server = [ { product_name: "Search" }, { product_name: "Search" }, { product_name: "Audio" } ];
await onRecommend();
R.kept = { asked: __confirms.length, notes: state.lineItems.map(function (l) { return l.line_modifier; }) };
return JSON.stringify(R);
"""
    out = json.loads(_node_async(script))
    assert out["declined"] == {"asked": 1, "mentions": True, "same": True, "renders": 0, "btn": False}
    assert out["accepted"]["asked"] == 1 and out["accepted"]["btn"] is False
    assert out["accepted"]["lines"] == [["Search", "Hispanic", True], ["Display", "", True]]
    assert out["kept"] == {"asked": 0, "notes": ["Hispanic", "Brand", ""]}


def test_the_step_04_input_and_every_reference_spot_are_wired_in_the_markup():
    render = _extract_fn("renderLineItems")
    assert 'data-line-note-input' in render and 'data-key="line_modifier"' in render
    assert 'maxlength="${LINE_MODIFIER_MAX}"' in render and 'placeholder="e.g. Hispanic, Phoenix only"' in render
    assert "Line note (optional)" in render and "shown under the product name in the export" in render
    assert "escapeHtml(_lineNote(li))" in render                                       # the value is escaped, not interpolated raw
    assert "notes_override" in render                                                  # the existing Planner note input is still there
    # the generic line-edit handler cleans it as typed; reopen cleans it too
    assert 'key === "line_modifier"' in _extract_fn("onLineItemEdit") and "_cleanLineModifier" in _extract_fn("onLineItemEdit")
    assert "_cleanLineModifier(item.line_modifier)" in _extract_fn("_migrateReopenedTier")
    # the read-only reference spots
    assert "_lineNoteHtml(li" in _extract_fn("renderAvailsGrid")
    assert "_lineNoteHtml(li)" in _extract_fn("_mbLineItemBlockHtml")
    assert "_lineNoteHtml(li)" in _extract_fn("_mbRenderByMonthView")
    assert "data-lines-for" in _extract_fn("renderRoadblocks") and "_refreshRoadblockLineNotes();" in _extract_fn("renderRoadblocks")
    assert "_refreshRoadblockLineNotes();" in _extract_fn("_enterRoadblocksStep")
    assert "_generateSummaryLineNotesHtml(" in _extract_fn("renderGenerateSummary")
    assert "— Note: " in _extract_fn("buildGammaOutline")
    # a brand-new line starts with an empty note (first visit to Step 04 AND "+ Add")
    assert APP_JS.count('line_modifier: ""') >= 2
    # duplicating / adding an option spread the whole line (the note travels); Suggest Mix keeps notes by product
    assert "..." in _extract_fn("onDuplicateLineItem") and "_carryLineNotes(state.lineItems, data.line_items)" in _extract_fn("onRecommend")


def test_no_line_note_is_ever_interpolated_into_html_unescaped():
    """Every read of `line_modifier` in app.js goes through the cleaner / escaper helpers — nothing writes the raw field
    into a template string."""
    code = re.sub(r"^[ \t]*//[^\n]*", "", APP_JS, flags=re.M)
    for m in re.finditer(r"\$\{[^}]*line_modifier[^}]*\}", code):
        raise AssertionError(f"raw line_modifier interpolated into a template: {m.group(0)}")
    for m in re.finditer(r"innerHTML\s*[+]?=\s*[\w.]*line_modifier\b", code):
        raise AssertionError(f"raw line_modifier assigned into innerHTML: {m.group(0)}")
    owners = set()
    for m in re.finditer(r"line_modifier", code):
        starts = list(re.finditer(r"^(?:async )?function (\w+)\(", code[:m.start()], re.M))
        owners.add(starts[-1].group(1) if starts else None)
    # the field is touched only by the helpers, the line factories, the edit handler, the reopen migration, Suggest Mix
    allowed = {"_cleanLineModifier", "_lineNote", "onNext", "onAddProduct", "onLineItemEdit", "_migrateReopenedTier",
               "onRecommend", "renderLineItems", None}       # (renderLineItems: only the input's data-key attribute name)
    assert owners <= allowed, sorted(str(o) for o in owners - allowed)


def test_the_note_styles_exist():
    for selector in (".line-note", ".line-note-block", ".line-note-field-label", ".roadblock-line-notes", ".sum-line-notes"):
        assert re.search(rf"^{re.escape(selector)}[ ,{{:]", STYLES_CSS, re.M), selector
    assert re.search(r"^\.line-items input\.line-note-input \{", STYLES_CSS, re.M)


def test_docx_reports_do_not_list_lines_so_they_carry_no_note():
    """docx_builder writes the strategy brief, the roadblocks report (one section per PRODUCT, platform rules don't depend
    on the line) and the client email: none lists a line with its product + objective, so none needs the note."""
    src = (ROOT / "app" / "services" / "docx_builder.py").read_text(encoding="utf-8")
    assert "line_modifier" not in src and "objective_override" not in src
