"""
Covers the Monthly Breakdown export-layout change: Planner Notes/AdOps
must land AFTER the (variable-width) Monthly Breakdown block instead of
before it, and the inline block needs its own per-month total row aligned
with the sheet's "TOTAL DIGITAL MONTHLY" row.

Deliberately does NOT go through generate_proposal()/catalog.by_name() —
those hit a real Postgres connection (see app/db.py), which isn't
reachable from a plain local dev machine. Product/LineItem are plain
dataclasses, constructible directly without the DB at all, which is all
this needs: the column-position arithmetic and the inline total row are
both pure openpyxl/dict logic once given a worksheet and a product list.
"""
from __future__ import annotations

from openpyxl import Workbook

from app.catalog import Product
from app.services.proposal_generator import LineItem, _populate_monthly_breakdown_inline
from app import excel_template as et


def _product(name: str, family: str = "Display") -> Product:
    return Product(
        family=family, name=name, short_label=name, proposal_description="",
        sizes="", buying_model="CPM", base_rate=10.0, estimated_impressions=False,
        discloses_impressions=True, minimum_spend=0.0, minimum_flight_days=(0, 0),
        sla_data_days=None, sla_creative_days=None, sla_activate_days=None, sla_total_days=None,
        media_allocation_pct=1.0, margin_upper=0.5, margin_lower=0.5,
        tech_platform="", wide_orbit_code="",
    )


def _months(*keys: str) -> list[dict]:
    return [{"key": k, "label": k, "days_in_month": 30, "active_days": 30} for k in keys]


def test_reposition_notes_adops_unused_keeps_original_columns():
    wb = Workbook()
    ws = wb.active
    notes_col, adops_col = et.reposition_notes_adops(ws, 17, gross=False, mb_width=0)
    assert (notes_col, adops_col) == ("T", "V")
    notes_col, adops_col = et.reposition_notes_adops(ws, 17, gross=True, mb_width=0)
    assert (notes_col, adops_col) == ("W", "Y")


def test_reposition_notes_adops_moves_past_monthly_breakdown():
    wb = Workbook()
    ws = wb.active
    # Net: MB starts at X(24). A 3-month block occupies X,Y,Z(24-26);
    # Notes must land one spacer past that (28=AB), AdOps 2 past Notes (AD).
    notes_col, adops_col = et.reposition_notes_adops(ws, 17, gross=False, mb_width=3)
    assert notes_col == "AB"
    assert adops_col == "AD"
    assert ws[f"{notes_col}17"].value == "Planner Notes — internal guidance"
    assert ws[f"{adops_col}17"].value == "AdOps (Internal Use)"
    # The columns Notes/AdOps used to occupy are narrowed to a plain gap,
    # not left at their old wide (60/22) widths.
    assert ws.column_dimensions["T"].width == 4
    assert ws.column_dimensions["V"].width == 4

    # Gross: MB starts at AA(27).
    notes_col, adops_col = et.reposition_notes_adops(ws, 17, gross=True, mb_width=2)
    assert notes_col == "AD"
    assert adops_col == "AF"


def test_reposition_notes_adops_does_not_reset_row_height():
    # _set_header (called earlier by build_proposal_a/_gross) already bumps
    # row 17 to height 40 for the 2-line wrapped headers — the new
    # dynamic Notes/AdOps writer must not silently reset that back to 32.
    wb = Workbook()
    ws = wb.active
    ws.row_dimensions[17].height = 40
    et.reposition_notes_adops(ws, 17, gross=False, mb_width=3)
    assert ws.row_dimensions[17].height == 40


def test_inline_monthly_breakdown_writes_total_row_aligned_with_sheet_total():
    wb = Workbook()
    ws = wb.active
    products = [_product("Display A"), _product("Display B")]
    months = _months("2026-01", "2026-02")
    line_items = [
        LineItem(product_name="Display A", monthly_budget=1000.0, months=2,
                 monthly_allocations={"2026-01": 600.0, "2026-02": 1400.0}),
        LineItem(product_name="Display B", monthly_budget=500.0, months=2),  # never customized -> falls back to the mode estimate
    ]
    _populate_monthly_breakdown_inline(ws, products, line_items, months, gross=False, distribution_mode="even")

    start_col = et.MONTHLY_BREAKDOWN_START_COL[False]
    # Row math must match build_proposal_a's own: start_row(19) + one row
    # per product (2) = 21 is one past the last product row; total_row = row+1 = 22.
    total_row = 22
    jan_col = et.get_column_letter(start_col)
    feb_col = et.get_column_letter(start_col + 1)
    # Display A: 600 + Display B's even-split half of $1000 total (500 each month) = 1100
    assert ws[f"{jan_col}{total_row}"].value == 1100.0
    # Display A: 1400 + Display B's 500 = 1900
    assert ws[f"{feb_col}{total_row}"].value == 1900.0
    # Sums to the real combined total (2000 + 1000 = 3000) either way —
    # the per-month split changes, the grand sum never should.
    assert ws[f"{jan_col}{total_row}"].value + ws[f"{feb_col}{total_row}"].value == 3000.0


def test_inline_monthly_breakdown_respects_prorated_mode_for_uncustomized_line():
    wb = Workbook()
    ws = wb.active
    products = [_product("Display A")]
    months = _months("2026-01", "2026-02")
    # Uneven active_days -> prorated mode should NOT split 50/50.
    months[0]["active_days"] = 10
    months[1]["active_days"] = 20
    line_items = [LineItem(product_name="Display A", monthly_budget=300.0, months=2)]  # $600 total, never customized

    _populate_monthly_breakdown_inline(ws, products, line_items, months, gross=False, distribution_mode="prorated")
    start_col = et.MONTHLY_BREAKDOWN_START_COL[False]
    # start_row=19, one product -> row becomes 20 after the loop; total_row = 21
    total_row = 21
    jan_col = et.get_column_letter(start_col)
    feb_col = et.get_column_letter(start_col + 1)
    assert ws[f"{jan_col}{total_row}"].value == 200.0   # 600 * 10/30
    assert ws[f"{feb_col}{total_row}"].value == 400.0   # 600 * 20/30
