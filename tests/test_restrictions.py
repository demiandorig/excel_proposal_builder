"""
The Restricted Verticals sheet as the Roadblocks step's source of truth
(app/services/restrictions.py + roadblocks.py). The headline regression: political ads ARE
allowed on Display and OLV per the sheet, yet the old catalog flags marked both not_allowed and
nothing checked the model's verdict.
"""
from __future__ import annotations

import asyncio
import base64
import io
import json
import os

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import pytest
from fastapi import HTTPException

from app.services import restrictions as rs
from app.services import roadblocks as rb
from app.services.proposal_generator import LineItem
from app.services.notion_parser import ProposalRequest


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())


@pytest.fixture
def builtin_state(monkeypatch):
    monkeypatch.setattr(rs, "load_state", lambda: {"categories": rs.builtin_categories(), "using_builtin": True, "sync": None})


# --- a workbook shaped like the real sheet ---------------------------------------------------

def _sheet_bytes() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)

    cannabis = wb.create_sheet("Cannabis")
    for r in ["", "", "** All Entravision Digital Products are currently on hold and will not accept cannabis ads due to legal roadblocks **",
              "", "From our EVC c-level suite:", "We spoke to our outside legal counsel on advertising issues.", "",
              "Previously approved products (currently on hold as well):", "CTV/OTT", "eDigital OLV"]:
        cannabis.append([r])

    blood = wb.create_sheet("Blood Donation")
    for r in ["", "", "* **Interest Targeting**", "", "* Interests related to **charity** can be used.",
              "* Limit the audience to **adults 18+**."]:
        blood.append([r])

    casino = wb.create_sheet("Casino_Gambling")     # Excel forbids "/": the sheet's "Casino/Gambling" tab exports like this
    casino.append(["Platform", "Casinos / Gambling", "Guideline/ Resources URLs", "Guideline/ Resources URLs"])
    casino.append(["TTD", "We do accept brick+mortar ads, but there are some guidelines.", "https://desk.thetradedesk.com/x", None])
    casino.append(["Meta", "Client's Preparation for Agency-Managed Ads: obtain licenses.", "Meta's Gambling Policy", None])
    casino.append([None, "Agency's Preparation: apply for written permission.", "Meta's Gambling Policy", None])
    casino.append(["TikTok", "May be allowed if requirements are met.", None, None])

    sexual = wb.create_sheet("Sexual Enchancement Procedures")
    sexual.append(["Platform", "Sexual Enchancement Procedures", "URLs"])
    sexual.append(["TTD", "May fall into a banned category.", "Ad Content Guidelines"])
    sexual.append(["MadHive", "Fully allowed.", None])

    political = wb.create_sheet("Political")
    political.append([None])
    political.append(["** Political Ads are only accepted for the below products **"])
    for item in rs._POLITICAL_ITEMS:
        political.append([item])
    political.append([None]); political.append([None]); political.append([None])
    political.append(["2024 Political Campaigns - What to Expect?"])
    political.append(["2024 Political Digital Rate Card"])

    idx = wb.create_sheet("Policy Document Index")
    idx.append(["Document Name", "Document Link"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- parsing --------------------------------------------------------------------------------------

def test_every_tab_is_classified_by_what_it_says_and_the_index_tab_is_skipped():
    cats = {c.name: c for c in rs.parse_workbook(_sheet_bytes())}
    assert set(cats) == {"Cannabis", "Blood Donation", "Casino Gambling", "Sexual Enchancement Procedures", "Political"}
    assert cats["Cannabis"].kind == "blanket_hold"
    assert "will not accept cannabis ads" in cats["Cannabis"].summary
    assert cats["Political"].kind == "allowlist"
    assert cats["Casino Gambling"].kind == "guidance" and cats["Blood Donation"].kind == "guidance"


def test_the_political_allowlist_stops_at_the_blank_row_and_keeps_every_product():
    political = {c.name: c for c in rs.parse_workbook(_sheet_bytes())}["Political"]
    assert political.items == rs._POLITICAL_ITEMS            # all 11, none of the trailing link text
    assert "2024 Political Digital Rate Card" not in political.items


def test_platform_grids_become_readable_lines_and_merged_platform_cells_carry_down():
    casino = {c.name: c for c in rs.parse_workbook(_sheet_bytes())}["Casino Gambling"]
    lines = casino.body.splitlines()
    assert lines[0].startswith("TTD: We do accept brick+mortar ads")
    assert "[https://desk.thetradedesk.com/x]" in lines[0]
    assert lines[1].startswith("Meta:") and lines[2].startswith("Meta: Agency's Preparation")   # blank cell inherits "Meta"
    assert lines[3].startswith("TikTok: May be allowed")


def test_markdown_noise_is_stripped_from_cell_text():
    blood = {c.name: c for c in rs.parse_workbook(_sheet_bytes())}["Blood Donation"]
    assert "**" not in blood.body and "Interest Targeting" in blood.body
    assert not any(line.startswith("* ") for line in blood.body.splitlines())


def test_parse_workbook_rejects_non_workbooks_and_empty_ones():
    with pytest.raises(ValueError, match="xlsx"):
        rs.parse_workbook(b"not a workbook")
    from openpyxl import Workbook
    wb = Workbook(); wb.active.title = "Policy Document Index"
    buf = io.BytesIO(); wb.save(buf)
    with pytest.raises(ValueError, match="No restricted-vertical"):
        rs.parse_workbook(buf.getvalue())


def test_default_keywords_cover_the_tab_name_and_known_synonyms():
    assert {"political", "election", "candidate for office"} <= set(rs.default_keywords("Political"))
    assert {"casino", "gambling", "sportsbook"} <= set(rs.default_keywords("Casino Gambling"))
    assert "marijuana" in rs.default_keywords("Cannabis")
    assert "something new" in rs.default_keywords("Something New")     # unknown tab: its own name still works


def test_suggest_categories_reads_the_plan_text():
    cats = rs.builtin_categories()
    assert rs.suggest_categories("Re-elect Maria Lopez — political campaign for city council", cats) == ["Political"]
    assert rs.suggest_categories("Green Leaf Dispensary and CBD shop", cats) == ["Cannabis"]
    assert rs.suggest_categories("Smith Family Dental", cats) == []


# --- sheet entry -> catalog product matching ------------------------------------------------

@pytest.mark.parametrize("item,expected", [
    ("eDigital Display", ["eDigital Network Display - Standard IAB"]),
    ("eDigital Display Geo Fence", ["Display - Geo Fence"]),
    ("eDigital OLV", ["Video - Pre-roll (OLV)"]),
    ("Entravision O&O Audio Streaming", ["Audio - EVC Audio Streaming"]),
    ("Audio Engage", ["AudioEngage"]),
    ("Email Marketing Campaigns", ["Email Campaigns and/or Email Campaigns - Re-Drop"]),
    ("Email Campaigns - Display Re-targeting", ["Email Campaigns - Display Re-targeting"]),
    ("Digital Out of Home", ["Digital Out of Home"]),
])
def test_sheet_entries_map_to_the_right_catalog_products(item, expected):
    assert rs.match_catalog_products(item) == expected


def test_run_of_network_ctv_excludes_named_property_buys():
    names = rs.match_catalog_products(
        "CTV / OTT -- Entravision Plus Run-Of-Network (Does not include Netflix, Prime or Roku)")
    assert names and all("CTV/OTT" in n for n in names)
    assert not any(bad in n for n in names for bad in ("Netflix", "Prime", "Roku", "VIX"))


def test_meta_ads_covers_the_core_buys_not_page_or_branded_content():
    names = rs.match_catalog_products("Meta Ads (Only through client's handle, no branded content)")
    assert names and all(n.startswith("Facebook & Instagram Ads |") for n in names)
    assert not any("Branded" in n or "Noticias" in n for n in names)


def test_an_entry_with_no_catalog_product_matches_nothing_rather_than_guessing():
    assert rs.match_catalog_products("eDigital OLV Geo Fence") == []
    assert rs.match_catalog_products("Totally Unknown Product") == []


# --- verdicts: the political-on-Display/OLV regression ---------------------------------------

def _political():
    return [c for c in rs.builtin_categories() if c.name == "Political"]


def test_political_is_allowed_on_display_and_olv_per_the_sheet_even_though_the_old_catalog_flags_said_no():
    from app.catalog import by_name
    for name in ("Video - Pre-roll (OLV)", "eDigital Network Display - Standard IAB"):
        assert by_name(name).political_policy == "not_allowed"     # the stale flag this replaces
        verdict = rs.evaluate_product(name, by_name(name).family, _political())
        assert verdict["verdict"] == "allowed", name
        assert verdict["checks"][0]["status"] == "allowed"


def test_products_the_political_list_omits_are_not_accepted_and_services_are_exempt():
    from app.catalog import by_name
    for name in ("YouTube Ads", "Search - SEM", "Tiktok Ads", "LinkedIn", "Spotify"):
        v = rs.evaluate_product(name, by_name(name).family, _political())
        assert v["verdict"] == "not_allowed", name
        assert "Not on the approved product list for Political" in v["checks"][0]["note"]
    services = by_name("Web Services - Landing Pages")
    assert rs.evaluate_product(services.name, services.family, _political())["verdict"] == "n/a"


def test_the_allowed_note_carries_the_sheets_own_conditions():
    from app.catalog import by_name
    meta = by_name("Facebook & Instagram Ads | Awareness")
    note = rs.evaluate_product(meta.name, meta.family, _political())["checks"][0]["note"]
    assert "Only through client's handle, no branded content" in note


def test_cannabis_is_not_accepted_anywhere_and_the_worst_verdict_wins_across_verticals():
    from app.catalog import by_name
    cats = rs.builtin_categories()
    olv = by_name("Video - Pre-roll (OLV)")
    both = rs.evaluate_product(olv.name, olv.family, cats)
    assert both["verdict"] == "not_allowed"                          # allowed for Political, not for Cannabis
    assert {c["category"]: c["status"] for c in both["checks"]} == {"Political": "allowed", "Cannabis": "not_allowed"}


def test_guidance_verticals_surface_as_guidance_not_a_made_up_yes_or_no():
    cats = {c.name: c for c in rs.parse_workbook(_sheet_bytes())}
    v = rs.evaluate_product("AudioEngage", "Audio", [cats["Casino Gambling"]])
    assert v["verdict"] == "guidance"


def test_no_confirmed_verticals_means_no_verdicts(builtin_state):
    m = rs.build_matrix(["YouTube Ads"], [])
    assert m["categories"] == [] and m["verdicts"]["YouTube Ads"] == {"verdict": "none", "checks": []}


# --- fail-safe loading ---------------------------------------------------------------------------

def test_load_state_falls_back_to_the_builtin_snapshot_when_the_tables_are_unreachable(monkeypatch):
    def boom():
        raise RuntimeError("relation \"restriction_categories\" does not exist")
    monkeypatch.setattr(rs, "get_connection", boom)
    state = rs.load_state()
    assert state["using_builtin"] is True
    assert {c.name for c in state["categories"]} == {"Political", "Cannabis"}


# --- the prompt ----------------------------------------------------------------------------------

def _req():
    return ProposalRequest(client_name="Re-elect Maria", campaign_goal="Win the election", geo="Fresno")


def _items(*names):
    return [LineItem(product_name=n, monthly_budget=1000, months=3, id=f"i{i}") for i, n in enumerate(names)]


def test_prompt_states_the_sheets_verdicts_as_facts_and_drops_the_stale_catalog_flags(builtin_state):
    items = _items("Video - Pre-roll (OLV)", "YouTube Ads")
    matrix = rs.build_matrix([li.product_name for li in items], ["Political"])
    prompt = rb._build_prompt(_req(), items, None, matrix)
    assert "RESTRICTED VERTICALS MATRIX — AUTHORITATIVE" in prompt
    assert "- Video - Pre-roll (OLV): Political — ALLOWED" in prompt
    assert "- YouTube Ads: Political — NOT ACCEPTED" in prompt
    assert "Never contradict the Restricted Verticals Matrix" in prompt
    assert "KNOWN INTERNAL POLICY FLAGS" not in prompt
    assert "political: not_allowed" not in prompt


def test_prompt_without_confirmed_verticals_forbids_inferring_restrictions_from_the_industry(builtin_state):
    items = _items("Video - Pre-roll (OLV)")
    prompt = rb._build_prompt(_req(), items, None, rs.build_matrix(["Video - Pre-roll (OLV)"], []))
    assert "NONE CONFIRMED" in prompt and "never cite Entravision internal flags" in prompt


# --- enforcement on the model's answer -----------------------------------------------------------

def _result(*roadblocks):
    return {"overall_summary": "s", "product_roadblocks": list(roadblocks), "categories": [], "used_web_search": True, "error": None}


def test_a_model_risk_claiming_political_is_not_allowed_on_an_allowed_product_is_removed(builtin_state):
    items = _items("Video - Pre-roll (OLV)")
    matrix = rs.build_matrix(["Video - Pre-roll (OLV)"], ["Political"])
    result = _result({
        "product_name": "Video - Pre-roll (OLV)", "risk_level": "high",
        "risks": [
            {"issue": "Political ads not feasible", "detail": "Political advertising is not allowed on OLV.", "source": "x"},
            {"issue": "Additional: ID verification", "detail": "Political ads on OLV partners need advertiser verification before launch.", "source": "y"},
        ],
        "recommended_mitigation": "m",
    })
    rb._apply_matrix(result, matrix, items)
    entry = result["product_roadblocks"][0]
    assert [r["issue"] for r in entry["risks"]] == ["Additional: ID verification"]   # the contradiction went, the real caveat stayed
    assert entry["matrix"]["verdict"] == "allowed"
    assert result["categories"] == ["Political"]


def test_not_accepted_products_get_a_deterministic_high_risk_even_if_the_model_said_low(builtin_state):
    items = _items("YouTube Ads")
    matrix = rs.build_matrix(["YouTube Ads"], ["Political"])
    result = _result({"product_name": "YouTube Ads", "risk_level": "low", "risks": [], "recommended_mitigation": ""})
    rb._apply_matrix(result, matrix, items)
    entry = result["product_roadblocks"][0]
    assert entry["risk_level"] == "high"
    assert entry["risks"][0]["issue"] == "Not accepted for Political"
    assert "Restricted Verticals sheet" in entry["risks"][0]["source"]


def test_a_product_the_model_never_mentioned_still_gets_its_verdict(builtin_state):
    items = _items("Tiktok Ads")
    matrix = rs.build_matrix(["Tiktok Ads"], ["Political"])
    result = _result()
    rb._apply_matrix(result, matrix, items)
    assert [r["product_name"] for r in result["product_roadblocks"]] == ["Tiktok Ads"]
    assert result["product_roadblocks"][0]["risk_level"] == "high"


def test_guidance_raises_a_low_risk_to_medium_but_never_lowers_a_high_one(builtin_state):
    cats = {c.name: c for c in rs.parse_workbook(_sheet_bytes())}
    matrix = {"categories": [cats["Casino Gambling"]], "using_builtin": False, "synced_at": None,
              "verdicts": {"AudioEngage": rs.evaluate_product("AudioEngage", "Audio", [cats["Casino Gambling"]])}}
    low = _result({"product_name": "AudioEngage", "risk_level": "low", "risks": [], "recommended_mitigation": ""})
    high = _result({"product_name": "AudioEngage", "risk_level": "high", "risks": [], "recommended_mitigation": ""})
    items = _items("AudioEngage")
    rb._apply_matrix(low, matrix, items)
    rb._apply_matrix(high, matrix, items)
    assert low["product_roadblocks"][0]["risk_level"] == "medium"
    assert high["product_roadblocks"][0]["risk_level"] == "high"


def test_risks_that_dont_mention_the_allowed_vertical_are_untouched(builtin_state):
    items = _items("Video - Pre-roll (OLV)")
    matrix = rs.build_matrix(["Video - Pre-roll (OLV)"], ["Political"])
    unrelated = {"issue": "Skippable ads not allowed over 30s", "detail": "Pre-roll longer than 30 seconds is not allowed.", "source": "z"}
    result = _result({"product_name": "Video - Pre-roll (OLV)", "risk_level": "medium", "risks": [unrelated], "recommended_mitigation": ""})
    rb._apply_matrix(result, matrix, items)
    assert result["product_roadblocks"][0]["risks"] == [unrelated]


# --- generate_roadblocks end to end (model mocked) -------------------------------------------------

class _FakeOpenAI:
    def __init__(self, api_key=None):
        pass


def test_generate_roadblocks_never_returns_a_political_not_feasible_claim_for_an_allowed_product(monkeypatch, builtin_state):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    monkeypatch.setattr(rb, "_HAS_OPENAI", True)
    monkeypatch.setattr(rb, "_OpenAI", _FakeOpenAI, raising=False)
    model_says = json.dumps({
        "overall_summary": "Political is risky.",
        "product_roadblocks": [
            {"product_name": "eDigital Network Display - Standard IAB", "risk_level": "high",
             "risks": [{"issue": "Political ads not feasible on Display", "detail": "Political advertising is not allowed on Display.", "source": "s, 2025"}],
             "recommended_mitigation": "Avoid."},
        ],
    })
    monkeypatch.setattr(rb.llm_utils, "responses_text", lambda *a, **k: model_says)
    result = rb.generate_roadblocks(_req(), _items("eDigital Network Display - Standard IAB"), categories=["Political"])
    entry = result["product_roadblocks"][0]
    assert entry["risks"] == [] and entry["risk_level"] == "low"
    assert entry["matrix"]["verdict"] == "allowed"
    assert result["categories"] == ["Political"]


def test_when_the_model_is_unavailable_the_sheets_verdicts_still_come_back(monkeypatch, builtin_state):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.setattr(rb, "_HAS_OPENAI", True)
    result = rb.generate_roadblocks(_req(), _items("YouTube Ads", "Video - Pre-roll (OLV)"), categories=["Political"])
    assert "OPENAI_API_KEY" in result["error"]
    verdicts = {r["product_name"]: r["matrix"]["verdict"] for r in result["product_roadblocks"]}
    assert verdicts == {"YouTube Ads": "not_allowed", "Video - Pre-roll (OLV)": "allowed"}


# --- endpoints ------------------------------------------------------------------------------------

def test_planner_endpoint_lists_verticals_and_flags_the_ones_the_plan_suggests(builtin_state):
    import app.main as m
    body = m.RestrictionCategoriesRequest(request={"client_name": "Re-elect Maria", "campaign_goal": "Win the election"},
                                          product_names=["YouTube Ads"])
    out = asyncio.run(m.restriction_categories(body))
    by = {c["name"]: c for c in out["categories"]}
    assert by["Political"]["suggested"] is True and by["Cannabis"]["suggested"] is False
    assert out["using_builtin"] is True


def test_admin_sync_parses_the_fetched_workbook_and_stores_it(monkeypatch):
    import app.main as m
    stored = {}
    monkeypatch.setattr(rs, "fetch_sheet_xlsx", lambda file_id: (_sheet_bytes(), "Restricted Verticals by Platforms"))
    monkeypatch.setattr(rs, "replace_all", lambda cats, **kw: stored.update(cats=cats, **kw))
    monkeypatch.setattr(rs, "load_state", lambda: {"categories": stored.get("cats", []), "using_builtin": False,
                                                    "sync": {"sheet_url": "u", "source_title": "t", "synced_at": None, "synced_by": "a@x"}})

    class _Req:
        class state:
            user = {"email": "admin@entravision.com"}
    out = asyncio.run(m.admin_sync_restrictions(m.RestrictionSyncRequest(sheet_url=rs.DEFAULT_SHEET_URL), _Req()))
    assert out["synced"] is True
    assert {c.name for c in stored["cats"]} >= {"Political", "Cannabis", "Casino Gambling"}
    assert stored["synced_by"] == "admin@entravision.com" and stored["source_title"] == "Restricted Verticals by Platforms"


def test_admin_sync_and_upload_turn_problems_into_400s(monkeypatch):
    import app.main as m

    class _Req:
        class state:
            user = {"email": "a@x"}

    def no_access(file_id):
        raise ValueError("The connected Google account couldn't open that sheet")
    monkeypatch.setattr(rs, "fetch_sheet_xlsx", no_access)
    with pytest.raises(HTTPException) as e1:
        asyncio.run(m.admin_sync_restrictions(m.RestrictionSyncRequest(sheet_url=rs.DEFAULT_SHEET_URL), _Req()))
    assert e1.value.status_code == 400 and "couldn't open" in e1.value.detail
    with pytest.raises(HTTPException) as e2:
        asyncio.run(m.admin_upload_restrictions(m.RestrictionUploadRequest(filename="x.xlsx", data_base64=base64.b64encode(b"junk").decode()), _Req()))
    assert e2.value.status_code == 400
    with pytest.raises(HTTPException) as e3:
        asyncio.run(m.admin_sync_restrictions(m.RestrictionSyncRequest(sheet_url="not a link"), _Req()))
    assert e3.value.status_code == 400


def test_extract_file_id_accepts_links_and_bare_ids():
    sid = "1FpWkPQRXim8_CGkUxagmQ19JuKtBSn29XoRqzoquAp0"
    assert rs.extract_file_id(f"https://docs.google.com/spreadsheets/d/{sid}/edit?gid=198512657#gid=198512657") == sid
    assert rs.extract_file_id(sid) == sid
    with pytest.raises(ValueError):
        rs.extract_file_id("hello")


def test_update_requires_a_keyword(monkeypatch):
    with pytest.raises(ValueError, match="at least one keyword"):
        rs.update_category("Political", keywords=" , ", product_map={})


def test_schema_declares_the_restriction_tables():
    from pathlib import Path
    sql = (Path(__file__).resolve().parents[1] / "schema.sql").read_text(encoding="utf-8")
    assert "CREATE TABLE IF NOT EXISTS restriction_categories" in sql
    assert "CHECK (kind IN ('allowlist', 'blanket_hold', 'guidance'))" in sql
    assert "CREATE TABLE IF NOT EXISTS restriction_sync" in sql


# --- the Word report ------------------------------------------------------------------------------

def test_word_report_shows_each_products_sheet_verdict(tmp_path):
    docx = pytest.importorskip("docx")
    from app.services import docx_builder
    out = tmp_path / "rb.docx"
    ok = docx_builder.build_roadblocks_docx(
        out, "Acme", "summary",
        [{"product_name": "YouTube Ads", "risk_level": "high", "risks": [],
          "matrix": {"verdict": "not_allowed", "checks": [{"category": "Political", "status": "not_allowed", "note": "Not on the approved product list for Political ads."}]}}],
        used_web_search=True, categories=["Political"])
    assert ok
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "Restricted verticals checked: Political" in text
    assert "Political: Not accepted" in text and "Not on the approved product list" in text
