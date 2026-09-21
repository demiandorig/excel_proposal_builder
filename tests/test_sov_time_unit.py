"""
Covers a real gap found after the Week/Month/Quarter toggle shipped: SOV
("Share of Voice") compares the curated budget against the catalog's
avails ceiling, which stays a flat MONTHLY figure by design (see
excel_template._UNIT_LABELS' own comment on why that ceiling doesn't get
relabeled). Without converting the budget to a monthly-equivalent first,
SOV% would read wildly low in Weekly mode (~1/4 the real rate) or wildly
high in Quarterly mode (~3x), even though real pacing hasn't changed.
Planner-specified factors: week x4, quarter /3 (simpler than the 12/52
ratio used for minimum-spend scaling elsewhere — explicit choice, not an
oversight, see excel_template.SOV_MONTHLY_EQUIVALENT_SCALE's own comment).
"""
import dataclasses
import os

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import pytest

from app.catalog import Product, by_name
from app import excel_template as et


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())


def _cpm_product(**overrides) -> Product:
    # Start from a real catalog product (Product has ~20 required fields
    # with no defaults) and override just what this test cares about.
    base = by_name("Search - SEM")
    fields = dict(buying_model="CPM", base_rate=10.0, minimum_spend=500.0, estimated_cpm_for_imps=None)
    fields.update(overrides)
    return dataclasses.replace(base, **fields)


def test_month_granularity_matches_pre_toggle_behavior_exactly():
    p = _cpm_product()
    avail = {"max_spend": 4000.0}
    pct_no_arg = et.compute_sov_pct(p, 1000.0, avail)
    pct_explicit_month = et.compute_sov_pct(p, 1000.0, avail, time_unit="month")
    assert pct_no_arg == pct_explicit_month == 25.0  # 1000/4000 * 100, unchanged


def test_weekly_budget_scales_up_four_x_before_comparing():
    p = _cpm_product()
    avail = {"max_spend": 4000.0}
    # $1000/week ~ $4000/month-equivalent -> 100% of a $4000 monthly ceiling,
    # NOT 25% (which is what the raw unscaled $1000 would naively give).
    pct = et.compute_sov_pct(p, 1000.0, avail, time_unit="week")
    assert pct == 100.0


def test_quarterly_budget_scales_down_by_three_before_comparing():
    p = _cpm_product()
    avail = {"max_spend": 4000.0}
    # $12000/quarter ~ $4000/month-equivalent -> 100%, not 300%.
    pct = et.compute_sov_pct(p, 12000.0, avail, time_unit="quarter")
    assert pct == 100.0


def test_scaling_also_applies_to_the_imps_derived_fallback_path():
    p = _cpm_product(base_rate=10.0)  # CPM $10 -> 100k imps costs $1000
    avail = {"max_imps": 400_000}  # implied monthly ceiling spend = $4000
    pct = et.compute_sov_pct(p, 1000.0, avail, time_unit="week")
    assert pct == 100.0  # same $1000/week -> $4000-equivalent logic as the max_spend test


def test_live_sov_formula_multiplies_by_four_for_weekly():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["L5"] = 1000  # budget cell
    ws["O5"] = 4000  # spend cell (default Net-sheet column; already-derived/typed numeric spend)
    et.write_avails_cells(
        ws, 5, {"max_spend": 4000.0, "basis": "spend"}, _cpm_product(),
        cols=None, sov_col="Q", budget_col="L", time_unit="week",
    )
    assert ws["Q5"].value == '=IFERROR((L5*4)/O5,"")'


def test_live_sov_formula_divides_by_three_for_quarterly():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["L5"] = 12000
    ws["O5"] = 4000
    et.write_avails_cells(
        ws, 5, {"max_spend": 4000.0, "basis": "spend"}, _cpm_product(),
        cols=None, sov_col="Q", budget_col="L", time_unit="quarter",
    )
    assert ws["Q5"].value == '=IFERROR((L5/3)/O5,"")'


def test_live_sov_formula_unchanged_for_monthly():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["L5"] = 1000
    ws["O5"] = 4000
    et.write_avails_cells(
        ws, 5, {"max_spend": 4000.0, "basis": "spend"}, _cpm_product(),
        cols=None, sov_col="Q", budget_col="L",
    )
    assert ws["Q5"].value == '=IFERROR(L5/O5,"")'
