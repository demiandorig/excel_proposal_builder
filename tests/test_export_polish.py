"""
Export polish + admin FAQ + a live restrictions sheet:
  * Full Flight shows no "Months" label/cell and calls its first total "SUBTOTAL DIGITAL";
  * sheet banners read "<option or client name> - Proposal" / "- Avails" (never "Proposal A - Proposal");
  * the Gross sheet carries an "Agency Comm:" legend and its eCPM comes from the NET budget;
  * the Process FAQs tab is built from admin-managed rows (real hyperlinks) with a header link to the guidelines;
  * the restricted-verticals sheet is re-read from Google when the stored copy is stale;
 """
from __future__ import annotations

import asyncio
import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import openpyxl
import pytest
from fastapi import HTTPException

from app import excel_template as et
from app.services import faq_entries as faq
from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import LineItem, _banner_title, generate_proposal

EMAIL = "Email Campaigns - Display Re-targeting"
IAB = "eDigital Network Display - Standard IAB"


@pytest.fixture(autouse=True)
def _no_market_config_db(monkeypatch):
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


def _req(**overrides) -> ProposalRequest:
    fields = dict(client_name="Acme Corp", request_type="New Business", start_date="2026-11-24",
                  end_date="2026-12-31", total_months=2, geo="Los Angeles", agency_fee=0.15)
    fields.update(overrides)
    return ProposalRequest(**fields)


def _export(tmp_path, *, time_unit="full_flight", tiers=None, force_tabs=None, lines=None):
    out = tmp_path / "x.xlsx"
    lines = lines if lines is not None else [
        LineItem(product_name=EMAIL, monthly_budget=680.0, months=1 if time_unit == "full_flight" else 2, id="a"),
        LineItem(product_name=IAB, monthly_budget=5470.0, months=1 if time_unit == "full_flight" else 2, id="b")]
    summary = generate_proposal(_req(), lines, out, tiers=tiers, time_unit=time_unit, force_tabs=force_tabs)
    return summary, openpyxl.load_workbook(out)


def _column(ws, col="C", start=1):
    return [ws[f"{col}{r}"].value for r in range(start, ws.max_row + 1)]


# --- Full Flight: no Months, SUBTOTAL DIGITAL --------------------------------------------------------------------------------

def test_full_flight_shows_no_months_and_calls_the_first_total_subtotal(tmp_path):
    _, wb = _export(tmp_path)
    for name in ("Proposal A", "Proposal A (Gross)"):
        ws = wb[name]
        assert ws["H10"].value is None and ws["I10"].value is None, name
        labels = _column(ws)
        assert "SUBTOTAL DIGITAL" in labels and "TOTAL DIGITAL — CAMPAIGN" in labels, name
        assert "TOTAL DIGITAL" not in labels, name


def test_other_billing_periods_keep_their_months_cell_and_total_label(tmp_path):
    _, wb = _export(tmp_path, time_unit="month")
    ws = wb["Proposal A"]
    assert ws["H10"].value == "Months:" and ws["I10"].value == 2
    assert "TOTAL DIGITAL MONTHLY" in _column(ws) and "SUBTOTAL DIGITAL" not in _column(ws)


# --- banners ------------------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("name,kind,expected", [
    ("ViX Only", "Proposal", "ViX Only - Proposal"),
    ("Option A", "Proposal", "Option A - Proposal"),
    ("Spanish Networks", "Proposal", "Spanish Networks - Proposal"),
    ("Proposal A", "Proposal", "Proposal A"),                       # already says it: never "Proposal A - Proposal"
    ("My PROPOSAL (rev 2)", "Proposal", "My PROPOSAL (rev 2)"),
    ("Acme Corp", "Avails", "Acme Corp - Avails"),
    ("Avails Only", "Avails", "Avails Only"),
    ("Proposal A", "Avails", "Proposal A - Avails"),                # the Avails sheet still says what it is
    ("  Acme   Corp  ", "Proposal", "Acme Corp - Proposal"),
    ("", "Proposal", "TBD - Proposal"),
    (None, "Avails", "TBD - Avails"),
])
def test_banner_title(name, kind, expected):
    assert _banner_title(name, kind) == expected


def test_a_single_option_banner_is_the_client_name_plus_proposal(tmp_path):
    _, wb = _export(tmp_path, force_tabs={"net": True, "wsections": True, "gross": True, "avails_only": True})
    assert wb["Proposal A"]["E2"].value == "Acme Corp - Proposal"
    assert wb["Proposal A (wsections)"]["E2"].value == "Acme Corp - Proposal"
    assert wb["Proposal A (Gross)"]["E2"].value == "Acme Corp - Proposal"
    assert wb["Avails-Only"]["E2"].value == "Acme Corp - Avails"


def test_each_option_banner_uses_that_options_own_name(tmp_path):
    lines = [LineItem(product_name=EMAIL, monthly_budget=500.0, months=1, id="x")]
    tiers = [{"label": "A", "name": "ViX Only", "line_items": lines, "avails_data": {}},
             {"label": "B", "name": "Proposal B", "line_items": lines, "avails_data": {}},
             {"label": "C", "line_items": lines, "avails_data": {}}]
    _, wb = _export(tmp_path, tiers=tiers, lines=lines, force_tabs={"net": True, "avails_only": True})
    banners = {ws.title: ws["E2"].value for ws in wb.worksheets if ws["E2"].value}
    assert "ViX Only - Proposal" in banners.values()
    assert "Proposal B" in banners.values() and "Proposal B - Proposal" not in banners.values()
    assert "Option C - Proposal" in banners.values()
    assert "ViX Only - Avails" in banners.values() and "Option C - Avails" in banners.values()


# --- Gross: Agency Comm legend, eCPM from NET -------------------------------------------------------------------------------

def test_gross_sheet_has_the_agency_comm_legend_next_to_the_percent(tmp_path):
    _, wb = _export(tmp_path)
    ws = wb["Proposal A (Gross)"]
    assert ws["H14"].value == "Agency Comm:" and ws["H14"].font.bold
    assert ws["H14"].alignment.horizontal == "right"
    assert ws["I14"].value == 0.15 and ws["I14"].number_format == "0%"
    assert ws["I14"].fill.start_color.rgb.endswith("FFFF00") and ws["I14"].font.bold      # the highlighted input


def test_a_fractional_agency_fee_still_shows_its_decimals(tmp_path):
    out = tmp_path / "frac.xlsx"
    generate_proposal(_req(agency_fee=0.125), [LineItem(product_name=EMAIL, monthly_budget=680.0, months=1, id="a")], out,
                      time_unit="full_flight")
    ws = openpyxl.load_workbook(out)["Proposal A (Gross)"]
    assert ws["I14"].value == 0.125 and ws["I14"].number_format == "0.00%"


def test_gross_ecpm_is_calculated_from_the_net_budget_like_the_net_sheet(tmp_path):
    _, wb = _export(tmp_path)
    gross, net = wb["Proposal A (Gross)"], wb["Proposal A"]
    g_total = next(r for r in range(19, gross.max_row) if gross[f"C{r}"].value == "SUBTOTAL DIGITAL")
    n_total = next(r for r in range(19, net.max_row) if net[f"C{r}"].value == "SUBTOTAL DIGITAL")
    assert gross[f"J{g_total}"].value == '="eCPM:"'
    assert gross[f"K{g_total}"].value == f'=IFERROR(L{g_total}/I{g_total}*1000,"")'          # L = NET BUDGET, not N = GROSS
    assert net[f"K{n_total}"].value == f'=IFERROR(L{n_total}/I{n_total}*1000,"")'            # the Net sheet's own formula


# --- the Process FAQs tab ------------------------------------------------------------------------------------------------------

def _wb_with_faqs(entries):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return et.build_process_faqs(wb, entries=entries)


def test_the_default_faq_rows_are_the_ones_the_tab_always_had():
    rows = faq.default_entries()
    assert [r["section"] for r in rows].count("faq") == 6 and [r["section"] for r in rows].count("restriction") == 5
    assert rows[0]["services"] == "All Products" and rows[0]["link_title"] == "Entravision's Digital Specs and Best Practices"
    assert all(r["link_url"] == "" for r in rows)                  # their URLs were never recorded - an admin adds them


def test_the_sheet_writes_real_hyperlinks_and_a_header_link_to_the_guidelines():
    ws = _wb_with_faqs([
        {"section": "faq", "services": "All Products", "question": "Where are the specs?",
         "link_title": "Specs", "link_url": "https://example.com/specs"},
        {"section": "faq", "services": "TikTok Ads", "question": "How do I give access?", "link_title": "No link yet", "link_url": ""},
        {"section": "restriction", "services": "Alcohol — All products", "question": "Related to: alcohol.",
         "link_title": "FTC Guidelines", "link_url": "https://ftc.example/alcohol"},
    ])
    assert ws["C4"].value == "See latest information about Entravision's Ad Guidelines"
    assert ws["C4"].hyperlink.target == "https://entravision-digital-ad-guidelines.netlify.app/"
    assert [ws[f"{c}6"].value for c in "CDE"] == ["SERVICES", "Frequently Asked Question", "Resource Link"]
    assert (ws["C8"].value, ws["D8"].value, ws["E8"].value) == ("All Products", "Where are the specs?", "Specs")
    assert ws["E8"].hyperlink.target == "https://example.com/specs" and ws["E8"].font.underline == "single"
    assert ws["E9"].value == "No link yet" and ws["E9"].hyperlink is None and not ws["E9"].font.underline
    assert ws["C11"].value == "CATEGORY AND PRODUCT RESTRICTIONS"
    assert ws["C12"].value == "Alcohol — All products" and ws["E12"].hyperlink.target == "https://ftc.example/alcohol"


def test_text_that_looks_like_a_formula_stays_text_and_tall_rows_get_room():
    long_question = "x " * 200
    ws = _wb_with_faqs([{"section": "faq", "services": "=SUM(A1:A2)", "question": long_question, "link_title": "", "link_url": ""}])
    assert ws["C8"].value == "=SUM(A1:A2)" and ws["C8"].data_type == "s"
    assert ws.row_dimensions[8].height > 60


def test_a_workbook_without_a_restriction_block_has_no_empty_heading():
    ws = _wb_with_faqs([{"section": "faq", "services": "A", "question": "Q", "link_title": "", "link_url": ""}])
    assert "CATEGORY AND PRODUCT RESTRICTIONS" not in [c for row in ws.iter_rows(values_only=True) for c in row]


def test_generate_builds_the_faq_tab_from_the_stored_rows(monkeypatch, tmp_path):
    monkeypatch.setattr(faq, "load_entries", lambda: [
        {"id": 1, "section": "faq", "services": "Stored Service", "question": "Stored question?",
         "link_title": "Stored link", "link_url": "https://stored.example/x", "sort_order": 0}])
    _, wb = _export(tmp_path)
    ws = wb["Process FAQs"]
    assert ws["C8"].value == "Stored Service" and ws["E8"].hyperlink.target == "https://stored.example/x"
    assert ws["C9"].value is None


# --- faq_entries: validation -----------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("raw,expected", [
    ("", ""), ("   ", ""), ("https://example.com/a?b=1", "https://example.com/a?b=1"),
    ("http://example.com", "http://example.com"), ("example.com/page", "https://example.com/page"),
    ("  www.example.com  ", "https://www.example.com"),
])
def test_clean_url_accepts_web_addresses(raw, expected):
    assert faq.clean_url(raw) == expected


@pytest.mark.parametrize("raw", ["javascript:alert(1)", "data:text/html,x", "file:///etc/passwd", "mailto:a@b.com",
                                 "ftp://example.com/x", "https://", "https://nodots", "two words.com", "https://a b.com"])
def test_clean_url_rejects_anything_that_is_not_a_web_link(raw):
    with pytest.raises(ValueError):
        faq.clean_url(raw)


def test_clean_entry_requires_services_and_question_and_defaults_the_link_title():
    with pytest.raises(ValueError, match="SERVICES"):
        faq.clean_entry({"question": "Q"})
    with pytest.raises(ValueError, match="question"):
        faq.clean_entry({"services": "S"})
    with pytest.raises(ValueError, match="Section"):
        faq.clean_entry({"services": "S", "question": "Q", "section": "other"})
    entry = faq.clean_entry({"services": " S\x00 ", "question": "Q", "link_url": "example.com/x"})
    assert entry == {"section": "faq", "services": "S", "question": "Q", "link_title": "https://example.com/x",
                     "link_url": "https://example.com/x"}
    with pytest.raises(ValueError, match="too long"):
        faq.clean_entry({"services": "S" * 400, "question": "Q"})
    # an update keeps what it doesn't mention
    base = {"section": "restriction", "services": "S", "question": "Q", "link_title": "T", "link_url": "https://a.example/x"}
    assert faq.clean_entry({"question": "New"}, base) == {**base, "question": "New"}


# --- faq_entries: storage (a small in-memory stand-in for the table) --------------------------------------------------------------

class _Result:
    def __init__(self, rows):
        self._rows = rows

    def fetchone(self):
        return self._rows[0] if self._rows else None

    def fetchall(self):
        return list(self._rows)


class FakeFaqDB:
    def __init__(self):
        self.rows, self.next_id = [], 1

    def __call__(self):
        return self

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def _sorted(self, rows=None):
        return sorted(rows if rows is not None else self.rows,
                      key=lambda r: (0 if r["section"] == "faq" else 1, r["sort_order"], r["id"]))

    def execute(self, sql, params=()):
        q = re.sub(r"\s+", " ", sql).strip()
        p = list(params)
        if q.startswith("SELECT id, section, services, question, link_title, link_url, sort_order FROM faq_entries ORDER BY"):
            return _Result([dict(r) for r in self._sorted()])
        if q.startswith("SELECT id, section, services, question, link_title, link_url, sort_order FROM faq_entries WHERE id"):
            return _Result([dict(r) for r in self.rows if r["id"] == p[0]])
        if q.startswith("INSERT INTO faq_entries"):
            row = dict(zip(["section", "services", "question", "link_title", "link_url", "sort_order"], p), id=self.next_id)
            self.next_id += 1
            self.rows.append(row)
            return _Result([dict(row)] if "RETURNING" in q else [])
        if q.startswith("SELECT COALESCE(MAX(sort_order)"):
            return _Result([{"n": max([r["sort_order"] for r in self.rows if r["section"] == p[0]], default=-1) + 1}])
        if q.startswith("UPDATE faq_entries SET section"):
            row = next(r for r in self.rows if r["id"] == p[-1])
            row.update(dict(zip(["section", "services", "question", "link_title", "link_url", "sort_order"], p[:6])))
            return _Result([dict(row)])
        if q.startswith("DELETE FROM faq_entries"):
            hit = [r for r in self.rows if r["id"] == p[0]]
            self.rows = [r for r in self.rows if r["id"] != p[0]]
            return _Result([{"id": hit[0]["id"]}] if hit else [])
        if q.startswith("SELECT id, section FROM faq_entries WHERE id"):
            return _Result([{"id": r["id"], "section": r["section"]} for r in self.rows if r["id"] == p[0]])
        if q.startswith("SELECT id FROM faq_entries WHERE section"):
            return _Result([{"id": r["id"]} for r in sorted((r for r in self.rows if r["section"] == p[0]),
                                                           key=lambda r: (r["sort_order"], r["id"]))])
        if q.startswith("UPDATE faq_entries SET sort_order"):
            next(r for r in self.rows if r["id"] == p[1])["sort_order"] = p[0]
            return _Result([])
        raise AssertionError(f"FakeFaqDB doesn't understand: {q[:100]}")


@pytest.fixture
def faq_db(monkeypatch):
    fake = FakeFaqDB()
    monkeypatch.setattr(faq, "get_connection", fake)
    return fake


def test_the_first_admin_read_seeds_the_built_in_rows_once(faq_db):
    first = faq.list_entries()
    assert len(first) == 11 and first[0]["id"] and first[0]["section"] == "faq" and first[-1]["section"] == "restriction"
    assert len(faq.list_entries()) == 11                           # not seeded again


def test_create_update_move_and_delete_keep_each_block_in_order(faq_db):
    faq.list_entries()
    added = faq.create_entry({"section": "faq", "services": "New Service", "question": "New?", "link_url": "example.com/n"})
    assert added["sort_order"] == 6 and added["link_url"] == "https://example.com/n"
    assert [e["services"] for e in faq.list_entries() if e["section"] == "faq"][-1] == "New Service"

    faq.move_entry(added["id"], "up")
    names = [e["services"] for e in faq.list_entries() if e["section"] == "faq"]
    assert names[-2] == "New Service" and names[-1] == "TikTok Ads"
    faq.move_entry(added["id"], "up")
    faq.move_entry(added["id"], "up")
    faq.move_entry(added["id"], "up")
    faq.move_entry(added["id"], "up")
    faq.move_entry(added["id"], "up")
    faq.move_entry(added["id"], "up")                              # past the top: a no-op, never an error
    assert [e["services"] for e in faq.list_entries() if e["section"] == "faq"][0] == "New Service"
    assert sorted(e["sort_order"] for e in faq.list_entries() if e["section"] == "faq") == list(range(7))   # no ties

    changed = faq.update_entry(added["id"], {"section": "restriction", "question": "Now a restriction"})
    assert changed["section"] == "restriction" and changed["question"] == "Now a restriction"
    others = [e["sort_order"] for e in faq.list_entries() if e["section"] == "restriction" and e["id"] != added["id"]]
    assert changed["sort_order"] > max(others)                     # appended to the end of its new block

    assert faq.delete_entry(added["id"]) is True and faq.delete_entry(added["id"]) is False
    with pytest.raises(LookupError):
        faq.update_entry(added["id"], {"question": "x"})
    with pytest.raises(LookupError):
        faq.move_entry(added["id"], "up")
    with pytest.raises(ValueError):
        faq.move_entry(1, "sideways")


def test_a_write_on_an_empty_table_seeds_it_first(faq_db):
    faq.create_entry({"services": "Only One", "question": "Q"})
    assert len(faq_db.rows) == 12


def test_the_export_falls_back_to_the_built_in_rows_when_the_table_is_empty_or_unreadable(faq_db, monkeypatch):
    load = faq._real_load_entries                                   # (conftest swaps load_entries for the built-ins)
    assert [r["services"] for r in load()] == [r["services"] for r in faq.default_entries()]

    def down():
        raise ConnectionError("database is down")
    monkeypatch.setattr(faq, "get_connection", down)
    assert len(load()) == 11                                       # never raises: generating a proposal must not depend on this
    with pytest.raises(RuntimeError, match="Couldn't read"):
        faq.list_entries()


def test_a_missing_table_gets_the_apply_schema_message(monkeypatch):
    class UndefinedTable(Exception):
        pass

    def missing():
        raise UndefinedTable("relation faq_entries does not exist")
    monkeypatch.setattr(faq, "get_connection", missing)
    with pytest.raises(RuntimeError, match="schema.sql"):
        faq.list_entries()


# --- the admin endpoints ----------------------------------------------------------------------------------------------------------

def test_the_faq_endpoints_map_errors_to_http_statuses(faq_db):
    import app.main as m
    out = asyncio.run(m.admin_list_faq())
    assert len(out["entries"]) == 11 and out["guidelines_url"] == "https://entravision-digital-ad-guidelines.netlify.app/"
    created = asyncio.run(m.admin_create_faq(m.FaqEntryRequest(services="S", question="Q", link_url="example.com")))
    assert created["added"] and created["entry"]["link_url"] == "https://example.com"
    with pytest.raises(HTTPException) as exc:
        asyncio.run(m.admin_create_faq(m.FaqEntryRequest(services="S", question="Q", link_url="javascript:alert(1)")))
    assert exc.value.status_code == 400
    with pytest.raises(HTTPException) as exc:
        asyncio.run(m.admin_update_faq(9999, m.FaqEntryRequest(question="x")))
    assert exc.value.status_code == 404
    with pytest.raises(HTTPException) as exc:
        asyncio.run(m.admin_delete_faq(9999))
    assert exc.value.status_code == 404
    with pytest.raises(HTTPException) as exc:
        asyncio.run(m.admin_move_faq(created["entry"]["id"], m.FaqMoveRequest(direction="left")))
    assert exc.value.status_code == 400
    assert asyncio.run(m.admin_move_faq(created["entry"]["id"], m.FaqMoveRequest(direction="up")))["moved"]
    assert asyncio.run(m.admin_delete_faq(created["entry"]["id"]))["deleted"]


def test_the_faq_tab_is_wired_into_the_admin_page():
    root = Path(__file__).resolve().parent.parent / "app"
    html = (root / "templates" / "admin.html").read_text(encoding="utf-8")
    js = (root / "static" / "admin.js").read_text(encoding="utf-8")
    assert 'data-tab="faq"' in html and 'id="tab-faq"' in html and 'id="faq-form"' in html
    assert 'id="faq-link-title"' in html and 'id="faq-link-url"' in html and 'id="faq-services"' in html
    assert "https://entravision-digital-ad-guidelines.netlify.app/" in html
    for needle in ('tab === "faq"', "function loadFaq", "function onSubmitFaq", "/api/admin/faq"):
        assert needle in js
    index = (root / "templates" / "index.html").read_text(encoding="utf-8")
    assert "See latest information about Entravision's Ad Guidelines" in index         # the Roadblocks step's link


# --- a live restrictions sheet ----------------------------------------------------------------------------------------------------

@pytest.fixture
def refresh(monkeypatch):
    import app.main as m
    calls = {"fetch": [], "ingest": []}
    state = {"sync": {"synced_at": datetime.now(timezone.utc) - timedelta(hours=3),
                      "sheet_url": "https://docs.google.com/spreadsheets/d/AAAAAAAAAAAAAAAAAAAAAAAAAAAA/edit",
                      "source_title": "Restricted Verticals", "synced_by": "admin@x"},
             "using_builtin": False, "error": False, "categories": []}
    monkeypatch.setattr(m.restrictions_svc, "load_state", lambda: state)
    monkeypatch.setattr(m.restrictions_svc, "fetch_sheet_xlsx", lambda file_id: calls["fetch"].append(file_id) or (b"xlsx", "Sheet Title"))
    monkeypatch.setattr(m, "_ingest_restrictions", lambda data, url, title, by: calls["ingest"].append((data, url, title, by)) or [])
    monkeypatch.setattr(m, "_RESTRICTIONS_REFRESH_MINUTES", 60.0)
    m._restrictions_refresh["attempted_at"] = None
    return m, state, calls


def test_a_stale_synced_copy_is_re_read_from_the_sheet(refresh):
    m, state, calls = refresh
    assert m._refresh_restrictions_if_stale() is True
    assert calls["fetch"] == ["AAAAAAAAAAAAAAAAAAAAAAAAAAAA"] and calls["ingest"][0][3] == "auto-sync"
    assert calls["ingest"][0][1].endswith("/edit") and calls["ingest"][0][2] == "Sheet Title"


@pytest.mark.parametrize("change", [
    {"synced_at": datetime.now(timezone.utc) - timedelta(minutes=10)},        # still fresh
    {"source_title": "Uploaded: restrictions.xlsx"},                          # a deliberate upload is never overridden
])
def test_a_fresh_or_uploaded_copy_is_left_alone(refresh, change):
    m, state, calls = refresh
    state["sync"].update(change)
    assert m._refresh_restrictions_if_stale() is False and not calls["fetch"]


def test_the_built_in_snapshot_a_failed_read_and_a_never_synced_table_are_never_replaced(refresh):
    m, state, calls = refresh
    for change in ({"using_builtin": True}, {"error": True}, {"sync": None}):
        saved = dict(state)
        state.update(change)
        assert m._refresh_restrictions_if_stale() is False
        state.clear()
        state.update(saved)
    assert not calls["fetch"]


def test_a_failed_refresh_keeps_the_stored_copy_and_is_not_retried_every_request(refresh, monkeypatch):
    m, state, calls = refresh

    def boom(file_id):
        calls["fetch"].append(file_id)
        raise ValueError("Google Drive isn't connected yet")
    monkeypatch.setattr(m.restrictions_svc, "fetch_sheet_xlsx", boom)
    assert m._refresh_restrictions_if_stale() is False
    assert m._refresh_restrictions_if_stale() is False
    assert len(calls["fetch"]) == 1                                # the second call was throttled
    assert not calls["ingest"]


def test_the_refresh_can_be_switched_off(refresh, monkeypatch):
    m, state, calls = refresh
    monkeypatch.setattr(m, "_RESTRICTIONS_REFRESH_MINUTES", 0.0)
    assert m._refresh_restrictions_if_stale() is False and not calls["fetch"]


def test_the_chooser_and_the_check_both_refresh_first(refresh, monkeypatch):
    m, state, calls = refresh
    order = []
    monkeypatch.setattr(m, "_refresh_restrictions_if_stale", lambda: order.append("refresh") or True)
    state["categories"] = []

    class _Req:
        request = {"client_name": "Acme"}
        product_names = []
    monkeypatch.setattr(m, "_reconstruct_proposal_request", lambda d: ProposalRequest())
    asyncio.run(m.restriction_categories(_Req()))
    assert order == ["refresh"]
