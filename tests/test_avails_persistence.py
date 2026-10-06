"""
Avails persistence end to end: what a planner enters in Step 05 must reach
the exported workbook — per line, per option, including two lines of the same
product — and the paths that used to drop it silently (an all-off tab set,
uniques-only entries, Added Value lines, "500k" typed as 500) must hold.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
from pathlib import Path

os.environ.setdefault("DATABASE_URL", "postgresql://fake:fake@localhost/fake")

import openpyxl
import pytest

from app.services.notion_parser import ProposalRequest
from app.services.proposal_generator import LineItem, generate_proposal

APP_JS = Path(__file__).resolve().parents[1] / "app" / "static" / "app.js"


@pytest.fixture(autouse=True)
def _mock_catalog_db(monkeypatch):
    monkeypatch.setattr("app.catalog.load_rate_overrides", lambda: {})
    monkeypatch.setattr("app.catalog.load_custom_products", lambda: [])
    monkeypatch.setattr("app.catalog.load_deleted_builtin_names", lambda: set())
    monkeypatch.setattr("app.market_config.load_market_config", lambda: {})


DISPLAY = "eDigital Network Display - Standard IAB"


def _req(**kw):
    base = dict(client_name="Acme", request_type="New Business", start_date="2026-10-01",
                end_date="2026-12-31", total_months=3)
    base.update(kw)
    return ProposalRequest(**base)


def _line_rows(ws, product_prefix):
    """Row numbers of the product lines (column C starts with the name)."""
    return [r for r in range(1, ws.max_row + 1)
            if isinstance(ws.cell(row=r, column=3).value, str) and ws.cell(row=r, column=3).value.startswith(product_prefix)]


def _tier(label, lines, avails):
    return {"label": label, "line_items": lines, "avails_data": avails}


# --- per line / per option ------------------------------------------------------------

def test_two_lines_of_the_same_product_keep_their_own_avails(tmp_path):
    a = LineItem(product_name=DISPLAY, monthly_budget=1000.0, months=3, id="a", target_override="A18+")
    b = LineItem(product_name=DISPLAY, monthly_budget=2000.0, months=3, id="b", target_override="Hispanic A25-54")
    avails = {"a": {"max_imps": 100000, "est_uniques": 40000, "basis": "imps"},
              "b": {"max_imps": 500000, "est_uniques": 90000, "basis": "imps"}}
    out = tmp_path / "dup.xlsx"
    generate_proposal(_req(), [], out, tiers=[_tier("A", [a, b], avails)])
    ws = openpyxl.load_workbook(out)["Proposal A"]
    rows = _line_rows(ws, DISPLAY)
    assert len(rows) == 2
    assert [ws[f"N{r}"].value for r in rows] == [100000, 500000]
    assert [ws[f"P{r}"].value for r in rows] == [40000, 90000]


def test_each_option_exports_its_own_avails(tmp_path):
    a = LineItem(product_name=DISPLAY, monthly_budget=1000.0, months=3, id="a1")
    b = LineItem(product_name=DISPLAY, monthly_budget=1000.0, months=3, id="b1")
    tiers = [
        _tier("A", [a], {"a1": {"max_imps": 111111, "basis": "imps"}}),
        _tier("B", [b], {"b1": {"max_imps": 222222, "basis": "imps"}}),
    ]
    out = tmp_path / "tiers.xlsx"
    generate_proposal(_req(), [], out, tiers=tiers)
    wb = openpyxl.load_workbook(out)
    ws_a = next(w for w in wb.worksheets if "Option A" in w.title)
    ws_b = next(w for w in wb.worksheets if "Option B" in w.title)
    assert ws_a[f"N{_line_rows(ws_a, DISPLAY)[0]}"].value == 111111
    assert ws_b[f"N{_line_rows(ws_b, DISPLAY)[0]}"].value == 222222


def test_freeform_text_avails_reach_the_net_sheet(tmp_path):
    li = LineItem(product_name="Search - SEM", monthly_budget=1000.0, months=3, id="s")
    avails = {"s": {"freeform": True, "max_imps_text": "50 to 100", "max_spend_text": "TBD", "est_uniques_text": "n/a"}}
    out = tmp_path / "ff.xlsx"
    generate_proposal(_req(), [], out, tiers=[_tier("A", [li], avails)])
    ws = openpyxl.load_workbook(out)["Proposal A"]
    r = _line_rows(ws, "Search - SEM")[0]
    assert (ws[f"N{r}"].value, ws[f"O{r}"].value, ws[f"P{r}"].value) == ("50 to 100", "TBD", "n/a")


# --- the all-tabs-off path --------------------------------------------------------------

@pytest.mark.parametrize("override", [
    {"net": False, "wsections": False, "gross": False, "avails_only": False, "dooh_summary": False, "dooh_screenlist": False},
])
def test_an_all_off_tab_override_falls_back_to_the_suggested_tabs_and_says_so(tmp_path, override):
    li = LineItem(product_name=DISPLAY, monthly_budget=1000.0, months=3, id="a")
    out = tmp_path / "alloff.xlsx"
    summary = generate_proposal(_req(), [], out, tiers=[_tier("A", [li], {"a": {"max_imps": 123456, "basis": "imps"}})],
                                force_tabs=override)
    assert summary["tabs_built"], "an all-off override must never produce an empty workbook"
    assert any("suggested tabs were generated instead" in w for w in summary["warnings"])
    ws = openpyxl.load_workbook(out)["Proposal A"]
    assert ws[f"N{_line_rows(ws, DISPLAY)[0]}"].value == 123456


def test_a_partial_override_changes_only_what_the_planner_toggled(tmp_path):
    li = LineItem(product_name=DISPLAY, monthly_budget=1000.0, months=3, id="a")
    out = tmp_path / "partial.xlsx"
    summary = generate_proposal(_req(agency_fee=0.15), [], out, tiers=[_tier("A", [li], {})],
                                force_tabs={"gross": False})   # net/wsections/etc. stay on the server's own suggestion
    assert "Proposal A" in summary["tabs_built"]
    assert not any("Gross" in t for t in summary["tabs_built"])


def test_the_server_suggestion_follows_the_final_plan_not_a_stale_parse_time_one(tmp_path):
    li = LineItem(product_name=DISPLAY, monthly_budget=1000.0, months=3, id="a")
    # No force_tabs at all (nothing touched): an agency fee set later in the
    # flow still yields the Gross tab, which a stale all-false override used to drop.
    out = tmp_path / "fresh.xlsx"
    summary = generate_proposal(_req(agency_fee=0.15), [], out, tiers=[_tier("A", [li], {})], force_tabs={})
    assert any("Gross" in t for t in summary["tabs_built"])


def test_suggest_tabs_endpoint_matches_the_generate_classification():
    import asyncio
    import app.main as m
    body = m.SuggestTabsRequest(request={"request_type": "New Business", "agency_fee": 0.15},
                                product_names=[DISPLAY])
    out = asyncio.run(m.suggest_tabs(body))["tabs"]
    assert out["net"] is True and out["gross"] is True and out["avails_only"] is False
    plain = asyncio.run(m.suggest_tabs(m.SuggestTabsRequest(request={"request_type": "New Business"}, product_names=[DISPLAY])))["tabs"]
    assert plain["gross"] is False


# --- Avails-Only tab ----------------------------------------------------------------------

def test_avails_only_keeps_a_line_whose_only_entry_is_est_uniques(tmp_path):
    li = LineItem(product_name=DISPLAY, monthly_budget=1000.0, months=3, id="a")
    out = tmp_path / "ao.xlsx"
    generate_proposal(_req(), [], out, force_tabs={"avails_only": True},
                      tiers=[_tier("A", [li], {"a": {"est_uniques": 50000}})])
    wb = openpyxl.load_workbook(out)
    ws = next(w for w in wb.worksheets if "Avails" in w.title)
    r = _line_rows(ws, DISPLAY)[0]
    assert ws[f"L{r}"].value == 50000


# --- Added Value lines --------------------------------------------------------------------

def test_added_value_lines_export_the_stored_derived_value_not_a_zero_rate_formula(tmp_path):
    li = LineItem(product_name=DISPLAY, monthly_budget=0.0, months=3, id="av", is_added_value=True)
    # the planner typed SPEND; the UI derived imps at the catalog rate and stored both
    avails = {"av": {"max_spend": 4500, "max_imps": 500000, "basis": "spend"}}
    out = tmp_path / "av.xlsx"
    generate_proposal(_req(), [], out, tiers=[_tier("A", [li], avails)])
    ws = openpyxl.load_workbook(out)["Proposal A"]
    r = _line_rows(ws, DISPLAY)[0]
    assert ws[f"N{r}"].value == 500000          # a number, not "=IFERROR(O19/K19*1000,0)" evaluating to 0
    assert ws[f"O{r}"].value == 4500


def test_normal_lines_still_use_live_formulas_for_the_derived_side(tmp_path):
    li = LineItem(product_name=DISPLAY, monthly_budget=1000.0, months=3, id="n")
    avails = {"n": {"max_spend": 4500, "max_imps": 500000, "basis": "spend"}}
    out = tmp_path / "n.xlsx"
    generate_proposal(_req(), [], out, tiers=[_tier("A", [li], avails)])
    ws = openpyxl.load_workbook(out)["Proposal A"]
    r = _line_rows(ws, DISPLAY)[0]
    assert str(ws[f"N{r}"].value).startswith("=")


# --- the Step 05 / Step 04 input parsing (runs the real app.js function in Node) ---------

def _node_parse(values):
    node = shutil.which("node")
    if not node:
        pytest.skip("node not installed")
    src = APP_JS.read_text(encoding="utf-8")
    start = src.index("function parseFormattedInput(s) {")
    depth, i = 0, src.index("{", start)
    for j in range(i, len(src)):
        depth += src[j] == "{"
        depth -= src[j] == "}"
        if depth == 0:
            fn = src[start:j + 1]
            break
    script = fn + f"\nconsole.log(JSON.stringify({json.dumps(values)}.map(parseFormattedInput)));"
    return json.loads(subprocess.run([node, "-e", script], capture_output=True, text=True, check=True).stdout)


def test_parse_formatted_input_understands_k_m_b_suffixes_instead_of_corrupting_them():
    out = _node_parse(["500k", "1.5M", "2b", "$2.5K", "750 imps", "1,500,000", "3.5m impressions"])
    assert out[:4] == [500000, 1500000, 2000000000, 2500]
    assert out[5] == 1500000 and out[6] == 3500000


def test_parse_formatted_input_keeps_its_old_behavior_for_plain_and_formatted_numbers():
    out = _node_parse(["35,000", "35,000.00", "Est. 333,462", "$1,200", "", "abc", "5.5", "12%"])
    assert out == [35000, 35000, 333462, 1200, None, None, 5.5, 12]
