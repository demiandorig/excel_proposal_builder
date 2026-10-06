"""
Restricted-verticals matrix — Entravision's own "Restricted Verticals by
Platforms" Google Sheet as the single source of truth for whether a client
vertical (Political, Cannabis, Casino/Gambling, ...) is allowed on each
product, used by the Roadblocks step.

Why this exists: the Roadblocks check used to rely on two hard-coded flags in
catalog.py (political_policy / cannabis_policy) plus whatever the model found
by searching the web. Those flags drifted from reality — e.g. Video Pre-roll
(OLV) and eDigital Display were both marked "not_allowed" for political ads
even though the sheet's Political tab lists them as accepted — and nothing
checked the model's verdict against anything.

HOW THE SHEET IS READ. Each tab is one vertical, and the tabs aren't uniform,
so each is classified by what it actually says:
  - allowlist     "Political Ads are only accepted for the below products" +
                  the list. A product is ALLOWED if the sheet lists it,
                  NOT ACCEPTED if it doesn't (non-media services excepted).
  - blanket_hold  "All Entravision Digital Products are currently on hold and
                  will not accept cannabis ads" — NOT ACCEPTED everywhere.
  - guidance      Anything else (platform-by-platform grids like Casino/
                  Gambling, bullet notes like Blood Donation). Free-text
                  advice isn't reduced to a yes/no — it's passed to the model
                  verbatim as authoritative context and shown to the planner.
The deterministic verdicts (allowlist / blanket_hold) are computed here, handed
to the model as facts it may not contradict, and enforced again in code after
the model answers (see roadblocks._apply_matrix).

The sheet is ingested, not read live: an admin syncs it (through the app's
existing Google Drive connection) or uploads an .xlsx export, and the result
is stored in Postgres. Until then a built-in snapshot of the two tabs with
clear-cut product rules (Political, Cannabis) applies, and DB trouble never
breaks Roadblocks — see load_state().
"""
from __future__ import annotations

import base64
import io
import logging
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any, Iterable, Optional

from app.db import get_connection

log = logging.getLogger(__name__)

DEFAULT_SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/1FpWkPQRXim8_CGkUxagmQ19JuKtBSn29XoRqzoquAp0/edit"
)
KINDS = ("allowlist", "blanket_hold", "guidance")
# Services and measurement aren't ad inventory, so an "approved products" list
# for a vertical says nothing about them.
NON_MEDIA_FAMILIES = {"Services", "Measurement"}
MAX_BODY_IN_PROMPT = 8000
# Hard caps on what one workbook may contribute (an uploaded file is admin-supplied, but a stray cell far
# out in a sheet would otherwise make the parse walk millions of empty cells).
MAX_SHEETS = 40
MAX_ROWS = 400
MAX_COLS = 30
MAX_ITEM_CHARS = 160

STATUS_ORDER = {"not_allowed": 0, "guidance": 1, "allowed": 2, "n/a": 3}


@dataclass
class Category:
    name: str
    kind: str
    keywords: list[str] = field(default_factory=list)
    items: list[str] = field(default_factory=list)          # allowlist: product names exactly as the sheet lists them
    summary: str = ""
    body: str = ""                                          # the tab's text, verbatim — what the model reads
    product_map: dict[str, list[str]] = field(default_factory=dict)   # sheet item -> catalog product names

    def to_dict(self) -> dict[str, Any]:
        return {
            "name": self.name, "kind": self.kind, "keywords": list(self.keywords), "items": list(self.items),
            "summary": self.summary, "body": self.body, "product_map": {k: list(v) for k, v in self.product_map.items()},
        }


# ---------------------------------------------------------------------------
# Keywords (what in a plan suggests a vertical applies)
# ---------------------------------------------------------------------------

_KEYWORD_SEEDS: dict[str, list[str]] = {
    # Whole phrases / unambiguous words only: a keyword only SUGGESTS a vertical, but a suggestion the planner
    # has to untick is noise ("Pac-12 Network" is not a PAC, "plasma TV" is not plasma donation).
    "political": ["political", "election", "ballot", "ballot measure", "campaign committee",
                  "political action committee", "political candidate", "candidate for office", "running for office", "for congress", "for senate",
                  "for governor", "for mayor", "for sheriff", "for city council", "committee to elect", "re-elect",
                  "reelect", "vote yes on", "vote no on", "yes on prop", "no on prop", "get out the vote"],
    "cannabis": ["cannabis", "marijuana", "dispensary", "cbd", "thc"],
    "casino": ["casino", "gambling", "sportsbook", "sports betting", "lottery", "bingo", "poker"],
    "gambling": ["casino", "gambling", "sportsbook", "sports betting", "lottery", "bingo", "poker"],
    "blood": ["blood donation", "blood drive", "plasma donation", "plasma center"],
    "sexual": ["sexual enhancement", "male enhancement", "erectile"],
}
_STOPWORDS = {"and", "or", "the", "of", "for", "a", "an", "procedures", "procedure", "services", "service"}
# Tab-name words too generic to be a trigger on their own ("donation" alone is not blood donation).
_GENERIC_TAB_WORDS = {
    "donation", "donations", "ads", "advertising", "products", "product", "general", "other", "health", "medical",
    "financial", "policy", "policies", "guidelines", "restrictions", "restricted", "verticals", "industry", "content",
    "enhancement",
}


def default_keywords(tab_name: str) -> list[str]:
    tokens = [t for t in re.split(r"[^a-z0-9]+", tab_name.lower()) if t and t not in _STOPWORDS]
    out: list[str] = []
    for tok in tokens:
        if tok in _KEYWORD_SEEDS:
            candidates = _KEYWORD_SEEDS[tok]
        elif len(tokens) == 1 or (tok not in _GENERIC_TAB_WORDS and len(tok) >= 4):
            candidates = [tok]
        else:
            candidates = []
        for kw in candidates:
            if kw not in out:
                out.append(kw)
    full = tab_name.strip().lower()
    if full and full not in out:
        out.append(full)
    return out


def suggest_categories(plan_text: str, categories: list[Category]) -> list[str]:
    """Category names whose keywords appear (whole-word) in `plan_text`."""
    from app.disclaimers import text_matches
    return [c.name for c in categories if c.keywords and text_matches(c.keywords, plan_text)]


# ---------------------------------------------------------------------------
# Sheet-name -> catalog-product matching
# ---------------------------------------------------------------------------

def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", (s or "").lower())).strip()


def _named(products, pred):
    return [p.name for p in products if pred(p)]


def match_catalog_products(item: str, products: Optional[Iterable] = None) -> list[str]:
    """Catalog products a sheet entry ("eDigital Display Geo Fence", "CTV / OTT
    -- Entravision Plus Run-Of-Network (Does not include Netflix, Prime or
    Roku)", ...) stands for. Explicit rules for how this sheet names things,
    then a conservative token-overlap fallback. Deliberately errs toward
    matching FEWER products: an unmatched product reads "not on the approved
    list" (visible, fixable in Admin), whereas a wrongly matched one would
    read ALLOWED."""
    from app.catalog import CATALOG
    products = list(products if products is not None else CATALOG)
    t = _norm(item)
    low = lambda p: p.name.lower()
    geo = bool(re.search(r"geo ?fenc", t))      # "Geo Fence", "Geofencing", "Geo-Fencing", "Geo Fencing"

    if geo and ("display" in t or "olv" in t or "video" in t):
        # "Geofencing Display and Geofencing OLV" names both buys
        out: list[str] = []
        if "display" in t:
            out += _named(products, lambda p: p.family == "Display" and "geo fence" in low(p))
        if "olv" in t or "video" in t:
            out += _named(products, lambda p: p.family == "Online Video" and "geo fence" in low(p))
        return out
    if "display re targeting" in t:
        return _named(products, lambda p: "display re-targeting" in low(p))
    if "youtube" in t:
        return _named(products, lambda p: "youtube" in low(p))
    if re.search(r"\bedigital display\b|^display\b", t):
        return _named(products, lambda p: p.family == "Display" and "geo fence" not in low(p))
    if re.search(r"\bedigital olv\b|^olv\b|pre roll", t):
        return _named(products, lambda p: p.family == "Online Video" and "youtube" not in low(p) and "geo fence" not in low(p))
    if re.search(r"\bctv\b|\bott\b", t) and ("entravision plus" in t or "run of network" in t):
        # Generic run-of-network CTV/OTT only; named-property buys (Netflix, Prime,
        # Roku, VIX) and OLV Blend aren't "run of network" — and the sheet itself
        # names Netflix, Prime and Roku as excluded.
        return _named(products, lambda p: p.family == "Entravision Plus" and "ctv/ott" in low(p))
    if "audio streaming" in t and not re.search(r"spotify|pandora|iheart|siriusxm|tunein", t):
        return _named(products, lambda p: low(p) == "audio - evc audio streaming")
    if "audio engage" in t or "audioengage" in t:
        return _named(products, lambda p: low(p) == "audioengage")
    if t.startswith("meta ads") or "meta ads" in t:
        # "Only through client's handle, no branded content": the three core
        # Facebook & Instagram buys, not the Entravision-page or branded-content ones.
        return _named(products, lambda p: low(p).startswith("facebook & instagram ads |"))
    if re.fullmatch(r"email (?:marketing )?campaigns?", t):
        # only the plain entry — "Email Campaigns - Matchback" and friends are different products
        return _named(products, lambda p: low(p).startswith("email campaigns and/or"))
    if "digital out of home" in t or "dooh" in t:
        return _named(products, lambda p: p.family == "DOOH")

    tokens = set(t.split())
    best: list[str] = []
    for p in products:
        ptokens = set(_norm(p.name).split())
        if not tokens or not ptokens:
            continue
        overlap = len(tokens & ptokens) / len(tokens | ptokens)
        if overlap >= 0.8:                       # a near-exact name only; anything fuzzier is the admin's call
            best.append(p.name)
    return best


def auto_product_map(items: list[str], products: Optional[Iterable] = None) -> dict[str, list[str]]:
    return {item: match_catalog_products(item, products) for item in items}


# ---------------------------------------------------------------------------
# Built-in snapshot (applies until an admin syncs the sheet, and if the DB is unreachable)
# ---------------------------------------------------------------------------

_POLITICAL_ITEMS = [
    "eDigital Display",
    "eDigital Display Geo Fence",
    "eDigital OLV Geo Fence",
    "eDigital OLV",
    "CTV / OTT -- Entravision Plus Run-Of-Network (Does not include Netflix, Prime or Roku)",
    "Entravision O&O Audio Streaming",
    "Audio Engage",
    "Meta Ads (Only through client's handle, no branded content)",
    "Email Marketing Campaigns",
    "Email Campaigns - Display Re-targeting",
    "Digital Out of Home",
]
_CANNABIS_HEADLINE = "All Entravision Digital Products are currently on hold and will not accept cannabis ads due to legal roadblocks"


def builtin_categories() -> list[Category]:
    political = Category(
        name="Political", kind="allowlist", keywords=default_keywords("Political"), items=list(_POLITICAL_ITEMS),
        summary="Political Ads are only accepted for the products listed below.",
        body="Political Ads are only accepted for the below products:\n" + "\n".join(_POLITICAL_ITEMS),
    )
    political.product_map = auto_product_map(political.items)
    cannabis = Category(
        name="Cannabis", kind="blanket_hold", keywords=default_keywords("Cannabis"), items=[],
        summary=_CANNABIS_HEADLINE,
        body=(_CANNABIS_HEADLINE + ".\nOutside legal counsel advises that accepting cannabis ads on digital platforms "
              "remains risky, and Entravision has paused accepting them on digital until the issues settle.\n"
              "Previously approved (also on hold): CTV/OTT, eDigital OLV, DOOH (only through Wilkins Media), Email Marketing "
              "(LeadMe and SiteImpact), AudioEngage (AdsWizz and Triton), Landing Pages, GeoFencing Display and GeoFencing OLV "
              "(only through simpli.fi)."),
    )
    return [political, cannabis]


# ---------------------------------------------------------------------------
# Parsing the sheet (an .xlsx export of it)
# ---------------------------------------------------------------------------

_SKIP_TABS = {"policy document index"}
_LISTED = r"(?:the )?(?:below|following|these|listed)"
_ALLOWLIST_HEADER = re.compile(
    rf"\bonly\b.{{0,40}}\b(?:accept\w*|run|available|allowed|permitted|approved|eligible)\b.{{0,40}}\b(?:for|on|with|in)\b.{{0,20}}\b{_LISTED}\b|"
    rf"\b(?:accepted|approved|allowed|permitted)\b.{{0,20}}\b(?:for|on)\b.{{0,20}}\b{_LISTED}\b|"
    rf"\blimited to\b.{{0,20}}\b{_LISTED}\b", re.I)
# A header that says the OPPOSITE ("NOT accepted for the below", "Previously approved for the following ... on hold")
# is not an approved-products list.
_NEGATED = re.compile(
    r"\b(?:not|never|no longer|previously|used to|on hold|paused|suspended|cannot|prohibit\w*|formerly)\b|n[\u2019']t\b|\bwere\b|\bwas\b",
    re.I)
# "All Entravision Digital Products are currently on hold and will not accept cannabis ads" — and only a headline
# about Entravision's DIGITAL products as a whole (a note about one product family isn't a blanket hold).
_HOLD_PHRASE = r"\b(?:on hold|will not accept|do not accept|does not accept|not accepting|paused|suspended)\b"
_BLANKET_HOLD = re.compile(
    r"(?:\b(?:all|every)\b.{0,30}|\bentravision\b.{0,20})\bdigital\b.{0,30}\b(?:products?|platforms?|inventory)\b.{0,80}" + _HOLD_PHRASE + "|"
    r"\b(?:all|every)\b.{0,12}\bentravision\b.{0,25}\b(?:products?|platforms?|inventory)\b.{0,80}" + _HOLD_PHRASE + "|"
    + _HOLD_PHRASE + r".{0,60}\b(?:all|every|entravision)\b.{0,30}\bdigital\b", re.I)
_LIST_POINTER = re.compile(r"\b(?:below|following)\b", re.I)


def _is_allowlist_header(cell: str) -> bool:
    """An approved-products header — judged clause by clause, so "...only accepted for the below products. All other
    products are not accepting political ads." is still a header, while "NOT accepted for the below" is not."""
    for clause in re.split(r"[.;\n]+", re.sub(r"\([^)]*\)", " ", cell)):
        if _ALLOWLIST_HEADER.search(clause) and not _NEGATED.search(clause):
            return True
    return False


def _ends_list(cell: str) -> bool:
    """A row that can't be another approved product: a lead-in to another list ("...NOT accepted for the below:"),
    a heading ending in ":", or a sentence — unless it names a catalog product (an entry that happens to end with
    a period or say "see restrictions below" is still an entry)."""
    if cell.endswith(":") or _ALLOWLIST_HEADER.search(cell):
        return True
    if _LIST_POINTER.search(cell) and _NEGATED.search(cell):
        return True
    return cell.endswith(".") and len(cell.split()) > 6 and not match_catalog_products(cell)


def _clean(v: Any) -> str:
    if v is None:
        return ""
    s = str(v).replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"\*\*|__", "", s)                      # markdown bold left in the cells
    s = re.sub(r"^\s*[*•\-]\s+", "", s)           # a leading bullet marker
    return re.sub(r"[ \t]+", " ", s).strip()


_TRAILING_LINK = re.compile(r"\s*<https?://[^>]*>\s*$")


def _strip_link(text: str) -> str:
    """A cell's text without the " <url>" _cell_text appended for a hyperlink (the link stays in the body)."""
    return _TRAILING_LINK.sub("", text or "").strip()


def _cell_text(cell: Any) -> str:
    text = _clean(cell.value)
    link = getattr(cell, "hyperlink", None)
    target = getattr(link, "target", None) if link else None
    if target and target not in text:
        text = f"{text} <{target}>".strip()
    return text


def _has_content_beyond_caps(ws: Any) -> bool:
    """True only if a cell with a VALUE sits past the read caps — max_row/max_column also count styled-but-empty
    cells (a column's formatting running down to row 1000), which isn't worth a warning."""
    if ws.max_row > MAX_ROWS:
        for row in ws.iter_rows(min_row=MAX_ROWS + 1, max_row=min(ws.max_row, MAX_ROWS + 5000), values_only=True):
            if any(v not in (None, "") for v in row):
                return True
    if ws.max_column > MAX_COLS:
        for row in ws.iter_rows(max_row=MAX_ROWS, min_col=MAX_COLS + 1, values_only=True):
            if any(v not in (None, "") for v in row):
                return True
    return False


def parse_workbook_ex(data: bytes) -> tuple[list[Category], list[str]]:
    """(categories, warnings): every visible vertical tab of an .xlsx export of
    the sheet, classified, plus anything an admin should look at — a tab that
    couldn't be classified as expected, entries that matched no catalog product,
    text trimmed for the model. Raises ValueError if the file isn't a readable
    workbook, holds no verticals, or two tabs would collide."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"That doesn't look like an .xlsx workbook ({exc}).") from exc
    cats: list[Category] = []
    warnings: list[str] = []
    seen: dict[str, str] = {}
    if len(wb.worksheets) > MAX_SHEETS:
        warnings.append(f"The workbook has {len(wb.worksheets)} tabs; only the first {MAX_SHEETS} were read.")
    for ws in wb.worksheets[:MAX_SHEETS]:
        # Excel sheet names can't contain "/" (the sheet's "Casino/Gambling" tab), so an
        # export swaps forbidden characters for something else — read them as spaces.
        name = re.sub(r"[_]+", " ", ws.title or "").strip()
        if not name or name.lower() in _SKIP_TABS:
            continue
        if getattr(ws, "sheet_state", "visible") != "visible":
            warnings.append(f"Skipped the hidden tab \"{ws.title}\".")
            continue
        key = re.sub(r"\s+", " ", name).casefold()
        if key in seen:
            raise ValueError(f"Two tabs read as \"{name}\" (\"{seen[key]}\" and \"{ws.title}\") — rename one in the sheet.")
        seen[key] = ws.title
        if _has_content_beyond_caps(ws):
            warnings.append(f"\"{name}\" has content past {MAX_ROWS} rows x {MAX_COLS} columns; anything beyond that was not read.")
        rows = [[_cell_text(c) for c in row] for row in ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS)]
        cat = _parse_tab(name, rows, warnings)
        if cat:
            cats.append(cat)
            if len(cat.body) > MAX_BODY_IN_PROMPT:
                warnings.append(f"\"{cat.name}\" is {len(cat.body):,} characters; the model reads the first {MAX_BODY_IN_PROMPT:,}.")
    if not cats:
        raise ValueError("No restricted-vertical tabs found in that workbook.")
    warnings.extend(mapping_warnings(cats))
    return cats, warnings


def parse_workbook(data: bytes) -> list[Category]:
    return parse_workbook_ex(data)[0]


def mapping_warnings(cats: list[Category]) -> list[str]:
    """Sheet entries of an allowlist that matched no catalog product — those
    products read NOT ACCEPTED until an admin maps them."""
    out: list[str] = []
    for cat in cats:
        if cat.kind != "allowlist":
            continue
        unmapped = [item for item in cat.items if not (cat.product_map or {}).get(item)]
        if unmapped:
            out.append(
                f"{cat.name}: {len(unmapped)} of {len(cat.items)} listed entries matched no catalog product "
                f"({'; '.join(unmapped[:4])}{' …' if len(unmapped) > 4 else ''}). Map them in Admin → Restrictions "
                "or those products will read NOT ACCEPTED.")
    return out


def _parse_tab(name: str, rows: list[list[str]], warnings: Optional[list[str]] = None) -> Optional[Category]:
    warnings = warnings if warnings is not None else []
    if not any(any(r) for r in rows):
        return None

    # A platform grid: header row with a "Platform(s)" cell in its first two columns (Casino/Gambling, ...).
    grid_i, pcol = None, 0
    for i, r in enumerate(rows):
        hit = next((j for j, c in enumerate(r[:2]) if c.lower() in ("platform", "platforms")), None)
        if hit is not None:
            grid_i, pcol = i, hit
            break

    flat = [c for r in rows for c in r if c]
    body = "\n".join(flat)
    above = rows[:grid_i] if grid_i is not None else rows        # a headline ABOVE a grid still decides the tab's kind

    # The tab's two decisive lines: an approved-products header, and a blanket-hold headline (one of the first
    # three lines). Whichever comes first wins; a tab that has both says so.
    header_row = header_cell = None
    for i, r in enumerate(above):
        hit = next((c for c in r if c and _is_allowlist_header(c)), None)
        if hit:
            header_row, header_cell = i, hit
            break
    hold_row = hold_cell = None
    seen_rows = 0
    for i, r in enumerate(above):
        if not any(r):
            continue
        seen_rows += 1
        if seen_rows > 3:
            break                                       # the headline is one of a tab's first three LINES
        hit = next((c for c in r if c and _BLANKET_HOLD.search(c)), None)
        if hit:
            hold_row, hold_cell = i, hit
            break

    if hold_row is not None and (header_row is None or hold_row <= header_row):
        if header_row is not None:
            warnings.append(f"\"{name}\" has both a hold headline and an approved-products header; it is treated as not accepted anywhere.")
        return Category(name=name, kind="blanket_hold", keywords=default_keywords(name), summary=hold_cell, body=body)

    if header_row is not None:
        limit = grid_i if grid_i is not None else len(rows)
        items: list[str] = []
        skipped: list[str] = []
        header_cells = [c for c in rows[header_row] if c]
        for c in header_cells[header_cells.index(header_cell) + 1:]:            # entries sharing the header's row
            c = _strip_link(c)
            if re.search(r"[A-Za-z]", c) and not _ends_list(c) and match_catalog_products(c):   # not a bare column title
                items.append(c)
        stopped_at = None
        blank_run = 0
        for r in rows[header_row + 1:limit]:
            cells = [c for c in r if c]
            if not cells:
                blank_run += 1
                if items and blank_run >= 2:
                    break                               # two blank rows end the list; one spacer row doesn't
                continue
            blank_run = 0
            item = next((c for c in cells if re.search(r"[A-Za-z]", c)), None)   # skips a "1." numbering column
            if item is None:
                continue
            item = _strip_link(item)
            if _ends_list(item):
                if items:
                    stopped_at = item                   # a second list, a heading or a note: not more approved products
                    break
                continue
            if len(item) > MAX_ITEM_CHARS or len(item.split()) > 24:
                skipped.append(item)
                continue
            items.append(item)
        if skipped:
            warnings.append(f"\"{name}\": {len(skipped)} entr{'y was' if len(skipped) == 1 else 'ies were'} too long to read as a product "
                            f"name and skipped ({skipped[0][:60]}…).")
        if items:
            product_map = auto_product_map(items)
            if any(product_map.values()):
                if stopped_at:
                    warnings.append(f"\"{name}\": the approved list was read up to the row \"{stopped_at[:70]}\"; "
                                    "anything after it was not treated as an approved product.")
                cat = Category(name=name, kind="allowlist", keywords=default_keywords(name), items=items,
                               summary=header_cell, body=body)
                cat.product_map = product_map
                return cat
            warnings.append(f"\"{name}\" looks like an approved-products list but none of its entries match a catalog product "
                            "({}), so it is treated as free-text guidance instead of marking every product not accepted.".format(
                                "; ".join(i[:40] for i in items[:3])))
        else:
            warnings.append(f"\"{name}\" looks like an approved-products list but no entries were found under its header, "
                            "so it is treated as free-text guidance (no yes/no verdicts).")

    if grid_i is not None:
        lines = [c for r in rows[:grid_i] for c in r if c]       # text above the grid is guidance too
        platform = ""
        for r in rows[grid_i + 1:]:
            if not any(r):
                continue
            platform = (r[pcol] if len(r) > pcol else "") or platform   # merged platform cells only carry their value once
            guidance = r[pcol + 1] if len(r) > pcol + 1 else ""
            urls = ", ".join(u for u in r[pcol + 2:] if u)
            extras = ", ".join(u for u in r[:pcol] if u)
            if guidance or urls:
                lines.append(f"{platform}: {guidance}" + (f"  [{urls}]" if urls else ""))
            elif extras or (r[pcol] if len(r) > pcol else ""):
                lines.append(extras or r[pcol])                           # a note that sits under the grid
        if not lines:
            return None
        head = rows[grid_i][pcol + 1] if len(rows[grid_i]) > pcol + 1 and rows[grid_i][pcol + 1] else name
        return Category(name=name, kind="guidance", keywords=default_keywords(name),
                        summary=f"Platform-by-platform guidance for {head}.", body="\n".join(lines))

    first = next((ln for ln in flat if len(ln) >= 20), flat[0] if flat else "")
    warnings.append(f"\"{name}\" is free text, so it is shown to the model and the planner as guidance, not as a yes/no verdict.")
    return Category(name=name, kind="guidance", keywords=default_keywords(name),
                    summary=first[:200], body=body)


# ---------------------------------------------------------------------------
# Persistence (fail-safe: Roadblocks must never break because of this table)
# ---------------------------------------------------------------------------

def _row_to_category(row: dict) -> Category:
    return Category(
        name=row["name"], kind=row["kind"], keywords=list(row["keywords"] or []), items=list(row["items"] or []),
        summary=row["summary"] or "", body=row["body"] or "",
        product_map={k: list(v) for k, v in (row["product_map"] or {}).items()},
    )


def load_state() -> dict[str, Any]:
    """{categories, using_builtin, sync, error}. Never raises: on a missing
    table, DB error or an empty table the built-in snapshot applies. `error`
    is True when the read FAILED (as opposed to the table simply being empty),
    so a caller about to overwrite the stored rows can tell the two apart."""
    try:
        with get_connection() as conn:
            rows = conn.execute(
                "SELECT name, kind, keywords, items, summary, body, product_map FROM restriction_categories ORDER BY name"
            ).fetchall()
            sync = conn.execute(
                "SELECT sheet_url, source_title, synced_at, synced_by FROM restriction_sync WHERE id = 'default'"
            ).fetchone()
        if rows:
            return {"categories": [_row_to_category(r) for r in rows], "using_builtin": False,
                    "sync": dict(sync) if sync else None, "error": False}
    except Exception as exc:  # noqa: BLE001
        log.warning("restrictions tables unavailable, using built-in snapshot: %s", exc)
        return {"categories": builtin_categories(), "using_builtin": True, "sync": None, "error": True,
                "missing_table": type(exc).__name__ == "UndefinedTable"}
    return {"categories": builtin_categories(), "using_builtin": True, "sync": None, "error": False}


def carry_over_admin_edits(new: list[Category], previous: list[Category]) -> None:
    """A re-sync replaces every row, so the trigger keywords an admin tuned and
    the sheet-item -> product mappings they corrected would be lost. Keep them
    wherever the same tab / the same sheet entry is still there."""
    old_by_name = {c.name.casefold(): c for c in previous}
    for cat in new:
        old = old_by_name.get(cat.name.casefold())
        if not old:
            continue
        if old.keywords:
            cat.keywords = list(old.keywords)
        if cat.kind == "allowlist" and old.kind == "allowlist":
            for item in cat.items:
                if item in old.product_map and old.product_map[item]:
                    cat.product_map[item] = list(old.product_map[item])


def replace_all(categories: list[Category], *, sheet_url: str, source_title: str, synced_by: str) -> None:
    from psycopg.types.json import Jsonb
    with get_connection() as conn:
        conn.execute("DELETE FROM restriction_categories")
        for c in categories:
            conn.execute(
                """INSERT INTO restriction_categories (name, kind, keywords, items, summary, body, product_map)
                   VALUES (%s, %s, %s, %s, %s, %s, %s)""",
                (c.name, c.kind, c.keywords, c.items, c.summary, c.body, Jsonb(c.product_map)),
            )
        conn.execute(
            """INSERT INTO restriction_sync (id, sheet_url, source_title, synced_at, synced_by)
               VALUES ('default', %s, %s, now(), %s)
               ON CONFLICT (id) DO UPDATE SET sheet_url = EXCLUDED.sheet_url, source_title = EXCLUDED.source_title,
                                              synced_at = now(), synced_by = EXCLUDED.synced_by""",
            (sheet_url, source_title, synced_by),
        )


def update_category(name: str, *, keywords: Any, product_map: dict[str, Any]) -> Optional[Category]:
    """Admin edits: the trigger keywords and the sheet-item -> catalog-product
    mapping (the two things the sync can only guess at)."""
    from app.disclaimers import normalize_keywords
    from psycopg.types.json import Jsonb
    kws = normalize_keywords(keywords)
    if not kws:
        raise ValueError("Keep at least one keyword — otherwise this vertical could never be suggested.")
    clean_map = {str(k): [str(n).strip() for n in (v or []) if str(n).strip()] for k, v in (product_map or {}).items()}
    with get_connection() as conn:
        row = conn.execute(
            """UPDATE restriction_categories SET keywords = %s, product_map = %s, updated_at = now()
               WHERE name = %s RETURNING name, kind, keywords, items, summary, body, product_map""",
            (kws, Jsonb(clean_map), name),
        ).fetchone()
    return _row_to_category(row) if row else None


def seed_builtin_into_db(synced_by: str) -> None:
    """Lets an admin edit keywords/mappings of the built-in snapshot before the
    first real sync, by persisting it as rows."""
    replace_all(builtin_categories(), sheet_url=DEFAULT_SHEET_URL, source_title="Built-in snapshot", synced_by=synced_by)


# ---------------------------------------------------------------------------
# Getting the sheet: Drive export or an uploaded .xlsx
# ---------------------------------------------------------------------------

def extract_file_id(url_or_id: str) -> str:
    s = (url_or_id or "").strip()
    m = re.search(r"/d/([a-zA-Z0-9_-]{20,})", s)
    if m:
        return m.group(1)
    if re.fullmatch(r"[a-zA-Z0-9_-]{20,}", s):
        return s
    raise ValueError("That doesn't look like a Google Sheets link.")


_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def fetch_sheet_xlsx(file_id: str) -> tuple[bytes, str]:
    """(xlsx bytes, sheet title) via the app's existing Google Drive
    connection. ValueError with a planner-readable reason on any failure."""
    from app.services import drive_uploader
    if not drive_uploader._is_configured():
        raise ValueError("Google Drive isn't configured for this app — upload an .xlsx export of the sheet instead.")
    creds = drive_uploader._load_credentials()
    if creds is None:
        raise ValueError("Google Drive isn't connected yet — connect it from the Generate step, or upload an .xlsx export instead.")
    try:
        from googleapiclient.discovery import build
        service = build("drive", "v3", credentials=creds, cache_discovery=False)
        title = service.files().get(fileId=file_id, fields="name", supportsAllDrives=True).execute().get("name", "")
        data = service.files().export(fileId=file_id, mimeType=_XLSX).execute()
    except Exception as exc:  # noqa: BLE001
        raise ValueError(
            "The connected Google account couldn't open that sheet "
            f"({getattr(exc, 'reason', None) or exc}). Share it with that account, or upload an .xlsx export instead."
        ) from exc
    return bytes(data), title


def decode_upload(data_base64: str) -> bytes:
    try:
        return base64.b64decode(data_base64, validate=False)
    except Exception as exc:  # noqa: BLE001
        raise ValueError("Couldn't read the uploaded file.") from exc


# ---------------------------------------------------------------------------
# Verdicts
# ---------------------------------------------------------------------------

def _worst(statuses: list[str]) -> str:
    return min(statuses, key=lambda s: STATUS_ORDER.get(s, 9)) if statuses else "none"


def evaluate_product(product_name: str, family: Optional[str], categories: list[Category],
                     aliases: Optional[Iterable[str]] = None, custom: bool = False) -> dict[str, Any]:
    """{"verdict": allowed|not_allowed|guidance|n/a|none, "checks": [...]} for
    one product against the confirmed verticals. `aliases` are the other names
    the product answers to (its renamed display name / stable raw name) — the
    sheet mapping is keyed by raw catalog names, so a renamed product is looked
    up under all of them. `custom` (a product the catalog has no built-in entry
    for) or `family` None (not in the catalog at all) can't be judged against
    the sheet's list, so it reads as guidance — "ask Sales Planning" — rather
    than a hard NOT ACCEPTED the sheet never actually issued."""
    names = {product_name, *(aliases or [])}
    checks: list[dict[str, str]] = []
    for cat in categories:
        if cat.kind == "blanket_hold":
            checks.append({"category": cat.name, "status": "not_allowed", "note": cat.summary})
        elif family in NON_MEDIA_FAMILIES:
            checks.append({"category": cat.name, "status": "n/a",
                           "note": f"Not ad inventory, so {cat.name}'s product rules don't apply."})
        elif cat.kind == "allowlist":
            product_map = cat.product_map or auto_product_map(cat.items)
            listed = next((item for item, mapped in product_map.items() if names & set(mapped)), None)
            if listed:
                checks.append({"category": cat.name, "status": "allowed", "note": f"Approved on the sheet as: {listed}"})
            elif custom or family is None:
                checks.append({"category": cat.name, "status": "guidance",
                               "note": f"A custom product the {cat.name} sheet doesn't list — confirm with Sales Planning "
                                       "(or map it in Admin → Restrictions)."})
            else:
                checks.append({"category": cat.name, "status": "not_allowed",
                               "note": f"Not on the approved product list for {cat.name} ads."})
        else:
            checks.append({"category": cat.name, "status": "guidance", "note": cat.summary})
    return {"verdict": _worst([c["status"] for c in checks]), "checks": checks}


def build_matrix(product_names: list[str], category_names: Optional[list[str]]) -> dict[str, Any]:
    """Everything Roadblocks needs: the confirmed Category objects, a verdict
    per product, and where the data came from."""
    from app.catalog import CATALOG, by_name, name_variants
    state = load_state()
    wanted = {n.strip().lower() for n in (category_names or []) if n and n.strip()}
    selected = [c for c in state["categories"] if c.name.lower() in wanted]
    found = {c.name.lower() for c in selected}
    unresolved = sorted(n for n in (category_names or []) if n and n.strip() and n.strip().lower() not in found)
    builtin_names = {p.name for p in CATALOG}
    verdicts: dict[str, dict] = {}
    for name in dict.fromkeys(product_names):
        product = by_name(name)
        variants = name_variants(name)
        verdicts[name] = evaluate_product(
            name, product.family if product else None, selected,
            aliases=variants, custom=product is not None and not (variants & builtin_names))
    sync = state["sync"] or {}
    return {
        "categories": selected, "verdicts": verdicts, "using_builtin": state["using_builtin"],
        "synced_at": sync.get("synced_at").isoformat() if sync.get("synced_at") else None,
        "unresolved": unresolved,
    }


def matrix_prompt_block(matrix: dict[str, Any]) -> str:
    """The authoritative section of the Roadblocks prompt."""
    cats: list[Category] = matrix["categories"]
    if not cats:
        return (
            "## RESTRICTED VERTICALS — NONE CONFIRMED\n"
            "The planner did not confirm that this client falls under any of Entravision's restricted verticals. "
            "Do NOT claim a product is unavailable, not allowed or not feasible because of the client's industry "
            "unless a platform's own current policy says so, and never cite Entravision internal flags."
        )
    when = f" (synced {matrix['synced_at'][:10]})" if matrix.get("synced_at") else " (built-in snapshot — not yet synced)" if matrix.get("using_builtin") else ""
    lines = [
        f"## RESTRICTED VERTICALS MATRIX — AUTHORITATIVE{when}",
        "Source: Entravision's own \"Restricted Verticals by Platforms\" sheet. The planner confirmed this client falls under: "
        + ", ".join(c.name for c in cats) + ".",
        "This sheet is the ONLY source of truth for whether those verticals are allowed on a product. RULES:",
        "- A product marked ALLOWED below IS allowed. Never say or imply it is not allowed, not feasible, prohibited or restricted for that vertical.",
        "- A product marked NOT ACCEPTED is not accepted by Entravision for that vertical — state that plainly with the reason given; do not debate or soften it.",
        "- GUIDANCE means the sheet gives platform-specific advice instead of a yes/no — apply the sheet text below, and don't contradict it.",
        "- Web research may add ADDITIONAL practical caveats (ad-review specifics, creative requirements, approval lead times). Start each such issue with \"Additional:\".",
        "",
        "Per-product verdicts (computed from the sheet):",
    ]
    label = {"allowed": "ALLOWED", "not_allowed": "NOT ACCEPTED", "guidance": "SEE GUIDANCE", "n/a": "n/a (not ad inventory)"}
    for name, v in matrix["verdicts"].items():
        parts = [f"{c['category']} — {label.get(c['status'], c['status'])} ({c['note']})" for c in v["checks"]]
        lines.append(f"- {name}: " + "; ".join(parts))
    lines.append("")
    lines.append("Sheet text per vertical (verbatim; very long tabs are trimmed):")
    for c in cats:
        body = c.body if len(c.body) <= MAX_BODY_IN_PROMPT else c.body[:MAX_BODY_IN_PROMPT].rsplit(" ", 1)[0] + " …"
        lines.append(f"### {c.name}\n{body}")
    return "\n".join(lines)
